# -*- coding: utf-8 -*-
"""本地数据检查的无网络单元测试。"""
import hashlib
import json
import socket
import tempfile
import unittest
import zipfile
from datetime import date, datetime, time, timedelta
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from app.task_page import TaskPage
from tools.data_check import core
from tools.data_check.page import DataCheckPage


def comment_record(rpid=1, message="正常评论"):
    return {
        "rpid": rpid, "parent": 0, "root": 0, "is_main": True,
        "uname": "测试用户", "mid": 100, "sex": "保密", "level": 6,
        "vip": False, "message": message, "like": 0, "rcount": 0,
        "ctime": 1700000000, "location": "", "is_top": False,
    }


def video_record(bvid="BV1Demo", fetched_at=1700000000, **extra):
    value = {
        "bvid": bvid, "aid": 1, "title": "测试视频", "owner": "测试UP主",
        "owner_mid": 2, "tname": "测试", "pubdate": 1690000000,
        "duration": 60, "view": 0, "danmaku": 0, "reply": 0,
        "favorite": 0, "coin": 0, "share": 0, "like": False,
        "fetched_at": fetched_at,
    }
    value.update(extra)
    return value


class _FakeItem:
    def __init__(self, text):
        self.text = text
        self.tooltip = ""

    def setToolTip(self, value):
        self.tooltip = value


class _FakeList:
    def __init__(self):
        self.items = []
        self.focused = False

    def clear(self):
        self.items.clear()

    def addItem(self, value):
        self.items.append(_FakeItem(str(value)))

    def count(self):
        return len(self.items)

    def item(self, index):
        return self.items[index]

    def setFocus(self):
        self.focused = True


class _FakePathRow:
    def __init__(self):
        self.value = ""

    def set_value(self, value):
        self.value = str(value)


class DataCheckCoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="data-check-test-")
        self.root = Path(self.temp.name)
        self.out = self.root / "reports"

    def tearDown(self):
        self.temp.cleanup()

    def write_jsonl(self, name, values, encoding="utf-8"):
        path = self.root / name
        with path.open("w", encoding=encoding, newline="\n") as handle:
            for value in values:
                if isinstance(value, str):
                    handle.write(value)
                else:
                    handle.write(json.dumps(value, ensure_ascii=False))
                handle.write("\n")
        return path

    def write_jsonl_bytes(self, name, value):
        path = self.root / name
        path.write_bytes(value)
        return path

    def write_workbook(self, name, sheets):
        path = self.root / name
        workbook = Workbook()
        first = workbook.active
        workbook.remove(first)
        for sheet_name, rows in sheets:
            sheet = workbook.create_sheet(sheet_name)
            for row in rows:
                sheet.append(row)
        workbook.save(path)
        return path

    def run_check(self, paths, **kwargs):
        return core.run_check([str(path) for path in paths], str(self.out), **kwargs)

    def categories(self, result):
        return [item["category"] for item in result["issues"]]

    def snapshot_source(self, path):
        stat = path.stat()
        return hashlib.sha256(path.read_bytes()).hexdigest(), stat.st_size, stat.st_mtime_ns

    def assert_source_unchanged(self, path, before):
        after = path.stat()
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before[0])
        self.assertEqual(after.st_size, before[1])
        self.assertEqual(after.st_mtime_ns, before[2])

    def test_normal_comment_jsonl_has_no_false_positive(self):
        path = self.write_jsonl("comments.jsonl", [comment_record()])
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["stats"]["record_count"], 1)
        self.assertEqual(result["stats"]["total_issues"], 0)
        self.assertEqual(result["files"][0]["structure"], "已知评论 JSONL")
        self.assertEqual(result["source_files"][0], {
            "path": str(path.resolve(strict=False)),
            "sha256": before[0], "size": before[1], "mtime_ns": before[2],
        })
        self.assert_source_unchanged(path, before)

    def test_comment_rpid_duplicate(self):
        path = self.write_jsonl("duplicate-comments.jsonl", [
            comment_record(1), comment_record("1"),
        ])
        result = self.run_check([path])
        self.assertEqual(result["stats"]["duplicates"], 1)
        self.assertIn("rpid=1", result["issues"][0]["key"])

    def test_comment_missing_fields_and_blank_message(self):
        value = comment_record(message="   ")
        value.pop("location")
        path = self.write_jsonl("missing-comments.jsonl", [value])
        result = self.run_check([path])
        self.assertEqual(result["stats"]["missing"], 1)
        self.assertEqual(result["stats"]["blanks"], 1)
        self.assertIn("缺失字段", self.categories(result))
        self.assertIn("空白内容", self.categories(result))

    def test_jsonl_blank_bad_json_and_non_object(self):
        path = self.root / "bad-lines.jsonl"
        path.write_text(
            json.dumps(comment_record(), ensure_ascii=False) + "\n\n{broken\n[]\n",
            encoding="utf-8",
        )
        result = self.run_check([path])
        self.assertEqual(result["stats"]["record_count"], 1)
        self.assertEqual(result["stats"]["blanks"], 1)
        self.assertEqual(result["stats"]["formats"], 2)
        self.assertIn("JSON语法错误", self.categories(result))
        self.assertIn("非对象记录", self.categories(result))

    def test_jsonl_non_utf8_is_damaged(self):
        path = self.write_jsonl_bytes("not-utf8.jsonl", b'{"x": 1}\n\xff')
        result = self.run_check([path])
        self.assertEqual(result["stats"]["damaged_files"], 1)
        self.assertEqual(result["files"][0]["conclusion"], "文件损坏，未能完整读取")

    def test_bom_only_jsonl_is_one_blank_issue_not_damaged(self):
        path = self.write_jsonl_bytes("bom-only.jsonl", b"\xef\xbb\xbf")
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["stats"]["blanks"], 1)
        self.assertEqual(result["stats"]["total_issues"], 1)
        self.assertEqual(result["stats"]["damaged_files"], 0)
        self.assertEqual(result["issues"][0]["category"], "空白内容")
        self.assertEqual(result["issues"][0]["location"], "文件")
        self.assertEqual(result["issues"][0]["field"], "整文件")
        self.assert_source_unchanged(path, before)

    def test_bom_with_three_blank_lines_counts_only_physical_lines(self):
        path = self.write_jsonl_bytes("bom-and-blanks.jsonl", b"\xef\xbb\xbf\n  \n\t\n")
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["stats"]["blanks"], 3)
        self.assertEqual(result["stats"]["total_issues"], 3)
        self.assertEqual(result["stats"]["damaged_files"], 0)
        self.assertFalse(any(
            item["location"] == "文件" and item["field"] == "整文件"
            for item in result["issues"]
        ))
        self.assert_source_unchanged(path, before)

    def test_zero_byte_jsonl_is_one_blank_issue_not_damaged(self):
        path = self.write_jsonl_bytes("zero-byte.jsonl", b"")
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["stats"]["blanks"], 1)
        self.assertEqual(result["stats"]["total_issues"], 1)
        self.assertEqual(result["stats"]["damaged_files"], 0)
        self.assertEqual(result["files"][0]["conclusion"], "发现 1 项问题")
        self.assertEqual(result["issues"][0]["category"], "空白内容")
        self.assertEqual(result["issues"][0]["location"], "文件")
        self.assertEqual(result["issues"][0]["field"], "整文件")
        self.assert_source_unchanged(path, before)

    def test_whitespace_only_jsonl_counts_lines_without_whole_file_issue(self):
        path = self.write_jsonl_bytes("whitespace-only.jsonl", b"\n  \n\t\n")
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["stats"]["blanks"], 3)
        self.assertEqual(result["stats"]["total_issues"], 3)
        self.assertEqual(result["stats"]["damaged_files"], 0)
        self.assertFalse(any(
            item["location"] == "文件" and item["field"] == "整文件"
            for item in result["issues"]
        ))
        self.assert_source_unchanged(path, before)

    def test_video_different_timestamps_are_not_duplicates(self):
        path = self.write_jsonl("snapshots.jsonl", [
            video_record(fetched_at=1), video_record(fetched_at=2),
        ])
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["stats"]["duplicates"], 0)
        self.assertEqual(result["files"][0]["structure"], "已知视频快照 JSONL")
        self.assert_source_unchanged(path, before)

    def test_video_same_bvid_and_timestamp_is_duplicate(self):
        path = self.write_jsonl("duplicate-snapshots.jsonl", [
            video_record(fetched_at=1), video_record(fetched_at=1),
        ])
        result = self.run_check([path])
        self.assertEqual(result["stats"]["duplicates"], 1)
        self.assertIn("bvid=BV1Demo", result["issues"][0]["key"])

    def test_invalid_fetched_at_values_are_format_issues(self):
        invalid_values = ["not-a-time", "ordinary text", 0, -1, True, False, [], {}, 1.5]
        path = self.write_jsonl("invalid-times.jsonl", [
            video_record(bvid=f"BVInvalid{index}", fetched_at=value)
            for index, value in enumerate(invalid_values, 1)
        ])
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["stats"]["formats"], len(invalid_values))
        self.assertEqual(result["stats"]["blanks"], 0)
        self.assertEqual(result["stats"]["duplicates"], 0)
        self.assertTrue(all(
            item["field"] == "fetched_at" and item["category"] == "异常格式"
            for item in result["issues"]
        ))
        self.assert_source_unchanged(path, before)

    def test_positive_integer_timestamp_and_digit_string_share_duplicate_key(self):
        path = self.write_jsonl("normalised-times.jsonl", [
            video_record(fetched_at=1700000000),
            video_record(fetched_at="1700000000"),
        ])
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["stats"]["formats"], 0)
        self.assertEqual(result["stats"]["duplicates"], 1)
        self.assert_source_unchanged(path, before)

    def test_partial_first_comment_record_does_not_lock_unknown_structure(self):
        path = self.write_jsonl("neutral-comments-data.jsonl", [
            {"rpid": 1}, comment_record(2),
        ])
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["files"][0]["structure"], "已知评论 JSONL")
        self.assertEqual(result["stats"]["missing"], len(core.COMMENT_FIELDS) - 1)
        self.assert_source_unchanged(path, before)

    def test_partial_first_video_record_does_not_lock_unknown_structure(self):
        path = self.write_jsonl("neutral-video-data.jsonl", [
            {"bvid": "BVPartial"}, video_record(bvid="BVComplete"),
        ])
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["files"][0]["structure"], "已知视频快照 JSONL")
        self.assertEqual(result["stats"]["missing"], len(core.VIDEO_FIELDS) - 1)
        self.assert_source_unchanged(path, before)

    def test_split_video_markers_across_objects_stay_unknown(self):
        path = self.write_jsonl("split-markers.jsonl", [
            {"title": "A"}, {"view": 1},
        ])
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["files"][0]["structure"], "未知结构，仅完成通用检查")
        self.assertEqual(result["stats"]["missing"], 0)
        self.assertNotIn("缺失字段", self.categories(result))
        self.assert_source_unchanged(path, before)

    def test_mixed_comment_and_video_strong_evidence_stays_unknown(self):
        path = self.write_jsonl("mixed-known-structures.jsonl", [
            comment_record(), video_record(),
        ])
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["files"][0]["structure"], "未知结构，仅完成通用检查")
        self.assertEqual(result["stats"]["missing"], 0)
        self.assertNotIn("缺失字段", self.categories(result))
        self.assert_source_unchanged(path, before)

    def test_single_video_strong_record_cannot_hijack_mostly_unknown_file(self):
        path = self.write_jsonl("mostly-unknown-with-video.jsonl", [
            *({"x": index} for index in range(9)),
            {"bvid": "BV1Demo", "fetched_at": 1},
        ])
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["files"][0]["structure"], "未知结构，仅完成通用检查")
        self.assertEqual(result["stats"]["missing"], 0)
        self.assertNotIn("缺失字段", self.categories(result))
        self.assert_source_unchanged(path, before)

    def test_single_comment_strong_record_cannot_hijack_mostly_unknown_file(self):
        path = self.write_jsonl("mostly-unknown-with-comment.jsonl", [
            *({"x": index} for index in range(9)),
            comment_record(),
        ])
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["files"][0]["structure"], "未知结构，仅完成通用检查")
        self.assertEqual(result["stats"]["missing"], 0)
        self.assertNotIn("缺失字段", self.categories(result))
        self.assert_source_unchanged(path, before)

    def test_title_and_view_only_object_stays_unknown(self):
        path = self.write_jsonl("weak-video-object.jsonl", [
            {"title": "A", "view": 1},
        ])
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["files"][0]["structure"], "未知结构，仅完成通用检查")
        self.assertEqual(result["stats"]["missing"], 0)
        self.assertEqual(result["stats"]["total_issues"], 0)
        self.assert_source_unchanged(path, before)

    def test_overlong_string_rpid_is_format_issue(self):
        path = self.write_jsonl("overlong-rpid.jsonl", [
            comment_record(rpid="9" * 5000),
        ])
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["stats"]["formats"], 1)
        self.assertEqual(result["stats"]["damaged_files"], 0)
        self.assertEqual(result["issues"][0]["field"], "rpid")
        self.assert_source_unchanged(path, before)

    def test_overlong_string_fetched_at_is_format_issue(self):
        path = self.write_jsonl("overlong-fetched-at.jsonl", [
            video_record(fetched_at="9" * 5000),
        ])
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["stats"]["formats"], 1)
        self.assertEqual(result["stats"]["damaged_files"], 0)
        self.assertEqual(result["issues"][0]["field"], "fetched_at")
        self.assert_source_unchanged(path, before)

    def test_overlong_json_integer_is_syntax_issue_and_next_line_is_checked(self):
        overlong_line = '{"rpid": ' + ("9" * 5000) + "}\n"
        normal_line = json.dumps(comment_record(2), ensure_ascii=False) + "\n"
        path = self.write_jsonl_bytes(
            "overlong-json-number.jsonl", (overlong_line + normal_line).encode("utf-8")
        )
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["files"][0]["structure"], "已知评论 JSONL")
        self.assertEqual(result["stats"]["formats"], 1)
        self.assertEqual(result["stats"]["record_count"], 1)
        self.assertEqual(result["stats"]["damaged_files"], 0)
        self.assertEqual(result["issues"][0]["category"], "JSON语法错误")
        with zipfile.ZipFile(result["report"]) as archive:
            report_xml = b"".join(
                archive.read(name) for name in archive.namelist()
                if name.startswith("xl/worksheets/")
            )
        self.assertNotIn(("9" * 100).encode(), report_xml)
        self.assert_source_unchanged(path, before)

    def test_zero_and_false_are_not_blank(self):
        path = self.write_jsonl("zero-values.jsonl", [video_record()])
        result = self.run_check([path])
        self.assertEqual(result["stats"]["blanks"], 0)
        self.assertEqual(result["stats"]["total_issues"], 0)

    def test_unknown_jsonl_is_conservative_but_checks_duplicates_and_shape(self):
        path = self.write_jsonl("unknown.jsonl", [
            {"x": 1}, {"x": 1}, {"y": 2},
        ])
        result = self.run_check([path])
        self.assertEqual(result["files"][0]["structure"], "未知结构，仅完成通用检查")
        self.assertEqual(result["stats"]["duplicates"], 1)
        self.assertEqual(result["stats"]["formats"], 1)
        self.assertNotIn("缺失字段", self.categories(result))

    def test_normal_comment_excel_with_leading_margin_column(self):
        headers = ["rpid", "评论内容", "用户昵称"]
        path = self.write_workbook("comments.xlsx", [
            ("全量评论", [[None, *headers], [None, 1, "正常", "用户"]]),
        ])
        result = self.run_check([path])
        self.assertEqual(result["stats"]["total_issues"], 0)
        self.assertEqual(result["files"][0]["structure"], "已知评论 Excel")

    def test_normal_video_excel(self):
        path = self.write_workbook("videos.xlsx", [
            ("视频总表", [[None, "BV号", "标题"], [None, "BV1Demo", "视频"]]),
        ])
        result = self.run_check([path])
        self.assertEqual(result["stats"]["total_issues"], 0)
        self.assertEqual(result["files"][0]["structure"], "已知视频 Excel")

    def test_excel_missing_key_header(self):
        path = self.write_workbook("missing-header.xlsx", [
            ("全量评论", [[None, "rpid"], [None, 1]]),
        ])
        result = self.run_check([path])
        self.assertEqual(result["stats"]["missing"], 1)
        self.assertEqual(result["stats"]["blanks"], 0)

    def test_excel_duplicate_record(self):
        path = self.write_workbook("duplicate.xlsx", [
            ("视频总表", [[None, "BV号", "标题"],
                         [None, "BV1Demo", "视频"],
                         [None, "BV1Demo", "视频"]]),
        ])
        result = self.run_check([path])
        self.assertEqual(result["stats"]["duplicates"], 1)

    def test_damaged_xlsx_is_result_and_does_not_abort(self):
        path = self.root / "damaged.xlsx"
        path.write_bytes(b"not an xlsx")
        result = self.run_check([path])
        self.assertEqual(result["stats"]["damaged_files"], 1)
        self.assertTrue(Path(result["report"]).is_file())

    def test_unknown_excel_uses_generic_duplicate_check(self):
        path = self.write_workbook("unknown.xlsx", [
            ("数据", [["列A", "列B"], [1, "x"], [1, "x"]]),
        ])
        result = self.run_check([path])
        self.assertEqual(result["files"][0]["structure"], "未知结构，仅完成通用检查")
        self.assertEqual(result["stats"]["duplicates"], 1)

    def test_unknown_excel_accepts_common_date_and_time_cell_types(self):
        path = self.write_workbook("dates.xlsx", [
            ("数据", [
                ["日期时间", "日期", "时间", "时长", "空值", "文本", "数字", "布尔值"],
                [datetime(2026, 9, 8, 12, 30, 45), date(2026, 9, 8),
                 time(12, 30, 45), timedelta(days=2, seconds=3), None, "正常", 42, True],
            ]),
        ])
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["files"][0]["structure"], "未知结构，仅完成通用检查")
        self.assertEqual(result["stats"]["damaged_files"], 0)
        self.assertEqual(result["stats"]["total_issues"], 0)
        self.assertTrue(Path(result["report"]).is_file())
        self.assert_source_unchanged(path, before)

    def test_unknown_excel_duplicate_rows_with_dates_are_detected(self):
        values = [date(2026, 9, 8), time(12, 30), timedelta(hours=1)]
        path = self.write_workbook("duplicate-dates.xlsx", [
            ("数据", [["日期", "时间", "时长"], values, list(values)]),
        ])
        before = self.snapshot_source(path)
        result = self.run_check([path])
        self.assertEqual(result["stats"]["damaged_files"], 0)
        self.assertEqual(result["stats"]["duplicates"], 1)
        self.assert_source_unchanged(path, before)

    def test_multiple_files_continue_after_damaged_file(self):
        damaged = self.root / "broken.xlsx"
        damaged.write_bytes(b"broken")
        good = self.write_jsonl("good.jsonl", [comment_record()])
        result = self.run_check([damaged, good])
        self.assertEqual(result["stats"]["file_count"], 2)
        self.assertEqual(result["stats"]["damaged_files"], 1)
        self.assertEqual(result["stats"]["readable_files"], 1)
        self.assertEqual(result["stats"]["record_count"], 1)

    def test_cancellation_is_deterministic_and_leaves_no_report(self):
        path = self.write_jsonl("cancel.jsonl", [comment_record()])
        result = self.run_check([path], cancel=lambda: True)
        self.assertTrue(result["stats"]["cancelled"])
        self.assertEqual(list(self.out.glob("*.xlsx")), [])
        self.assertEqual(list(self.out.glob("*.tmp")), [])

    def test_issue_detail_cap_preserves_full_counts(self):
        path = self.write_jsonl("many-issues.jsonl", [{"x": 1}] * 16)
        result = self.run_check([path], max_issue_details=3)
        self.assertEqual(result["stats"]["total_issues"], 15)
        self.assertEqual(len(result["issues"]), 3)
        self.assertTrue(result["stats"]["details_truncated"])
        workbook = load_workbook(result["report"], read_only=True, data_only=True)
        overview = list(workbook["检查概览"].iter_rows(values_only=True))
        workbook.close()
        self.assertTrue(any("仅展示前 3 条" in str(row[1]) for row in overview))

    def test_report_has_three_sheets_and_correct_numeric_statistics(self):
        path = self.write_jsonl("report.jsonl", [comment_record(1), comment_record(1)])
        result = self.run_check([path])
        workbook = load_workbook(result["report"], read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["检查概览", "文件汇总", "问题明细"])
        overview = list(workbook["检查概览"].iter_rows(values_only=True))
        workbook.close()
        values = {row[0]: row[1] for row in overview[1:]}
        self.assertEqual(values["文件数量"], 1)
        self.assertEqual(values["重复数量"], 1)
        self.assertEqual(values["总问题数量"], 1)

    def test_report_does_not_copy_sensitive_or_full_record_text(self):
        secret = "cookie-secret-should-not-appear"
        record = comment_record(message=f"Cookie={secret}; Authorization=Bearer auth-secret")
        path = self.write_jsonl("sensitive.jsonl", [record])
        result = self.run_check([path])
        with zipfile.ZipFile(result["report"]) as archive:
            raw = b"".join(archive.read(name) for name in archive.namelist()
                           if name.startswith("xl/worksheets/"))
        self.assertNotIn(secret.encode(), raw)
        self.assertNotIn(b"auth-secret", raw)
        self.assertNotIn(b"Cookie=", raw)
        self.assertNotIn(b"Authorization", raw)

    def test_source_sha_size_and_mtime_are_unchanged(self):
        path = self.write_jsonl("immutable.jsonl", [comment_record()])
        before_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        before_stat = path.stat()
        result = self.run_check([path])
        after_stat = path.stat()
        self.assertTrue(Path(result["report"]).exists())
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before_hash)
        self.assertEqual(after_stat.st_size, before_stat.st_size)
        self.assertEqual(after_stat.st_mtime_ns, before_stat.st_mtime_ns)

    def test_report_does_not_overwrite_existing_or_input_named_report(self):
        class FrozenDateTime:
            @classmethod
            def now(cls):
                return datetime(2026, 9, 8, 19, 30, 0)

        expected = self.out / "数据检查报告_20260908_193000.xlsx"
        self.out.mkdir(parents=True)
        expected.write_bytes(b"existing-report")
        source = self.write_jsonl("source.jsonl", [comment_record()])
        with patch.object(core, "datetime", FrozenDateTime):
            result = self.run_check([source])
        self.assertNotEqual(Path(result["report"]), expected)
        self.assertEqual(expected.read_bytes(), b"existing-report")

        source_xlsx = self.out / "数据检查报告_20260908_193000_1.xlsx"
        workbook = Workbook()
        workbook.save(source_xlsx)
        before = source_xlsx.read_bytes()
        with patch.object(core, "datetime", FrozenDateTime):
            result = self.run_check([source_xlsx])
        self.assertNotEqual(Path(result["report"]), source_xlsx)
        self.assertEqual(source_xlsx.read_bytes(), before)

    def test_validation_rejects_wrong_extension_and_missing_file(self):
        txt = self.root / "data.txt"
        txt.write_text("x", encoding="utf-8")
        with self.assertRaises(ValueError):
            core.validate_inputs([str(txt)], str(self.out))
        with self.assertRaises(ValueError):
            core.validate_inputs([str(self.root / "missing.jsonl")], str(self.out))

    def test_no_network_calls_are_possible_in_core_check(self):
        path = self.write_jsonl("offline.jsonl", [comment_record()])
        with patch.object(socket, "create_connection", side_effect=AssertionError("network")), \
                patch("urllib.request.urlopen", side_effect=AssertionError("network")):
            result = self.run_check([path])
        self.assertEqual(result["stats"]["total_issues"], 0)


class DataCheckPageTests(unittest.TestCase):
    def test_reusable_params_only_fill_fields_and_do_not_start(self):
        page = DataCheckPage.__new__(DataCheckPage)
        page.file_list = _FakeList()
        page.out_row = _FakePathRow()
        with patch.object(TaskPage, "on_start") as start:
            page.apply_reusable_params({
                "files": ["D:\\input\\one.jsonl", "D:\\input\\two.xlsx"],
                "out_dir": "D:\\reports",
            })
        start.assert_not_called()
        self.assertEqual([item.text for item in page.file_list.items], [
            "D:\\input\\one.jsonl", "D:\\input\\two.xlsx",
        ])
        self.assertEqual(page.out_row.value, "D:\\reports")
        self.assertTrue(page.file_list.focused)


if __name__ == "__main__":
    unittest.main()
