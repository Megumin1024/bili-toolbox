# -*- coding: utf-8 -*-
"""关系分析工具的无网络单元测试。

覆盖任务卡要求的场景：列名映射 / 编码回退 / 去重 / 缺列报错 / 超限报错 /
空文件 / 坏行跳过计数；集合运算边界（空清单 / 全互关 / 无交集）；差异
（新增/取关/未变）；按月分布；单时点退化（时点 2 留空）；取消不落半截 Excel。
"""
import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from tools import TOOLS
from tools.relation_analysis import core
from tools.relation_analysis.core import (
    Cancelled,
    _s,
    compute_analysis,
    load_roster,
    parse_time_value,
    validate_inputs,
)
from tools.relation_analysis.page import RelationAnalysisPage
from tools.relation_analysis.pipeline import run_pipeline


def _csv_text(rows, header="mid,uname,mtime"):
    # 用 \n 连接：write_text 在 Windows 上会翻译成 \r\n，避免写出 \r\r\n
    lines = [header] + rows
    return "\n".join(lines) + "\n"


class BaseTestCase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory(prefix="relation-analysis-test-")
        self.root = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def write_file(self, name, text, encoding="utf-8"):
        path = self.root / name
        if isinstance(text, bytes):
            path.write_bytes(text)
        else:
            path.write_text(text, encoding=encoding)
        return path


class TimeParsingTests(BaseTestCase):
    def test_epoch_int_and_float(self):
        self.assertEqual(parse_time_value(1_700_000_000),
                         datetime.fromtimestamp(1_700_000_000))
        self.assertEqual(parse_time_value(1_700_000_000.5),
                         datetime.fromtimestamp(1_700_000_000.5))

    def test_epoch_numeric_string(self):
        self.assertEqual(parse_time_value(" 1700000000 "),
                         datetime.fromtimestamp(1_700_000_000))

    def test_date_and_datetime_strings(self):
        self.assertEqual(parse_time_value("2024-05-01"), datetime(2024, 5, 1))
        self.assertEqual(parse_time_value("2024-05-01 12:30:45"),
                         datetime(2024, 5, 1, 12, 30, 45))

    def test_invalid_values_return_none(self):
        for value in (None, True, "", "   ", "not-a-time", "2024/05/01",
                      "9" * 40):
            self.assertIsNone(parse_time_value(value), repr(value))

    def test_s_dash_fallback(self):
        self.assertEqual(_s(None), "—")
        self.assertEqual(_s("   "), "—")
        self.assertEqual(_s("张三"), "张三")


class CsvParsingTests(BaseTestCase):
    def load(self, text, encoding="utf-8", name="fans.csv"):
        path = self.write_file(name, text, encoding=encoding)
        return load_roster(path, "时点 1 · 粉丝清单")

    def test_basic_column_mapping(self):
        roster = self.load(_csv_text([
            "1001,张三,1700000000",
            "1002,李四,2024-05-01 08:00:00",
        ]))
        self.assertEqual(len(roster.rows), 2)
        self.assertEqual(roster.rows[0]["mid"], 1001)
        self.assertEqual(roster.rows[0]["name"], "张三")
        self.assertEqual(roster.rows[0]["time"],
                         datetime.fromtimestamp(1_700_000_000))
        self.assertEqual(roster.rows[1]["time"], datetime(2024, 5, 1, 8, 0, 0))
        self.assertTrue(roster.has_time)

    def test_synonym_columns(self):
        roster = self.load(_csv_text([
            "1001,昵称甲,2024-01-02",
            "1002,昵称乙,2024-01-03",
        ], header="uid,昵称,关注时间"))
        self.assertEqual([row["mid"] for row in roster.rows], [1001, 1002])
        self.assertEqual(roster.rows[0]["name"], "昵称甲")
        self.assertEqual(roster.rows[0]["time"], datetime(2024, 1, 2))
        self.assertTrue(roster.has_time)

    def test_time_column_absent(self):
        roster = self.load(_csv_text(["1001,张三"], header="mid,name"))
        self.assertFalse(roster.has_time)
        self.assertIsNone(roster.rows[0]["time"])

    def test_encoding_fallback_utf8_sig_and_gbk(self):
        bom_text = _csv_text(["1001,张三"])
        roster = self.load(bom_text.encode("utf-8-sig"))
        self.assertEqual(roster.rows[0]["name"], "张三")
        roster = self.load(_csv_text(["1001,李四"]).encode("gbk"),
                           name="gbk.csv")
        self.assertEqual(roster.rows[0]["name"], "李四")

    def test_undecodable_bytes_raise(self):
        with self.assertRaises(ValueError) as ctx:
            self.load(b"\xff\xfe\x81\x81", name="bad.csv")
        self.assertIn("编码", str(ctx.exception))

    def test_missing_mid_column_raises(self):
        with self.assertRaises(ValueError) as ctx:
            self.load(_csv_text(["1001,张三"], header="uid2,name"))
        self.assertIn("mid", str(ctx.exception))

    def test_dedup_keeps_first(self):
        roster = self.load(_csv_text([
            "1001,第一,1700000000",
            "1001,第二,1700000001",
        ]))
        self.assertEqual(len(roster.rows), 1)
        self.assertEqual(roster.rows[0]["name"], "第一")
        self.assertEqual(roster.duplicates, 1)

    def test_bad_rows_skipped_and_counted(self):
        roster = self.load(_csv_text([
            "1001,张三,1700000000",
            ",缺mid,1700000000",
            "abc,非数字,1700000000",
            "1002,李四,1700000000",
        ]))
        self.assertEqual([row["mid"] for row in roster.rows], [1001, 1002])
        self.assertEqual(roster.bad_rows, 2)
        self.assertEqual(roster.data_rows, 4)

    def test_time_parse_failure_keeps_row(self):
        roster = self.load(_csv_text(["1001,张三,乱七八糟"]))
        self.assertEqual(len(roster.rows), 1)
        self.assertIsNone(roster.rows[0]["time"])

    def test_blank_lines_skipped_silently(self):
        roster = self.load(_csv_text(["", "1001,张三", "", "1002,李四"]))
        self.assertEqual(len(roster.rows), 2)
        self.assertEqual(roster.bad_rows, 0)

    def test_empty_file_raises(self):
        with self.assertRaises(ValueError) as ctx:
            self.load("")
        self.assertIn("空", str(ctx.exception))

    def test_header_only_file_is_valid_empty(self):
        roster = self.load(_csv_text([], header="mid,uname,mtime"))
        self.assertEqual(roster.rows, [])

    def test_row_limit_no_partial_processing(self):
        rows = [f"{1000 + i},用户{i}" for i in range(4)]
        with patch.object(core, "MAX_ROWS", 3):
            with self.assertRaises(ValueError) as ctx:
                self.load(_csv_text(rows, header="mid,uname"))
        self.assertIn("上限", str(ctx.exception))

    def test_file_size_limit_no_partial_processing(self):
        with patch.object(core, "MAX_FILE_BYTES", 8):
            with self.assertRaises(ValueError) as ctx:
                self.load(_csv_text(["1001,张三"]))
        self.assertIn("上限", str(ctx.exception))

    def test_mid_as_float_string(self):
        roster = self.load(_csv_text(["1001.0,张三"], header="mid,uname"))
        self.assertEqual(roster.rows[0]["mid"], 1001)


class JsonParsingTests(BaseTestCase):
    def load(self, payload, name="fans.json"):
        path = self.write_file(name, json.dumps(payload, ensure_ascii=False))
        return load_roster(path, "时点 1 · 粉丝清单")

    def test_basic_object_array(self):
        roster = self.load([
            {"mid": 1001, "uname": "张三", "mtime": 1_700_000_000},
            {"mid": 1002, "uname": "李四", "mtime": "2024-05-01"},
        ])
        self.assertEqual(len(roster.rows), 2)
        self.assertEqual(roster.rows[0]["time"],
                         datetime.fromtimestamp(1_700_000_000))
        self.assertEqual(roster.rows[1]["time"], datetime(2024, 5, 1))
        self.assertTrue(roster.has_time)

    def test_synonym_keys(self):
        roster = self.load([
            {"uid": 1001, "nickname": "昵称甲", "follow_time": "2024-01-02"},
        ])
        self.assertEqual(roster.rows[0]["mid"], 1001)
        self.assertEqual(roster.rows[0]["name"], "昵称甲")
        self.assertEqual(roster.rows[0]["time"], datetime(2024, 1, 2))

    def test_mid_priority_over_synonyms(self):
        roster = self.load([{"mid": 1001, "id": 999999, "uname": "张三"}])
        self.assertEqual(roster.rows[0]["mid"], 1001)

    def test_non_list_top_level_raises(self):
        with self.assertRaises(ValueError):
            self.load({"list": []})

    def test_malformed_json_raises(self):
        path = self.write_file("broken.json", "{not json")
        with self.assertRaises(ValueError):
            load_roster(path, "测试")

    def test_missing_mid_key_raises(self):
        with self.assertRaises(ValueError) as ctx:
            self.load([{"uname": "张三"}, {"nickname": "李四"}])
        self.assertIn("mid", str(ctx.exception))

    def test_non_dict_elements_are_bad_rows(self):
        roster = self.load(["字符串", 42, {"mid": 1001, "uname": "张三"}])
        self.assertEqual(len(roster.rows), 1)
        self.assertEqual(roster.bad_rows, 2)

    def test_empty_array_is_valid(self):
        roster = self.load([])
        self.assertEqual(roster.rows, [])
        self.assertFalse(roster.has_time)

    def test_bool_mid_is_bad_row(self):
        roster = self.load([{"mid": True, "uname": "怪值"}])
        self.assertEqual(roster.rows, [])
        self.assertEqual(roster.bad_rows, 1)

    def test_float_mid_kept_as_int(self):
        roster = self.load([{"mid": 1001.0, "uname": "张三"}])
        self.assertEqual(roster.rows[0]["mid"], 1001)

    def test_dedup_and_bad_count(self):
        roster = self.load([
            {"mid": 1001, "uname": "第一"},
            {"mid": 1001, "uname": "第二"},
            {"mid": "bad", "uname": "坏行"},
        ])
        self.assertEqual(len(roster.rows), 1)
        self.assertEqual(roster.rows[0]["name"], "第一")
        self.assertEqual(roster.duplicates, 1)
        self.assertEqual(roster.bad_rows, 1)


def _make_roster(label, mids_names_times, has_time=True):
    roster = core.Roster(label=label, has_time=has_time)
    for mid, name, time_value in mids_names_times:
        roster.rows.append({"mid": mid, "name": name, "time": time_value})
    return roster


class SetOperationTests(BaseTestCase):
    def test_mutual_only_fans_only_follows(self):
        fans = _make_roster("粉丝", [(3, "丙", None), (1, "甲", None), (2, "乙", None)])
        follows = _make_roster("关注", [(2, "乙", None), (4, "丁", None)])
        analysis = compute_analysis({"fans_t1": fans, "follows_t1": follows})
        self.assertEqual([row["mid"] for row in analysis["mutual_rows"]], [2])
        self.assertEqual([row["mid"] for row in analysis["only_fans_rows"]], [1, 3])
        self.assertEqual([row["mid"] for row in analysis["only_follows_rows"]], [4])

    def test_all_mutual(self):
        fans = _make_roster("粉丝", [(1, "甲", None), (2, "乙", None)])
        follows = _make_roster("关注", [(1, "甲", None), (2, "乙", None)])
        analysis = compute_analysis({"fans_t1": fans, "follows_t1": follows})
        self.assertEqual(analysis["mutual_rows"], fans.rows)
        self.assertEqual(analysis["only_fans_rows"], [])
        self.assertEqual(analysis["only_follows_rows"], [])

    def test_no_intersection(self):
        fans = _make_roster("粉丝", [(1, "甲", None)])
        follows = _make_roster("关注", [(2, "乙", None)])
        analysis = compute_analysis({"fans_t1": fans, "follows_t1": follows})
        self.assertEqual([row["mid"] for row in analysis["only_fans_rows"]], [1])
        self.assertEqual([row["mid"] for row in analysis["only_follows_rows"]], [2])
        self.assertEqual(analysis["mutual_rows"], [])

    def test_empty_rosters_still_compute(self):
        fans = _make_roster("粉丝", [])
        follows = _make_roster("关注", [])
        analysis = compute_analysis({"fans_t1": fans, "follows_t1": follows})
        self.assertEqual(analysis["mutual_rows"], [])
        self.assertEqual(analysis["only_fans_rows"], [])
        self.assertEqual(analysis["only_follows_rows"], [])

    def test_missing_follows_roster_skips_set_ops(self):
        fans = _make_roster("粉丝", [(1, "甲", None)])
        analysis = compute_analysis({"fans_t1": fans})
        self.assertIsNone(analysis["mutual_rows"])
        self.assertIsNone(analysis["only_fans_rows"])
        self.assertIsNone(analysis["only_follows_rows"])

    def test_diff_added_removed_unchanged(self):
        fans1 = _make_roster(
            "时点1粉丝",
            [(1, "留", datetime(2024, 1, 1)), (2, "取关", datetime(2024, 1, 2))])
        fans2 = _make_roster(
            "时点2粉丝",
            [(1, "留", datetime(2024, 1, 1)), (3, "新增", datetime(2024, 3, 1))])
        analysis = compute_analysis({"fans_t1": fans1, "fans_t2": fans2})
        self.assertEqual([row["mid"] for row in analysis["diff_added_rows"]], [3])
        self.assertEqual(analysis["diff_added_rows"][0]["name"], "新增")
        self.assertEqual(analysis["diff_added_rows"][0]["time"], datetime(2024, 3, 1))
        self.assertEqual([row["mid"] for row in analysis["diff_removed_rows"]], [2])
        self.assertEqual(analysis["diff_removed_rows"][0]["time"], datetime(2024, 1, 2))
        self.assertEqual(analysis["diff_unchanged"], 1)

    def test_no_diff_without_two_fans_snapshots(self):
        fans1 = _make_roster("时点1粉丝", [(1, "甲", None)])
        analysis = compute_analysis({"fans_t1": fans1})
        self.assertIsNone(analysis["diff_added_rows"])
        self.assertIsNone(analysis["diff_removed_rows"])
        self.assertIsNone(analysis["diff_unchanged"])

    def test_primary_snapshot_prefers_t1(self):
        fans1 = _make_roster("时点1粉丝", [(1, "甲", None)])
        fans2 = _make_roster("时点2粉丝", [(2, "乙", None)])
        analysis = compute_analysis({"fans_t1": fans1, "fans_t2": fans2})
        self.assertEqual(analysis["primary_fans_key"], "fans_t1")

    def test_t2_only_becomes_primary(self):
        fans2 = _make_roster("时点2粉丝", [(2, "乙", None)])
        follows2 = _make_roster("时点2关注", [(2, "乙", None)])
        analysis = compute_analysis({"fans_t2": fans2, "follows_t2": follows2})
        self.assertEqual(analysis["primary_fans_key"], "fans_t2")
        self.assertEqual(analysis["primary_follows_key"], "follows_t2")
        self.assertEqual([row["mid"] for row in analysis["mutual_rows"]], [2])
        self.assertIsNone(analysis["diff_added_rows"])

    def test_monthly_distribution(self):
        fans1 = _make_roster("时点1粉丝", [
            (1, "甲", datetime(2024, 1, 15)),
            (2, "乙", datetime(2024, 1, 20)),
            (3, "丙", datetime(2024, 3, 1)),
            (4, "丁", None),
        ])
        analysis = compute_analysis({"fans_t1": fans1})
        self.assertEqual(analysis["monthly"], [("2024-01", 2), ("2024-03", 1)])
        self.assertEqual(analysis["monthly_unparsed"], 1)

    def test_monthly_requires_time_column(self):
        fans1 = _make_roster("时点1粉丝", [(1, "甲", None)], has_time=False)
        analysis = compute_analysis({"fans_t1": fans1})
        self.assertIsNone(analysis["monthly"])

    def test_monthly_only_from_t1(self):
        fans1 = _make_roster("时点1粉丝", [(1, "甲", None)], has_time=False)
        fans2 = _make_roster("时点2粉丝", [(2, "乙", datetime(2024, 5, 1))])
        analysis = compute_analysis({"fans_t1": fans1, "fans_t2": fans2})
        self.assertIsNone(analysis["monthly"])


class ValidateInputsTests(BaseTestCase):
    def make_csv(self, name="fans.csv"):
        return self.write_file(name, _csv_text(["1001,张三"]))

    def test_no_files_raises(self):
        with self.assertRaises(ValueError):
            validate_inputs("", "", "", "", str(self.root))

    def test_missing_file_raises(self):
        with self.assertRaises(ValueError):
            validate_inputs(str(self.root / "不存在.csv"), "", "", "", str(self.root))

    def test_wrong_extension_raises(self):
        path = self.write_file("清单.txt", _csv_text(["1001,张三"]))
        with self.assertRaises(ValueError):
            validate_inputs(str(path), "", "", "", str(self.root))

    def test_empty_out_dir_raises(self):
        path = self.make_csv()
        with self.assertRaises(ValueError):
            validate_inputs(str(path), "", "", "", "   ")

    def test_valid_paths_resolved(self):
        path = self.make_csv()
        paths, out_dir = validate_inputs(str(path), "", "", "", str(self.root / "导出"))
        self.assertEqual(set(paths), {"fans_t1"})
        self.assertEqual(paths["fans_t1"].name, "fans.csv")
        self.assertTrue(out_dir.is_dir())

    def test_t2_only_is_valid(self):
        path = self.make_csv("t2.csv")
        paths, _out_dir = validate_inputs("", "", "", str(path), str(self.root))
        self.assertEqual(set(paths), {"follows_t2"})


class PipelineTests(BaseTestCase):
    def write_fixtures(self):
        fans1 = self.write_file("fans1.csv", _csv_text([
            "1001,留甲,2024-01-10",
            "1002,取关乙,2024-01-20",
            "1004,仅粉丝丁,",
        ]))
        follows1 = self.write_file("follows1.csv", _csv_text([
            "1001,留甲,",
            "1003,仅关注丙,",
        ], header="mid,uname,ctime"))
        fans2 = self.write_file("fans2.csv", _csv_text([
            "1001,留甲,2024-02-01",
            "1005,新增戊,2024-02-15",
        ]))
        return fans1, follows1, fans2

    def run_ok(self, **overrides):
        kwargs = {
            "fans_t1": "", "follows_t1": "", "fans_t2": "", "follows_t2": "",
            "out_dir": str(self.root / "导出"),
        }
        kwargs.update(overrides)
        return run_pipeline(**kwargs)

    def test_end_to_end_full_sheets(self):
        fans1, follows1, fans2 = self.write_fixtures()
        result = self.run_ok(fans_t1=str(fans1), follows_t1=str(follows1),
                             fans_t2=str(fans2))
        stats = result["stats"]
        self.assertFalse(stats["cancelled"])
        self.assertEqual(stats["added"], 1)
        self.assertEqual(stats["removed"], 2)
        self.assertEqual(stats["unchanged"], 1)
        self.assertEqual(stats["mutual"], 1)
        self.assertEqual(stats["only_fans"], 2)
        self.assertEqual(stats["only_follows"], 1)
        self.assertTrue(stats["monthly"])
        self.assertTrue(Path(result["excel"]).is_file())

        wb = load_workbook(result["excel"], read_only=True)
        self.addCleanup(wb.close)
        self.assertEqual(
            wb.sheetnames,
            ["汇总", "粉丝清单", "关注清单", "互相关注", "仅粉丝", "仅关注",
             "粉丝差异", "按月新增分布"])
        # 粉丝差异表：状态列在前，新增/取关分组，mid 升序；时点 1 独有的
        # 1004 同样计入取关（时点 2 粉丝相对时点 1 粉丝）
        diff_ws = wb["粉丝差异"]
        rows = [row for row in diff_ws.iter_rows(
            min_col=2, max_col=5, values_only=True)][3:]
        data = [row for row in rows if row[0] in ("新增", "取关")]
        self.assertEqual(data, [
            ("新增", 1005, "新增戊", "2024-02-15 00:00:00"),
            ("取关", 1002, "取关乙", "2024-01-20 00:00:00"),
            ("取关", 1004, "仅粉丝丁", "—"),
        ])
        # 汇总：来源声明必须落在表内
        summary_ws = wb["汇总"]
        values = [str(row[0]) for row in summary_ws.iter_rows(
            min_col=2, max_col=2, values_only=True) if row[0]]
        self.assertTrue(any("不发起任何网络请求" in value for value in values))
        # 明细表 mid 为整型，昵称不出现空串/None
        fans_ws = wb["粉丝清单"]
        fan_rows = [row for row in fans_ws.iter_rows(
            min_col=2, max_col=4, values_only=True)][3:]
        fan_data = [row for row in fan_rows if isinstance(row[0], int)]
        self.assertEqual([row[0] for row in fan_data], [1001, 1002, 1004])
        for row in fan_data:
            self.assertTrue(str(row[1]).strip())
            self.assertTrue(str(row[2]).strip())
        wb.close()

    def test_single_timepoint_degrades(self):
        fans1, _follows1, _fans2 = self.write_fixtures()
        result = self.run_ok(fans_t1=str(fans1))
        wb = load_workbook(result["excel"], read_only=True)
        self.addCleanup(wb.close)
        # 时点 1 粉丝清单带时间列 → 按月分布照常输出；无差异/集合表
        self.assertEqual(wb.sheetnames, ["汇总", "粉丝清单", "按月新增分布"])
        stats = result["stats"]
        self.assertIsNone(stats["mutual"])
        self.assertIsNone(stats["added"])
        self.assertTrue(stats["monthly"])

    def test_cancel_returns_clean_result_without_excel(self):
        fans1, follows1, fans2 = self.write_fixtures()
        result = self.run_ok(fans_t1=str(fans1), follows_t1=str(follows1),
                             fans_t2=str(fans2), cancel=lambda: True)
        self.assertTrue(result["stats"]["cancelled"])
        self.assertEqual(result["excel"], "")
        self.assertEqual(list((self.root / "导出").glob("*.xlsx")), [])

    def test_cancel_during_parse_keeps_no_excel(self):
        fans1, follows1, fans2 = self.write_fixtures()
        calls = {"n": 0}

        def cancel():
            calls["n"] += 1
            return calls["n"] > 1

        with patch.object(core, "_CANCEL_GAP", 1):
            result = self.run_ok(fans_t1=str(fans1), follows_t1=str(follows1),
                                 fans_t2=str(fans2), cancel=cancel)
        self.assertTrue(result["stats"]["cancelled"])
        self.assertEqual(list((self.root / "导出").glob("*.xlsx")), [])
        self.assertEqual(list((self.root / "导出").glob("*.tmp")), [])

    def test_over_limit_raises_and_writes_nothing(self):
        fans1, _follows1, _fans2 = self.write_fixtures()
        with patch.object(core, "MAX_ROWS", 2):
            with self.assertRaises(ValueError):
                self.run_ok(fans_t1=str(fans1))
        self.assertEqual(list((self.root / "导出").glob("*.xlsx")), [])

    def test_output_dir_created(self):
        fans1, _follows1, _fans2 = self.write_fixtures()
        nested = self.root / "层级" / "导出"
        result = self.run_ok(fans_t1=str(fans1), out_dir=str(nested))
        self.assertTrue(nested.is_dir())
        self.assertEqual(result["dir"], str(nested))


class RegistryAndBoundaryTests(unittest.TestCase):
    def test_registered_as_last_tool(self):
        # 2026-09-12 直播追踪（live_room）按任务卡以 append 一行注册到末位，
        # relation_analysis 不再是 TOOLS 的最后一项。这是预期变更：改测试
        # 而不是改注册顺序——本用例改为锁 relation_analysis 的既有位置
        # （第 8 位）与注册表总数。
        spec = next(item for item in TOOLS if item.id == "relation_analysis")
        self.assertEqual(TOOLS.index(spec), 7)
        self.assertEqual(len(TOOLS), 9)
        self.assertEqual(spec.name, "关系分析")
        self.assertEqual(TOOLS[-1].id, "live_room")

    def test_subtitle_consistent_between_registry_and_page(self):
        spec = next(item for item in TOOLS if item.id == "relation_analysis")
        self.assertEqual(spec.subtitle, RelationAnalysisPage.tool_subtitle)

    def test_tool_modules_are_network_free(self):
        """任务卡硬约束：新工具目录零网络引用（与验收 grep 同口径）。"""
        import tools.relation_analysis as package

        forbidden = ("session", "client", "transport", "gate", "wbi",
                     "budget", "urllib", "requests", "socket", "httpx",
                     "http")
        base = Path(package.__file__).parent
        for py_file in sorted(base.glob("*.py")):
            text = py_file.read_text(encoding="utf-8").lower()
            for token in forbidden:
                self.assertNotIn(
                    token, text,
                    f"{py_file.name} 出现零网络禁用词：{token}")


if __name__ == "__main__":
    unittest.main()
