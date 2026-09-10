# -*- coding: utf-8 -*-
"""P0 报告中心：真实导出字段、键契约、时间契约、导出保护和 UI 状态。"""
from __future__ import annotations

import json
import os
import socket
import tempfile
import threading
import time
import unittest
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from tools.comments import core as comments_core
from tools.collector import core as collector_core
from tools.report_center import core
from tools.report_center.pipeline import run_export, run_file_comparison, run_refresh, run_search


BASE_TS = 1_700_000_000


def comment_row(rpid, ctime=BASE_TS, like=1, message="正常评论"):
    return {
        "rpid": rpid, "parent": 0, "root": 0, "is_main": True,
        "uname": f"user{rpid}", "mid": rpid + 100, "sex": "保密", "level": 6,
        "vip": False, "message": message, "like": like, "rcount": 0,
        "ctime": ctime, "location": "", "is_top": False,
    }


def video_row(bvid="BV1Report0001", fetched_at=BASE_TS, view=100):
    return {
        "bvid": bvid, "aid": 1, "title": "测试视频", "owner": "测试UP",
        "owner_mid": 2, "tname": "测试", "pubdate": BASE_TS, "duration": 10,
        "view": view, "danmaku": 2, "reply": 3, "favorite": 4, "coin": 5,
        "share": 6, "like": 7, "fetched_at": fetched_at,
    }


def monitor_row(ts=BASE_TS, view=100):
    return {
        "ts": ts, "bvid": "BV1Monitor0001", "view": view, "danmaku": 2,
        "reply": 3, "favorite": 4, "coin": 5, "share": 6, "like": 7,
        "online": 8,
    }


class ReportCenterCoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def write_jsonl(self, name, rows):
        path = self.root / name
        path.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")
        return path

    def make_comments_xlsx(self):
        path = self.root / "comments-export.xlsx"
        rows = [comment_row(11)]
        kpi = {
            "coverage": "100%", "main": 1, "sub": 0, "users": 1, "per_user": "1.00",
            "hours": 1, "zero_like_pct": "0%", "vip_pct": "0%", "build": 0,
            "build_pct": "0%", "apology": 0, "defense": 0, "stop_pay": 0,
            "genshin": 0, "spam": 0, "top_words": [], "emo_top": [], "uni_top": [],
            "by_day": [("11-14", 1)], "by_hour": [(22, 1)],
        }
        comments_core.export_xlsx(
            rows,
            {"pub_ts": BASE_TS, "claimed_comment_count": 1, "author": "测试UP", "title": "测试"},
            kpi,
            path,
        )
        workbook = load_workbook(path)
        sheet = workbook["全量评论"]
        # The existing exporter places 点赞数 in K2. Formula text must remain
        # visible because report-center reads with data_only=False.
        sheet["K2"] = "=1+1"
        workbook.save(path)
        workbook.close()
        return path

    def make_video_xlsx(self, name="video-export.xlsx"):
        path = self.root / name
        collector_core.export_xlsx(
            [video_row()], [],
            {"title": "测试", "source": "本地夹具", "ok": 1, "attempted": 1,
             "monitor": "否", "rounds": 1, "interval_min": 1},
            path,
        )
        return path

    def test_epoch_units_are_strict_and_timezone_explicit(self):
        value = core.parse_epoch_seconds(BASE_TS)
        self.assertIsNotNone(value)
        self.assertEqual(core.LOCAL_TZ.tzname(None), "Asia/Shanghai")
        self.assertEqual(value.utcoffset(), core.LOCAL_TZ.utcoffset(value))
        self.assertIsNone(core.parse_epoch_seconds(BASE_TS * 1000))
        self.assertIsNone(core.parse_epoch_seconds(BASE_TS * 1_000_000))
        self.assertIn("+", core.format_datetime(value))

    def test_discovery_uses_history_and_manual_files_and_keeps_missing_output(self):
        comment = self.write_jsonl("comments.jsonl", [comment_row(1)])
        video = self.write_jsonl("snapshots.jsonl", [video_row()])
        missing = self.root / "deleted.xlsx"
        records = [{"outputs": [str(comment), str(missing)]}]
        sources = core.discover_sources(records, manual_files=[video], manual_dirs=[], monitor_roots=[])
        paths = {Path(item.path).name for item in sources}
        self.assertEqual(paths, {"comments.jsonl", "snapshots.jsonl", "deleted.xlsx"})
        self.assertEqual(next(item for item in sources if item.file_name == "deleted.xlsx").status, core.STATUS_UNREADABLE)

    def test_damaged_file_does_not_block_valid_file(self):
        damaged = self.root / "damaged.jsonl"
        damaged.write_text('{"rpid": 1}\n{bad json}\n', encoding="utf-8")
        valid = self.write_jsonl("valid.jsonl", [comment_row(1)])
        sources = core.discover_sources(manual_files=[damaged, valid], manual_dirs=[], monitor_roots=[])
        by_name = {item.file_name: item for item in sources}
        self.assertEqual(by_name["damaged.jsonl"].status, core.STATUS_DAMAGED)
        self.assertEqual(by_name["valid.jsonl"].status, core.STATUS_READABLE)

    def test_cancelled_discovery_stops_large_file_and_discards_partial_source(self):
        large = self.root / "large.jsonl"
        rows = [comment_row(index, message=f"row-{index}") for index in range(1, 2201)]
        large.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")
        later = self.write_jsonl("later.jsonl", [comment_row(9000)])
        cancelled = False
        read_calls = []

        def cancel():
            return cancelled

        def progress(**values):
            nonlocal cancelled
            if values.get("row", 0) >= 500:
                cancelled = True

        real_read_source = core.read_source

        def tracked_read_source(path, *args, **kwargs):
            read_calls.append(Path(path).name)
            return real_read_source(path, *args, **kwargs)

        with patch.object(core, "read_source", side_effect=tracked_read_source):
            sources = core.discover_sources(
                history_records=[], manual_files=[large, later], manual_dirs=[], monitor_roots=[],
                cancel=cancel, progress=progress,
            )
        self.assertEqual(read_calls, ["large.jsonl"])
        self.assertEqual(sources, [])
        self.assertFalse(any(item.status == core.STATUS_INCOMPLETE for item in sources))

    def test_existing_exporters_create_readable_xlsx_fixtures_and_formula_is_not_evaluated(self):
        comment_path = self.make_comments_xlsx()
        video_path = self.make_video_xlsx()
        comment_result = core.read_source(comment_path)
        video_result = core.read_source(video_path)
        self.assertEqual(comment_result.source.kind, core.KIND_COMMENT)
        self.assertIn("ctime", comment_result.source.fields)
        self.assertIsNotNone(comment_result.source.time_start)
        self.assertEqual(comment_result.source.time_start.utcoffset(), core.LOCAL_TZ.utcoffset(comment_result.source.time_start))
        self.assertEqual(comment_result.preview_rows[0]["_raw_fields"]["点赞数"], "=1+1")
        self.assertEqual(comment_result.source.record_count, 1)
        self.assertEqual(video_result.source.kind, core.KIND_VIDEO)
        self.assertIn("view", video_result.source.fields)
        self.assertEqual(video_result.preview_rows[0]["_raw_fields"]["播放"], 100)
        # Existing 视频总表 has 发布时间 (publication time), not fetched_at;
        # report-center must not guess that it is a snapshot time.
        self.assertNotIn("fetched_at", video_result.source.fields)
        self.assertFalse(video_result.source.can_compare)
        self.assertIsNone(video_result.source.time_start)

    def test_comment_file_comparison_uses_rpid_and_reports_metric_changes(self):
        left = self.write_jsonl("comments-a.jsonl", [comment_row(1, like=1), comment_row(2, like=2)])
        right = self.write_jsonl("comments-b.jsonl", [comment_row(1, like=4), comment_row(3, like=3)])
        result = core.compare_sources(left, right)
        self.assertTrue(result.compatible)
        self.assertEqual(result.matching_mode, "key")
        self.assertEqual((result.a_records, result.b_records, result.common_records), (2, 2, 1))
        self.assertEqual((result.only_a_records, result.only_b_records, result.changed_records), (1, 1, 1))
        like = next(item for item in result.metric_changes if item.field == "like")
        self.assertEqual((like.a_value, like.b_value, like.absolute_change), (1, 4, 3))
        self.assertAlmostEqual(like.percent_change, 300.0)
        self.assertEqual(len(result.missing_fields_a), 0)

    def test_video_comparison_uses_bvid_and_fetched_at(self):
        left = self.write_jsonl("video-a.jsonl", [video_row(fetched_at=BASE_TS), video_row(bvid="BV1Report0002")])
        right = self.write_jsonl("video-b.jsonl", [video_row(fetched_at=BASE_TS, view=120), video_row(bvid="BV1Report0002", fetched_at=BASE_TS + 1)])
        result = core.compare_sources(left, right)
        self.assertEqual(result.matching_mode, "key")
        self.assertEqual(result.common_records, 1)
        self.assertEqual(result.only_a_records, 1)
        self.assertEqual(result.only_b_records, 1)

    def test_incompatible_types_are_rejected_and_unknown_key_is_summary_only(self):
        comments = self.write_jsonl("comments.jsonl", [comment_row(1)])
        monitor = self.write_jsonl("history_BV1.jsonl", [monitor_row()])
        result = core.compare_sources(comments, monitor)
        self.assertFalse(result.compatible)
        summary = core.compare_sources(
            self.make_video_xlsx("video-a.xlsx"),
            self.make_video_xlsx("video-b.xlsx"),
        )
        self.assertEqual(summary.matching_mode, "summary")
        self.assertIsNone(summary.common_records)
        self.assertIn("无法逐条匹配", summary.message)

    def test_period_comparison_reports_all_stats_and_missing_values(self):
        source_path = self.write_jsonl("history_BV1.jsonl", [
            monitor_row(BASE_TS, 10), monitor_row(BASE_TS + 10, 20),
            {**monitor_row(BASE_TS + 20, 30), "view": None},
            monitor_row(BASE_TS + 30, 40),
        ])
        start_a = core.parse_epoch_seconds(BASE_TS)
        end_a = core.parse_epoch_seconds(BASE_TS + 10)
        start_b = core.parse_epoch_seconds(BASE_TS + 20)
        end_b = core.parse_epoch_seconds(BASE_TS + 30)
        result = core.compare_periods(source_path, (start_a, end_a), (start_b, end_b), "view")
        self.assertEqual((result.period_a.sample_count, result.period_a.start_value, result.period_a.end_value), (2, 10, 20))
        self.assertEqual((result.period_b.sample_count, result.period_b.valid_count, result.period_b.missing_data), (2, 1, 1))
        self.assertEqual(result.period_b.end_value, 40)
        self.assertEqual(result.period_a.absolute_change, 10)
        with self.assertRaises(ValueError):
            core.compare_periods(source_path, (start_a, end_a), (start_b, end_b), "missing_metric")

    def test_missing_time_field_blocks_period_compare(self):
        source_path = self.write_jsonl("no-time.jsonl", [{"bvid": "BV1", "view": 1}])
        result = core.read_source(source_path).source
        with self.assertRaises(ValueError) as context:
            core.compare_periods(result, ("2023-01-01T00:00:00+08:00", "2023-01-01T01:00:00+08:00"),
                                  ("2023-01-02T00:00:00+08:00", "2023-01-02T01:00:00+08:00"), "view")
        self.assertIn("可靠时间", str(context.exception))

    def test_exports_never_overwrite_sources_and_preserve_fingerprint(self):
        source_path = self.write_jsonl("comments.jsonl", [comment_row(1, message="safe")])
        source = core.read_source(source_path).source
        before_bytes = source_path.read_bytes()
        before_fp = core.snapshot_source(source_path)
        output_one = core.export_filtered(source, self.root / "reports", "JSONL", query="safe")
        output_two = core.export_filtered(source, self.root / "reports", "JSONL", query="safe")
        self.assertNotEqual(output_one, output_two)
        self.assertNotEqual(output_one.resolve(), source_path.resolve())
        self.assertEqual(source_path.read_bytes(), before_bytes)
        self.assertEqual(core.snapshot_source(source_path), before_fp)
        comparison = core.compare_sources(source_path, source_path)
        compare_output = core.export_file_comparison(comparison, self.root / "reports", "XLSX")
        self.assertTrue(compare_output.is_file())

    def test_comparison_exports_reject_sources_changed_after_generation(self):
        left = self.write_jsonl("comparison-left.jsonl", [comment_row(1, like=1)])
        right = self.write_jsonl("comparison-right.jsonl", [comment_row(1, like=2)])
        comparison = core.compare_sources(left, right)
        self.assertIsNotNone(comparison.fingerprint_a)
        left.write_text(left.read_text(encoding="utf-8") + "\n", encoding="utf-8")
        with self.assertRaisesRegex(core.SourceChangedError, "源文件已变化，请重新生成对比"):
            core.export_file_comparison(comparison, self.root / "reports", "JSONL")

        period_path = self.write_jsonl("period-source.jsonl", [monitor_row(BASE_TS, 10)])
        period = core.compare_periods(
            period_path,
            ("2023-11-14T00:00:00+08:00", "2023-11-15T00:00:00+08:00"),
            ("2023-11-16T00:00:00+08:00", "2023-11-17T00:00:00+08:00"),
            "view",
        )
        self.assertIsNotNone(period.fingerprint)
        period_path.write_text(period_path.read_text(encoding="utf-8").replace('"view": 100', '"view": 101'), encoding="utf-8")
        with self.assertRaisesRegex(core.SourceChangedError, "源文件已变化，请重新生成对比"):
            core.export_period_comparison(period, self.root / "reports", "JSONL")

    def test_full_record_search_finds_record_201(self):
        path = self.write_jsonl(
            "search.jsonl",
            [comment_row(index, message="needle-201" if index == 201 else "ordinary")
             for index in range(1, 202)],
        )
        source = core.read_source(path).source
        self.assertEqual(len(source.preview_rows), core.PREVIEW_LIMIT)
        matches = core.search_sources([source], "needle-201")
        self.assertEqual([item.path for item in matches], [source.path])
        self.assertEqual(core.search_source_matches([source], "needle-201")[source.path.casefold()], "content")

    def test_metadata_hit_wins_over_partial_content_hit_and_exports_all_rows(self):
        path = self.write_jsonl("needle.jsonl", [
            comment_row(1, message="needle in first row"), comment_row(2, message="second"),
        ])
        source = core.read_source(path).source
        before = core.snapshot_source(path)
        result = run_search([source], "needle")
        self.assertEqual(result["match_modes"][source.path.casefold()], "metadata")
        output = run_export(
            "filtered", {"source": source, "query": ""}, self.root / "reports", "JSONL"
        )
        exported = [json.loads(line) for line in Path(output["path"]).read_text(encoding="utf-8").splitlines()]
        self.assertEqual(len(exported), 2)
        self.assertEqual({row["rpid"] for row in exported}, {1, 2})
        self.assertEqual(core.snapshot_source(path), before)

    def test_content_only_hit_exports_only_matching_rows(self):
        path = self.write_jsonl("comments.jsonl", [
            comment_row(1, message="needle in first row"), comment_row(2, message="second"),
        ])
        source = core.read_source(path).source
        before = core.snapshot_source(path)
        result = run_search([source], "needle")
        self.assertEqual(result["match_modes"][source.path.casefold()], "content")
        output = run_export(
            "filtered", {"source": source, "query": "needle"}, self.root / "reports", "JSONL"
        )
        exported = [json.loads(line) for line in Path(output["path"]).read_text(encoding="utf-8").splitlines()]
        self.assertEqual(len(exported), 1)
        self.assertEqual(exported[0]["rpid"], 1)
        self.assertEqual(core.snapshot_source(path), before)

    def test_comparison_worker_cancel_returns_no_partial_result(self):
        left = self.write_jsonl("cancel-left.jsonl", [comment_row(index) for index in range(1, 2201)])
        right = self.write_jsonl("cancel-right.jsonl", [comment_row(index, like=2) for index in range(1, 2201)])
        cancelled = False

        def cancel():
            return cancelled

        def progress(**values):
            nonlocal cancelled
            if values.get("row", 0) >= 500:
                cancelled = True

        result = run_file_comparison(left, right, cancel=cancel, progress=progress)
        self.assertTrue(result["cancelled"])
        self.assertNotIn("comparison", result)

    def test_xlsx_export_cleans_controls_escapes_formula_text_and_preserves_types(self):
        source_path = self.write_jsonl(
            "xlsx-source.jsonl",
            [{**comment_row(index, message=text), "is_main": index % 2 == 0,
              "vip": index % 2 == 1, "like": -5 if index == 1 else index}
             for index, text in enumerate(
                 ("\x01control", "=1+1", "+SUM(A1)", "-cmd", "@cmd"), 1)],
        )
        source = core.read_source(source_path).source
        output = core.export_filtered(source, self.root / "xlsx-reports", "XLSX")
        exported = load_workbook(output, data_only=False)
        sheet = exported["数据"]
        headers = {cell.value: index for index, cell in enumerate(sheet[1])}
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        messages = [row[headers["message"]] for row in rows]
        self.assertEqual(messages, ["control", "'=1+1", "'+SUM(A1)", "'-cmd", "'@cmd"])
        self.assertEqual(rows[0][headers["like"]], -5)
        self.assertIsInstance(rows[1][headers["like"]], int)
        self.assertIs(rows[1][headers["is_main"]], True)
        exported.close()

        date_source = core.read_source(self.make_comments_xlsx()).source
        date_output = core.export_filtered(date_source, self.root / "xlsx-reports", "XLSX")
        date_book = load_workbook(date_output, data_only=False)
        date_sheet = date_book["数据"]
        date_headers = {cell.value: index for index, cell in enumerate(date_sheet[1])}
        date_row = next(date_sheet.iter_rows(min_row=2, values_only=True))
        self.assertIsInstance(date_row[date_headers["发布时间"]], datetime)
        date_book.close()

    def test_sensitive_fields_are_not_exported(self):
        source_path = self.write_jsonl("sensitive.jsonl", [
            {**comment_row(1), "Cookie": "SESSDATA=secret", "access_token": "secret"},
        ])
        source = core.read_source(source_path).source
        output = core.export_filtered(source, self.root / "reports", "JSONL")
        exported = json.loads(output.read_text(encoding="utf-8"))
        self.assertNotIn("Cookie", exported)
        self.assertNotIn("access_token", exported)
        self.assertNotIn("secret", output.read_text(encoding="utf-8"))

    def test_refresh_export_have_no_network_or_taskrunner_boundary(self):
        source_path = self.write_jsonl("comments.jsonl", [comment_row(1)])
        with patch.object(socket, "create_connection", side_effect=AssertionError("network")), \
                patch("app.task_runner.TaskRunner", side_effect=AssertionError("TaskRunner")) as runner:
            result = run_refresh(history_records=[], manual_files=[source_path], manual_dirs=[], monitor_roots=[])
            self.assertEqual(result["count"], 1)
            output = run_export("filtered", {"source": result["sources"][0], "query": ""},
                                self.root / "reports", "JSONL")
            self.assertTrue(Path(output["path"]).is_file())
            runner.assert_not_called()


class ReportCenterPageTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try:
            from PySide6.QtWidgets import QApplication
        except ImportError:
            cls.app = None
            return
        cls.app = QApplication.instance() or QApplication([])

    def wait_for(self, predicate, timeout=4000):
        deadline = time.monotonic() + timeout / 1000
        while time.monotonic() < deadline:
            self.app.processEvents()
            if predicate():
                return True
            time.sleep(0.01)
        self.app.processEvents()
        return bool(predicate())

    def test_page_has_required_file_button_tabs_and_compact_no_horizontal_overflow(self):
        if self.app is None:
            self.skipTest("PySide6 unavailable in this interpreter")
        from PySide6.QtCore import Qt
        from PySide6.QtWidgets import QScrollArea
        from tools.report_center.page import ReportCenterPage
        page = ReportCenterPage({"out_dir": tempfile.gettempdir()})
        page.resize(960, 640)
        self.app.processEvents()
        self.assertEqual(page.btn_add_files.text(), "添加文件")
        self.assertEqual([page.tabs.tabText(i) for i in range(page.tabs.count())], ["浏览报告", "文件对比", "时间段对比", "导出结果"])
        page._on_refresh_done({"cancelled": True})
        self.assertIn("取消", page.status_pill.text())
        page._on_worker_failed("损坏文件")
        self.assertEqual(page.status_pill.property("state"), "error")
        scroll = page.findChild(QScrollArea, "pageScroll")
        self.assertIsNotNone(scroll)
        self.assertEqual(scroll.horizontalScrollBar().maximum(), 0)
        page.close()

    def test_page_states_expose_cancel_damage_incompatible_missing_and_empty_messages(self):
        if self.app is None:
            self.skipTest("PySide6 unavailable in this interpreter")
        from tools.report_center.page import ReportCenterPage

        page = ReportCenterPage({"out_dir": tempfile.gettempdir()})
        page.on_app_close()
        self.assertTrue(self.wait_for(lambda: page._thread is None))
        page._on_refresh_done({"cancelled": True})
        self.assertIn("取消", page.status_pill.text())
        page._on_worker_failed("损坏文件")
        self.assertIn("损坏文件", page.source_detail.text())
        page._comparison = object()
        page._period_comparison = object()
        page._on_compare_done({"cancelled": True})
        page._on_period_done({"cancelled": True})
        self.assertIsNone(page._comparison)
        self.assertIsNone(page._period_comparison)
        page._on_worker_failed("SourceChangedError: 源文件已变化，请重新生成对比")
        self.assertIsNone(page._comparison)
        self.assertIsNone(page._period_comparison)
        self.assertIn("源文件已变化，请重新生成对比", page.export_status.text())

        page.compare_files()
        self.assertIn("请选择文件", page.compare_summary.text())
        page.export_filtered(core.FORMAT_JSONL)
        self.assertIn("导出失败", page.export_status.text())

        temp = tempfile.TemporaryDirectory()
        root = Path(temp.name)
        comments = root / "comments.jsonl"
        monitor = root / "history_BV1.jsonl"
        no_time = root / "no-time.jsonl"
        comments.write_text(json.dumps(comment_row(1), ensure_ascii=False) + "\n", encoding="utf-8")
        monitor.write_text(json.dumps(monitor_row(), ensure_ascii=False) + "\n", encoding="utf-8")
        no_time.write_text('{"bvid":"BV1","view":1}\n', encoding="utf-8")
        comment_source = core.read_source(comments).source
        monitor_source = core.read_source(monitor).source
        no_time_source = core.read_source(no_time).source
        page.sources = [comment_source, monitor_source, no_time_source]
        page._source_map = {item.path.casefold(): item for item in page.sources}
        page._populate_source_controls()
        page.compare_a.setCurrentIndex(page.compare_a.findData(comment_source.path))
        page.compare_b.setCurrentIndex(page.compare_b.findData(monitor_source.path))
        page.compare_files()
        self.assertIn("已阻止对比", page.compare_summary.text())

        page.period_source.clear()
        page.period_source.addItem("无可靠时间", no_time_source.path)
        page.period_metric.clear()
        page.period_metric.addItem("view")
        page.compare_periods()
        self.assertIn("可靠时间", page.period_summary.text())

        page.period_source.clear()
        page.period_source.addItem("监控", monitor_source.path)
        page.period_metric.clear()
        page.period_metric.addItem("not_a_metric")
        page.compare_periods()
        self.assertIn("不是可计算数值", page.period_summary.text())

        page.period_metric.clear()
        page.period_metric.addItem("view")
        page.period_a_start.setText("2030-01-01T00:00:00+08:00")
        page.period_a_end.setText("2030-01-01T01:00:00+08:00")
        page.period_b_start.setText("2030-01-02T00:00:00+08:00")
        page.period_b_end.setText("2030-01-02T01:00:00+08:00")
        page.compare_periods()
        self.assertTrue(self.wait_for(lambda: page._period_comparison is not None))
        self.assertIn("暂无可绘制数据", page.period_chart._empty_text)
        page.close()
        temp.cleanup()

    def test_comparisons_execute_in_local_worker_thread(self):
        if self.app is None:
            self.skipTest("PySide6 unavailable in this interpreter")
        from tools.report_center import core as report_core
        from tools.report_center import page as page_module

        temp = tempfile.TemporaryDirectory()
        root = Path(temp.name)
        left = root / "left.jsonl"
        right = root / "right.jsonl"
        rows = [comment_row(index, like=index) for index in range(1, 1201)]
        left.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")
        right.write_text("".join(json.dumps({**row, "like": row["like"] + 1}, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")
        left_source = report_core.read_source(left).source
        right_source = report_core.read_source(right).source
        page = page_module.ReportCenterPage({"out_dir": root})
        page.on_app_close()
        self.assertTrue(self.wait_for(lambda: page._thread is None))
        page.sources = [left_source, right_source]
        page._source_map = {item.path.casefold(): item for item in page.sources}
        page._populate_source_controls()
        page.compare_a.setCurrentIndex(page.compare_a.findData(left_source.path))
        page.compare_b.setCurrentIndex(page.compare_b.findData(right_source.path))
        gui_thread = threading.get_ident()
        compare_threads = []
        period_threads = []
        real_compare = report_core.compare_sources
        real_period = report_core.compare_periods

        def tracked_compare(*args, **kwargs):
            compare_threads.append(threading.get_ident())
            return real_compare(*args, **kwargs)

        def tracked_period(*args, **kwargs):
            period_threads.append(threading.get_ident())
            return real_period(*args, **kwargs)

        with patch.object(page_module.core, "compare_sources", side_effect=tracked_compare), \
                patch.object(page_module.core, "compare_periods", side_effect=tracked_period):
            page.compare_files()
            self.assertTrue(
                self.wait_for(lambda: page._comparison is not None and page._thread is None),
                f"comparison={page._comparison!r} thread={page._thread!r} status={page.status_pill.text()!r} detail={page.source_detail.text()!r} count={page.compare_a.count()}/{page.compare_b.count()} data={page.compare_a.currentData()!r}/{page.compare_b.currentData()!r}",
            )
            self.assertTrue(compare_threads)
            self.assertNotEqual(compare_threads[0], gui_thread)

            page.period_source.setCurrentIndex(page.period_source.findData(left_source.path))
            page.period_metric.setCurrentText("like")
            page.period_a_start.setText("2023-11-14T00:00:00+08:00")
            page.period_a_end.setText("2023-11-15T00:00:00+08:00")
            page.period_b_start.setText("2023-11-15T00:00:01+08:00")
            page.period_b_end.setText("2023-11-16T00:00:00+08:00")
            page.compare_periods()
            self.assertTrue(self.wait_for(lambda: page._period_comparison is not None and page._thread is None))
            self.assertTrue(period_threads)
            self.assertNotEqual(period_threads[0], gui_thread)
        page.on_app_close()
        page.close()
        temp.cleanup()

    def test_filename_metadata_hit_page_export_uses_full_source(self):
        if self.app is None:
            self.skipTest("PySide6 unavailable in this interpreter")
        from tools.report_center.page import ReportCenterPage

        temp = tempfile.TemporaryDirectory()
        root = Path(temp.name)
        source_path = root / "visible-keyword.jsonl"
        source_path.write_text(
            "".join(json.dumps(comment_row(index, message="ordinary"), ensure_ascii=False) + "\n" for index in (1, 2)),
            encoding="utf-8",
        )
        source = core.read_source(source_path).source
        page = ReportCenterPage({"out_dir": root})
        page.on_app_close()
        self.assertTrue(self.wait_for(lambda: page._thread is None))
        page.sources = [source]
        page._source_map = {source.path.casefold(): source}
        page._fill_source_table([source])
        page.source_table.selectRow(0)
        page.search_edit.blockSignals(True)
        page.search_edit.setText("visible-keyword")
        page.search_edit.blockSignals(False)
        page._search_match_modes = {source.path.casefold(): "metadata"}
        self.assertEqual(page._export_query_for_source(source), "")
        page.export_dir.set_value(str(root / "reports"))
        page.export_filtered(core.FORMAT_JSONL)
        self.assertTrue(
            self.wait_for(lambda: "导出成功" in page.export_status.text()),
            f"status={page.export_status.text()!r} worker_present={page._worker is not None} thread_present={page._thread is not None} selected={page._selected_source()!r}",
        )
        exported = list((root / "reports").glob("*.jsonl"))
        self.assertEqual(len(exported), 1)
        self.assertEqual(len(exported[0].read_text(encoding="utf-8").splitlines()), 2)
        page.on_app_close()
        page.close()
        temp.cleanup()


if __name__ == "__main__":
    unittest.main()
