# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path
import tempfile
import unittest
import zipfile

from openpyxl import load_workbook

from core import xlsx
from core.xlsx_metadata import QualityItem
from core.xlsx_presentation import (
    TableLayout,
    append_sheet_directory,
    apply_quality_conditional_formats,
    configure_table,
    external_link_cell,
    finish_table,
    wrap_cell,
)


class XlsxPresentationTests(unittest.TestCase):
    def save(self, workbook: object, directory: Path) -> Path:
        path = directory / "presentation.xlsx"
        xlsx.save_workbook_atomic(workbook, path)
        self.assertTrue(path.exists())
        return path

    def test_real_workbook_preserves_a2_a4_and_finite_filters(self):
        with tempfile.TemporaryDirectory(prefix="xlsx_presentation_") as raw:
            root = Path(raw)
            workbook = xlsx.new_workbook()
            writer = xlsx.SheetWriter(workbook)

            first = workbook.create_sheet("header-one")
            first_layout = TableLayout(1, 2, 3)
            configure_table(first, first_layout)
            writer.header_row(first, ["编号", "正文"])
            writer.append([None, writer.wc("1", kind=xlsx.CellKind.ID),
                           wrap_cell(writer, "第一条正文")])
            finish_table(first, first_layout, 1)

            third = workbook.create_sheet("header-three")
            third_layout = TableLayout(3, 2, 3)
            configure_table(third, third_layout)
            writer.title_row(third, "标题", 2)
            writer.header_row(third, ["编号", "正文"])
            finish_table(third, third_layout, 0)

            path = self.save(workbook, root)
            workbook.close()
            reopened = load_workbook(path, read_only=False, data_only=False)
            try:
                self.assertEqual(reopened["header-one"].freeze_panes, "A2")
                self.assertEqual(reopened["header-one"].auto_filter.ref, "B1:C2")
                self.assertEqual(reopened["header-three"].freeze_panes, "A4")
                self.assertEqual(reopened["header-three"].auto_filter.ref, "B3:C3")
                self.assertEqual(reopened["header-one"].sheet_format.defaultRowHeight, 32.0)
                self.assertTrue(reopened["header-one"]["C2"].alignment.wrap_text)
                self.assertEqual(reopened["header-one"]._charts, [])
            finally:
                reopened.close()

    def test_internal_directory_uses_location_without_external_relationship(self):
        with tempfile.TemporaryDirectory(prefix="xlsx_directory_") as raw:
            root = Path(raw)
            workbook = xlsx.new_workbook()
            writer = xlsx.SheetWriter(workbook)
            overview = workbook.create_sheet("概览")
            target = workbook.create_sheet("综'述")
            append_sheet_directory(writer, overview, (("友好名称", target.title),))
            writer.ws = target
            writer.append([writer.wc("目标")])
            path = self.save(workbook, root)
            workbook.close()

            reopened = load_workbook(path, read_only=False, data_only=False)
            try:
                link = reopened["概览"]["B3"].hyperlink
                self.assertIsNotNone(link)
                self.assertEqual(link.location, "'综''述'!A1")
                self.assertIsNone(link.target)
            finally:
                reopened.close()
            with zipfile.ZipFile(path) as archive:
                rels = [name for name in archive.namelist()
                        if name.endswith("worksheets/_rels/sheet1.xml.rels")]
                self.assertEqual(rels, [])

    def test_external_link_validation_and_formula_safe_rejection(self):
        values = (
            "https://WWW.BILIBILI.COM/video/BV1x",
            "//space.bilibili.com/1",
            "javascript:alert(1)",
            "http://www.bilibili.com/video/BV1x",
            "https://bilibili.com.evil.example/",
            "https://user:pass@www.bilibili.com/",
            "https://www.bilibili.com:8443/",
            "=HYPERLINK(\"https://evil.example/\")",
            "https://www.bilibili.com/?token=secret-value",
        )
        with tempfile.TemporaryDirectory(prefix="xlsx_external_") as raw:
            root = Path(raw)
            workbook = xlsx.new_workbook()
            writer = xlsx.SheetWriter(workbook)
            sheet = workbook.create_sheet("链接")
            writer.ws = sheet
            cells = [external_link_cell(writer, value) for value in values]
            writer.append(cells)
            path = self.save(workbook, root)
            workbook.close()
            reopened = load_workbook(path, read_only=False, data_only=False)
            try:
                self.assertEqual(reopened["链接"]["A1"].hyperlink.target,
                                 "https://www.bilibili.com/video/BV1x")
                self.assertEqual(reopened["链接"]["B1"].hyperlink.target,
                                 "https://space.bilibili.com/1")
                for column in "CDEFGI":
                    self.assertIsNone(reopened["链接"][f"{column}1"].hyperlink)
                self.assertTrue(reopened["链接"]["H1"].quotePrefix)
                self.assertEqual(reopened["链接"]["I1"].value, "[已隐藏不安全链接]")
            finally:
                reopened.close()
            with zipfile.ZipFile(path) as archive:
                content = b"".join(archive.read(name) for name in archive.namelist())
                self.assertNotIn(b"secret-value", content)

    def test_external_link_rejects_control_and_overlong_targets(self):
        control = "https://www.bilibili.com/video/BV1\nunsafe"
        overlong = "https://www.bilibili.com/" + ("x" * 2048)
        with tempfile.TemporaryDirectory(prefix="xlsx_external_negative_") as raw:
            root = Path(raw)
            workbook = xlsx.new_workbook()
            writer = xlsx.SheetWriter(workbook)
            sheet = workbook.create_sheet("链接")
            writer.ws = sheet
            writer.append([external_link_cell(writer, control),
                           external_link_cell(writer, overlong)])
            path = self.save(workbook, root)
            workbook.close()
            reopened = load_workbook(path, read_only=False, data_only=False)
            try:
                self.assertIsNone(reopened["链接"]["A1"].hyperlink)
                self.assertIsNone(reopened["链接"]["B1"].hyperlink)
                self.assertEqual(reopened["链接"]["A1"].value,
                                 "[已隐藏不安全链接]")
                self.assertEqual(reopened["链接"]["B1"].value,
                                 "[已隐藏不安全链接]")
            finally:
                reopened.close()
            with zipfile.ZipFile(path) as archive:
                content = b"".join(archive.read(name) for name in archive.namelist())
                self.assertNotIn(b"x" * 128, content)

    def test_quality_rules_are_narrow_and_zero_is_not_colored(self):
        with tempfile.TemporaryDirectory(prefix="xlsx_quality_") as raw:
            root = Path(raw)
            workbook = xlsx.new_workbook()
            writer = xlsx.SheetWriter(workbook)
            sheet = workbook.create_sheet("数据质量")
            writer.ws = sheet
            writer.append([writer.wc("类别"), writer.wc("项目"), writer.wc("说明"),
                           writer.wc("数值"), writer.wc("状态"), writer.wc("备注")])
            writer.append([writer.wc("字段"), writer.wc("警告"), writer.wc(""),
                           writer.wc(None), writer.wc("状态"), writer.wc("")])
            writer.append([writer.wc("字段"), writer.wc("计数"), writer.wc(""),
                           writer.wc(0, kind=xlsx.CellKind.INTEGER), writer.wc("条"), writer.wc("")])
            items = (
                QualityItem("字段", "警告", xlsx.cell_value(None, xlsx.CellKind.NOT_RETURNED), "状态", "未返回"),
                QualityItem("字段", "计数", xlsx.cell_value(0, xlsx.CellKind.INTEGER), "条", "真实零值", "error"),
            )
            apply_quality_conditional_formats(sheet, 2, items)
            path = self.save(workbook, root)
            workbook.close()
            reopened = load_workbook(path, read_only=False, data_only=False)
            try:
                rules = reopened["数据质量"].conditional_formatting
                self.assertEqual(len(rules), 2)
                self.assertEqual({str(rule.sqref) for rule in rules}, {"B2:F2", "B3:F3"})
            finally:
                reopened.close()


if __name__ == "__main__":
    unittest.main()
