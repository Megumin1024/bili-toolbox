# -*- coding: utf-8 -*-
"""弹幕分析（热词 / 高频弹幕 / 时间轴热点）的离线测试。

全部离线，连模拟字节都不需要：分析是纯计算，喂 dict 就有结果。唯一碰 I/O 的
是最后那组导出/流水线用例，它们验的是"分析有没有真的进到产物里"。

这里钉死的几条容易在重构里丢掉的约束：
- 词频按"出现在几条弹幕里"算，同一条里重复不累加（否则 "哈哈哈" 会靠重叠
  n-gram 压过 "哈哈哈哈" 排到前面）；
- 热点分钟按 (分P, 分钟) 分开，两个分P的 00:00 不是同一分钟；
- 多分P不合并用户统计——mid_hash 跨分P是否同一个盐未经证实；
- 用户可见的文本里不许出现 None/nan。
"""
from __future__ import annotations

import tempfile
import unittest
from collections import Counter
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from core import text as text_mod
from tools.danmaku import analysis, core, danmaku_pb2, pipeline
from tools.danmaku.page import DanmakuPage

META = {"bvid": "BV1GJ411x7h7", "aid": 12345, "cid": 999, "duration": 720,
        "title": "测试视频", "owner": "测试UP", "page": 1, "page_count": 1,
        "claimed_danmaku": 100}


def row(content, page=1, progress_ms=0, mode=1, pool=0, color="#FFFFFF",
        ctime=1600000000, mid="h1"):
    """一条弹幕行，字段与 core.row_from_elem 的输出对齐。"""
    return {"id": "1", "page": page, "progress_ms": progress_ms,
            "time": core.hhmmss(progress_ms), "mode": mode,
            "mode_label": {1: "滚动", 4: "底部", 5: "顶部", 6: "逆向"}.get(mode, "滚动"),
            "content": content, "ctime": ctime, "sent_time": "",
            "mid_hash": mid, "color": 0, "color_hex": color, "fontsize": 25,
            "weight": 9, "pool": pool,
            "pool_label": {0: "普通", 1: "字幕", 2: "特殊"}.get(pool, "普通"),
            "attr": 0}


def repeat(content, n, **kw):
    return [row(content, **kw) for _ in range(n)]


# ============================ 分词 ============================

class NormalizeTests(unittest.TestCase):
    def test_collapses_inner_whitespace(self):
        self.assertEqual(text_mod.normalize("  a \n\t b  "), "a b")

    def test_none_and_empty_become_empty_string(self):
        self.assertEqual(text_mod.normalize(None), "")
        self.assertEqual(text_mod.normalize("   "), "")

    def test_meaningful_rejects_punctuation_and_single_char(self):
        for bad in ("", " ", "!!", "。。。", "6", "??"):
            self.assertFalse(text_mod.is_meaningful(bad), f"{bad!r} 不该算一个词")

    def test_meaningful_accepts_digits_and_words(self):
        for good in ("66", "666", "yyds", "前方", "哈哈"):
            self.assertTrue(text_mod.is_meaningful(good), f"{good!r} 应该算一个词")


class TokenizeTests(unittest.TestCase):
    def test_digit_runs_survive_whole(self):
        """666 是弹幕里最有信息量的串之一，不能被当噪音切掉。"""
        self.assertEqual(text_mod.tokenize("666")["666"], 1)

    def test_latin_runs_are_lowercased_whole(self):
        tokens = text_mod.tokenize("YYDS")
        self.assertEqual(tokens["yyds"], 1)

    def test_emoji_markers_are_not_words(self):
        self.assertNotIn("doge", text_mod.tokenize("[doge]"))
        self.assertEqual(text_mod.tokenize("[doge]"), {})

    def test_stopwords_are_dropped(self):
        self.assertNotIn("什么", text_mod.tokenize("什么"))

    def test_filler_char_ngrams_are_dropped(self):
        """滑窗会切出"的觉"这种半截词，靠字符黑名单挡掉。"""
        self.assertEqual(text_mod.tokenize("我觉得"), {})

    def test_repeated_word_within_one_text_counts_once(self):
        """同一条弹幕里重复出现不累加，否则重叠 n-gram 会把次数刷上天。"""
        self.assertEqual(text_mod.tokenize("哈哈哈哈")["哈哈哈"], 1)
        self.assertEqual(text_mod.tokenize("哈哈哈哈")["哈哈哈哈"], 1)

    def test_cjk_windows_of_all_three_sizes(self):
        tokens = text_mod.tokenize("前方高能")
        for expected in ("前方", "方高", "高能", "前方高", "方高能", "前方高能"):
            self.assertIn(expected, tokens)

    def test_none_input_is_empty(self):
        self.assertEqual(text_mod.tokenize(None), {})


class TopPhrasesTests(unittest.TestCase):
    def test_fragments_are_suppressed_by_the_longer_phrase(self):
        counter = Counter()
        for _ in range(5):
            counter.update(text_mod.tokenize("前方高能"))
        top = [w for w, _ in text_mod.top_phrases(counter, limit=10, min_count=2)]
        self.assertIn("前方高能", top)
        self.assertNotIn("方高", top, "碎片不该和整词一起霸榜")
        self.assertNotIn("方高能", top, "碎片不该和整词一起霸榜")

    def test_min_count_filters_singletons(self):
        counter = text_mod.tokenize("牛逼")
        self.assertEqual(text_mod.top_phrases(counter, min_count=2), [])

    def test_limit_is_respected(self):
        """造一批都过门槛、互不包含的词，再验只会留下 limit 个。"""
        words = ["甲乙", "丙丁", "戊己", "庚辛", "壬癸", "子丑", "寅卯",
                 "辰巳", "午未", "申酉", "戌亥"]
        counter = Counter()
        for word in words:
            for _ in range(3):
                counter.update(text_mod.tokenize(word))
        self.assertEqual(len(counter), len(words), "先确认候选确实多于 limit")
        self.assertEqual(len(text_mod.top_phrases(counter, limit=5)), 5)

    def test_ties_are_ordered_deterministically(self):
        """并列项的顺序不能取决于谁先出现，否则同一份数据结果会飘。"""
        a = text_mod.tokenize("牛逼 厉害")
        b = text_mod.tokenize("厉害 牛逼")
        self.assertEqual(text_mod.top_phrases(a), text_mod.top_phrases(b))

    def test_ranked_keeps_higher_count_first(self):
        counter = Counter()
        for _ in range(3):
            counter.update(text_mod.tokenize("666"))
        counter.update(text_mod.tokenize("不错"))
        self.assertEqual(text_mod.ranked(counter)[0], ("666", 3))


# ============================ 分析主体 ============================

class AnalyzeEmptyInputTests(unittest.TestCase):
    def test_empty_rows_yield_a_complete_empty_result(self):
        out = analysis.analyze([])
        self.assertEqual(out["total"], 0)
        self.assertFalse(out["multi"])
        self.assertEqual(out["top_danmaku"], [])
        self.assertEqual(out["top_words"], [])
        self.assertEqual(out["hot_minutes"], [])
        self.assertEqual(out["kpi"]["modes"], [])

    def test_none_rows_do_not_crash(self):
        self.assertEqual(analysis.analyze(None)["total"], 0)

    def test_rows_missing_every_optional_field_do_not_crash(self):
        out = analysis.analyze([{"content": "666"}, {"content": "666"},
                                {"content": "666"}])
        self.assertEqual(out["total"], 3)
        self.assertEqual(out["top_danmaku"][0][0], "666")


class AnalyzeKpiTests(unittest.TestCase):
    def test_total_and_percentages(self):
        rows = repeat("666", 3) + repeat("哈哈", 1)
        out = analysis.analyze(rows)
        self.assertEqual(out["total"], 4)
        self.assertAlmostEqual(out["top_danmaku"][0][2], 75.0)

    def test_colored_counts_only_non_white(self):
        rows = [row("a", color="#FFFFFF"), row("b", color="#FF0000"),
                row("c", color="#00FF00")]
        kpi = analysis.analyze(rows)["kpi"]
        self.assertEqual(kpi["colored"], 2)
        self.assertAlmostEqual(kpi["colored_pct"], 200 / 3)

    def test_avg_length(self):
        kpi = analysis.analyze([row("1234"), row("12")])["kpi"]
        self.assertAlmostEqual(kpi["avg_len"], 3.0)

    def test_modes_and_pools_are_counted(self):
        rows = [row("a", mode=1), row("b", mode=5), row("c", mode=5, pool=1)]
        kpi = analysis.analyze(rows)["kpi"]
        self.assertEqual(dict((k, v) for k, v, _ in kpi["modes"]),
                         {"滚动": 1, "顶部": 2})
        self.assertEqual(dict((k, v) for k, v, _ in kpi["pools"]),
                         {"普通": 2, "字幕": 1})

    def test_span_uses_first_and_last_send_time(self):
        rows = [row("a", ctime=1600000000), row("b", ctime=1600003600)]
        kpi = analysis.analyze(rows)["kpi"]
        self.assertAlmostEqual(kpi["span_hours"], 1.0)
        self.assertNotEqual(kpi["first_sent"], analysis.DASH)

    def test_single_row_has_zero_span_not_none(self):
        kpi = analysis.analyze([row("a")])["kpi"]
        self.assertEqual(kpi["span_hours"], 0.0)
        self.assertNotIn("None", str(kpi))


class HighFrequencyTests(unittest.TestCase):
    def test_threshold_is_three(self):
        self.assertEqual(analysis.analyze(repeat("666", 2))["top_danmaku"], [])
        self.assertEqual(analysis.analyze(repeat("666", 3))["top_danmaku"][0][0], "666")

    def test_single_char_danmaku_is_not_ranked(self):
        self.assertEqual(analysis.analyze(repeat("6", 10))["top_danmaku"], [])

    def test_punctuation_only_danmaku_is_not_ranked(self):
        self.assertEqual(analysis.analyze(repeat("！！！", 10))["top_danmaku"], [])

    def test_whitespace_variants_count_as_the_same_danmaku(self):
        rows = [row("前方高能"), row(" 前方高能 "), row("前方高能")]
        self.assertEqual(analysis.analyze(rows)["top_danmaku"][0][1], 3)

    def test_sorted_by_count_desc(self):
        rows = repeat("666", 5) + repeat("哈哈", 3) + repeat("厉害", 4)
        top = [c for c, _, _ in analysis.analyze(rows)["top_danmaku"]]
        self.assertEqual(top, ["666", "厉害", "哈哈"])


class HotMinuteTests(unittest.TestCase):
    def test_multi_part_keeps_the_same_minute_apart(self):
        """P1 的第 0 分钟和 P2 的第 0 分钟是两档，不能合并。"""
        rows = (repeat("p1", 5, page=1, progress_ms=0)
                + repeat("p2", 3, page=2, progress_ms=0))
        hot = analysis.analyze(rows)["hot_minutes"]
        keys = [(m["page"], m["minute"]) for m in hot]
        self.assertEqual(keys, [(1, 0), (2, 0)])

    def test_single_part_has_page_one(self):
        hot = analysis.analyze(repeat("a", 3, progress_ms=120_000))["hot_minutes"]
        self.assertEqual(hot[0]["page"], 1)
        self.assertEqual(hot[0]["minute"], 2)

    def test_examples_are_the_most_repeated_in_that_minute(self):
        rows = (repeat("666", 5, progress_ms=1_000)
                + repeat("哈哈", 2, progress_ms=1_000)
                + repeat("别的", 4, progress_ms=200_000))
        hot = analysis.analyze(rows)["hot_minutes"]
        self.assertEqual(hot[0]["examples"][0], "666（5次）")

    def test_minute_without_repeats_has_no_examples(self):
        """一分钟里全是各说各的，就不该硬凑"刷得最多"的那句话。"""
        rows = [row("甲", progress_ms=1_000), row("乙", progress_ms=2_000)]
        hot = analysis.analyze(rows)["hot_minutes"]
        self.assertTrue(all(m["examples"] == [] for m in hot))

    def test_capped_at_the_limit(self):
        rows = [row("a", progress_ms=i * 60_000) for i in range(40)]
        self.assertEqual(len(analysis.analyze(rows)["hot_minutes"]),
                         analysis.HOT_MINUTE_LIMIT)

    def test_sorted_by_count_desc(self):
        rows = (repeat("a", 5, progress_ms=0)
                + repeat("b", 2, progress_ms=600_000))
        hot = analysis.analyze(rows)["hot_minutes"]
        self.assertEqual([m["count"] for m in hot], [5, 2])

    def test_pct_is_share_of_all_danmaku(self):
        rows = repeat("a", 3, progress_ms=0) + repeat("b", 1, progress_ms=600_000)
        hot = analysis.analyze(rows)["hot_minutes"]
        self.assertAlmostEqual(hot[0]["pct"], 75.0)


class UserStatsTests(unittest.TestCase):
    def test_single_part_counts_distinct_hashes(self):
        rows = [row("a", mid="h1"), row("b", mid="h1"), row("c", mid="h2")]
        out = analysis.analyze(rows)
        self.assertEqual(out["users"], 2)
        self.assertEqual(out["users_note"], "")

    def test_top10_share(self):
        rows = repeat("a", 5, mid="big") + repeat("b", 5, mid="small")
        self.assertAlmostEqual(analysis.analyze(rows)["top10_user_pct"], 100.0)

    def test_multi_part_leaves_users_empty_with_a_reason(self):
        rows = repeat("a", 3, page=1) + repeat("b", 3, page=2)
        out = analysis.analyze(rows)
        self.assertIsNone(out["users"])
        self.assertIsNone(out["top10_user_pct"])
        self.assertIn("盐", out["users_note"], "留空必须写明原因，不能只给个横杠")

    def test_multi_part_derived_from_rows_even_without_parts_arg(self):
        """没传 parts 也要认出这是多分P，不能默认去合并用户。"""
        rows = repeat("a", 3, page=1) + repeat("b", 3, page=2)
        self.assertTrue(analysis.analyze(rows)["multi"])

    def test_rows_without_hashes_give_zero_users(self):
        rows = [row("a", mid=""), row("b", mid=None)]
        self.assertEqual(analysis.analyze(rows)["users"], 0)


# ============================ Markdown ============================

class MarkdownTests(unittest.TestCase):
    def _md(self, rows=None, meta=None, parts=None):
        ana = analysis.analyze(rows or [], meta or META, parts)
        return analysis.render_markdown(ana, meta or META)

    def test_never_prints_none_or_nan(self):
        """上一轮 _completion_note 把 "None" 印进用户文案的坑，这里双向钉死。"""
        for rows in ([], repeat("666", 3), [{"content": "x"}]):
            md = self._md(rows)
            self.assertNotIn("None", md)
            self.assertNotIn("nan", md)

    def test_has_all_sections(self):
        md = self._md(repeat("666", 3, progress_ms=1_000))
        for section in ("# 弹幕分析报告", "## 高频弹幕", "## 热词",
                        "## 热点分钟", "## 构成"):
            self.assertIn(section, md)

    def test_empty_input_says_so_instead_of_an_empty_table(self):
        md = self._md([])
        self.assertIn("没有出现达到 3 次的重复弹幕", md)
        self.assertIn("没有弹幕，无法给出热点分钟", md)

    def test_pipes_in_content_do_not_break_the_table(self):
        md = self._md(repeat("a|b", 3))
        self.assertIn("a\\|b", md)
        self.assertNotIn("| a|b |", md)

    def test_multi_part_adds_the_page_column(self):
        rows = repeat("a", 3, page=1, progress_ms=0) + repeat("b", 3, page=2)
        self.assertIn("| 分P | 分钟区间 |", self._md(rows))

    def test_single_part_has_no_page_column(self):
        md = self._md(repeat("a", 3, progress_ms=0))
        self.assertNotIn("| 分P | 分钟区间 |", md)
        self.assertIn("| 分钟区间 | 条数 |", md)

    def test_multi_part_states_that_words_are_merged(self):
        rows = repeat("a", 3, page=1) + repeat("b", 3, page=2)
        self.assertIn("跨全部分P合并统计", self._md(rows))

    def test_users_dash_when_not_merged(self):
        rows = repeat("a", 3, page=1) + repeat("b", 3, page=2)
        self.assertIn(f"| 独立发送者 | {analysis.DASH} |", self._md(rows))


# ============================ 导出集成 ============================

class AnalysisSheetTests(unittest.TestCase):
    def _wb(self, rows, analysis_arg=True, parts=None, multi=False):
        tmp = Path(tempfile.mkdtemp(prefix="danmaku_ana_"))
        path = tmp / "out.xlsx"
        stats = {"segments": 1, "expected_segments": 1, "rows": len(rows),
                 "truncated": False, "cancelled": False, "duplicates": 0}
        ana = (analysis.analyze(rows, META, parts) if analysis_arg else None)
        core.export_xlsx(rows, META, stats, path, parts=parts, analysis=ana)
        return load_workbook(path)

    def test_sheet_is_absent_when_analysis_is_not_passed(self):
        """保持向后兼容：不传分析就不出这张表。"""
        wb = self._wb(repeat("666", 3), analysis_arg=False)
        self.assertNotIn("弹幕分析", wb.sheetnames)

    def test_sheet_present_when_analysis_is_passed(self):
        wb = self._wb(repeat("666", 3))
        self.assertIn("弹幕分析", wb.sheetnames)

    def test_sheet_sits_right_after_the_overview(self):
        wb = self._wb(repeat("666", 3))
        self.assertEqual(wb.sheetnames.index("弹幕分析"),
                         wb.sheetnames.index("概览") + 1)

    def test_header_row_covers_both_rankings(self):
        ws = self._wb(repeat("666", 3))["弹幕分析"]
        headers = [c.value for row in ws.iter_rows() for c in row if c.value]
        self.assertIn("高频弹幕", headers)
        self.assertIn("热词", headers)
        self.assertIn("出现在几条弹幕", headers)
        self.assertIn("该分钟刷得最多", headers)

    def test_ranked_content_actually_lands_in_the_sheet(self):
        ws = self._wb(repeat("666", 3) + repeat("牛逼", 4))["弹幕分析"]
        cells = [c.value for row in ws.iter_rows() for c in row if c.value]
        self.assertIn("666", cells)
        self.assertIn("牛逼", cells)

    def test_values_contain_no_none(self):
        """空单元格本来就是 None，这里要抓的是把 None 当**文本**写进去。"""
        ws = self._wb(repeat("666", 3, progress_ms=1_000))["弹幕分析"]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    self.assertNotIn("None", str(cell.value))

    def test_hot_minute_header_gains_a_page_column_only_when_multi(self):
        rows = repeat("a", 3, page=1, progress_ms=0) + repeat("b", 3, page=2)
        parts = [{"page": 1, "part": "P1", "duration": 60, "rows": 3,
                  "segments": 1, "expected_segments": 1, "duplicates": 0,
                  "cancelled": False, "truncated": False},
                 {"page": 2, "part": "P2", "duration": 60, "rows": 3,
                  "segments": 1, "expected_segments": 1, "duplicates": 0,
                  "cancelled": False, "truncated": False}]
        ws = self._wb(rows, parts=parts)["弹幕分析"]
        cells = [c.value for row in ws.iter_rows() for c in row if c.value]
        self.assertIn("分P", cells)
        self.assertIn("P1", cells)

    def test_single_part_has_no_page_label_in_the_analysis_sheet(self):
        ws = self._wb(repeat("a", 3, progress_ms=0))["弹幕分析"]
        cells = [c.value for row in ws.iter_rows() for c in row if c.value]
        self.assertNotIn("P1", cells)

    def test_small_sample_says_so(self):
        ws = self._wb([row("各家各话")])["弹幕分析"]
        text = " ".join(str(c.value) for row in ws.iter_rows() for c in row
                        if c.value)
        self.assertIn("样本太少", text)


# ============================ 流水线 ============================

def _api_payload(claimed=10):
    return {"code": 0, "data": {
        "bvid": "BV1GJ411x7h7", "aid": 12345, "cid": 999, "duration": 720,
        "title": "测试视频", "owner": {"name": "测试UP"}, "pubdate": 1600000000,
        "stat": {"danmaku": claimed}, "pages": [{"cid": 999, "duration": 720,
                                                 "page": 1, "part": "P1"}]}}


def _segment(*contents):
    reply = danmaku_pb2.DmSegMobileReply()
    for i, content in enumerate(contents, 1):
        e = reply.elems.add()
        e.id = i
        e.idStr = str(i)
        e.progress = 1000 * i
        e.mode = 1
        e.color = 16777215
        e.ctime = 1600000000
        e.pool = 0
        e.midHash = f"h{i}"
        e.content = content
    return reply.SerializeToString()


class PipelineAnalysisTests(unittest.TestCase):
    def _run(self, contents, **kw):
        out = tempfile.mkdtemp(prefix="danmaku_ana_pipe_")
        fetch = lambda url, **k: _segment(*contents) if "segment_index=1" in url else b""
        with patch("core.session.http_get_json",
                   return_value=_api_payload(claimed=len(contents))), \
                patch("core.session.http_get_bytes", side_effect=fetch):
            return pipeline.run_pipeline("BV1GJ411x7h7", out, sleep=0,
                                         cancel=lambda: False, **kw)

    def test_report_file_is_written(self):
        result = self._run(["666", "666", "666"])
        self.assertTrue(result["report"])
        report = Path(result["report"])
        self.assertTrue(report.exists())
        self.assertIn("弹幕分析报告", report.read_text(encoding="utf-8"))

    def test_result_carries_the_analysis(self):
        result = self._run(["666", "666", "666"])
        self.assertEqual(result["analysis"]["top_danmaku"][0][0], "666")

    def test_report_failure_does_not_sink_the_task(self):
        """报告是附加产物，写不出来不能把已经落盘的 Excel 一起废掉。"""
        logs = []
        with patch("tools.danmaku.analysis.render_markdown",
                   side_effect=OSError("磁盘满了")):
            result = self._run(["666", "666", "666"],
                               progress=lambda **kw: logs.append(kw))
        self.assertEqual(result["report"], "")
        self.assertTrue(Path(result["xlsx"]).exists())
        self.assertTrue(any(kw.get("level") == "warn" for kw in logs),
                        "降级必须说出来，不能假装成功")

    def test_jsonl_is_untouched_by_the_analysis(self):
        """分析是只读的，不能把已经落盘的原始数据改了。"""
        result = self._run(["666", "666", "666"])
        raw = Path(result["jsonl"]).read_text(encoding="utf-8")
        self.assertIn("666", raw)


# ============================ 页面 ============================

class AnalysisPageTests(unittest.TestCase):
    """沿用仓库既有做法（`__new__` + 假控件），不启动 GUI。"""

    class _Card:
        def __init__(self):
            self.links = None

        def show_result(self, title, links):
            self.links = list(links)

    class _Edit:
        def __init__(self, text=""):
            self._t = text

        def setText(self, value):
            self._t = value

        def text(self):
            return self._t

        def setFocus(self):
            pass

    class _Check:
        def __init__(self, checked=False):
            self._c = checked

        def setChecked(self, value):
            self._c = bool(value)

        def isChecked(self):
            return self._c

    class _Row:
        def __init__(self, value=""):
            self._v = value

        def set_value(self, value):
            self._v = value

        def value(self):
            return self._v

    def _page(self):
        page = DanmakuPage.__new__(DanmakuPage)
        page.target_edit = self._Edit("BV1GJ411x7h7")
        page.out_row = self._Row("D:\\out")
        page.segments_edit = self._Edit("20")
        page.sleep_edit = self._Edit("0.8")
        page.all_pages = self._Check(False)
        page.auto_open = self._Check(True)
        page.result_card = self._Card()
        return page

    def test_history_paths_include_the_report(self):
        page = self._page()
        self.assertEqual(page.history_output_paths(
            {"xlsx": "a.xlsx", "report": "r.md", "jsonl": "b.jsonl"}),
            ["a.xlsx", "r.md", "b.jsonl"])

    def test_history_paths_drop_an_empty_report(self):
        """空路径会被 Path("") 当成当前目录，点开就落到别处。"""
        page = self._page()
        self.assertEqual(page.history_output_paths(
            {"xlsx": "a.xlsx", "report": "", "jsonl": "b.jsonl"}),
            ["a.xlsx", "b.jsonl"])

    def test_finished_card_offers_the_report(self):
        page = self._page()
        page.on_finished({"rows": 3, "stats": {}, "parts": [],
                          "xlsx": "a.xlsx", "report": "r.md",
                          "jsonl": "b.jsonl", "dir": "d"})
        labels = [label for label, _ in page.result_card.links]
        self.assertIn("分析报告(MD)", labels)

    def test_finished_card_hides_the_report_when_it_failed(self):
        page = self._page()
        page.on_finished({"rows": 3, "stats": {}, "parts": [],
                          "xlsx": "a.xlsx", "report": "",
                          "jsonl": "b.jsonl", "dir": "d"})
        labels = [label for label, _ in page.result_card.links]
        self.assertNotIn("分析报告(MD)", labels)
        self.assertNotIn(("分析报告(MD)", ""), page.result_card.links)


if __name__ == "__main__":
    unittest.main()
