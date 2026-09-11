# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import os
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import Mock, patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from core.risk import RiskChallengeError
from tools.collector import pipeline as collector_pipeline
from tools.collector.comparison import SessionComparison, numeric_value, sort_rows


def snapshot(bvid, view=100, like=10, fetched_at=1_700_000_000, title=None, owner=None):
    return {
        "bvid": bvid,
        "aid": 1,
        "title": title or f"标题-{bvid}",
        "owner": owner or f"UP-{bvid}",
        "owner_mid": 2,
        "tname": "测试",
        "pubdate": 1_600_000_000,
        "duration": 60,
        "view": view,
        "danmaku": 1,
        "reply": 2,
        "favorite": 3,
        "coin": 4,
        "share": 5,
        "like": like,
        "fetched_at": fetched_at,
    }


def write_snapshot(path, value):
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(value, ensure_ascii=False) + "\n")


class ComparisonAggregationTests(unittest.TestCase):
    def rows(self, payload):
        return {row["bvid"]: row for row in payload["rows"]}

    def test_multiple_videos_total_delta_growth_and_engagement(self):
        comparison = SessionComparison(["BV1", "BV2", "BV3"])
        first = comparison.publish_round([
            snapshot("BV1", 100, 10, 1_700_000_000),
            snapshot("BV2", 200, 20, 1_700_000_000),
            snapshot("BV3", 300, 0, 1_700_000_000),
        ])
        self.assertEqual(first["summary"]["total_view"], 600)
        self.assertIsNone(self.rows(first)["BV1"]["delta_view"])
        second = comparison.publish_round([
            snapshot("BV1", 130, 13, 1_700_003_600),
            snapshot("BV2", 150, 15, 1_700_003_600),
            snapshot("BV3", 330, 33, 1_700_003_600),
        ], round_number=2, state="completed")
        rows = self.rows(second)
        self.assertEqual(rows["BV1"]["delta_view"], 30)
        self.assertEqual(rows["BV2"]["delta_view"], -50)
        self.assertAlmostEqual(rows["BV1"]["view_per_hour"], 30.0)
        self.assertAlmostEqual(rows["BV2"]["view_per_hour"], -50.0)
        self.assertAlmostEqual(rows["BV1"]["engagement_rate"], 10.0)
        self.assertEqual(second["summary"]["total_delta_view"], 10)

    def test_single_snapshot_has_no_delta_or_growth(self):
        comparison = SessionComparison(["BV1", "BV2"])
        payload = comparison.publish_round([snapshot("BV1")], failed=["BV2"])
        rows = self.rows(payload)
        self.assertIsNone(rows["BV1"]["delta_view"])
        self.assertIsNone(rows["BV1"]["delta_like"])
        self.assertIsNone(rows["BV1"]["view_per_hour"])
        self.assertEqual(rows["BV2"]["status"], "本轮失败")

    def test_same_backwards_and_missing_time_have_no_growth(self):
        for first_time, last_time in ((100, 100), (200, 100), (None, 200)):
            with self.subTest(first_time=first_time, last_time=last_time):
                comparison = SessionComparison(["BV1"])
                payload = comparison.publish_round([
                    snapshot("BV1", 100, 10, first_time),
                    snapshot("BV1", 120, 12, last_time),
                ])
                self.assertIsNone(self.rows(payload)["BV1"]["view_per_hour"])

    def test_numeric_safety_for_zero_missing_bool_strings_and_long_values(self):
        cases = (
            (0, 10, None),
            (None, 10, None),
            (True, 10, None),
            ("100", "10", 10.0),
            ("9" * 21, "10", None),
        )
        for view, like, expected_engagement in cases:
            with self.subTest(view_type=type(view).__name__):
                comparison = SessionComparison(["BV1"])
                payload = comparison.publish_round([snapshot("BV1", view=view, like=like)])
                row = self.rows(payload)["BV1"]
                if isinstance(view, str) and len(view) > 20:
                    self.assertIsNone(row["view"])
                    self.assertIsNone(row["engagement_rate"])
                else:
                    self.assertEqual(row["engagement_rate"], expected_engagement)

    def test_numeric_value_rejects_overlong_decimal_int_and_timestamp_values(self):
        self.assertEqual(numeric_value(10 ** 20 - 1), 10 ** 20 - 1)
        self.assertEqual(numeric_value(-(10 ** 20 - 1)), -(10 ** 20 - 1))
        self.assertIsNone(numeric_value(10 ** 20))
        self.assertEqual(numeric_value("9" * 20), int("9" * 20))
        for value in (
            "9" * 21,
            "9" * 5000,
            Decimal("9" * 21),
            10 ** 21,
            10 ** 4999,
            -(10 ** 4999),
            "1e21",
            Decimal("1e999999999999"),
            float("inf"),
            float("nan"),
        ):
            with self.subTest(value_type=type(value).__name__):
                self.assertIsNone(numeric_value(value))

    def test_huge_int_snapshot_fields_are_ignored_without_crashing(self):
        huge = 10 ** 4999
        comparison = SessionComparison(["BV1", "BV2"])
        first = comparison.publish_round([
            snapshot("BV1", huge, huge, huge),
            snapshot("BV2", 100, 10, 1_700_000_000),
        ])
        rows = self.rows(first)
        huge_row = rows["BV1"]
        self.assertIsNone(huge_row["view"])
        self.assertIsNone(huge_row["like"])
        self.assertIsNone(huge_row["fetched_at"])
        self.assertIsNone(huge_row["delta_view"])
        self.assertIsNone(huge_row["delta_like"])
        self.assertIsNone(huge_row["view_per_hour"])
        self.assertIsNone(huge_row["engagement_rate"])
        self.assertEqual(first["summary"]["total_view"], 100)

        second = comparison.publish_round([
            snapshot("BV1", 50, 5, 1_700_003_600),
            snapshot("BV2", 130, 13, 1_700_003_600),
        ], round_number=2)
        rows = self.rows(second)
        self.assertEqual(rows["BV1"]["view"], 50)
        self.assertEqual(rows["BV1"]["engagement_rate"], 10.0)
        self.assertIsNone(rows["BV1"]["delta_view"])
        self.assertIsNone(rows["BV1"]["view_per_hour"])
        self.assertEqual(rows["BV2"]["delta_view"], 30)
        self.assertEqual(rows["BV2"]["view_per_hour"], 30.0)

    def test_overlong_values_are_excluded_but_later_legal_samples_work(self):
        overlong = "9" * 21
        comparison = SessionComparison(["BV1", "BV2", "BV3"])
        first = comparison.publish_round([
            snapshot("BV1", overlong, overlong, overlong),
            snapshot("BV2", 100, 10, 1_700_000_000),
            snapshot("BV3", 200, 20, 1_700_000_000),
        ])
        first_rows = self.rows(first)
        self.assertIsNone(first_rows["BV1"]["view"])
        self.assertIsNone(first_rows["BV1"]["like"])
        self.assertIsNone(first_rows["BV1"]["fetched_at"])
        self.assertEqual(first["summary"]["total_view"], 300)
        self.assertEqual([row["bvid"] for row in sort_rows(first["rows"], "view")],
                         ["BV3", "BV2", "BV1"])

        second = comparison.publish_round([
            snapshot("BV1", 50, 5, 1_700_003_600),
            snapshot("BV2", 130, 13, 1_700_003_600),
            snapshot("BV3", 230, 23, 1_700_003_600),
        ], round_number=2)
        rows = self.rows(second)
        self.assertEqual(rows["BV1"]["view"], 50)
        self.assertEqual(rows["BV1"]["engagement_rate"], 10.0)
        self.assertIsNone(rows["BV1"]["delta_view"])
        self.assertIsNone(rows["BV1"]["view_per_hour"])
        self.assertEqual(rows["BV2"]["delta_view"], 30)
        self.assertEqual(rows["BV3"]["delta_view"], 30)

    def test_view_decrease_keeps_negative_delta_and_growth(self):
        comparison = SessionComparison(["BV1"])
        comparison.publish_round([snapshot("BV1", 100, 10, 1_700_000_000)])
        payload = comparison.publish_round(
            [snapshot("BV1", 90, 9, 1_700_003_600)], round_number=2
        )
        row = self.rows(payload)["BV1"]
        self.assertEqual(row["delta_view"], -10)
        self.assertEqual(row["delta_like"], -1)
        self.assertEqual(row["view_per_hour"], -10.0)

    def test_old_snapshots_are_not_used_as_session_baseline(self):
        comparison = SessionComparison(["BV1"])
        comparison.publish_round([snapshot("BV1", 100, 10, 1_700_000_000)])
        payload = comparison.publish_round([snapshot("BV1", 130, 13, 1_700_003_600)], round_number=2)
        self.assertEqual(self.rows(payload)["BV1"]["delta_view"], 30)

    def test_first_round_failure_then_success_starts_new_baseline(self):
        comparison = SessionComparison(["BV1", "BV2"])
        comparison.publish_round([snapshot("BV2", 200, 20)], failed=["BV1"])
        payload = comparison.publish_round([
            snapshot("BV1", 50, 5, 1_700_003_600),
            snapshot("BV2", 220, 22, 1_700_003_600),
        ], round_number=2)
        rows = self.rows(payload)
        self.assertEqual(rows["BV1"]["sample_count"], 1)
        self.assertIsNone(rows["BV1"]["delta_view"])
        self.assertEqual(rows["BV2"]["delta_view"], 20)

    def test_round_failure_keeps_previous_successful_value(self):
        comparison = SessionComparison(["BV1", "BV2"])
        comparison.publish_round([snapshot("BV1", 100), snapshot("BV2", 200)])
        payload = comparison.publish_round([snapshot("BV1", 110)], failed=["BV2"], round_number=2)
        rows = self.rows(payload)
        self.assertEqual(rows["BV2"]["view"], 200)
        self.assertEqual(rows["BV2"]["status"], "本轮失败")
        self.assertEqual(rows["BV1"]["status"], "本轮成功")

    def test_one_video_failure_does_not_hide_other_videos(self):
        comparison = SessionComparison(["BV1", "BV2", "BV3"])
        payload = comparison.publish_round([snapshot("BV1", 100), snapshot("BV3", 300)], failed=["BV2"])
        rows = self.rows(payload)
        self.assertEqual(rows["BV1"]["view"], 100)
        self.assertEqual(rows["BV3"]["view"], 300)
        self.assertEqual(rows["BV2"]["status"], "本轮失败")

    def test_rounds_are_grouped_by_bvid(self):
        comparison = SessionComparison(["BV1", "BV2"])
        comparison.publish_round([snapshot("BV1", 100), snapshot("BV2", 200)])
        payload = comparison.publish_round([
            snapshot("BV2", 260, 26, 1_700_003_600),
            snapshot("BV1", 130, 13, 1_700_003_600),
        ], round_number=2)
        rows = self.rows(payload)
        self.assertEqual(rows["BV1"]["delta_view"], 30)
        self.assertEqual(rows["BV2"]["delta_view"], 60)

    def test_duplicate_bvid_in_one_round_has_one_row_and_uses_last_record(self):
        comparison = SessionComparison(["BV1", "BV2"])
        payload = comparison.publish_round([
            snapshot("BV1", 100),
            snapshot("BV1", 120),
            snapshot("BV2", 200),
        ])
        rows = self.rows(payload)
        self.assertEqual(len(payload["rows"]), 2)
        self.assertEqual(rows["BV1"]["view"], 120)

    def test_sorting_is_stable_by_bvid_for_equal_values(self):
        rows = [
            {"bvid": "BV2", "view": 100},
            {"bvid": "BV1", "view": 100},
            {"bvid": "BV3", "view": None},
        ]
        self.assertEqual([row["bvid"] for row in sort_rows(rows, "view")], ["BV1", "BV2", "BV3"])
        self.assertEqual([row["bvid"] for row in sort_rows(rows, "view", descending=False)],
                         ["BV1", "BV2", "BV3"])


class PipelineRoundBoundaryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.bvids = ["BV1", "BV2", "BV3", "BV4"]

    def tearDown(self):
        self.temp.cleanup()

    def expand(self, _line, progress=None):
        return [(bvid, "") for bvid in self.bvids]

    def run_pipeline_case(self, collect, capture_export=None, **kwargs):
        events = []
        def export_xlsx(*args, **values):
            if capture_export is not None:
                capture_export.append((args, values))

        with patch.object(collector_pipeline.session, "ensure_ready"), \
                patch.object(collector_pipeline.links, "expand_source", side_effect=self.expand), \
                patch.object(collector_pipeline.core, "collect_snapshot", side_effect=collect), \
                patch.object(collector_pipeline.core, "export_xlsx", side_effect=export_xlsx), \
                patch.object(collector_pipeline.risk, "risk_recovery_flow", return_value=True), \
                patch.object(collector_pipeline.time, "sleep", return_value=None):
            result = collector_pipeline.run_pipeline(
                ["BV1"], self.root, sleep=0, progress=lambda **kw: events.append(kw), **kwargs
            )
        return result, events

    @staticmethod
    def write_progress(snapshot_path, progress, values, done, total, ok, fail):
        write_snapshot(snapshot_path, values)
        progress(done=done, total=total, ok=ok, fail=fail, text=f"采集 {done}/{total}")

    def test_risk_recovery_keeps_same_offset_and_includes_all_successes(self):
        calls = []
        export_calls = []
        count = 0

        def collect(values, progress, snapshot_path, **_kwargs):
            nonlocal count
            count += 1
            calls.append(list(values))
            if count == 1:
                self.write_progress(snapshot_path, progress, snapshot("BV1", 100), 1, 4, 1, 0)
                progress(done=2, total=4, ok=1, fail=1, text="普通失败 BV2")
                self.write_progress(snapshot_path, progress, snapshot("BV3", 300), 3, 4, 2, 1)
                error = RiskChallengeError("voucher")
                error.resume_index = 3
                raise error
            if count == 2:
                error = RiskChallengeError("voucher-again")
                error.resume_index = 0
                raise error
            self.write_progress(snapshot_path, progress, snapshot("BV4", 400), 1, 1, 1, 0)
            return [snapshot("BV4", 400)], []

        result, events = self.run_pipeline_case(collect, capture_export=export_calls)
        rows = {row["bvid"]: row for row in result["dashboard"]["rows"]}
        self.assertEqual(calls, [self.bvids, ["BV4"], ["BV4"]])
        self.assertEqual(set(rows), set(self.bvids))
        self.assertEqual(rows["BV1"]["view"], 100)
        self.assertEqual(rows["BV3"]["view"], 300)
        self.assertEqual(rows["BV4"]["view"], 400)
        self.assertEqual(rows["BV2"]["status"], "本轮失败")
        self.assertEqual(result["fail"], 1)
        self.assertEqual(result["dashboard"]["state"], "partial_failure")
        self.assertTrue(any("成功 3，失败 1" in event.get("text", "") for event in events))
        self.assertEqual(export_calls[-1][0][2]["ok"], 3)
        self.assertEqual(export_calls[-1][0][2]["attempted"], 4)
        self.assertNotIn("voucher", json.dumps(events, ensure_ascii=False).lower())

    def test_cancel_mid_round_does_not_publish_partial_dashboard(self):
        def collect(values, progress, snapshot_path, **_kwargs):
            self.assertEqual(list(values), self.bvids)
            self.write_progress(snapshot_path, progress, snapshot("BV1", 100), 1, 4, 1, 0)
            return [snapshot("BV1", 100)], []

        result, events = self.run_pipeline_case(collect, monitor=True, rounds=3, cancel=lambda: True)
        self.assertEqual(result["dashboard"]["rounds"], 0)
        dashboard_events = [event for event in events if "dashboard" in event]
        self.assertTrue(dashboard_events)
        self.assertTrue(all(event["dashboard"]["rounds"] == 0 for event in dashboard_events))

    def test_cancel_after_last_video_allows_complete_round(self):
        def collect(values, progress, snapshot_path, **_kwargs):
            out = []
            for index, bvid in enumerate(values, 1):
                value = snapshot(bvid, index * 100)
                out.append(value)
                self.write_progress(snapshot_path, progress, value, index, len(values), index, 0)
            return out, []

        result, _events = self.run_pipeline_case(collect, cancel=lambda: True)
        self.assertEqual(result["dashboard"]["rounds"], 1)
        self.assertEqual(result["dashboard"]["summary"]["total_view"], 1000)

    def test_old_full_and_partial_records_do_not_mix(self):
        snapshot_path = self.root / "snapshots.jsonl"
        write_snapshot(snapshot_path, snapshot("BV1", 1, 1, 1))
        calls = 0
        cancel_calls = 0

        def cancel():
            nonlocal cancel_calls
            cancel_calls += 1
            return cancel_calls >= 3

        def collect(values, progress, snapshot_path, **_kwargs):
            nonlocal calls
            calls += 1
            if calls == 1:
                for index, bvid in enumerate(("BV1", "BV2", "BV3", "BV4"), 1):
                    value = snapshot(bvid, index * 100, index, 1_700_000_000)
                    self.write_progress(snapshot_path, progress, value, index, 4, index, 0)
                return [snapshot(bvid, index * 100, index, 1_700_000_000)
                        for index, bvid in enumerate(self.bvids, 1)], []
            value = snapshot("BV1", 999, 99, 1_700_003_600)
            self.write_progress(snapshot_path, progress, value, 1, 4, 1, 0)
            return [value], []

        result, _events = self.run_pipeline_case(collect, monitor=True, rounds=2, interval_min=0, cancel=cancel)
        rows = {row["bvid"]: row for row in result["dashboard"]["rows"]}
        self.assertEqual(result["dashboard"]["rounds"], 1)
        self.assertEqual(rows["BV1"]["view"], 100)
        self.assertEqual(rows["BV2"]["view"], 200)
        self.assertNotEqual(rows["BV1"]["view"], 1)
        self.assertNotEqual(rows["BV1"]["view"], 999)

    def test_pipeline_is_serial_and_fetches_each_video_once_per_round(self):
        calls = []

        def fake_fetch(bvid, cancel=None):
            calls.append(bvid)
            return snapshot(bvid, 100 + len(calls), 10 + len(calls), 1_700_000_000 + len(calls))

        with patch.object(collector_pipeline.session, "ensure_ready"), \
                patch.object(collector_pipeline.links, "expand_source", side_effect=self.expand), \
                patch.object(collector_pipeline.core, "fetch_view", side_effect=fake_fetch), \
                patch.object(collector_pipeline.core, "export_xlsx"), \
                patch.object(collector_pipeline.core.time, "sleep", return_value=None):
            collector_pipeline.run_pipeline(["BV1"], self.root, sleep=0, monitor=True,
                                            rounds=2, interval_min=0, cancel=lambda: False)
        self.assertEqual(calls, self.bvids + self.bvids)

    def test_progress_dashboard_has_no_credentials_or_paths(self):
        def collect(values, progress, snapshot_path, **_kwargs):
            out = []
            for index, bvid in enumerate(values, 1):
                value = snapshot(bvid, index * 100)
                out.append(value)
                self.write_progress(snapshot_path, progress, value, index, len(values), index, 0)
            return out, []

        _result, events = self.run_pipeline_case(collect)
        forbidden = ("cookie", "token", "proxy", "password", "authorization",
                     str(self.root).lower().replace("\\", ""))
        for event in events:
            serialized = json.dumps(event, ensure_ascii=False).lower().replace("\\", "")
            self.assertFalse(any(value in serialized for value in forbidden), serialized)
        for event in events:
            if "dashboard" in event:
                self.assertNotIn("path", json.dumps(event["dashboard"], ensure_ascii=False).lower())

    def test_snapshots_keys_and_excel_structure_stay_unchanged(self):
        def fake_fetch(bvid, cancel=None):
            return snapshot(bvid, 100, 10)

        with patch.object(collector_pipeline.session, "ensure_ready"), \
                patch.object(collector_pipeline.links, "expand_source", side_effect=self.expand), \
                patch.object(collector_pipeline.core, "fetch_view", side_effect=fake_fetch), \
                patch.object(collector_pipeline.core.time, "sleep", return_value=None):
            result = collector_pipeline.run_pipeline(["BV1"], self.root, sleep=0)
        snapshot_lines = [json.loads(line) for line in (self.root / "snapshots.jsonl").read_text(encoding="utf-8").splitlines()]
        self.assertEqual(set(snapshot_lines[0]), set(snapshot("BV1")))
        from openpyxl import load_workbook
        workbook = load_workbook(result["xlsx"], read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["采集概览", "视频总表", "增速榜"])
        headers = [cell.value for cell in next(workbook["视频总表"].iter_rows(min_row=1, max_row=1))]
        self.assertEqual(headers[1:], ["排名", "BV号", "标题", "UP主", "分区", "时长(秒)", "发布时间",
                                       "播放", "弹幕", "评论", "点赞", "投币", "收藏", "分享"])
        workbook.close()


class CollectorPageTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try:
            from PySide6.QtWidgets import QApplication
        except ImportError:
            cls.app = None
            return
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        if self.app is None:
            self.skipTest("PySide6 unavailable in this interpreter")
        from tools.collector.page import CollectorPage
        self.page = CollectorPage({"out_dir": tempfile.gettempdir()})
        self.page.resize(960, 640)
        self.page.show()
        self.app.processEvents()

    def tearDown(self):
        self.page.close()
        self.page.deleteLater()
        self.app.processEvents()
        from PySide6.QtCore import QCoreApplication, QEvent
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
        self.app.processEvents()

    def _scroll(self):
        from PySide6.QtWidgets import QScrollArea
        scroll = self.page.findChild(QScrollArea, "pageScroll")
        self.assertIsNotNone(scroll)
        return scroll

    def _open_dialog(self):
        self.page._open_dashboard()
        self.app.processEvents()
        self.assertIsNotNone(self.page._dashboard_dialog)
        self.assertTrue(self.page._dashboard_dialog.isVisible())
        return self.page._dashboard_dialog

    def _completed_payload(self):
        comparison = SessionComparison(["BV1", "BV2"])
        comparison.publish_round([
            snapshot("BV1", 100, 10),
            snapshot("BV2", 200, 20),
        ])
        return comparison.publish_round([
            snapshot("BV1", 150, 15, 1_700_003_600),
            snapshot("BV2", 260, 26, 1_700_003_600),
        ], round_number=2, state="completed")

    def test_default_startup_content_fills_page_scroll_viewport(self):
        scroll = self._scroll()
        self.page.resize(904, 640)
        self.app.processEvents()
        self.assertGreaterEqual(scroll.widget().width(), scroll.viewport().width() - 1)
        self.assertGreaterEqual(scroll.widget().maximumWidth(), scroll.viewport().width())
        self.assertEqual(scroll.horizontalScrollBar().maximum(), 0)

    def test_page_fills_without_horizontal_overflow_at_supported_sizes(self):
        scroll = self._scroll()
        for size in ((1120, 760), (1440, 900), (960, 640)):
            with self.subTest(size=size):
                self.page.resize(*size)
                self.app.processEvents()
                self.assertGreaterEqual(scroll.widget().width(), scroll.viewport().width() - 1)
                self.assertEqual(scroll.horizontalScrollBar().maximum(), 0)

    def test_resizing_does_not_leave_a_stale_maximum_width(self):
        scroll = self._scroll()
        for size in ((960, 640), (1440, 900), (1120, 760), (960, 640)):
            self.page.resize(*size)
            self.app.processEvents()
            self.assertGreaterEqual(scroll.widget().maximumWidth(), scroll.viewport().width())
            self.assertGreaterEqual(scroll.widget().width(), scroll.viewport().width() - 1)

    def test_main_page_keeps_only_compact_dashboard_entry(self):
        from PySide6.QtWidgets import QTableWidget
        from tools.collector.comparison_chart import ComparisonChart

        scroll = self._scroll()
        self.assertIsNone(scroll.findChild(QTableWidget))
        self.assertIsNone(scroll.findChild(ComparisonChart))
        self.assertIsNone(self.page._dashboard_dialog)
        self.assertTrue(self.page.dashboard_open_button.isVisible())
        self.assertLess(self.page.dashboard_open_button.parentWidget().height(), 100)

    def test_open_button_creates_and_shows_one_non_modal_dialog(self):
        from PySide6.QtWidgets import QDialog

        with patch("app.task_page.TaskRunner") as runner, \
                patch.object(collector_pipeline.session, "ensure_ready") as ensure_ready:
            dialog = self._open_dialog()
            self.assertIsInstance(dialog, QDialog)
            self.assertFalse(dialog.isModal())
            runner.assert_not_called()
            ensure_ready.assert_not_called()

            self.page.dashboard_open_button.click()
            self.app.processEvents()
            self.assertIs(self.page._dashboard_dialog, dialog)
            self.assertTrue(dialog.isVisible())
            runner.assert_not_called()
            ensure_ready.assert_not_called()

    def test_metric_switch_does_not_start_runner_or_network(self):
        with patch("app.task_page.TaskRunner") as runner, \
                patch.object(collector_pipeline.session, "ensure_ready") as ensure_ready:
            dialog = self._open_dialog()
            dialog.dashboard_metric.setCurrentIndex(1)
            self.app.processEvents()
            self.assertEqual(self.page._dashboard_metric_key, "delta_view")
            runner.assert_not_called()
            ensure_ready.assert_not_called()
        self.assertEqual(self._scroll().horizontalScrollBar().maximum(), 0)

    def test_initial_view_sort_indicator_matches_play_count_descending(self):
        from PySide6.QtCore import Qt

        self.page._render_dashboard(self._completed_payload())
        dialog = self._open_dialog()
        header = dialog.dashboard_table.horizontalHeader()
        self.assertEqual(self.page._dashboard_sort_field, "view")
        self.assertTrue(self.page._dashboard_sort_descending)
        self.assertEqual(header.sortIndicatorSection(), 4)
        self.assertEqual(header.sortIndicatorOrder(), Qt.DescendingOrder)
        self.assertFalse(dialog.dashboard_table.isSortingEnabled())
        self.assertEqual(dialog.dashboard_table.item(0, 0).text(), "BV2")

    def test_metric_switch_updates_sort_field_column_and_descending_arrow(self):
        from PySide6.QtCore import Qt

        self.page._render_dashboard(self._completed_payload())
        dialog = self._open_dialog()
        dialog.dashboard_metric.setCurrentIndex(dialog.dashboard_metric.findData("delta_view"))
        self.app.processEvents()
        header = dialog.dashboard_table.horizontalHeader()
        self.assertEqual(self.page._dashboard_sort_field, "delta_view")
        self.assertTrue(self.page._dashboard_sort_descending)
        self.assertEqual(header.sortIndicatorSection(), 5)
        self.assertEqual(header.sortIndicatorOrder(), Qt.DescendingOrder)
        self.assertEqual(dialog.dashboard_table.item(0, 0).text(), "BV2")

    def test_header_clicks_keep_custom_sort_and_indicator_in_sync(self):
        from PySide6.QtCore import Qt

        self.page._render_dashboard(self._completed_payload())
        dialog = self._open_dialog()
        header = dialog.dashboard_table.horizontalHeader()

        header.sectionClicked.emit(0)
        self.app.processEvents()
        self.assertEqual(self.page._dashboard_sort_field, "bvid")
        self.assertFalse(self.page._dashboard_sort_descending)
        self.assertEqual(header.sortIndicatorSection(), 0)
        self.assertEqual(header.sortIndicatorOrder(), Qt.AscendingOrder)
        self.assertEqual(dialog.dashboard_table.item(0, 0).text(), "BV1")

        header.sectionClicked.emit(0)
        self.app.processEvents()
        self.assertEqual(self.page._dashboard_sort_field, "bvid")
        self.assertTrue(self.page._dashboard_sort_descending)
        self.assertEqual(header.sortIndicatorSection(), 0)
        self.assertEqual(header.sortIndicatorOrder(), Qt.DescendingOrder)
        self.assertEqual(dialog.dashboard_table.item(0, 0).text(), "BV2")

        header.sectionClicked.emit(1)
        self.app.processEvents()
        self.assertEqual(self.page._dashboard_sort_field, "title")
        self.assertFalse(self.page._dashboard_sort_descending)
        self.assertEqual(header.sortIndicatorSection(), 1)
        self.assertEqual(header.sortIndicatorOrder(), Qt.AscendingOrder)

    def test_dialog_reopen_preserves_latest_payload_and_metric_selection(self):
        comparison = SessionComparison(["BV1", "BV2"])
        first = comparison.publish_round([snapshot("BV1"), snapshot("BV2")])
        second = comparison.publish_round([
            snapshot("BV1", 140, 14, 1_700_003_600),
            snapshot("BV2", 260, 26, 1_700_003_600),
        ], round_number=2, state="completed")
        self.page._render_dashboard(first)
        dialog = self._open_dialog()
        dialog.dashboard_metric.setCurrentIndex(1)
        dialog.close()
        self.assertFalse(dialog.isVisible())

        self.page._render_dashboard(second)
        self._open_dialog()
        self.assertIs(self.page._dashboard_dialog, dialog)
        self.assertEqual(dialog.dashboard_metric.currentData(), "delta_view")
        self.assertEqual(dialog.dashboard_table.rowCount(), 2)
        self.assertIn("已完成", dialog.dashboard_status.text())

    def test_progress_received_while_closed_is_visible_when_reopened(self):
        comparison = SessionComparison(["BV1", "BV2"])
        payload = comparison.publish_round(
            [snapshot("BV1", 120), snapshot("BV2", 240)], state="only_one_round")
        dialog = self._open_dialog()
        dialog.close()
        self.page._on_progress({"dashboard": payload})
        self.app.processEvents()
        self._open_dialog()
        self.assertEqual(dialog.dashboard_table.rowCount(), 2)
        self.assertIn("仅有一轮", dialog.dashboard_status.text())

    def test_page_hide_hides_dialog_and_reopen_restores_payload_metric_and_sort(self):
        from PySide6.QtCore import Qt

        payload = self._completed_payload()
        self.page._render_dashboard(payload)
        dialog = self._open_dialog()
        dialog.dashboard_metric.setCurrentIndex(dialog.dashboard_metric.findData("delta_view"))
        dialog.dashboard_table.horizontalHeader().sectionClicked.emit(5)
        self.app.processEvents()
        self.assertFalse(self.page._dashboard_sort_descending)

        self.page.hide()
        self.app.processEvents()
        self.assertFalse(dialog.isVisible())
        self.assertEqual(self.page._dashboard_payload, payload)

        self.page.show()
        self.app.processEvents()
        self.assertFalse(dialog.isVisible())
        self.page._open_dashboard()
        self.app.processEvents()
        self.assertIs(self.page._dashboard_dialog, dialog)
        self.assertTrue(dialog.isVisible())
        self.assertEqual(dialog.dashboard_metric.currentData(), "delta_view")
        self.assertEqual(self.page._dashboard_sort_field, "delta_view")
        self.assertFalse(self.page._dashboard_sort_descending)
        self.assertEqual(dialog.dashboard_table.horizontalHeader().sortIndicatorSection(), 5)
        self.assertEqual(dialog.dashboard_table.horizontalHeader().sortIndicatorOrder(), Qt.AscendingOrder)

    def test_page_hide_show_and_reopen_do_not_start_runner_or_network(self):
        with patch("app.task_page.TaskRunner") as runner, \
                patch.object(collector_pipeline.session, "ensure_ready") as ensure_ready:
            dialog = self._open_dialog()
            self.page.hide()
            self.app.processEvents()
            self.page.show()
            self.app.processEvents()
            self.page._open_dashboard()
            self.app.processEvents()
            self.assertIs(self.page._dashboard_dialog, dialog)
            runner.assert_not_called()
            ensure_ready.assert_not_called()

    def test_dashboard_screenshot_path_requires_native_size_without_scaling(self):
        source = (Path(__file__).resolve().parents[1] / "scripts" / "take_screenshots.py").read_text(
            encoding="utf-8")
        dashboard_source = source.split("def shoot_dashboard", 1)[1]
        self.assertNotIn("pixmap.scaled", dashboard_source)
        self.assertIn("dialog.resize(*dialog_size)", dashboard_source)
        self.assertIn("actual_size = (pixmap.width(), pixmap.height())", dashboard_source)

    def test_chart_is_top_ten_but_dialog_table_keeps_all_rows(self):
        from PySide6.QtCore import Qt

        records = [snapshot(f"BV{i:02d}", i * 100) for i in range(1, 13)]
        comparison = SessionComparison([f"BV{i:02d}" for i in range(1, 13)])
        payload = comparison.publish_round(records, state="completed")
        self.page._render_dashboard(payload)
        dialog = self._open_dialog()
        self.assertEqual(dialog.dashboard_table.rowCount(), 12)
        self.assertLessEqual(len(dialog.dashboard_chart._items), 10)
        self.assertEqual(dialog.dashboard_table.horizontalScrollBarPolicy(), Qt.ScrollBarAsNeeded)

    def test_empty_single_partial_failure_and_cancelled_states_are_visible(self):
        dialog = self._open_dialog()
        self.assertTrue(dialog.empty_hint.isVisible())
        self.assertIn("等待首轮", dialog.dashboard_status.text())

        comparison = SessionComparison(["BV1"])
        single = comparison.publish_round([snapshot("BV1")])
        self.page._render_dashboard(single)
        self.assertFalse(dialog.empty_hint.isVisible())
        self.assertEqual(dialog.dashboard_table.rowCount(), 1)
        self.assertIn("至少需要两个视频", dialog.dashboard_summary.text())

        comparison = SessionComparison(["BV1", "BV2"])
        partial = comparison.publish_round(
            [snapshot("BV1", 100)], failed=["BV2"], state="partial_failure")
        self.page._render_dashboard(partial)
        self.assertIn("部分失败", dialog.dashboard_status.text())
        self.assertIn("本轮失败", dialog.dashboard_table.item(1, 3).text())

        self.page._on_cancelled({"dashboard": partial})
        self.assertIn("已取消", dialog.dashboard_status.text())
        self.assertEqual(dialog.dashboard_table.rowCount(), 2)

    def test_new_valid_task_clears_dashboard_without_opening_dialog(self):
        comparison = SessionComparison(["BV1", "BV2"])
        self.page._render_dashboard(comparison.publish_round([snapshot("BV1"), snapshot("BV2")]))
        self.page.src_edit.setPlainText("BV1\nBV2")
        with patch("app.task_page.TaskRunner") as runner:
            runner.return_value.isRunning.return_value = False
            self.page.on_start()
            runner.assert_called_once()
        self.assertIsNone(self.page._dashboard_dialog)
        self.assertEqual(self.page._dashboard_payload["rows"], [])

    def test_invalid_parameters_preserve_previous_dashboard(self):
        comparison = SessionComparison(["BV1", "BV2"])
        payload = comparison.publish_round([snapshot("BV1"), snapshot("BV2")])
        self.page._render_dashboard(payload)
        self.page.src_edit.clear()
        with patch("app.task_page.QMessageBox.warning"), patch("app.task_page.TaskRunner") as runner:
            self.page.on_start()
        self.assertEqual(self.page._dashboard_payload, payload)
        runner.assert_not_called()

    def test_repeated_start_while_runner_is_running_preserves_dashboard(self):
        comparison = SessionComparison(["BV1", "BV2"])
        payload = comparison.publish_round([snapshot("BV1"), snapshot("BV2")])
        self.page._render_dashboard(payload)
        running_runner = Mock()
        running_runner.isRunning.return_value = True
        self.page.runner = running_runner
        self.page.src_edit.clear()
        self.page.on_start()
        self.assertEqual(self.page._dashboard_payload, payload)

    def test_dialog_is_clamped_to_available_geometry_at_compact_size(self):
        self.page.resize(960, 640)
        self.app.processEvents()
        dialog = self._open_dialog()
        available = dialog.screen().availableGeometry()
        frame = dialog.frameGeometry()
        self.assertGreaterEqual(frame.left(), available.left())
        self.assertGreaterEqual(frame.top(), available.top())
        self.assertLessEqual(frame.right(), available.right())
        self.assertLessEqual(frame.bottom(), available.bottom())
        self.assertLessEqual(dialog.width(), 1000)
        self.assertLessEqual(dialog.height(), 700)

    def test_page_close_hides_dialog_and_leaves_no_visible_dashboard_window(self):
        dialog = self._open_dialog()
        self.page.close()
        self.app.processEvents()
        self.assertFalse(dialog.isVisible())
        self.assertNotIn(dialog, [widget for widget in self.app.topLevelWidgets() if widget.isVisible()])

    def test_chart_horizontal_bars_handle_signed_zero_equal_single_and_missing_values(self):
        from tools.collector.comparison_chart import ComparisonChart

        chart = ComparisonChart()
        chart.resize(520, 170)
        try:
            chart.set_items("播放增量", [(f"BV-{index}-超长标签", value)
                                         for index, value in enumerate(
                                             (10, -8, 0, 4, -2, 3, -1, 6, -4, 2, 1, -3), 1)])
            self.assertEqual(len(chart._items), 10)
            self.assertTrue(any(value < 0 for _label, value in chart._items))
            self.assertTrue(any(value > 0 for _label, value in chart._items))
            chart.show()
            self.app.processEvents()
            self.assertFalse(chart.grab().isNull())

            for values in (((1, 1), (2, 1)), (("only", 0),), (("missing", None),)):
                chart.set_items("边界", [(str(label), value) for label, value in values])
                self.app.processEvents()
                self.assertFalse(chart.grab().isNull())
        finally:
            chart.close()


if __name__ == "__main__":
    unittest.main()
