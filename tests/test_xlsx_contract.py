"""core.xlsx 阶段一合同：类型、安全文本、时区和原子发布。"""

from datetime import datetime, tzinfo
from decimal import Decimal
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from core import xlsx


class XlsxContractTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix="xlsx_contract_"))
        self.addCleanup(lambda: self._cleanup_root())

    def _cleanup_root(self):
        for path in sorted(self.root.glob("**/*"), reverse=True):
            if path.is_file():
                path.unlink()
            elif path.is_dir():
                path.rmdir()
        if self.root.exists():
            self.root.rmdir()

    def _workbook(self, rows):
        workbook = xlsx.new_workbook()
        worksheet = workbook.create_sheet("数据")
        writer = xlsx.SheetWriter(workbook)
        writer.ws = worksheet
        for row in rows:
            writer.append(row)
        return workbook

    def _save_and_load(self, workbook, name="data.xlsx"):
        path = self.root / name
        xlsx.save_workbook_atomic(workbook, path)
        workbook.close()
        loaded = load_workbook(path, data_only=False)
        self.addCleanup(loaded.close)
        return path, loaded["数据"]

    def _assert_write_rejected(self, value, kind):
        workbook = xlsx.new_workbook()
        worksheet = workbook.create_sheet("数据")
        writer = xlsx.SheetWriter(workbook)
        writer.ws = worksheet
        try:
            with self.assertRaises((TypeError, ValueError)) as caught:
                writer.append([xlsx.cell_value(value, kind)])
        finally:
            workbook.close()
        return caught.exception

    def test_ids_keep_exact_text_and_text_format(self):
        values = [
            "123456789012345",
            "1234567890123456",
            "12345678901234567890",
            "9" * 21,
            12345678901234567890,
        ]
        workbook = self._workbook([
            [xlsx.cell_value(value, xlsx.CellKind.ID) for value in values]
        ])
        _path, worksheet = self._save_and_load(workbook)

        self.assertEqual([cell.value for cell in worksheet[1]],
                         [str(value) for value in values])
        for cell in worksheet[1]:
            self.assertEqual(cell.data_type, "s")
            self.assertEqual(cell.number_format, "@")

    def test_integer_percent_and_datetime_keep_real_excel_types(self):
        when = xlsx.unix_seconds_to_excel_datetime(1_700_000_000)
        self.assertIsInstance(when, datetime)
        self.assertIsNone(when.tzinfo)
        workbook = self._workbook([[
            xlsx.cell_value(0, xlsx.CellKind.INTEGER),
            xlsx.cell_value(1234.5, xlsx.CellKind.DECIMAL),
            xlsx.cell_value(0.123, xlsx.CellKind.PERCENT),
            xlsx.cell_value(when, xlsx.CellKind.DATETIME,
                            number_format="yyyy-mm-dd hh:mm"),
        ]])
        _path, worksheet = self._save_and_load(workbook)

        self.assertEqual(worksheet.cell(1, 1).value, 0)
        self.assertIsInstance(worksheet.cell(1, 1).value, int)
        self.assertEqual(worksheet.cell(1, 1).number_format, "#,##0")
        self.assertEqual(worksheet.cell(1, 2).value, 1234.5)
        self.assertEqual(worksheet.cell(1, 2).number_format, "#,##0.00")
        self.assertAlmostEqual(worksheet.cell(1, 3).value, 0.123)
        self.assertEqual(worksheet.cell(1, 3).number_format, "0.0%")
        self.assertIsInstance(worksheet.cell(1, 4).value, datetime)
        self.assertEqual(worksheet.cell(1, 4).number_format,
                         "yyyy-mm-dd hh:mm")

    def test_numeric_limits_reject_huge_values_before_openpyxl(self):
        self.assertEqual(
            xlsx._integer_value(xlsx.MAX_INTEGER_VALUE),
            xlsx.MAX_INTEGER_VALUE,
        )
        self._assert_write_rejected(xlsx.MAX_INTEGER_VALUE + 1,
                                    xlsx.CellKind.INTEGER)
        self._assert_write_rejected("9" * 5000, xlsx.CellKind.INTEGER)
        self._assert_write_rejected(10 ** 4999, xlsx.CellKind.INTEGER)
        self._assert_write_rejected(True, xlsx.CellKind.INTEGER)
        self._assert_write_rejected(1.25, xlsx.CellKind.INTEGER)

        for kind in (xlsx.CellKind.DECIMAL, xlsx.CellKind.PERCENT):
            self._assert_write_rejected("9" * 5000, kind)
            self._assert_write_rejected(10 ** 4999, kind)
            self._assert_write_rejected(Decimal("9" * 5000), kind)
            self._assert_write_rejected(Decimal("1e+5000"), kind)
            self._assert_write_rejected(Decimal("1e-5000"), kind)
            for value in (float("nan"), float("inf"), "NaN", "Infinity", True):
                self._assert_write_rejected(value, kind)

            workbook = self._workbook([[xlsx.cell_value(
                xlsx.MAX_DECIMAL_ABS, kind
            )]])
            _path, worksheet = self._save_and_load(workbook,
                                                    name=f"{kind.value}.xlsx")
            self.assertTrue(worksheet.cell(1, 1).value is not None)
            self.assertEqual(worksheet.cell(1, 1).number_format,
                             "0.0%" if kind is xlsx.CellKind.PERCENT
                             else "#,##0.00")
            self._assert_write_rejected(xlsx.MAX_DECIMAL_ABS + 1, kind)

        percent_workbook = self._workbook([[xlsx.cell_value(
            Decimal("1.5"), xlsx.CellKind.PERCENT
        )]])
        _path, percent_sheet = self._save_and_load(
            percent_workbook, name="growth.xlsx"
        )
        self.assertAlmostEqual(percent_sheet.cell(1, 1).value, 1.5)

    def test_unix_seconds_are_fixed_to_asia_shanghai(self):
        value = xlsx.unix_seconds_to_excel_datetime(0)
        self.assertEqual(value, datetime(1970, 1, 1, 8, 0, 0))
        precise = xlsx.unix_seconds_to_excel_datetime(
            Decimal("1700000000.123456")
        )
        self.assertEqual(precise, datetime(2023, 11, 15, 6, 13, 20, 123456))
        self.assertEqual(
            xlsx.unix_seconds_to_excel_datetime("1700000000.12345").microsecond,
            123450,
        )
        with self.assertRaises(ValueError):
            xlsx.unix_seconds_to_excel_datetime("1700000000.1234567")
        self.assertEqual(
            xlsx.unix_seconds_to_excel_datetime(xlsx.MAX_UNIX_SECONDS),
            datetime(9999, 12, 31, 23, 59, 59),
        )
        for invalid in (
            -1,
            xlsx.MAX_UNIX_SECONDS + 1,
            1_700_000_000_000,
            1_700_000_000_000_000,
            "9" * 5000,
            Decimal("1e+5000"),
            Decimal("1e-5000"),
            Decimal("NaN"),
            float("inf"),
            True,
        ):
            with self.assertRaises((TypeError, ValueError)):
                xlsx.unix_seconds_to_excel_datetime(invalid)

        with self.assertRaises(TypeError):
            xlsx.local_text_to_excel_datetime(10 ** 4999)
        with self.assertRaises(TypeError):
            xlsx.safe_text(10 ** 4999)

        for platform_error in (TypeError("platform"),
                               ValueError("platform"),
                               OverflowError("platform"),
                               OSError("platform")):
            with patch.object(xlsx, "_unix_datetime_from_parts",
                              side_effect=platform_error):
                with self.assertRaises(ValueError):
                    xlsx.unix_seconds_to_excel_datetime(0)

        parsed = xlsx.local_text_to_excel_datetime("2024-01-02 03:04:05")
        self.assertEqual(parsed, datetime(2024, 1, 2, 3, 4, 5))
        self.assertIsNone(parsed.tzinfo)

    def test_integer_text_counts_effective_digits_after_leading_zeros(self):
        cases = (
            ("0000000000000001", 1),
            ("-0000000000000001", -1),
            ("0000000000000000", 0),
            ("0" * 64, 0),
            ("000000000000000" + "9" * 15, 999999999999999),
        )
        for text, expected in cases:
            self.assertEqual(xlsx._integer_value(text), expected)

        self._assert_write_rejected("9" * 16, xlsx.CellKind.INTEGER)
        self._assert_write_rejected("0" * 5000 + "1", xlsx.CellKind.INTEGER)
        self._assert_write_rejected("9" * 5000, xlsx.CellKind.INTEGER)

    def test_datetime_platform_errors_are_normalized(self):
        class BrokenTimezone(tzinfo):
            def utcoffset(self, _value):
                raise OSError("platform timezone failure")

        with self.assertRaises(ValueError):
            xlsx._datetime_value(datetime(2024, 1, 1,
                                           tzinfo=BrokenTimezone()))

    def test_text_safety_is_idempotent_and_explicit_formula_remains_formula(self):
        for value in ("=1+1", "+cmd", "-cmd", "@cmd", "a\x00b"):
            safe = xlsx.safe_text(value)
            self.assertEqual(xlsx.safe_text(safe), safe)
            self.assertNotIn("\x00", safe)

        workbook = xlsx.new_workbook()
        worksheet = workbook.create_sheet("数据")
        writer = xlsx.SheetWriter(workbook)
        writer.ws = worksheet
        cleaned = xlsx.safe_text("=1+1\x00")
        worksheet.append([cleaned])
        writer.append([
            xlsx.cell_value(cleaned, xlsx.CellKind.TEXT),
            xlsx.cell_value(-1, xlsx.CellKind.INTEGER),
            xlsx.cell_value("=SUM(1,2)", xlsx.CellKind.FORMULA),
        ])
        _path, worksheet = self._save_and_load(workbook)
        self.assertEqual(worksheet.cell(1, 1).value, "=1+1")
        self.assertEqual(worksheet.cell(1, 1).data_type, "f")
        self.assertEqual(worksheet.cell(2, 1).value, "=1+1")
        self.assertEqual(worksheet.cell(2, 1).data_type, "s")
        self.assertTrue(worksheet.cell(2, 1).quotePrefix)
        self.assertEqual(worksheet.cell(2, 2).value, -1)
        self.assertEqual(worksheet.cell(2, 3).value, "=SUM(1,2)")
        self.assertEqual(worksheet.cell(2, 3).data_type, "f")

    def test_state_cells_are_blank_with_explicit_comments(self):
        workbook = self._workbook([[
            xlsx.cell_value(None, xlsx.CellKind.MISSING),
            xlsx.cell_value(None, xlsx.CellKind.NOT_RETURNED),
            xlsx.cell_value(None, xlsx.CellKind.NOT_APPLICABLE),
        ]])
        _path, worksheet = self._save_and_load(workbook)
        self.assertEqual([cell.value for cell in worksheet[1]], [None, None, None])
        self.assertEqual([cell.comment.text for cell in worksheet[1]],
                         ["字段缺失", "接口未返回", "不适用"])

    def test_atomic_new_file_and_overwrite(self):
        path = self.root / "result.xlsx"
        xlsx.save_workbook_atomic(self._workbook([[xlsx.cell_value("old")]]), path)
        old_bytes = path.read_bytes()
        xlsx.save_workbook_atomic(self._workbook([[xlsx.cell_value("new")]]), path)
        self.assertTrue(path.exists())
        self.assertNotEqual(path.read_bytes(), old_bytes)
        self.assertEqual(list(self.root.glob("*.xlsx.tmp")), [])
        self.assertEqual(list(self.root.glob(".*.xlsx.tmp")), [])

    def test_save_failure_preserves_existing_target_and_cleans_temp(self):
        path = self.root / "result.xlsx"
        xlsx.save_workbook_atomic(self._workbook([[xlsx.cell_value("old")]]), path)
        old_bytes = path.read_bytes()
        old_stat = path.stat()
        workbook = self._workbook([])
        try:
            with patch.object(workbook, "save", side_effect=OSError("save failed")):
                with self.assertRaises(OSError):
                    xlsx.save_workbook_atomic(workbook, path)
        finally:
            workbook.close()
        self.assertEqual(path.read_bytes(), old_bytes)
        self.assertEqual(path.stat().st_mtime_ns, old_stat.st_mtime_ns)
        self.assertEqual(list(self.root.glob(".*.xlsx.tmp")), [])

    def test_replace_failure_preserves_existing_target_and_cleans_temp(self):
        path = self.root / "result.xlsx"
        xlsx.save_workbook_atomic(self._workbook([[xlsx.cell_value("old")]]), path)
        old_bytes = path.read_bytes()
        old_stat = path.stat()
        workbook = self._workbook([[xlsx.cell_value("new")]])
        try:
            with patch.object(xlsx, "_os_replace", side_effect=OSError("replace failed")):
                with self.assertRaises(OSError):
                    xlsx.save_workbook_atomic(workbook, path)
        finally:
            workbook.close()
        self.assertEqual(path.read_bytes(), old_bytes)
        self.assertEqual(path.stat().st_mtime_ns, old_stat.st_mtime_ns)
        self.assertEqual(list(self.root.glob(".*.xlsx.tmp")), [])

    def test_before_replace_failure_preserves_existing_target_and_cleans_temp(self):
        path = self.root / "result.xlsx"
        xlsx.save_workbook_atomic(self._workbook([[xlsx.cell_value("old")]]), path)
        old_bytes = path.read_bytes()
        old_stat = path.stat()
        workbook = self._workbook([[xlsx.cell_value("new")]])

        def cancel_before_publish():
            raise RuntimeError("cancelled")

        try:
            with self.assertRaises(RuntimeError):
                xlsx.save_workbook_atomic(
                    workbook, path, before_replace=cancel_before_publish
                )
        finally:
            workbook.close()
        self.assertEqual(path.read_bytes(), old_bytes)
        self.assertEqual(path.stat().st_mtime_ns, old_stat.st_mtime_ns)
        self.assertEqual(list(self.root.glob(".*.xlsx.tmp")), [])


if __name__ == "__main__":
    unittest.main()
