# -*- coding: utf-8 -*-
"""第一方 XLSX 元数据共享契约的离线测试。"""

from datetime import datetime
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook

from core.xlsx import (
    ASIA_SHANGHAI,
    CellKind,
    cell_value,
    new_workbook,
    save_workbook_atomic,
)
from core.xlsx_metadata import (
    FieldDefinition,
    QualityItem,
    WorkbookMetadata,
    XLSX_METADATA_MARKER,
    XLSX_SCHEMA_VERSION,
    build_parameter_summary,
    is_complete_metadata_workbook,
    is_standard_metadata_sheet,
    percent_cell,
    write_metadata_sheets,
)


class XlsxMetadataTests(unittest.TestCase):
    def make_metadata(self):
        return WorkbookMetadata(
            tool="评论",
            report_type="评论报告",
            generated_at=datetime(2026, 9, 14, 12, 34, 56, tzinfo=ASIA_SHANGHAI),
            timezone="Asia/Shanghai / UTC+08:00",
            parameter_summary=build_parameter_summary(
                {
                    "目标类型": cell_value("dynamic", CellKind.TEXT),
                    "规范化 oid": cell_value("123456", CellKind.ID),
                },
                ("目标类型", "规范化 oid"),
            ),
            quality_items=(
                QualityItem("记录", "真实零值", cell_value(0, CellKind.INTEGER), "条", "0 是真实值"),
                QualityItem("记录", "字段缺失", cell_value(None, CellKind.MISSING), "状态", "输入未提供"),
                QualityItem("接口", "声明数量", cell_value(None, CellKind.NOT_RETURNED), "状态", "接口未返回"),
                QualityItem("覆盖", "覆盖率", percent_cell(0, 0), "状态", "分母为 0 时不适用"),
            ),
            fields=(
                FieldDefinition(
                    worksheet="全量评论",
                    display_name="评论 ID",
                    stable_name="rpid",
                    data_type="ID",
                    unit="",
                    nullable="否",
                    source="评论接口归一化",
                    metric="按 rpid 去重",
                    example=cell_value("123456789", CellKind.ID),
                    missing_meaning="接口未返回",
                ),
                FieldDefinition(
                    worksheet="全量评论",
                    display_name="用户内容",
                    stable_name="message",
                    data_type="文本",
                    unit="",
                    nullable="是",
                    source="用户来源文本",
                    metric="不参与计数",
                    example=cell_value("=不是公式", CellKind.TEXT),
                    missing_meaning="字段缺失",
                ),
            ),
        )

    def test_appends_two_marked_sheets_and_preserves_typed_cells(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "report.xlsx"
            workbook = new_workbook()
            sheet = workbook.create_sheet("业务表")
            sheet.append(["旧列", "值"])
            sheet.append(["a", 1])
            write_metadata_sheets(workbook, self.make_metadata())
            save_workbook_atomic(workbook, path)

            loaded = load_workbook(path, data_only=False)
            self.assertEqual(loaded.sheetnames, ["业务表", "数据质量", "字段说明"])
            quality = loaded["数据质量"]
            fields = loaded["字段说明"]
            quality_values = [cell.value for row in quality.iter_rows() for cell in row]
            field_values = [cell.value for row in fields.iter_rows() for cell in row]
            self.assertGreaterEqual(quality_values.count(XLSX_METADATA_MARKER), 1)
            self.assertGreaterEqual(field_values.count(XLSX_METADATA_MARKER), 1)
            self.assertIn(XLSX_SCHEMA_VERSION, quality_values)
            self.assertIn(XLSX_SCHEMA_VERSION, field_values)

            # 第一列是 SheetWriter 的布局占位列；按固定表头寻找数据区，
            # 不把顶部生成信息行号当成对外契约。
            header_row = next(
                row[0].row
                for row in quality.iter_rows()
                if row[1].value == "类别" and row[5].value == "说明"
            )
            headers = [quality.cell(header_row, column).value for column in range(1, 7)]
            self.assertEqual(headers[1:], ["类别", "项目", "数值", "单位或状态", "说明"])
            item_rows = {
                quality.cell(row_index, 3).value: row_index
                for row_index in range(header_row + 1, quality.max_row + 1)
                if quality.cell(row_index, 3).value
            }
            zero_cell = quality.cell(item_rows["真实零值"], 4)
            self.assertEqual(zero_cell.value, 0)
            self.assertEqual(zero_cell.data_type, "n")
            missing_cell = quality.cell(item_rows["字段缺失"], 4)
            self.assertIsNone(missing_cell.value)
            self.assertIn("字段缺失", missing_cell.comment.text)
            not_returned_cell = quality.cell(item_rows["声明数量"], 4)
            self.assertIsNone(not_returned_cell.value)
            self.assertIn("接口未返回", not_returned_cell.comment.text)
            not_applicable_cell = quality.cell(item_rows["覆盖率"], 4)
            self.assertIsNone(not_applicable_cell.value)
            self.assertIn("不适用", not_applicable_cell.comment.text)

            formula_row = next(
                row[0].row
                for row in fields.iter_rows()
                if row[3].value == "message"
            )
            formula_text = fields.cell(formula_row, 10)
            self.assertEqual(formula_text.value, "=不是公式")
            self.assertEqual(formula_text.data_type, "s")
            self.assertTrue(formula_text.quotePrefix)
            self.assertEqual(loaded["业务表"].cell(2, 2).value, 1)

    def test_percent_cell_is_numeric_and_zero_denominator_is_not_applicable(self):
        usable = percent_cell(25, 100)
        self.assertEqual(usable.kind, CellKind.PERCENT)
        self.assertEqual(usable.value, 0.25)
        self.assertEqual(usable.number_format, "0.0%")
        unusable = percent_cell(0, 0)
        self.assertEqual(unusable.kind, CellKind.NOT_APPLICABLE)
        self.assertIsNone(unusable.value)
        self.assertIn("分母为 0", unusable.note)

    def test_repeating_percentages_reopen_as_numeric_excel_cells(self):
        items = (
            QualityItem("覆盖", "三分之一", percent_cell(1, 3), "百分比", "可写数值"),
            QualityItem("覆盖", "三分之二", percent_cell(2, 3), "百分比", "可写数值"),
            QualityItem("覆盖", "零除三", percent_cell(0, 3), "百分比", "真实零值"),
            QualityItem("覆盖", "零除零", percent_cell(0, 0), "状态", "分母为 0"),
        )
        metadata = WorkbookMetadata(
            tool="测试", report_type="比例测试",
            generated_at=datetime(2026, 9, 14, 12, 34, 56, tzinfo=ASIA_SHANGHAI),
            timezone="Asia/Shanghai / UTC+08:00",
            parameter_summary=(), quality_items=items, fields=(),
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "percentages.xlsx"
            workbook = new_workbook()
            workbook.create_sheet("业务表")
            write_metadata_sheets(workbook, metadata)
            save_workbook_atomic(workbook, path)
            loaded = load_workbook(path, data_only=False)
            quality = loaded["数据质量"]
            rows = {
                quality.cell(row, 3).value: row
                for row in range(1, quality.max_row + 1)
                if quality.cell(row, 3).value
            }
            for label, expected in (("三分之一", 1 / 3), ("三分之二", 2 / 3)):
                cell = quality.cell(rows[label], 4)
                self.assertIsInstance(cell.value, (int, float))
                self.assertEqual(cell.data_type, "n")
                self.assertAlmostEqual(cell.value, expected, places=14)
                self.assertEqual(cell.number_format, "0.0%")
                self.assertIsNone(cell.comment)
            zero = quality.cell(rows["零除三"], 4)
            self.assertEqual(zero.value, 0)
            self.assertEqual(zero.data_type, "n")
            self.assertEqual(zero.number_format, "0.0%")
            not_applicable = quality.cell(rows["零除零"], 4)
            self.assertIsNone(not_applicable.value)
            self.assertIn("分母为 0", not_applicable.comment.text)
            loaded.close()

    def test_parameter_summary_is_explicit_allowlist_and_quarantines_sensitive_values(self):
        summary = build_parameter_summary(
            {"自由文本": cell_value("cookie=SESSDATA=secret", CellKind.TEXT)},
            ("自由文本",),
        )
        self.assertEqual(len(summary), 1)
        self.assertIsNone(summary[0][1].value)
        self.assertEqual(summary[0][1].kind, CellKind.NOT_APPLICABLE)
        self.assertEqual(summary[0][1].note, "已排除敏感值")
        # 普通词语 bearer 不应因缺少凭据而被误判为敏感值。
        ordinary = build_parameter_summary(
            {"说明": cell_value("bearer", CellKind.TEXT)},
            ("说明",),
        )
        self.assertEqual(ordinary[0][1].value, "bearer")
        with self.assertRaises(ValueError):
            build_parameter_summary(
                {
                    "安全键": cell_value("ok", CellKind.TEXT),
                    "未白名单键": cell_value("must not persist", CellKind.TEXT),
                },
                ("安全键",),
            )

    def test_metadata_write_failure_closes_writer_and_never_leaks_note(self):
        workbook = new_workbook()
        workbook.create_sheet("业务表")
        metadata = self.make_metadata()
        metadata = WorkbookMetadata(
            tool=metadata.tool, report_type=metadata.report_type,
            generated_at=metadata.generated_at, timezone=metadata.timezone,
            parameter_summary=metadata.parameter_summary,
            quality_items=metadata.quality_items,
            fields=(FieldDefinition(
                worksheet="业务表", display_name="安全字段", stable_name="safe",
                data_type="文本", unit="", nullable="是", source="本地",
                metric="测试", example=cell_value("安全", CellKind.TEXT,
                                                    note="Token=secret"),
                missing_meaning="不适用"),),
        )
        with self.assertRaises(ValueError):
            write_metadata_sheets(workbook, metadata)
        # write_only workbook must be closable after validation failure; this is
        # the Windows handle-leak regression guard.
        workbook.close()

    def test_sensitive_values_are_absent_from_cells_and_xlsx_xml(self):
        secret = "SendKey=top-secret-token"
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "safe.xlsx"
            workbook = new_workbook()
            workbook.create_sheet("业务表").append(["x"])
            metadata = self.make_metadata()
            quarantined = build_parameter_summary(
                {"不安全说明": cell_value(secret, CellKind.TEXT)},
                ("不安全说明",),
            )
            metadata = WorkbookMetadata(
                tool=metadata.tool,
                report_type=metadata.report_type,
                generated_at=metadata.generated_at,
                timezone=metadata.timezone,
                parameter_summary=(
                    ("安全说明", cell_value("没有凭据", CellKind.TEXT)),
                ) + quarantined,
                quality_items=metadata.quality_items,
                fields=metadata.fields,
            )
            write_metadata_sheets(workbook, metadata)
            save_workbook_atomic(workbook, path)
            loaded = load_workbook(path, data_only=False)
            all_values = "\n".join(
                str(cell.value)
                for sheet in loaded.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            )
            self.assertNotIn(secret, all_values)
            with ZipFile(path) as archive:
                xml = b"\n".join(archive.read(name) for name in archive.namelist())
            self.assertNotIn(secret.encode("utf-8"), xml)

    def test_standard_sheet_requires_fixed_name_and_exact_marker(self):
        self.assertTrue(is_standard_metadata_sheet("数据质量", [[XLSX_METADATA_MARKER]]))
        self.assertTrue(is_standard_metadata_sheet("字段说明", [["x", XLSX_METADATA_MARKER]]))
        self.assertFalse(is_standard_metadata_sheet("数据质量", [["other-marker"]]))
        self.assertFalse(is_standard_metadata_sheet("用户自建表", [[XLSX_METADATA_MARKER]]))

    def test_complete_metadata_requires_both_fixed_marked_sheets(self):
        marked_quality = [[XLSX_METADATA_MARKER]]
        marked_fields = [["说明", XLSX_METADATA_MARKER]]
        self.assertTrue(is_complete_metadata_workbook({
            "数据质量": marked_quality, "字段说明": marked_fields,
        }))
        self.assertFalse(is_complete_metadata_workbook({"数据质量": marked_quality}))
        self.assertFalse(is_complete_metadata_workbook({"字段说明": marked_fields}))
        self.assertFalse(is_complete_metadata_workbook({
            "数据质量": marked_quality, "字段说明": [["说明", "other"]],
        }))
        self.assertFalse(is_complete_metadata_workbook({
            "其他表": [[XLSX_METADATA_MARKER]],
            "数据质量": marked_quality, "字段说明": [["说明", "other"]],
        }))

    def test_danmaku_declared_count_distinguishes_missing_from_uncomparable(self):
        from tools.danmaku import core as danmaku_core

        base_meta = {
            "bvid": "BV1TEST",
            "aid": 1,
            "cid": 2,
            "title": "离线测试",
            "owner": "测试",
            "page": 1,
            "page_count": 1,
            "duration": 60,
        }
        stats = {
            "segments": 0,
            "expected_segments": 1,
            "duplicates": 0,
            "cancelled": False,
            "truncated": False,
            "stopped_reason": None,
            "max_segments": 20,
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            missing_path = Path(temp_dir) / "missing.xlsx"
            danmaku_core.export_xlsx([], dict(base_meta), dict(stats), missing_path)
            missing = load_workbook(missing_path, data_only=False)["数据质量"]
            missing_rows = {
                missing.cell(row, 3).value: row
                for row in range(1, missing.max_row + 1)
                if missing.cell(row, 3).value
            }
            missing_declared = missing.cell(missing_rows["接口声称弹幕数"], 4)
            self.assertIsNone(missing_declared.value)
            self.assertIn("接口未返回", missing_declared.comment.text)

            numeric_path = Path(temp_dir) / "numeric.xlsx"
            numeric_meta = dict(base_meta, claimed_danmaku=123)
            danmaku_core.export_xlsx([], numeric_meta, dict(stats), numeric_path)
            numeric = load_workbook(numeric_path, data_only=False)["数据质量"]
            numeric_rows = {
                numeric.cell(row, 3).value: row
                for row in range(1, numeric.max_row + 1)
                if numeric.cell(row, 3).value
            }
            declared = numeric.cell(numeric_rows["接口声称弹幕数"], 4)
            self.assertEqual(declared.value, 123)
            status = numeric.cell(numeric_rows["声明数量口径状态"], 4)
            self.assertIn("口径不可比较", status.comment.text)
            coverage = numeric.cell(numeric_rows["覆盖率"], 4)
            self.assertIsNone(coverage.value)
            self.assertIn("不计算覆盖率", coverage.comment.text)

    def test_danmaku_declared_count_boundaries_reopen_without_export_failure(self):
        from tools.danmaku import core as danmaku_core

        base_meta = {
            "bvid": "BV1TEST", "aid": 1, "cid": 2, "title": "离线测试",
            "owner": "测试", "page": 1, "page_count": 1, "duration": 60,
        }
        stats = {
            "segments": 0, "expected_segments": 1, "duplicates": 0,
            "cancelled": False, "truncated": False, "stopped_reason": None,
            "max_segments": 20,
        }
        cases = (
            ("missing", {}, "接口未返回", None),
            ("none", {"claimed_danmaku": None}, "接口未返回", None),
            ("zero", {"claimed_danmaku": 0}, "口径不可比较", 0),
            ("positive", {"claimed_danmaku": 3}, "口径不可比较", 3),
            ("negative", {"claimed_danmaku": -1}, "声明数量格式异常", None),
            ("bool", {"claimed_danmaku": True}, "声明数量格式异常", None),
            ("float-int", {"claimed_danmaku": 1.0}, "声明数量格式异常", None),
            ("float", {"claimed_danmaku": 1.5}, "声明数量格式异常", None),
            ("nan", {"claimed_danmaku": float("nan")}, "声明数量格式异常", None),
            ("infinity", {"claimed_danmaku": float("inf")}, "声明数量格式异常", None),
            ("over-limit", {"claimed_danmaku": 10 ** 16}, "声明数量格式异常", None),
            ("exponent-text", {"claimed_danmaku": "1e3"}, "声明数量格式异常", None),
            ("overlong-text", {"claimed_danmaku": "9" * 5000}, "声明数量格式异常", None),
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            for name, extra, note, expected in cases:
                with self.subTest(name=name):
                    path = Path(temp_dir) / f"{name}.xlsx"
                    meta = dict(base_meta, **extra)
                    danmaku_core.export_xlsx([], meta, dict(stats), path)
                    workbook = load_workbook(path, data_only=False)
                    quality = workbook["数据质量"]
                    by_item = {
                        quality.cell(row, 3).value: row
                        for row in range(1, quality.max_row + 1)
                        if quality.cell(row, 3).value
                    }
                    declared = quality.cell(by_item["接口声称弹幕数"], 4)
                    if expected is None:
                        self.assertIsNone(declared.value)
                        self.assertIn(note, declared.comment.text)
                    else:
                        self.assertEqual(declared.value, expected)
                        self.assertEqual(declared.data_type, "n")
                    status = quality.cell(by_item["声明数量口径状态"], 4)
                    self.assertIn(note, status.comment.text)
                    coverage = quality.cell(by_item["覆盖率"], 4)
                    self.assertIsNone(coverage.value)
                    if name in {"zero", "positive"}:
                        self.assertIn("不计算覆盖率", coverage.comment.text)
                    elif name not in {"missing", "none"}:
                        self.assertIn("不适用", coverage.comment.text)
                    workbook.close()

    def test_report_center_skips_only_name_and_exact_marker(self):
        from tools.report_center import core as report_core

        with tempfile.TemporaryDirectory() as temp_dir:
            unmarked_path = Path(temp_dir) / "unmarked.xlsx"
            unmarked = new_workbook()
            sheet = unmarked.create_sheet("数据质量")
            sheet.append(["rpid", "评论内容"])
            sheet.append(["123", "用户自建同名表"])
            save_workbook_atomic(unmarked, unmarked_path)
            layout = report_core._xlsx_layout(unmarked_path)
            self.assertEqual(layout[1], "数据质量")

            marked_path = Path(temp_dir) / "marked.xlsx"
            marked = new_workbook()
            business = marked.create_sheet("数据")
            business.append(["rpid", "评论内容"])
            business.append(["123", "业务记录"])
            write_metadata_sheets(marked, self.make_metadata())
            save_workbook_atomic(marked, marked_path)
            marked_layout = report_core._xlsx_layout(marked_path)
            self.assertEqual(marked_layout[1], "数据")

            partial_cases = (
                ("only-quality", (("数据质量", True),)),
                ("only-fields", (("字段说明", True),)),
                ("one-of-two", (("数据质量", True), ("字段说明", False))),
            )
            for filename, sheets in partial_cases:
                with self.subTest(filename=filename):
                    partial_path = Path(temp_dir) / f"{filename}.xlsx"
                    partial = new_workbook()
                    for sheet_name, marked in sheets:
                        sheet = partial.create_sheet(sheet_name)
                        sheet.append(["rpid", "评论内容"])
                        sheet.append(["123", "按普通工作表读取"])
                        sheet.append([XLSX_METADATA_MARKER] if marked else ["普通同名表"])
                    save_workbook_atomic(partial, partial_path)
                    partial_layout = report_core._xlsx_layout(partial_path)
                    self.assertEqual(partial_layout[1], sheets[0][0])


if __name__ == "__main__":
    unittest.main()
