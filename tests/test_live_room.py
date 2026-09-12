# -*- coding: utf-8 -*-
"""直播追踪的离线测试。

全程离线：fetch/sleeper/budget/notify 全部注入，不打网络、不真等待；
fixture 形状来自 2026-09-12 对 get_info 接口的探针实测（字段名一致，
内容全部合成，不把真实房间数据带进仓库）。

预算语义说明：业务请求数的强制点在 BiliClient._request（check_request +
observe_request）；离线注入的 fake 通道须自行模拟这两步，与既有预算
测试约定一致（见 tests/test_budget.py / test_user_dynamics.py）。
"""
from __future__ import annotations

import ast
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from core.budget import TaskBudget
from core.notify import webhook_payload
from core.transport import BiliApiError
from tools import TOOLS
from tools.live_room import core
from tools.live_room.page import LiveRoomPage
from tools.live_room.pipeline import run_pipeline


# ---------- 桩 ----------

class RecordingSleep:
    def __init__(self):
        self.calls = []
        self.total = 0.0

    def __call__(self, seconds):
        self.calls.append(seconds)
        self.total += seconds


class NotifyStub:
    def __init__(self, ok=True):
        self.ok = ok
        self.calls = []

    def __call__(self, event_type, title, text):
        self.calls.append((event_type, title, text))
        return self.ok


class FakeChannel:
    """按 URL 特征回包的假通道：记录每次请求，可编程逐轮快照与异常。"""

    def __init__(self, get_info_payloads=None, resolve_payload=None,
                 get_info_error=None):
        self.calls = []
        self.get_info_payloads = list(get_info_payloads or [])
        self.resolve_payload = resolve_payload
        self.get_info_error = get_info_error

    def __call__(self, url):
        self.calls.append(url)
        if "getRoomPlayInfo" in url:
            if self.resolve_payload is None:
                raise BiliApiError("fake: resolve failed")
            return self.resolve_payload
        if self.get_info_error is not None:
            # 仅首轮 get_info 注入业务错误（模拟"短号不认"）；归一重试放行。
            error, self.get_info_error = self.get_info_error, None
            raise error
        if len(self.get_info_payloads) > 1:
            return self.get_info_payloads.pop(0)
        return self.get_info_payloads[0]


def budgeted_fetch(channel, budget):
    """镜像 BiliClient._request 的预算契约：请求前检查、请求后记账。"""
    def _fetch(url):
        budget.check_request()
        budget.observe_request()
        return channel(url)
    return _fetch


# ---------- fixture（形状照抄探针实测的 get_info 响应） ----------

def make_get_info(room_id=6, live_status=1, title="示例直播间", online=12345,
                  area="唱歌", parent="虚拟主播", live_time="0", uid=42,
                  tags="示例标签"):
    # online 实测以字符串下发，故意保持字符串覆盖真实情况。
    return {"code": 0, "message": "0", "data": {
        "room_id": room_id, "uid": uid, "live_status": live_status,
        "title": title, "online": str(online), "area_name": area,
        "parent_area_name": parent, "live_time": live_time, "tags": tags}}


def make_resolve(room_id=22474988, uid=42):
    return {"code": 0, "message": "0", "data": {"room_id": room_id, "uid": uid}}


def statuses_channel(statuses, **kwargs):
    """第 n 轮 live_status 依次取 statuses[n-1]，之后维持最后一个。"""
    payloads = [make_get_info(live_status=s, **kwargs) for s in statuses]
    return FakeChannel(get_info_payloads=payloads)


class LiveRoomTestBase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.out_dir = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()


# ---------- 1. 输入归一 ----------

class ParseRoomInputTests(unittest.TestCase):
    def test_plain_digits(self):
        self.assertEqual(core.parse_room_input("6"), 6)
        self.assertEqual(core.parse_room_input(" 22474988 "), 22474988)

    def test_live_link(self):
        self.assertEqual(
            core.parse_room_input("https://live.bilibili.com/6"), 6)
        self.assertEqual(
            core.parse_room_input("live.bilibili.com/22474988?spm_id_from=x"),
            22474988)
        self.assertEqual(
            core.parse_room_input("看 https://live.bilibili.com/7777 直播"), 7777)

    def test_invalid(self):
        for bad in ("", None, "abc", "live.bilibili.com/", "b23.tv/xxx",
                    "https://live.bilibili.com/"):
            with self.assertRaises(ValueError):
                core.parse_room_input(bad)


class ParseGetInfoTests(unittest.TestCase):
    def test_full_payload(self):
        row = core.parse_get_info(make_get_info(live_time="1700000000"))
        self.assertEqual(row["room_id"], 6)
        self.assertEqual(row["uid"], 42)
        self.assertEqual(row["live_status"], 1)
        self.assertEqual(row["live_status_label"], "直播中")
        self.assertEqual(row["online"], 12345)
        self.assertEqual(row["area_name"], "唱歌")
        self.assertEqual(row["live_time"], "2023-11-15 06:13:20")
        self.assertEqual(row["tags"], "示例标签")

    def test_missing_fields_degrade(self):
        row = core.parse_get_info({"code": 0, "data": {}})
        self.assertEqual(row["live_status"], -1)
        self.assertEqual(row["live_status_label"], "未知(-1)")
        self.assertEqual(row["online"], 0)
        self.assertEqual(row["title"], "")

    def test_short_id_normalized_from_response(self):
        row = core.parse_get_info(make_get_info(room_id=22474988),
                                  fallback_room=6)
        self.assertEqual(row["room_id"], 22474988)

    def test_live_time_not_live(self):
        self.assertEqual(core.parse_get_info(make_get_info()).get("live_time"), "")
        self.assertEqual(core.parse_get_info(
            make_get_info(live_time="abc")).get("live_time"), "abc")


# ---------- 2. 快照 / 追踪输出 ----------

class SnapshotPipelineTests(LiveRoomTestBase):
    def test_snapshot_excel_overview_only(self):
        channel = statuses_channel([1])
        result = run_pipeline("6", self.out_dir, mode="snapshot", fetch=channel)
        self.assertEqual(result["rows"], 1)
        self.assertEqual(result["jsonl"], "")
        self.assertEqual(result["stats"]["requests"], 1)
        self.assertIsNone(result["stats"]["stopped_reason"])
        xlsx = Path(result["xlsx"])
        self.assertTrue(xlsx.exists())
        wb = load_workbook(xlsx)
        self.assertEqual(wb.sheetnames, ["概览"])

    def test_snapshot_no_jsonl_file(self):
        channel = statuses_channel([0])
        result = run_pipeline("6", self.out_dir, mode="snapshot", fetch=channel)
        self.assertFalse((self.out_dir / "直播_6" / "live_6.jsonl").exists())


class TrackPipelineTests(LiveRoomTestBase):
    def test_track_jsonl_and_excel_two_sheets(self):
        channel = statuses_channel([0, 1, 1])
        sleep = RecordingSleep()
        result = run_pipeline("6", self.out_dir, mode="track", rounds=3,
                              interval=30, fetch=channel, sleeper=sleep)
        self.assertEqual(result["rows"], 3)
        self.assertIsNone(result["stats"]["stopped_reason"])
        self.assertEqual(result["stats"]["rounds_done"], 3)
        jsonl = Path(result["jsonl"])
        self.assertTrue(jsonl.exists())
        lines = [json.loads(line) for line in
                 jsonl.read_text(encoding="utf-8").splitlines()]
        self.assertEqual(len(lines), 3)
        for n, row in enumerate(lines, 1):
            self.assertEqual(row["round"], n)
            self.assertIn("live_status", row)
            self.assertIn("live_status_label", row)
            self.assertIn("room_id", row)
            self.assertIn("title", row)
            self.assertIn("online", row)
            self.assertIn("ts", row)
        self.assertEqual([row["live_status"] for row in lines], [0, 1, 1])
        wb = load_workbook(Path(result["xlsx"]))
        self.assertEqual(wb.sheetnames, ["概览", "快照明细"])
        # 轮间等待只发生在轮与轮之间（3 轮 = 2 次等待），每次 30s。
        self.assertAlmostEqual(sleep.total, 60.0)

    def test_excel_never_says_watch_count(self):
        channel = statuses_channel([1, 0])
        result = run_pipeline("6", self.out_dir, mode="track", rounds=2,
                              interval=30, fetch=channel,
                              sleeper=RecordingSleep())
        wb = load_workbook(Path(result["xlsx"]))
        # 卡口径约束的是**表头/指标名**必须写「人气」不许写「观看人数」；
        # 概览说明列的「非精确观看人数」是要求的口径注释，不受此限。
        detail_headers = [str(c.value) for c in wb["快照明细"][3] if c.value]
        self.assertIn("人气", detail_headers)
        self.assertFalse(any("观看人数" in h for h in detail_headers))
        metric_names = [str(wb["概览"][r][1].value or "")
                        for r in range(4, 14)]
        self.assertIn("人气（最新）", metric_names)
        self.assertFalse(any("观看人数" in name for name in metric_names))


# ---------- 3. 状态翻转推送 ----------

class FlipPushTests(LiveRoomTestBase):
    def test_flip_pushes_exactly_once(self):
        channel = statuses_channel([0, 1])
        notify = NotifyStub()
        result = run_pipeline("6", self.out_dir, mode="track", rounds=2,
                              interval=30, fetch=channel, notify=notify,
                              sleeper=RecordingSleep())
        self.assertEqual(len(notify.calls), 1)
        self.assertEqual(result["stats"]["pushes"], 1)
        event_type, title, text = notify.calls[0]
        self.assertEqual(event_type, "live_status")
        self.assertEqual(title, "开播提醒")
        self.assertIn("房间 6", text)
        self.assertIn("人气 12345", text)
        # 白名单三字段：payload 只允许 title/text/event_type。
        payload = webhook_payload(event_type, title, text)
        self.assertEqual(set(payload.keys()),
                         {"title", "text", "event_type"})

    def test_no_flip_no_push(self):
        channel = statuses_channel([1, 1])
        notify = NotifyStub()
        result = run_pipeline("6", self.out_dir, mode="track", rounds=2,
                              interval=30, fetch=channel, notify=notify,
                              sleeper=RecordingSleep())
        self.assertEqual(notify.calls, [])
        self.assertEqual(result["stats"]["pushes"], 0)

    def test_flip_then_flip_back_two_pushes(self):
        channel = statuses_channel([0, 1, 0])
        notify = NotifyStub()
        run_pipeline("6", self.out_dir, mode="track", rounds=3, interval=30,
                     fetch=channel, notify=notify, sleeper=RecordingSleep())
        self.assertEqual(len(notify.calls), 2)
        self.assertEqual(notify.calls[0][1], "开播提醒")
        self.assertEqual(notify.calls[1][1], "下播提醒")

    def test_first_round_is_baseline_only(self):
        channel = statuses_channel([1, 1, 1])
        notify = NotifyStub()
        run_pipeline("6", self.out_dir, mode="track", rounds=3, interval=30,
                     fetch=channel, notify=notify, sleeper=RecordingSleep())
        self.assertEqual(notify.calls, [])

    def test_notify_none_silent(self):
        channel = statuses_channel([0, 1])
        result = run_pipeline("6", self.out_dir, mode="track", rounds=2,
                              interval=30, fetch=channel, notify=None,
                              sleeper=RecordingSleep())
        self.assertEqual(result["stats"]["pushes"], 0)


# ---------- 4. 预算 / 轮数 / 取消 ----------

class BudgetTests(LiveRoomTestBase):
    def test_request_budget_graceful_stop(self):
        channel = statuses_channel([1, 1, 1])
        budget = TaskBudget(max_requests=2)
        result = run_pipeline("6", self.out_dir, mode="track", rounds=3,
                              interval=30, fetch=budgeted_fetch(channel, budget),
                              sleeper=RecordingSleep(), budget=budget)
        self.assertEqual(result["stats"]["stopped_reason"], "budget_reached")
        self.assertEqual(result["rows"], 2)
        self.assertTrue(Path(result["xlsx"]).exists())

    def test_duration_budget_stops_before_first_request(self):
        channel = statuses_channel([1])
        budget = TaskBudget(max_seconds=0)
        result = run_pipeline("6", self.out_dir, mode="track", rounds=3,
                              interval=30, fetch=channel, budget=budget,
                              sleeper=RecordingSleep())
        self.assertEqual(result["stats"]["stopped_reason"], "budget_reached")
        self.assertEqual(result["rows"], 0)
        self.assertEqual(result["stats"]["requests"], 0)
        self.assertEqual(result["xlsx"], "")

    def test_rounds_completed_normally(self):
        channel = statuses_channel([1, 1, 1])
        result = run_pipeline("6", self.out_dir, mode="track", rounds=3,
                              interval=30, fetch=channel,
                              sleeper=RecordingSleep())
        self.assertIsNone(result["stats"]["stopped_reason"])
        self.assertFalse(result["stats"]["cancelled"])
        self.assertEqual(result["rows"], 3)

    def test_cancel_takes_priority_over_budget(self):
        channel = statuses_channel([1])
        budget = TaskBudget(max_seconds=0)
        result = run_pipeline("6", self.out_dir, mode="track", rounds=3,
                              interval=30, fetch=channel, budget=budget,
                              cancel=lambda: True, sleeper=RecordingSleep())
        self.assertTrue(result["stats"]["cancelled"])
        self.assertIsNone(result["stats"]["stopped_reason"])
        self.assertEqual(result["rows"], 0)

    def test_cancel_during_wait_keeps_rows(self):
        channel = statuses_channel([1, 1, 1])
        state = {"calls": 0}

        def cancel_after_first_round():
            state["calls"] += 1
            # 第 1 轮请求 + 轮间等待期间取消：等待在首轮完成后开始。
            return state["calls"] > 1

        result = run_pipeline("6", self.out_dir, mode="track", rounds=3,
                              interval=30, fetch=channel,
                              cancel=cancel_after_first_round,
                              sleeper=RecordingSleep())
        self.assertTrue(result["stats"]["cancelled"])
        self.assertEqual(result["rows"], 1)
        lines = Path(result["jsonl"]).read_text(
            encoding="utf-8").splitlines()
        self.assertEqual(len(lines), 1)


# ---------- 5. 轮询礼仪：间隔下限 / 轮数上限 ----------

class PolitenessTests(LiveRoomTestBase):
    def test_interval_floor_30s(self):
        channel = statuses_channel([1, 1, 1])
        sleep = RecordingSleep()
        result = run_pipeline("6", self.out_dir, mode="track", rounds=3,
                              interval=5, fetch=channel, sleeper=sleep)
        self.assertEqual(result["stats"]["interval"], 30)
        self.assertAlmostEqual(sleep.total, 60.0)

    def test_interval_ceiling(self):
        channel = statuses_channel([1, 1])
        result = run_pipeline("6", self.out_dir, mode="track", rounds=2,
                              interval=99999, fetch=channel,
                              sleeper=RecordingSleep())
        self.assertEqual(result["stats"]["interval"], 3600)

    def test_rounds_cap_500(self):
        channel = statuses_channel([1])
        result = run_pipeline("6", self.out_dir, mode="track", rounds=1000,
                              interval=30, fetch=channel,
                              sleeper=RecordingSleep())
        self.assertEqual(result["stats"]["rounds_requested"], 500)
        self.assertEqual(result["rows"], 500)

    def test_unknown_mode_falls_back_to_snapshot(self):
        channel = statuses_channel([1])
        result = run_pipeline("6", self.out_dir, mode="bogus", fetch=channel)
        self.assertEqual(result["mode"], "snapshot")
        self.assertEqual(result["jsonl"], "")


# ---------- 6. 输入归一兜底（getRoomPlayInfo 最多一次） ----------

class ResolveFallbackTests(LiveRoomTestBase):
    def test_resolve_once_then_retry(self):
        channel = FakeChannel(
            get_info_payloads=[make_get_info(room_id=22474988)],
            resolve_payload=make_resolve(room_id=22474988),
            get_info_error=BiliApiError("fake: bad room"))
        result = run_pipeline("6", self.out_dir, mode="snapshot",
                              fetch=channel)
        self.assertEqual(result["room_id"], 22474988)
        self.assertEqual(result["stats"]["resolve_requests"], 1)
        self.assertEqual(len(channel.calls), 3)
        self.assertIn("room_id=6", channel.calls[0])
        self.assertIn("getRoomPlayInfo", channel.calls[1])
        self.assertIn("room_id=22474988", channel.calls[2])

    def test_resolve_failure_reraises_original(self):
        channel = FakeChannel(
            get_info_payloads=[make_get_info()],
            resolve_payload=None,
            get_info_error=BiliApiError("fake: bad room"))
        with self.assertRaises(BiliApiError):
            run_pipeline("6", self.out_dir, mode="snapshot", fetch=channel)

    def test_resolve_same_id_reraises_original(self):
        channel = FakeChannel(
            get_info_payloads=[make_get_info()],
            resolve_payload=make_resolve(room_id=6),
            get_info_error=BiliApiError("fake: bad room"))
        with self.assertRaises(BiliApiError):
            run_pipeline("6", self.out_dir, mode="snapshot", fetch=channel)
        self.assertEqual(len(channel.calls), 2)

    def test_response_normalizes_short_id_without_extra_request(self):
        # get_info 响应自带真实 room_id：短号归一零额外请求（首轮即归一）。
        channel = statuses_channel([1], room_id=22474988)
        result = run_pipeline("6", self.out_dir, mode="track", rounds=2,
                              interval=30, fetch=channel,
                              sleeper=RecordingSleep())
        self.assertEqual(result["room_id"], 22474988)
        self.assertEqual(result["stats"]["resolve_requests"], 0)
        self.assertEqual(result["stats"]["requests"], 2)
        self.assertIn("room_id=22474988", channel.calls[1])


# ---------- 7. 上提结构锁 + 注册一致性 ----------

class NotifyLiftStructureTests(unittest.TestCase):
    def test_monitor_reexports_are_core_objects(self):
        import core.notify as n
        import tools.monitor.notifications as m
        for name in ("WEBHOOK_FORMATS", "WEBHOOK_RATE_LIMIT",
                     "WEBHOOK_RATE_WINDOW_SECONDS", "WEBHOOK_TIMEOUT_SECONDS",
                     "LogFunc", "WebhookAdapter", "WebhookRateLimiter",
                     "_post_json", "_post_serverchan", "serverchan_url",
                     "webhook_payload"):
            self.assertIs(getattr(m, name), getattr(n, name),
                          f"tools.monitor.notifications.{name} 应 re-export core.notify")

    def test_core_notify_has_no_tool_or_app_import(self):
        source = Path(core.__file__).parents[2] / "core" / "notify.py"
        tree = ast.parse(source.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                names = [alias.name for alias in node.names]
            elif isinstance(node, ast.ImportFrom):
                names = [node.module or ""]
            else:
                continue
            for name in names:
                self.assertFalse(
                    name == "tools" or name.startswith(("tools.", "app", "app.")),
                    f"core/notify.py 不得依赖上层：{name}")


class RegistryTests(unittest.TestCase):
    def test_registered_last_with_matching_subtitle(self):
        spec = TOOLS[-1]
        self.assertEqual(spec.id, "live_room")
        self.assertEqual(spec.name, "直播追踪")
        self.assertEqual(spec.subtitle, LiveRoomPage.tool_subtitle)
        self.assertEqual(len(TOOLS), 9)
        # 既有 8 个工具与注册顺序不变。
        self.assertEqual([s.id for s in TOOLS][:8],
                         ["comments", "collector", "monitor", "data_check",
                          "report_center", "user_dynamics", "danmaku",
                          "relation_analysis"])


class TransitionTests(unittest.TestCase):
    def test_transition_kinds(self):
        self.assertEqual(core.status_transition(None, 1), (False, ""))
        self.assertEqual(core.status_transition(1, 1), (False, ""))
        self.assertEqual(core.status_transition(0, 1), (True, "live"))
        self.assertEqual(core.status_transition(1, 0), (True, "offline"))
        self.assertEqual(core.status_transition(1, 2), (True, "change"))

    def test_push_message_offline_and_change(self):
        row = {"room_id": 6, "title": "示例直播间", "online": 7,
               "live_status_label": "未开播"}
        _et, title, text = core.push_message(row, "offline")
        self.assertEqual(title, "下播提醒")
        self.assertIn("已下播", text)
        _et, title, text = core.push_message(row, "change", "直播中")
        self.assertEqual(title, "直播状态变化提醒")
        self.assertIn("直播中 → 未开播", text)


if __name__ == "__main__":
    unittest.main()
