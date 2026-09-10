# -*- coding: utf-8 -*-
"""数据修复与合并的本地、只读回归测试。"""
import hashlib
import json
import os
import socket
import tempfile
import unittest
import zipfile
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from tools.data_check import repair_core


def comment(rpid=1, message="正常评论", **extra):
    value = {
        "rpid": rpid, "parent": 0, "root": 0, "is_main": True,
        "uname": "用户", "mid": 1, "sex": "保密", "level": 6,
        "vip": False, "message": message, "like": 0, "rcount": 0,
        "ctime": 1700000000, "location": "", "is_top": False,
    }
    value.update(extra)
    return value


def video(bvid="BV1Demo", fetched_at=1, **extra):
    value = {
        "bvid": bvid, "aid": 1, "title": "视频", "owner": "UP",
        "owner_mid": 2, "tname": "分区", "pubdate": 1, "duration": 60,
        "view": 0, "danmaku": 0, "reply": 0, "favorite": 0, "coin": 0,
        "share": 0, "like": False, "fetched_at": fetched_at,
    }
    value.update(extra)
    return value


class DataRepairTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="data-repair-test-")
        self.root = Path(self.temp.name)
        self.out = self.root / "out"

    def tearDown(self):
        self.temp.cleanup()

    def write_jsonl(self, name, values):
        path = self.root / name
        with path.open("w", encoding="utf-8", newline="\n") as handle:
            for value in values:
                if isinstance(value, str):
                    handle.write(value)
                else:
                    handle.write(json.dumps(value, ensure_ascii=False))
                if not (isinstance(value, str) and value.endswith("\n")):
                    handle.write("\n")
        return path

    def write_xlsx(self, name, sheets):
        path = self.root / name
        workbook = Workbook()
        workbook.remove(workbook.active)
        for title, rows in sheets:
            ws = workbook.create_sheet(title)
            for row in rows:
                ws.append(row)
        workbook.save(path)
        return path

    def snapshot(self, path):
        stat = path.stat()
        return hashlib.sha256(path.read_bytes()).hexdigest(), stat.st_size, stat.st_mtime_ns

    def assert_unchanged(self, path, before):
        stat = path.stat()
        self.assertEqual(self.snapshot(path), before)
        self.assertEqual(stat.st_mtime_ns, before[2])

    def run_repair(self, paths, **options):
        before = {path: self.snapshot(path) for path in paths}
        snapshots = repair_core.snapshot_sources(paths)
        result = repair_core.run_repair(
            [str(path) for path in paths], str(self.out), snapshots, **options)
        for path in paths:
            self.assert_unchanged(path, before[path])
        return result

    def read_jsonl(self, path):
        return [json.loads(line) for line in Path(path).read_text(encoding="utf-8").splitlines()
                if line.strip()]

    def test_jsonl_blank_cleanup_preserves_zero_false_and_empty_containers(self):
        path = self.write_jsonl("unknown.jsonl", ["\n", 0, False, [], {}, "   \n"])
        result = self.run_repair([path], normalize_fields=False)
        self.assertEqual(self.read_jsonl(result["copies"][0]), [0, False, [], {}])
        self.assertEqual(result["files"][0]["blanks_removed"], 2)

    def test_excel_blank_cleanup_preserves_dates_zero_and_false(self):
        path = self.write_xlsx("rows.xlsx", [("数据", [
            ["日期", "数字", "布尔"], [None, "  ", None],
            [date(2026, 9, 9), 0, False],
        ])])
        result = self.run_repair([path], normalize_fields=False)
        workbook = load_workbook(result["copies"][0], data_only=False)
        rows = list(workbook["数据"].iter_rows(values_only=True))
        workbook.close()
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][1:], (0, False))

    def test_comment_rpid_duplicate_and_conflict_keep_first(self):
        path = self.write_jsonl("comments.jsonl", [
            comment(1, "first"), comment("1", "first"), comment(1, "conflict"),
        ])
        result = self.run_repair([path])
        rows = self.read_jsonl(result["copies"][0])
        self.assertEqual([row["message"] for row in rows], ["first"])
        self.assertEqual(result["files"][0]["duplicates_removed"], 2)
        workbook = load_workbook(result["manifest"], read_only=True, data_only=True)
        details = list(workbook["修复明细"].iter_rows(values_only=True))
        workbook.close()
        self.assertTrue(any(row[2] == "冲突重复" and "rpid=1" in row[4] for row in details))

    def test_video_numeric_and_string_timestamp_are_same_key(self):
        path = self.write_jsonl("videos.jsonl", [video(fetched_at=1), video(fetched_at="1")])
        result = self.run_repair([path])
        self.assertEqual(len(self.read_jsonl(result["copies"][0])), 1)
        self.assertEqual(result["files"][0]["duplicates_removed"], 1)

    def test_cross_file_merge_is_stable_and_deduplicated(self):
        one = self.write_jsonl("one.jsonl", [comment(1), comment(2)])
        two = self.write_jsonl("two.jsonl", [comment("2"), comment(3)])
        result = self.run_repair([one, two], merge=True)
        self.assertEqual(len(result["copies"]), 2)
        self.assertEqual(len(result["merged"]), 1)
        self.assertEqual([row["rpid"] for row in self.read_jsonl(result["merged"][0])], [1, 2, 3])

    def test_unknown_exact_record_duplicate_is_removed(self):
        path = self.write_jsonl("unknown.jsonl", [{"x": 1}, {"x": 1}, {"x": 2}])
        result = self.run_repair([path])
        self.assertEqual(self.read_jsonl(result["copies"][0]), [{"x": 1}, {"x": 2}])

    def test_known_jsonl_missing_fields_and_field_order_with_extra(self):
        path = self.write_jsonl("partial.jsonl", [
            {"rpid": 1, "message": "x", "extra": "kept"}, comment(2),
        ])
        result = self.run_repair([path])
        rows = self.read_jsonl(result["copies"][0])
        self.assertEqual(list(rows[0])[:3], ["rpid", "parent", "root"])
        self.assertIsNone(rows[0]["parent"])
        self.assertEqual(rows[0]["extra"], "kept")
        self.assertGreater(result["files"][0]["missing_filled"], 0)

    def test_known_jsonl_uses_one_stable_extra_field_order(self):
        first = comment(1)
        first.update({"extra_b": 2})
        second = comment(2)
        second.update({"extra_a": 1})
        path = self.write_jsonl("stable-fields.jsonl", [first, second])
        result = self.run_repair([path])
        rows = self.read_jsonl(result["copies"][0])
        self.assertEqual(list(rows[0]), list(rows[1]))
        self.assertEqual(list(rows[0])[-2:], ["extra_b", "extra_a"])
        self.assertIsNone(rows[0]["extra_a"])
        self.assertIsNone(rows[1]["extra_b"])

    def test_known_excel_preserves_nondata_sheet_and_formula(self):
        path = self.write_xlsx("comments.xlsx", [
            ("统计概览", [["值", "公式"], [2, "=A2*2"]]),
            ("全量评论", [[None, "rpid", "评论内容", "额外列"],
                           [None, 1, "内容", "保留"]]),
        ])
        result = self.run_repair([path])
        workbook = load_workbook(result["copies"][0], data_only=False)
        self.assertEqual(workbook["统计概览"]["B2"].value, "=A2*2")
        headers = [cell.value for cell in workbook["全量评论"][1]]
        workbook.close()
        self.assertIn("额外列", headers)
        self.assertLess(headers.index("rpid"), headers.index("额外列"))

    def test_known_excel_missing_columns_are_added_as_blank(self):
        path = self.write_xlsx("partial-comments.xlsx", [
            ("全量评论", [[None, "rpid", "评论内容"], [None, 1, "内容"]]),
        ])
        result = self.run_repair([path])
        workbook = load_workbook(result["copies"][0], data_only=False)
        try:
            headers = [cell.value for cell in workbook["全量评论"][1]][1:]
            values = [cell.value for cell in workbook["全量评论"][2]][1:]
        finally:
            workbook.close()
        self.assertEqual(headers[:len(repair_core.COMMENT_EXCEL_FIELDS)],
                         list(repair_core.COMMENT_EXCEL_FIELDS))
        self.assertIsNone(values[headers.index("用户昵称")])

    def test_known_excel_invalid_business_key_is_retained_and_manifested(self):
        path = self.write_xlsx("invalid-rpid.xlsx", [
            ("全量评论", [[None, "rpid", "评论内容"], [None, "bad", "内容"]]),
        ])
        result = self.run_repair([path])
        workbook = load_workbook(result["copies"][0], read_only=True, data_only=True)
        try:
            self.assertEqual(workbook["全量评论"]["C2"].value, "bad")
        finally:
            workbook.close()
        self.assertEqual(result["files"][0]["unfixed"], 1)

    def test_compatible_excel_files_merge_and_cross_file_deduplicate(self):
        one = self.write_xlsx("one.xlsx", [
            ("全量评论", [[None, "rpid", "评论内容"], [None, 1, "一"]]),
        ])
        two = self.write_xlsx("two.xlsx", [
            ("全量评论", [[None, "rpid", "评论内容"],
                           [None, "1", "重复"], [None, 2, "二"]]),
        ])
        result = self.run_repair([one, two], merge=True)
        self.assertEqual(len(result["merged"]), 1)
        workbook = load_workbook(result["merged"][0], read_only=True, data_only=True)
        try:
            ws = workbook["全量评论"]
            values = list(ws.iter_rows(min_row=2, values_only=True))
        finally:
            workbook.close()
        self.assertEqual([row[2] for row in values], [1, 2])

    def test_comment_video_and_formats_are_never_cross_merged(self):
        comment_json = self.write_jsonl("comments.jsonl", [comment()])
        video_json = self.write_jsonl("videos.jsonl", [video()])
        comment_xlsx = self.write_xlsx("comments.xlsx", [
            ("全量评论", [[None, "rpid", "评论内容"], [None, 1, "内容"]])])
        result = self.run_repair([comment_json, video_json, comment_xlsx], merge=True)
        self.assertEqual(result["merged"], [])

    def test_unknown_complex_excel_is_skipped_without_blocking_good_file(self):
        complex_path = self.write_xlsx("complex.xlsx", [("数据", [["x"], ["=1+1"]])])
        good = self.write_jsonl("good.jsonl", [comment()])
        result = self.run_repair([complex_path, good])
        self.assertEqual(result["stats"]["skipped_files"], 1)
        self.assertEqual(result["stats"]["processed_files"], 1)
        self.assertEqual(len(result["copies"]), 1)

    def test_damaged_excel_is_skipped_without_blocking_good_file(self):
        damaged = self.root / "damaged.xlsx"
        damaged.write_bytes(b"not-an-xlsx")
        good = self.write_jsonl("good-after-damage.jsonl", [comment()])
        result = self.run_repair([damaged, good])
        self.assertEqual(result["stats"]["skipped_files"], 1)
        self.assertEqual(result["stats"]["processed_files"], 1)
        self.assertTrue(Path(result["manifest"]).is_file())

    def test_malformed_jsonl_is_skipped_and_manifested(self):
        path = self.write_jsonl("broken-lines.jsonl", ["{broken\n", comment()])
        result = self.run_repair([path])
        self.assertEqual(len(self.read_jsonl(result["copies"][0])), 1)
        self.assertEqual(result["files"][0]["unfixed"], 1)

    def test_existing_output_names_are_suffixed_and_unchanged(self):
        path = self.write_jsonl("same.jsonl", [comment()])
        self.out.mkdir()
        existing = self.out / "same_修复副本.jsonl"
        existing.write_text("old", encoding="utf-8")
        result = self.run_repair([path])
        self.assertEqual(existing.read_text(encoding="utf-8"), "old")
        self.assertNotEqual(Path(result["copies"][0]), existing)

    def test_output_directory_may_equal_source_directory_without_overwrite(self):
        path = self.write_jsonl("source.jsonl", [comment()])
        before = self.snapshot(path)
        result = repair_core.run_repair(
            [str(path)], str(self.root), repair_core.snapshot_sources([path]))
        self.assertTrue(Path(result["copies"][0]).is_file())
        self.assert_unchanged(path, before)

    def test_cancel_leaves_no_final_or_temp_files(self):
        path = self.write_jsonl("cancel.jsonl", [comment()])
        result = self.run_repair([path], cancel=lambda: True)
        self.assertTrue(result["stats"]["cancelled"])
        self.assertEqual(list(self.out.iterdir()), [])

    def test_write_failure_leaves_no_half_published_files(self):
        path = self.write_jsonl("failure.jsonl", [comment()])
        before = self.snapshot(path)
        with patch.object(repair_core, "_write_manifest", side_effect=OSError("disk")):
            with self.assertRaises(OSError):
                repair_core.run_repair(
                    [str(path)], str(self.out), repair_core.snapshot_sources([path]))
        self.assertEqual(list(self.out.iterdir()), [])
        self.assert_unchanged(path, before)

    def test_source_changed_before_repair_is_rejected(self):
        path = self.write_jsonl("changed.jsonl", [comment()])
        snapshots = repair_core.snapshot_sources([path])
        path.write_text("{}\n", encoding="utf-8")
        result = repair_core.run_repair([str(path)], str(self.out), snapshots)
        self.assertEqual(result["stats"]["status"], "interrupted")
        self.assertEqual(list(self.out.iterdir()), [])

    def test_source_changed_during_repair_stops_publication(self):
        path = self.write_jsonl("during.jsonl", [comment(1), comment(2)])
        snapshots = repair_core.snapshot_sources([path])
        changed = False

        def progress(**_kwargs):
            nonlocal changed
            if not changed:
                path.write_text(path.read_text(encoding="utf-8") + "{}\n", encoding="utf-8")
                changed = True

        result = repair_core.run_repair([str(path)], str(self.out), snapshots, progress=progress)
        self.assertEqual(result["stats"]["status"], "interrupted")
        self.assertEqual(list(self.out.iterdir()), [])

    def test_manifest_has_three_sheets_counts_and_no_record_secrets(self):
        secret = "Cookie=secret-value Authorization=Bearer abc"
        path = self.write_jsonl("privacy.jsonl", [comment(1, secret), comment(1, "other")])
        result = self.run_repair([path])
        workbook = load_workbook(result["manifest"], read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["修复概览", "文件汇总", "修复明细"])
        overview = {row[0]: row[1] for row in
                    workbook["修复概览"].iter_rows(min_row=2, values_only=True)}
        workbook.close()
        self.assertEqual(overview["删除重复记录"], 1)
        with zipfile.ZipFile(result["manifest"]) as archive:
            raw = b"".join(archive.read(name) for name in archive.namelist()
                           if name.startswith("xl/worksheets/"))
        self.assertNotIn(b"secret-value", raw)
        self.assertNotIn(b"Bearer abc", raw)

    def test_detail_cap_keeps_full_statistics(self):
        path = self.root / "many-blanks.jsonl"
        path.write_text("\n" * 10002, encoding="utf-8")
        result = self.run_repair([path], max_detail_rows=3)
        self.assertEqual(result["files"][0]["blanks_removed"], 10002)
        self.assertTrue(result["stats"]["details_truncated"])
        workbook = load_workbook(result["manifest"], read_only=True, data_only=True)
        try:
            self.assertEqual(len(list(workbook["修复明细"].iter_rows(values_only=True))), 4)
        finally:
            workbook.close()

    def test_repair_core_has_no_network_call(self):
        path = self.write_jsonl("offline.jsonl", [comment()])
        with patch.object(socket, "create_connection", side_effect=AssertionError("network")), \
                patch("urllib.request.urlopen", side_effect=AssertionError("network")):
            result = self.run_repair([path])
        self.assertEqual(result["stats"]["processed_files"], 1)

    def test_merge_options_false_preserve_jsonl_blanks_and_duplicates(self):
        one = self.write_jsonl("options-one.jsonl", [comment(1), "\n", comment(2)])
        two = self.write_jsonl("options-two.jsonl", [comment(1), "\n", comment(3)])
        result = self.run_repair(
            [one, two], deduplicate=False, clean_blanks=False, merge=True)
        merged = Path(result["merged"][0]).read_text(encoding="utf-8")
        self.assertEqual(merged.count("\n"), 6)
        self.assertEqual(merged.count('"rpid":1'), 2)
        self.assertEqual(merged.count('"rpid":2'), 1)
        self.assertEqual(merged.count('"rpid":3'), 1)
        self.assertEqual(merged.splitlines().count(""), 2)

    def test_merge_options_false_preserve_excel_blanks_and_duplicates(self):
        rows = [[None, "rpid", "评论内容"], [None, 1, "一"],
                [None, None, None], [None, 2, "二"]]
        one = self.write_xlsx("options-one.xlsx", [("全量评论", rows)])
        two = self.write_xlsx("options-two.xlsx", [("全量评论", rows)])
        result = self.run_repair(
            [one, two], deduplicate=False, clean_blanks=False, merge=True)
        workbook = load_workbook(result["merged"][0], read_only=True, data_only=True)
        try:
            values = list(workbook["全量评论"].iter_rows(min_row=2, values_only=True))
        finally:
            workbook.close()
        self.assertEqual(len(values), 6)
        self.assertEqual(sum(row[2] == 1 for row in values), 2)
        self.assertEqual(sum(row[2] == 2 for row in values), 2)
        self.assertEqual(sum(all(value is None for value in row) for row in values), 2)

    def test_known_jsonl_merge_normalizes_group_extension_union(self):
        first = comment(1, extra_a="A")
        second = comment(2, extra_b="B")
        one = self.write_jsonl("union-one.jsonl", [first])
        two = self.write_jsonl("union-two.jsonl", [second])
        result = self.run_repair([one, two], normalize_fields=True, merge=True)
        rows = self.read_jsonl(result["merged"][0])
        self.assertEqual(list(rows[0]), list(rows[1]))
        self.assertEqual(rows[0]["extra_a"], "A")
        self.assertIsNone(rows[0]["extra_b"])
        self.assertIsNone(rows[1]["extra_a"])
        self.assertEqual(rows[1]["extra_b"], "B")

    def test_known_excel_merge_preserves_group_extension_union(self):
        one = self.write_xlsx("union-one.xlsx", [
            ("全量评论", [[None, "rpid", "评论内容", "额外A"],
                           [None, 1, "一", "A"]]),
        ])
        two = self.write_xlsx("union-two.xlsx", [
            ("全量评论", [[None, "rpid", "评论内容", "额外B"],
                           [None, 2, "二", "B"]]),
        ])
        result = self.run_repair([one, two], normalize_fields=True, merge=True)
        workbook = load_workbook(result["merged"][0], read_only=True, data_only=True)
        try:
            rows = list(workbook["全量评论"].iter_rows(values_only=True))
        finally:
            workbook.close()
        headers = list(rows[0])
        self.assertIn("额外A", headers)
        self.assertIn("额外B", headers)
        values = {row[2]: row for row in rows[1:]}
        self.assertEqual(values[1][headers.index("额外A")], "A")
        self.assertEqual(values[2][headers.index("额外B")], "B")

    def test_known_excel_merge_preserves_later_nondata_sheets_and_formulas(self):
        one = self.write_xlsx("merge-one.xlsx", [
            ("统计A", [["值", "公式"], [2, "=A2*2"]]),
            ("全量评论", [[None, "rpid", "评论内容"], [None, 1, "一"]]),
        ])
        two = self.write_xlsx("merge-two.xlsx", [
            ("统计B", [["值", "公式"], [3, "=B2*3"]]),
            ("全量评论", [[None, "rpid", "评论内容"], [None, 2, "二"]]),
        ])
        result = self.run_repair([one, two], merge=True)
        workbook = load_workbook(result["merged"][0], data_only=False)
        try:
            self.assertIn("统计A", workbook.sheetnames)
            self.assertIn("统计B", workbook.sheetnames)
            self.assertEqual(workbook.sheetnames.count("全量评论"), 1)
            self.assertEqual(workbook["统计A"]["B2"].value, "=A2*2")
            self.assertEqual(workbook["统计B"]["B2"].value, "=B2*3")
            rows = list(workbook["全量评论"].iter_rows(min_row=2, values_only=True))
            self.assertEqual([row[2] for row in rows], [1, 2])
        finally:
            workbook.close()

    def test_known_excel_merge_suffixes_conflicting_nondata_sheet_names(self):
        one = self.write_xlsx("same-one.xlsx", [
            ("统计概览", [["来源", "值"], ["一", 1]]),
            ("全量评论", [[None, "rpid", "评论内容"], [None, 1, "一"]]),
        ])
        two = self.write_xlsx("same-two.xlsx", [
            ("统计概览", [["来源", "值"], ["二", 2]]),
            ("全量评论", [[None, "rpid", "评论内容"], [None, 2, "二"]]),
        ])
        result = self.run_repair([one, two], merge=True)
        workbook = load_workbook(result["merged"][0], data_only=False)
        try:
            self.assertIn("统计概览", workbook.sheetnames)
            self.assertIn("统计概览_2", workbook.sheetnames)
            self.assertEqual(workbook["统计概览"]["A2"].value, "一")
            self.assertEqual(workbook["统计概览_2"]["A2"].value, "二")
            self.assertEqual(workbook.sheetnames.count("全量评论"), 1)
        finally:
            workbook.close()

    def test_known_excel_merge_rewrites_renamed_nondata_formula_references(self):
        one = self.write_xlsx("refs-one.xlsx", [
            ("统计概览", [["值"], [1]]),
            ("引用", [["='统计概览'!A2"], ["=统计概览!A2"]]),
            ("全量评论", [[None, "rpid", "评论内容", "引用公式"],
                           [None, 1, "一", "='统计概览'!A2"]]),
        ])
        two = self.write_xlsx("refs-two.xlsx", [
            ("统计概览", [["值"], [2]]),
            ("引用", [["='统计概览'!A2"], ["=统计概览!A2"]]),
            ("全量评论", [[None, "rpid", "评论内容", "引用公式"],
                           [None, 2, "二", "='统计概览'!A2"]]),
        ])
        result = self.run_repair([one, two], merge=True)
        workbook = load_workbook(result["merged"][0], data_only=False)
        try:
            self.assertEqual(workbook["引用"]["A1"].value, "='统计概览'!A2")
            self.assertEqual(workbook["引用_2"]["A1"].value, "='统计概览_2'!A2")
            self.assertEqual(workbook["引用_2"]["A2"].value, "=统计概览_2!A2")
            headers = [cell.value for cell in workbook["全量评论"][1]]
            rows = list(workbook["全量评论"].iter_rows(min_row=2, values_only=True))
            formula_column = headers.index("引用公式")
            self.assertEqual(rows[0][formula_column], "='统计概览'!A2")
            self.assertEqual(rows[1][formula_column], "='统计概览_2'!A2")
        finally:
            workbook.close()

    def test_unknown_jsonl_signature_scans_after_first_hundred_rows(self):
        long_file = self.write_jsonl(
            "unknown-long.jsonl", [{"x": index} for index in range(100)] + [{"y": 100}])
        pure_file = self.write_jsonl("unknown-pure.jsonl", [{"x": 1000}])
        result = self.run_repair([long_file, pure_file], merge=True)
        self.assertEqual(result["merged"], [])

    def test_repair_history_reuse_restores_options_and_directory_without_starting(self):
        from app.task_page import TaskPage
        from tools.data_check.page import DataCheckPage
        from PySide6.QtWidgets import QApplication

        app = QApplication.instance() or QApplication([])

        with tempfile.TemporaryDirectory(prefix="data-repair-reuse-") as temp:
            root = Path(temp)
            source = root / "reuse.jsonl"
            source.write_text(json.dumps(comment()) + "\n", encoding="utf-8")
            page = DataCheckPage({"out_dir": str(root / "reports")})
            params = {
                "files": [str(source)],
                "out_dir": str(root / "reports"),
                "repair_out_dir": str(root / "custom-repair"),
                "deduplicate": False,
                "clean_blanks": False,
                "normalize_fields": False,
                "merge": True,
            }
            with patch.object(TaskPage, "on_start") as check_start, \
                    patch.object(page, "_start_repair") as repair_start:
                page.apply_reusable_params(params)
                check_start.assert_not_called()
                repair_start.assert_not_called()
            self.assertEqual(page.repair_out_row.value(), params["repair_out_dir"])
            self.assertFalse(page.repair_deduplicate.isChecked())
            self.assertFalse(page.repair_clean_blanks.isChecked())
            self.assertFalse(page.repair_normalize.isChecked())
            self.assertTrue(page.repair_merge.isChecked())
            page.on_finished({
                "report": str(root / "reports" / "report.xlsx"),
                "dir": str(root / "reports"),
                "stats": {"total_issues": 0},
                "source_files": repair_core.snapshot_sources([source]),
            })
            self.assertEqual(page.repair_out_row.value(), params["repair_out_dir"])
            app.processEvents()


class DataRepairPageTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtWidgets import QApplication
        cls.app = QApplication.instance() or QApplication([])

    def test_check_and_repair_runners_cannot_start_concurrently(self):
        from app.task_page import TaskPage
        from tools.data_check.page import DataCheckPage

        class Running:
            @staticmethod
            def isRunning():
                return True

        page = DataCheckPage.__new__(DataCheckPage)
        page.repair_runner = Running()
        with patch.object(TaskPage, "on_start") as parent_start:
            DataCheckPage.on_start(page)
        parent_start.assert_not_called()

        page.repair_runner = None
        page.runner = Running()
        DataCheckPage._start_repair(page)  # 返回时不应访问尚未构造的控件

    def test_success_exposes_repair_card_and_file_edit_invalidates_snapshot(self):
        from tools.data_check.page import DataCheckPage

        with tempfile.TemporaryDirectory(prefix="data-repair-page-") as temp:
            root = Path(temp)
            source = root / "data.jsonl"
            source.write_text(json.dumps(comment()) + "\n", encoding="utf-8")
            page = DataCheckPage({"out_dir": str(root)})
            page.file_list.addItem(str(source))
            report_dir = root / "reports"
            page.on_finished({
                "report": str(report_dir / "report.xlsx"),
                "dir": str(report_dir),
                "stats": {"total_issues": 0},
                "source_files": repair_core.snapshot_sources([source]),
            })
            self.assertTrue(page.repair_card.isVisibleTo(page))
            self.assertEqual(page.repair_out_row.value(), str(report_dir / "修复副本"))
            self.assertTrue(page._repair_snapshots)
            page.file_list.clear()
            self.app.processEvents()
            self.assertFalse(page._repair_snapshots)
            self.assertTrue(page.repair_card.isHidden())

    def test_repair_completion_updates_separate_history_with_all_outputs(self):
        from tools.data_check.page import DataCheckPage

        page = DataCheckPage({"out_dir": "."})
        page._repair_history_id = "repair-id"
        result = {
            "copies": ["D:/out/a.jsonl"],
            "merged": ["D:/out/merged.jsonl"],
            "manifest": "D:/out/manifest.xlsx",
            "dir": "D:/out",
            "stats": {"status": "completed", "skipped_files": 0},
        }
        with patch.object(page, "_update_repair_history") as update:
            page._on_repair_finished(result)
        update.assert_called_once_with(
            "completed",
            outputs=["D:/out/a.jsonl", "D:/out/merged.jsonl", "D:/out/manifest.xlsx"],
        )


if __name__ == "__main__":
    unittest.main()
