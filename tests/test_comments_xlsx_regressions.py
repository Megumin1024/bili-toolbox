# -*- coding: utf-8 -*-
"""评论 Excel 导出边界：VIP、日期显示和 UID 列表。"""

from datetime import datetime
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from tools.comments.core import analyze, export_xlsx


def meta():
    return {
        "author": "演示UP",
        "title": "演示视频",
        "pub_ts": 1_699_900_000,
        "claimed_comment_count": 10,
    }


def row(rpid, mid=100, vip=0, uname="用户", ctime=1_700_000_000):
    result = {
        "rpid": rpid,
        "is_main": True,
        "mid": mid,
        "uname": uname,
        "vip": vip,
        "sex": "保密",
        "level": 6,
        "message": f"评论{rpid}",
        "like": 0,
        "rcount": 0,
        "ctime": ctime,
        "location": "",
    }
    if mid == "__missing__":
        result.pop("mid")
    return result


def row_by_header(worksheet, header, header_row=1):
    for cell in worksheet[header_row]:
        if cell.value == header:
            return cell.row, cell.column
    raise AssertionError(f"未找到表头：{header}")


class CommentsXlsxRegressionTests(unittest.TestCase):
    def export_and_load(self, rows, metadata=None):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "comments.xlsx"
        metadata = metadata or meta()
        report, kpi = analyze(rows, metadata)
        export_xlsx(rows, metadata, kpi, path)
        workbook = load_workbook(path, data_only=False)
        self.addCleanup(workbook.close)
        return path, report, kpi, workbook

    def test_vip_statuses_reopen_as_boolean_or_explicit_annotated_state(self):
        missing_vip = row(5, 105, 0)
        missing_vip.pop("vip")
        rows = [
            row(1, 101, 0),
            row(2, 102, 1),
            row(3, 103, True),
            row(4, 104, False),
            missing_vip,
            row(6, 106, None),
            row(7, 107, 2),
            row(8, 108, 1.0),
            row(9, 109, "1"),
        ]
        _path, _report, kpi, workbook = self.export_and_load(rows)
        detail = workbook["全量评论"]
        _header_row, vip_col = row_by_header(detail, "大会员")
        expected = [False, True, True, False, None, None, None, None, None]
        for row_number, expected_value in enumerate(expected, 2):
            cell = detail.cell(row_number, vip_col)
            self.assertEqual(cell.value, expected_value)
            if expected_value is not None:
                self.assertEqual(cell.data_type, "b")
                self.assertIsNone(cell.comment)
            else:
                self.assertIsNotNone(cell.comment)
                if row_number == 6:
                    self.assertIn("未返回", cell.comment.text)
                elif row_number == 7:
                    self.assertIn("缺失", cell.comment.text)
                else:
                    self.assertIn("未知", cell.comment.text)

        self.assertEqual(kpi["vip_pct"], "22.2%")

    def test_analyze_counts_only_explicit_true_vip_states(self):
        rows = [
            row(1, 101, 0),
            row(2, 102, 1),
            row(3, 103, True),
            row(4, 104, 2),
            row(5, 105, "true"),
            row(6, 106, 1.0),
        ]
        _report, kpi = analyze(rows, meta())
        self.assertEqual(kpi["vip_ratio"], 2 / 6)
        self.assertEqual(kpi["vip_pct"], "33.3%")

    def test_dates_are_real_datetime_and_wide_enough_after_reopen(self):
        rows = [row(1, 101, 0), row(2, 102, 1, ctime=1_700_003_600)]
        _path, _report, _kpi, workbook = self.export_and_load(rows)

        overview = workbook["统计概览"]
        overview_label = next(
            cell for row_values in overview.iter_rows()
            for cell in row_values if cell.value == "发布时间"
        )
        overview_value = overview.cell(overview_label.row, overview_label.column + 1)
        self.assertIsInstance(overview_value.value, datetime)
        self.assertEqual(overview_value.number_format, "yyyy-mm-dd hh:mm:ss")
        self.assertGreaterEqual(overview.column_dimensions["C"].width, 21)

        detail = workbook["全量评论"]
        detail_header_row, date_col = row_by_header(detail, "发布时间")
        date_cell = detail.cell(detail_header_row + 1, date_col)
        self.assertIsInstance(date_cell.value, datetime)
        self.assertEqual(date_cell.number_format, "yyyy-mm-dd hh:mm:ss")
        self.assertGreaterEqual(
            detail.column_dimensions[get_column_letter(date_col)].width, 21
        )
        self.assertNotEqual(date_cell.value, "#####")
        for worksheet in workbook.worksheets:
            for row_values in worksheet.iter_rows():
                for cell in row_values:
                    if isinstance(cell.value, datetime):
                        self.assertGreaterEqual(
                            worksheet.column_dimensions[cell.column_letter].width,
                            21,
                            f"日期列过窄：{worksheet.title}!{cell.coordinate}",
                        )

    def test_uid_list_is_stable_deduplicated_exact_text_and_linked(self):
        first_uid = "12345678901234567890"
        second_uid = "98765432109876543210"
        rows = [
            row(1, first_uid, 0, "首个用户"),
            row(2, int(first_uid), 0, "重复用户"),
            row(3, second_uid, 1, "第二用户"),
            row(4, None, 0),
            row(5, "__missing__", 0),
            row(6, "not-a-mid", 0),
            row(7, 1.5, 0),
            row(8, True, 0),
            row(9, 0, 0),
        ]
        _path, _report, _kpi, workbook = self.export_and_load(rows)

        self.assertEqual(workbook.sheetnames, [
            "统计概览", "分布统计", "点赞Top100", "全量评论", "UID列表",
            "数据质量", "字段说明",
        ])

        uid_sheet = workbook["UID列表"]
        uid_header_row = next(
            cell.row for cell in uid_sheet[3] if cell.value == "UID"
        )
        uid_cols = {
            cell.value: cell.column for cell in uid_sheet[uid_header_row]
            if cell.value in {"UID", "用户昵称", "评论数"}
        }
        values = [
            (
                uid_sheet.cell(row_number, uid_cols["UID"]).value,
                uid_sheet.cell(row_number, uid_cols["用户昵称"]).value,
                uid_sheet.cell(row_number, uid_cols["评论数"]).value,
            )
            for row_number in range(uid_header_row + 1, uid_sheet.max_row + 1)
            if uid_sheet.cell(row_number, uid_cols["UID"]).value is not None
        ]
        self.assertEqual(values, [
            (first_uid, "首个用户", 2),
            (second_uid, "第二用户", 1),
        ])
        uid_cell = uid_sheet.cell(uid_header_row + 1, uid_cols["UID"])
        self.assertEqual(uid_cell.data_type, "s")
        self.assertEqual(uid_cell.number_format, "@")

        overview = workbook["统计概览"]
        linked = [
            cell for row_values in overview.iter_rows()
            for cell in row_values
            if cell.value == "UID列表" and cell.hyperlink is not None
        ]
        self.assertEqual(len(linked), 1)
        self.assertEqual(linked[0].hyperlink.location, "'UID列表'!A1")

        quality = workbook["数据质量"]
        quality_items = {
            quality.cell(row_number, 3).value: quality.cell(row_number, 4)
            for row_number in range(1, quality.max_row + 1)
            if quality.cell(row_number, 3).value
        }
        self.assertEqual(quality_items["有效 UID 数"].value, 2)
        self.assertEqual(quality_items["缺失 mid 数"].value, 2)
        self.assertEqual(quality_items["非法 mid 数"].value, 4)
        self.assertEqual(quality_items["排除 mid 总数"].value, 6)

        fields = workbook["字段说明"]
        uid_field = next(
            row_values for row_values in fields.iter_rows()
            if row_values[2].value == "UID"
        )
        self.assertEqual(uid_field[1].value, "UID列表")
        self.assertEqual(uid_field[3].value, "uid")
        self.assertEqual(uid_field[4].value, "ID")
        self.assertIn("mid", uid_field[7].value)
        self.assertIn("缺失", uid_field[10].value)


if __name__ == "__main__":
    unittest.main()
