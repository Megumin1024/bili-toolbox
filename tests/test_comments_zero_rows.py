# -*- coding: utf-8 -*-
"""comments 0 行评论的防御回归：启动即取消等路径拿到空行集也要能出报告。

历史缺陷（comments 预算验收备忘的预存项）：analyze([]) 在 0 赞占比、大会员
占比、盖楼占比上直接除以 n=0；export_xlsx 的「时间跨度」说明行对空 ctime
序列调 min()/max() 抛 ValueError。修复口径：占比与跨度说明降级为 "-"，
与 coverage 的既有兜底同口径，analyze + export 全程不抛异常。
"""
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from tools.comments.core import analyze, export_xlsx


def zero_meta():
    return {"author": "演示UP", "title": "演示视频", "claimed_comment_count": 0}


def sample_row(**overrides):
    row = {
        "rpid": 1, "is_main": True, "mid": 100, "uname": "用户", "vip": False,
        "sex": "男", "level": 6, "message": "普通评论内容", "like": 0,
        "rcount": 0, "ctime": 1_700_000_000, "location": "IP属地：北京",
    }
    row.update(overrides)
    return row


class ZeroRowCommentsTests(unittest.TestCase):
    def test_analyze_zero_rows_does_not_raise_and_degrades_kpi(self):
        report, kpi = analyze([], zero_meta())
        # 占比口径降级为 "-"（与 coverage 的既有兜底同口径）
        self.assertEqual(kpi["zero_like_pct"], "-")
        self.assertEqual(kpi["vip_pct"], "-")
        self.assertEqual(kpi["build_pct"], "-")
        self.assertEqual(kpi["hours"], "0")
        self.assertEqual(kpi["fetched"], 0)
        self.assertIn("0赞占比 -", report)

    def test_analyze_declared_count_missing_none_zero_positive_and_invalid(self):
        cases = (
            ("missing", "接口未返回", None),
            (None, "接口未返回", None),
            (0, "0", 0),
            (3, "3", 3),
            (-1, "声明数量格式异常", None),
            (True, "声明数量格式异常", None),
            (False, "声明数量格式异常", None),
            (1.0, "声明数量格式异常", None),
            (1.5, "声明数量格式异常", None),
            (float("nan"), "声明数量格式异常", None),
            (float("inf"), "声明数量格式异常", None),
            (10 ** 16, "声明数量格式异常", None),
            ("1e3", "声明数量格式异常", None),
            ("9" * 5000, "声明数量格式异常", None),
        )
        for declared, shown, kpi_value in cases:
            with self.subTest(declared=declared):
                meta = {"author": "演示UP", "title": "演示视频"}
                if declared != "missing":
                    meta["claimed_comment_count"] = declared
                report, kpi = analyze([], meta)
                self.assertIn(f"B站声称评论总数：{shown}", report)
                self.assertEqual(kpi["claimed"], kpi_value)
                if shown == "声明数量格式异常":
                    self.assertEqual(kpi["coverage"], "不适用")

    def test_declared_count_boundaries_reopen_with_distinct_comments(self):
        cases = (
            ("missing", {}, "接口未返回", None),
            ("none", {"claimed_comment_count": None}, "接口未返回", None),
            ("zero", {"claimed_comment_count": 0}, "", 0),
            ("positive", {"claimed_comment_count": 3}, "", 3),
            ("invalid", {"claimed_comment_count": 1.5}, "声明数量格式异常", None),
        )
        for name, extra, note, expected in cases:
            with self.subTest(name=name), tempfile.TemporaryDirectory() as tmp:
                meta = dict(zero_meta(), **extra)
                if name == "missing":
                    meta.pop("claimed_comment_count")
                _report, kpi = analyze([], meta)
                out = Path(tmp) / f"{name}.xlsx"
                export_xlsx([], meta, kpi, out)
                workbook = load_workbook(out, data_only=False)
                quality = workbook["数据质量"]
                by_item = {
                    quality.cell(row, 3).value: row
                    for row in range(1, quality.max_row + 1)
                    if quality.cell(row, 3).value
                }
                declared = quality.cell(by_item["接口声称评论数"], 4)
                if expected is None:
                    self.assertIsNone(declared.value)
                    self.assertIn(note, declared.comment.text)
                else:
                    self.assertEqual(declared.value, expected)
                    self.assertEqual(declared.data_type, "n")
                if name == "invalid":
                    coverage = quality.cell(by_item["覆盖率"], 4)
                    self.assertIsNone(coverage.value)
                    self.assertIn("不适用", coverage.comment.text)
                workbook.close()

    def test_export_xlsx_zero_rows_writes_report_with_placeholder(self):
        _report, kpi = analyze([], zero_meta())
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "zero-row.xlsx"
            export_xlsx([], zero_meta(), kpi, out)
            self.assertTrue(out.is_file())
            ws = load_workbook(str(out))["统计概览"]
            for row in ws.iter_rows(values_only=True):
                if "时间跨度" in [str(v) for v in row if v is not None]:
                    self.assertEqual(row[2], 0)
                    self.assertEqual(row[3], "-")
                    break
            else:
                self.fail("统计概览中未找到「时间跨度」行")

    def test_nonzero_rows_export_percentages_as_native_numbers(self):
        """报告文案保持不变，Excel 占比改为 0~1 数值并使用百分比格式。"""
        rows = [
            sample_row(rpid=1, like=0, ctime=1_700_000_000),
            sample_row(rpid=2, mid=200, like=5, ctime=1_700_003_600),
        ]
        meta = zero_meta()
        meta["pub_ts"] = 1_699_900_000
        _report, kpi = analyze(rows, meta)
        self.assertEqual(kpi["zero_like_pct"], "50.0%")
        self.assertEqual(kpi["vip_pct"], "0.0%")
        self.assertEqual(kpi["build_pct"], "0.0%")
        self.assertEqual(kpi["hours"], "1")
        # 时间跨度说明行保持原有 "~" 连接格式
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "rows.xlsx"
            export_xlsx(rows, meta, kpi, out)
            workbook = load_workbook(str(out), data_only=False)
            ws = workbook["统计概览"]
            rows_by_name = {row[0].value: row for row in ws.iter_rows(min_col=2, max_col=4)}
            self.assertEqual(rows_by_name["0赞占比"][1].value, 0.5)
            self.assertEqual(rows_by_name["0赞占比"][1].number_format, "0.0%")
            self.assertIn("~", rows_by_name["时间跨度"][2].value)
            # 1_700_000_000 / 1_700_003_600 在 Asia/Shanghai 均为 11-15。
            self.assertIn("11-15 06:13 ~ 11-15 07:13", rows_by_name["时间跨度"][2].value)

            detail = workbook["全量评论"]
            detail_row = detail[2]
            self.assertEqual(detail.freeze_panes, "A2")
            self.assertEqual(detail.auto_filter.ref, "B1:N3")
            self.assertEqual(detail_row[2].value, "1")
            self.assertEqual(detail_row[2].data_type, "s")
            self.assertEqual(detail_row[2].number_format, "@")
            self.assertIsInstance(detail_row[10].value, int)
            self.assertEqual(detail_row[10].number_format, "#,##0")
            self.assertIsInstance(detail_row[12].value, datetime)
            self.assertEqual(detail_row[12].number_format, "yyyy-mm-dd hh:mm:ss")
            self.assertTrue(detail_row[9].alignment.wrap_text)
            top = workbook["点赞Top100"]
            self.assertEqual(top.freeze_panes, "A4")
            self.assertEqual(top.auto_filter.ref, "B3:G5")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
