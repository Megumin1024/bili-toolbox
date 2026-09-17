# -*- coding: utf-8 -*-
"""openpyxl 导出基建：样式、类型安全、文本安全与原子发布。

评论与采集两套 Excel 报表共用的部分；各工具的报表内容见 tools/*/core.py。
"""
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from enum import Enum
import math
import os
from pathlib import Path
import re
import tempfile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import Cell, ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_os_replace = os.replace

FONT_NAME, HEADER_BOLD = "Microsoft YaHei", False
PRIMARY, NEUTRAL_0, NEUTRAL_100 = "1B2A4A", "FFFFFF", "F7F7F5"
NEUTRAL_200, NEUTRAL_600, NEUTRAL_900 = "E9E9E8", "8C8A84", "37352F"

F_TITLE = Font(name=FONT_NAME, size=16, bold=HEADER_BOLD, color=PRIMARY)
F_HEADER = Font(name=FONT_NAME, size=11, bold=HEADER_BOLD, color="FFFFFF")
F_BODY = Font(name=FONT_NAME, size=11, color=NEUTRAL_900)
F_CAPTION = Font(name=FONT_NAME, size=9, color=NEUTRAL_600)
FILL_HEADER = PatternFill("solid", fgColor=PRIMARY)
B_HEADER = Border(bottom=Side(style="thin", color=NEUTRAL_200))
A_HEADER = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_TEXT = Alignment(horizontal="left", vertical="center")

ASIA_SHANGHAI = timezone(timedelta(hours=8), "Asia/Shanghai")
_UNIX_EPOCH_UTC = datetime(1970, 1, 1, tzinfo=timezone.utc)
_MAX_LOCAL_DATETIME = datetime(9999, 12, 31, 23, 59, 59,
                               tzinfo=ASIA_SHANGHAI)


class CellKind(str, Enum):
    """Excel 单元格的显式业务类型。

    导出器必须为业务数据选择类型；本模块不根据字段名或字符串内容猜测
    类型。FORMULA 不是公式白名单，只是内部可信调用方的显式信任标记；
    用户文本必须使用 TEXT。
    """

    ID = "id"
    INTEGER = "integer"
    DECIMAL = "decimal"
    PERCENT = "percent"
    DATETIME = "datetime"
    TEXT = "text"
    BOOLEAN = "boolean"
    MISSING = "missing"
    NOT_RETURNED = "not_returned"
    NOT_APPLICABLE = "not_applicable"
    FORMULA = "formula"


@dataclass(frozen=True)
class CellValue:
    """供 SheetWriter.append()/kv() 使用的显式类型值。"""

    value: object
    kind: CellKind = CellKind.TEXT
    number_format: str | None = None
    note: str | None = None


_INTEGER_RE = re.compile(r"^[+-]?\d+$")
_DECIMAL_RE = re.compile(
    r"^[+-]?(?:(?:\d+(?:\.\d*)?)|(?:\.\d+))(?:[eE][+-]?\d+)?$"
)
MAX_SIGNIFICANT_DIGITS = 15
MAX_INTEGER_VALUE = 999_999_999_999_999
MIN_INTEGER_VALUE = -MAX_INTEGER_VALUE
MAX_DECIMAL_ABS = Decimal("999999999999999")
MAX_NUMERIC_TEXT_LENGTH = 64
MAX_DECIMAL_SCALE = 308
MAX_ID_DIGITS = 64
MAX_ID_TEXT_LENGTH = 128
MAX_UNIX_SECONDS = int(
    (_MAX_LOCAL_DATETIME.astimezone(timezone.utc) - _UNIX_EPOCH_UTC).total_seconds()
)
MAX_UNIX_FRACTION_DIGITS = 6
MAX_UNIX_SIGNIFICANT_DIGITS = 18
_DEFAULT_NUMBER_FORMATS = {
    CellKind.ID: "@",
    CellKind.INTEGER: "#,##0",
    CellKind.DECIMAL: "#,##0.00",
    CellKind.PERCENT: "0.0%",
    CellKind.DATETIME: "yyyy-mm-dd hh:mm:ss",
}
_STATE_NOTES = {
    CellKind.MISSING: "字段缺失",
    CellKind.NOT_RETURNED: "接口未返回",
    CellKind.NOT_APPLICABLE: "不适用",
}


def clean(v):
    """清洗 Excel 非法字符（B站文本常含控制字符，直接写入会抛错）。"""
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    return v


def clean_text(value):
    """只清理 Excel 非法控制字符，不负责公式注入防护。"""
    if value is None:
        return None
    if not isinstance(value, str):
        raise TypeError("文本清洗只接受 str 或 None")
    return clean(value)


def safe_text(value):
    """兼容旧名称：只清理控制字符，不保证可直接 worksheet.append()。

    公式防护必须通过 ``SheetWriter.wc(..., kind=CellKind.TEXT)`` 或
    ``SheetWriter.append([cell_value(..., CellKind.TEXT)])`` 完成。直接把本函数
    返回的字符串交给 openpyxl，仍可能被识别为公式；这是刻意保留的明确契约。
    """
    return clean_text(value)


def _needs_formula_quote(value):
    return isinstance(value, str) and value.startswith(("=", "+", "-", "@"))


def cell_value(value, kind=CellKind.TEXT, *, number_format=None, note=None):
    """创建一个带显式类型的值描述，不进行字段名推断。"""
    return CellValue(value, _coerce_kind(kind), number_format, note)


def checked_cell_value(value, kind=CellKind.TEXT, *, number_format=None,
                       note="格式异常"):
    """创建并预校验单元格值；外部坏值降级为空的缺失单元格。

    该入口只捕获类型/数值格式错误。保存、磁盘和发布错误仍由调用方看到，
    不会被导出器的单条数据容错吞掉。返回的 ``CellValue`` 已经携带清洗后的
    payload，因而业务导出器可以安全地继续写后续记录。
    """
    normalized_kind = _coerce_kind(kind)
    try:
        _, payload, normalized_format, normalized_note = _cell_payload(
            value, normalized_kind, number_format, note=None
        )
    except (TypeError, ValueError):
        return CellValue(None, CellKind.MISSING, note=note)
    return CellValue(payload, normalized_kind, normalized_format, normalized_note)


def unix_seconds_cell_value(value, *, note="格式异常"):
    """把 Unix 秒安全转换为无时区 Asia/Shanghai Excel datetime 单元格。"""
    try:
        converted = unix_seconds_to_excel_datetime(value)
    except (TypeError, ValueError):
        return CellValue(None, CellKind.MISSING, note=note)
    if converted is None:
        return CellValue(None, CellKind.MISSING, note=note)
    return CellValue(converted, CellKind.DATETIME,
                     _DEFAULT_NUMBER_FORMATS[CellKind.DATETIME])


def _coerce_kind(kind):
    if isinstance(kind, CellKind):
        return kind
    try:
        return CellKind(kind)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"未知 Excel 单元格类型：{kind!r}") from exc


def _id_text(value):
    if value is None:
        return None
    if isinstance(value, bool):
        raise TypeError("ID 不接受 bool")
    if isinstance(value, int):
        if abs(value) >= 10 ** MAX_ID_DIGITS:
            raise ValueError(f"ID 超出允许的 {MAX_ID_DIGITS} 位范围")
        return str(value)
    if isinstance(value, str):
        text = clean(value)
        if len(text) > MAX_ID_TEXT_LENGTH:
            raise ValueError(f"ID 文本超出允许的 {MAX_ID_TEXT_LENGTH} 个字符")
        numeric_text = text.strip()
        if _INTEGER_RE.fullmatch(numeric_text):
            digits = numeric_text.lstrip("+-")
            if len(digits) > MAX_ID_DIGITS:
                raise ValueError(f"ID 超出允许的 {MAX_ID_DIGITS} 位范围")
        return text
    raise TypeError("ID 必须是整数或字符串，禁止先转浮点")


def _integer_value(value):
    if value is None:
        return None
    if isinstance(value, bool):
        raise TypeError("整数不接受 bool")
    if isinstance(value, int):
        if value < MIN_INTEGER_VALUE or value > MAX_INTEGER_VALUE:
            raise ValueError(
                f"整数超出允许范围 [{MIN_INTEGER_VALUE}, {MAX_INTEGER_VALUE}]"
            )
        return value
    if isinstance(value, str):
        if len(value) > MAX_NUMERIC_TEXT_LENGTH:
            raise ValueError("整数数字文本过长")
        text = value.strip()
        if not _INTEGER_RE.fullmatch(text):
            raise ValueError("整数数值格式非法")
        digits = text.lstrip("+-").lstrip("0") or "0"
        if len(digits) > MAX_SIGNIFICANT_DIGITS:
            raise ValueError("整数有效数字超过 15 位")
        parsed = int(text)
        if parsed < MIN_INTEGER_VALUE or parsed > MAX_INTEGER_VALUE:
            raise ValueError(
                f"整数超出允许范围 [{MIN_INTEGER_VALUE}, {MAX_INTEGER_VALUE}]"
            )
        return parsed
    if isinstance(value, Decimal):
        parsed = _decimal_from_value(
            value, "整数", max_abs=MAX_DECIMAL_ABS,
        )
        if parsed != parsed.to_integral_value():
            raise ValueError("整数不接受小数")
        return int(parsed)
    if isinstance(value, float):
        raise TypeError("整数不接受 float，请显式传入整数")
    raise TypeError("整数类型必须是 int 或十进制整数字符串")


def _decimal_value(value):
    if value is None:
        return None
    parsed = _decimal_from_value(value, "小数", max_abs=MAX_DECIMAL_ABS)
    return float(parsed)


def _effective_digit_count(value):
    digits = value.as_tuple().digits
    count = len(digits)
    while count > 1 and digits[count - 1] == 0:
        count -= 1
    return count


def _decimal_from_value(value, name, *, max_abs,
                        max_scale=MAX_DECIMAL_SCALE,
                        max_significant_digits=MAX_SIGNIFICANT_DIGITS):
    """在任何 Decimal/float/int 转换前完成边界检查。"""
    if isinstance(value, bool):
        raise TypeError(f"{name} 不接受 bool")

    if isinstance(value, int):
        integer_limit = int(max_abs)
        if value < -integer_limit or value > integer_limit:
            raise ValueError(f"{name} 超出允许范围")
        parsed = Decimal(value)
    elif isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError(f"{name} 不接受 NaN 或 Infinity")
        text = repr(value)
        try:
            parsed = Decimal(text)
        except (InvalidOperation, ValueError, TypeError, OverflowError) as exc:
            raise ValueError(f"{name} 数值格式非法") from exc
    elif isinstance(value, Decimal):
        parsed = value
    elif isinstance(value, str):
        if len(value) > MAX_NUMERIC_TEXT_LENGTH:
            raise ValueError(f"{name} 数字文本过长")
        text = value.strip()
        if not _DECIMAL_RE.fullmatch(text):
            raise ValueError(f"{name} 数值格式非法")
        try:
            parsed = Decimal(text)
        except (InvalidOperation, ValueError, TypeError, OverflowError) as exc:
            raise ValueError(f"{name} 数值格式非法") from exc
    else:
        raise TypeError(f"{name} 必须是 int、float、Decimal 或十进制文本")

    if not parsed.is_finite():
        raise ValueError(f"{name} 不接受 NaN 或 Infinity")
    exponent = parsed.as_tuple().exponent
    if _effective_digit_count(parsed) > max_significant_digits:
        raise ValueError(f"{name} 有效数字超过 {max_significant_digits} 位")
    if exponent < -max_scale or exponent > max_scale:
        raise ValueError(f"{name} 指数超出允许范围")
    try:
        too_large = abs(parsed) > max_abs
    except (InvalidOperation, TypeError, OverflowError) as exc:
        raise ValueError(f"{name} 数值范围非法") from exc
    if too_large:
        raise ValueError(f"{name} 超出允许范围")
    return parsed


def _datetime_value(value):
    if value is None:
        return None
    if not isinstance(value, datetime):
        raise TypeError("datetime 类型必须传入 datetime")
    try:
        if value.tzinfo is None:
            return value
        return value.astimezone(ASIA_SHANGHAI).replace(tzinfo=None)
    except (TypeError, ValueError, OverflowError, OSError) as exc:
        raise ValueError("datetime 无法转换为 Asia/Shanghai 时间") from exc


def _unix_datetime_from_parts(seconds, microseconds):
    try:
        utc_value = _UNIX_EPOCH_UTC + timedelta(
            seconds=seconds, microseconds=microseconds
        )
        return utc_value.astimezone(ASIA_SHANGHAI).replace(tzinfo=None)
    except (TypeError, ValueError, OverflowError, OSError) as exc:
        raise ValueError("Unix 秒超出 datetime 范围") from exc


def unix_seconds_to_excel_datetime(value):
    """将明确的 Unix 秒转换为 Asia/Shanghai 的无时区 Excel datetime。"""
    if value is None:
        return None
    try:
        seconds = _decimal_from_value(
            value,
            "Unix 秒",
            max_abs=Decimal(MAX_UNIX_SECONDS),
            max_scale=MAX_DECIMAL_SCALE,
            max_significant_digits=MAX_UNIX_SIGNIFICANT_DIGITS,
        )
        if seconds < 0:
            raise ValueError("Unix 秒不接受负值")
        if -seconds.as_tuple().exponent > MAX_UNIX_FRACTION_DIGITS:
            raise ValueError("Unix 秒最多支持 6 位小数")
        whole_seconds = int(seconds)
        fraction = seconds - Decimal(whole_seconds)
        microseconds = int(fraction * Decimal(1_000_000))
        return _unix_datetime_from_parts(whole_seconds, microseconds)
    except (TypeError, ValueError, OverflowError, OSError) as exc:
        if isinstance(exc, ValueError):
            raise
        raise ValueError("Unix 秒格式或范围非法") from exc


def local_text_to_excel_datetime(value, fmt="%Y-%m-%d %H:%M:%S"):
    """按调用方明确给出的格式，把 Asia/Shanghai 本地时间文本转为 Excel datetime。"""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return _datetime_value(value)
    if not isinstance(value, str):
        raise TypeError("本地时间文本必须是 str 或 datetime")
    try:
        parsed = datetime.strptime(value, fmt).replace(
            tzinfo=ASIA_SHANGHAI
        )
        return parsed.astimezone(ASIA_SHANGHAI).replace(tzinfo=None)
    except (TypeError, ValueError, OverflowError, OSError) as exc:
        raise ValueError("本地时间格式或范围非法") from exc


def _cell_payload(value, kind, number_format=None, note=None):
    kind = _coerce_kind(kind)
    if kind in (CellKind.MISSING, CellKind.NOT_RETURNED, CellKind.NOT_APPLICABLE):
        payload = None
        note = note or _STATE_NOTES[kind]
    elif kind is CellKind.ID:
        payload = clean_text(_id_text(value))
    elif kind is CellKind.INTEGER:
        payload = _integer_value(value)
    elif kind is CellKind.DECIMAL or kind is CellKind.PERCENT:
        payload = _decimal_value(value)
    elif kind is CellKind.DATETIME:
        payload = _datetime_value(value)
    elif kind is CellKind.BOOLEAN:
        if value is not None and not isinstance(value, bool):
            raise TypeError("布尔类型必须是 bool")
        payload = value
    elif kind is CellKind.FORMULA:
        # 这里没有来源认证或公式白名单；调用方必须保证 value 来自内部可信构造。
        if not isinstance(value, str) or not value.startswith("="):
            raise ValueError("内部公式必须是以 = 开头的显式字符串")
        payload = clean(value)
    else:  # TEXT；兼容旧调用方的非字符串值，迁移阶段将显式声明类型。
        payload = clean_text(value) if isinstance(value, str) else clean(value)
    return kind, payload, number_format or _DEFAULT_NUMBER_FORMATS.get(kind), note


def new_workbook(creator="BiliToolbox"):
    wb = Workbook(write_only=True)
    wb.properties.creator = creator
    return wb


class SheetWriter:
    """write_only 模式的表辅助。ws 属性随切页更新，wc() 写入当前页。"""

    def __init__(self, wb):
        self.wb = wb
        self.ws = None

    def wc(self, value, font=None, fill=None, align=None, border=None, *,
           kind=CellKind.TEXT, number_format=None, note=None):
        kind, payload, fmt, note = _cell_payload(value, kind, number_format, note)
        c = WriteOnlyCell(self.ws, value=payload)
        if kind in (CellKind.ID, CellKind.TEXT) and _needs_formula_quote(payload):
            c.data_type = "s"
            c.quotePrefix = True
        c.font = font or F_BODY
        if fill:
            c.fill = fill
        c.alignment = align or A_TEXT
        if border:
            c.border = border
        if fmt:
            c.number_format = fmt
        if note:
            c.comment = Comment(clean(note), "BiliToolbox")
        return c

    def append(self, values):
        """追加一行；CellValue 用于要求显式类型的业务值。"""
        if self.ws is None:
            raise RuntimeError("SheetWriter.ws 尚未设置")
        cells = []
        for value in values:
            if isinstance(value, CellValue):
                cells.append(self.wc(value.value, kind=value.kind,
                                      number_format=value.number_format,
                                      note=value.note))
            elif isinstance(value, Cell) or value is None:
                cells.append(value)
            else:
                cells.append(self.wc(value))
        self.ws.append(cells)

    def title_row(self, ws, text, width):
        self.ws = ws
        self.append([None] + [self.wc(text, font=F_TITLE)] + [None] * (width - 1))
        ws.row_dimensions[1].height = 15
        ws.row_dimensions[2].height = 32
        self.append([None])

    def header_row(self, ws, headers):
        self.ws = ws
        self.append([None] + [self.wc(h, font=F_HEADER, fill=FILL_HEADER,
                                      align=A_HEADER, border=B_HEADER)
                              for h in headers])

    def kv(self, ws, pairs, width=4):
        self.ws = ws
        self.header_row(ws, ["指标", "数值", "说明"])
        for k, v, note in pairs:
            if isinstance(v, CellValue):
                value_cell = self.wc(v.value, kind=v.kind,
                                     number_format=v.number_format,
                                     note=v.note)
            else:
                value_cell = self.wc(v)
            self.append([None, self.wc(k), value_cell, self.wc(note)])
        self.append([None])


def save_workbook_atomic(workbook, path, before_replace=None):
    """同目录临时保存并原子发布工作簿。

    目标文件只在 save 成功且 replace 成功时被替换。失败路径不触碰旧目标，
    并尽力清理临时文件。before_replace 可用于 relation_analysis 的取消检查。
    """
    target = Path(path)
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(prefix=f".{target.stem}-",
                                         suffix=".xlsx.tmp",
                                         dir=str(target.parent), delete=False) as handle:
            temp_path = Path(handle.name)
        workbook.save(str(temp_path))
        if before_replace is not None:
            before_replace()
        _os_replace(str(temp_path), str(target))
    finally:
        if temp_path is not None and temp_path.exists():
            temp_path.unlink()
