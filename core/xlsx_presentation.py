# -*- coding: utf-8 -*-
"""First-party XLSX presentation helpers.

The helpers in this module are deliberately declarative: callers provide the
table coordinates and the semantic state of quality items.  This module never
scans a worksheet to infer business meaning and never reloads a workbook after
it has been written.
"""

from __future__ import annotations

from dataclasses import dataclass
import re
from typing import Any, Iterable, Sequence
from urllib.parse import urlsplit, urlunsplit

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.utils import get_column_letter

from .xlsx import (
    A_HEADER,
    A_TEXT,
    CellKind,
    CellValue,
    FILL_HEADER,
    F_HEADER,
    F_BODY,
    SheetWriter,
    clean_text,
)


MAX_SAFE_URL_LENGTH = 2048
DEFAULT_ROW_HEIGHT = 32.0

# These are presentation states, not a text classifier.  The caller selects a
# state from structured quality data; the fallback below only handles the
# existing explicit CellKind missing/not-returned states.
QUALITY_WARNING = "warning"
QUALITY_ERROR = "error"
QUALITY_STOP = "stop"
_QUALITY_STATES = frozenset({QUALITY_WARNING, QUALITY_ERROR, QUALITY_STOP})

WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
ERROR_FILL = PatternFill("solid", fgColor="F4CCCC")
STOP_FILL = PatternFill("solid", fgColor="D9EAF7")

A_WRAP_TEXT = Alignment(horizontal="left", vertical="top", wrap_text=True)

_CONTROL_RE = re.compile(r"[\x00-\x1f\x7f]")
_SENSITIVE_URL_RE = re.compile(
    r"(?i)(?:cookie|token|authorization|sessdata|sendkey|password|passwd|"
    r"secret|proxy|bili[_-]?jct|bearer|api[_-]?key)"
)


@dataclass(frozen=True)
class TableLayout:
    """Coordinates for one contiguous business table."""

    header_row: int
    first_col: int
    last_col: int
    default_row_height: float = DEFAULT_ROW_HEIGHT

    def __post_init__(self) -> None:
        if self.header_row < 1:
            raise ValueError("表头行必须从 1 开始")
        if self.first_col < 1 or self.last_col < self.first_col:
            raise ValueError("表格列范围非法")
        if self.default_row_height <= 0:
            raise ValueError("默认行高必须为正数")


def configure_table(worksheet: Any, layout: TableLayout) -> None:
    """Configure a table before the worksheet receives its first row."""
    # A-column freezing freezes rows only; it intentionally never freezes a
    # business column.  WriteOnlyWorksheet serializes freeze panes at close,
    # but only if this happens before its first append.
    worksheet.freeze_panes = f"A{layout.header_row + 1}"
    worksheet.sheet_format.defaultRowHeight = layout.default_row_height


def finish_table(worksheet: Any, layout: TableLayout, data_rows: int) -> str:
    """Set a finite auto-filter range and return it for assertions."""
    if isinstance(data_rows, bool) or not isinstance(data_rows, int) or data_rows < 0:
        raise ValueError("业务行数必须是非负整数")
    last_row = layout.header_row + data_rows
    ref = (
        f"{get_column_letter(layout.first_col)}{layout.header_row}:"
        f"{get_column_letter(layout.last_col)}{last_row}"
    )
    worksheet.auto_filter.ref = ref
    return ref


def _cell_value_args(value: CellValue) -> dict[str, Any]:
    return {
        "kind": value.kind,
        "number_format": value.number_format,
        "note": value.note,
    }


def styled_cell(
    writer: SheetWriter,
    value: object,
    *,
    align: Alignment,
    font=None,
    fill=None,
    border=None,
) -> Any:
    """Create a styled cell while preserving an explicit ``CellValue``."""
    if isinstance(value, CellValue):
        return writer.wc(
            value.value,
            font=font,
            fill=fill,
            align=align,
            border=border,
            **_cell_value_args(value),
        )
    return writer.wc(value, font=font, fill=fill, align=align, border=border)


def wrap_cell(writer: SheetWriter, value: object, **kwargs: Any) -> Any:
    """Create one explicitly declared long-text cell with shared alignment."""
    return styled_cell(writer, value, align=A_WRAP_TEXT, **kwargs)


def quote_sheet_name(sheet_name: str) -> str:
    """Return an Excel internal-reference-safe sheet-name literal."""
    if not isinstance(sheet_name, str) or not sheet_name:
        raise ValueError("工作表名称必须是非空文本")
    return "'" + sheet_name.replace("'", "''") + "'"


def internal_location(sheet_name: str, cell: str = "A1") -> str:
    """Build a native Excel hyperlink location, never an external URL."""
    if not isinstance(cell, str) or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", cell):
        raise ValueError("内部链接单元格地址非法")
    return f"{quote_sheet_name(sheet_name)}!{cell}"


def _hyperlink_cell(writer: SheetWriter, display: str, hyperlink: Hyperlink) -> Any:
    cell = writer.wc(display, kind=CellKind.TEXT)
    # WriteOnlyCell has no row/column until Worksheet._values_to_row receives
    # it.  Bypass the normal setter here; that writer later fills Hyperlink.ref
    # with the real coordinate before serializing the row.
    cell._hyperlink = hyperlink
    return cell


def append_sheet_directory(
    writer: SheetWriter,
    worksheet: Any,
    entries: Sequence[tuple[str, str]],
    *,
    prefix_columns: int = 1,
) -> None:
    """Append a compact directory after existing content.

    ``prefix_columns=1`` matches the first-party exporters that retain their
    historical blank A column.  The generated check/repair reports pass zero
    because their business tables start in column A.
    """
    if prefix_columns < 0:
        raise ValueError("目录前缀列数不能为负数")
    normalized = []
    for display, target_sheet in entries:
        if not isinstance(display, str) or not display:
            raise ValueError("目录显示值必须是非空文本")
        location = internal_location(target_sheet)
        normalized.append((display, location))
    if not normalized:
        return

    writer.ws = worksheet
    prefix = [None] * prefix_columns
    # Keep the separator wide enough for existing overview readers that
    # inspect the first two columns of every row. Empty strings survive
    # write-only serialization while remaining visually blank in Excel.
    writer.append(prefix + ["", ""])
    header_cells = [
        writer.wc("工作表目录", font=F_HEADER, fill=FILL_HEADER,
                  align=A_HEADER),
        writer.wc("目标地址", font=F_HEADER, fill=FILL_HEADER,
                  align=A_HEADER),
    ]
    writer.append(prefix + header_cells)
    for display, location in normalized:
        display_cell = _hyperlink_cell(
            writer,
            display,
            Hyperlink(ref="", location=location, display=display),
        )
        writer.append(prefix + [display_cell, writer.wc(location, align=A_TEXT)])


def _has_control(value: str) -> bool:
    return bool(_CONTROL_RE.search(value))


def _sensitive_url(value: str) -> bool:
    return bool(_SENSITIVE_URL_RE.search(value))


def safe_bilibili_url(value: object) -> str | None:
    """Validate one already-selected URL field without any network access."""
    if not isinstance(value, str):
        return None
    if _has_control(value) or len(value) > MAX_SAFE_URL_LENGTH:
        return None
    text = value.strip()
    if text.startswith("//"):
        text = "https:" + text
    if _sensitive_url(text):
        return None
    try:
        parts = urlsplit(text)
        scheme = parts.scheme.lower()
        host = parts.hostname
        port = parts.port
        username = parts.username
        password = parts.password
    except (TypeError, ValueError):
        return None
    if scheme != "https" or not host or username is not None or password is not None:
        return None
    host = host.lower()
    if not (host == "bilibili.com" or host.endswith(".bilibili.com")):
        return None
    if port not in (None, 443):
        return None
    netloc = host + (f":{port}" if port is not None else "")
    return urlunsplit(("https", netloc, parts.path, parts.query, parts.fragment))


def _safe_visible_url(value: object) -> str:
    if not isinstance(value, str):
        return ""
    if _has_control(value) or len(value) > MAX_SAFE_URL_LENGTH or _sensitive_url(value):
        return "[已隐藏不安全链接]"
    return clean_text(value)


def external_link_cell(writer: SheetWriter, value: object) -> Any:
    """Create a visible text cell and link it only when validation succeeds."""
    target = safe_bilibili_url(value)
    visible = _safe_visible_url(value)
    cell = writer.wc(visible, kind=CellKind.TEXT)
    if target is not None:
        cell._hyperlink = Hyperlink(ref="", target=target, display=visible)
    return cell


def _quality_predicate(row: int, value: object) -> str | None:
    kind = getattr(value, "kind", None)
    value_col = f"$D{row}"
    if kind is CellKind.BOOLEAN:
        return f"{value_col}=TRUE"
    if kind in (CellKind.INTEGER, CellKind.DECIMAL, CellKind.PERCENT):
        return f"{value_col}>0"
    if kind in (CellKind.MISSING, CellKind.NOT_RETURNED):
        return f"ISBLANK({value_col})"
    return None


def apply_quality_conditional_formats(
    worksheet: Any,
    quality_start_row: int,
    items: Iterable[object],
) -> None:
    """Apply narrow rules to explicit quality-item rows only."""
    if quality_start_row < 1:
        raise ValueError("质量项目起始行非法")
    fills = {
        QUALITY_WARNING: WARNING_FILL,
        QUALITY_ERROR: ERROR_FILL,
        QUALITY_STOP: STOP_FILL,
    }
    for offset, item in enumerate(items):
        value = getattr(item, "value", None)
        state = getattr(item, "presentation_state", None)
        if state is None and getattr(value, "kind", None) in (
            CellKind.MISSING,
            CellKind.NOT_RETURNED,
        ):
            state = QUALITY_WARNING
        if state not in _QUALITY_STATES:
            continue
        formula = _quality_predicate(quality_start_row + offset, value)
        if formula is None:
            continue
        row = quality_start_row + offset
        worksheet.conditional_formatting.add(
            f"B{row}:F{row}",
            FormulaRule(formula=[formula], fill=fills[state]),
        )
