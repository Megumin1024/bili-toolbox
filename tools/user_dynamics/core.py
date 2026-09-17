# -*- coding: utf-8 -*-
"""按 UID 抓取 B 站用户动态：解析、翻页、导出。

网络访问全部经注入的 fetch（默认走 core.session 四层风控通道，自动过全局闸门）。
本模块不自己建连接，因此翻页 / 解析 / 空结果重试都能离线测试。

两条来自实测的硬约束（依据见 core.wbi 与 core.transport 的模块注释）：

1. 动态接口必须带 WBI 签名，否则返回 code=0 但 items 为空——**静默清空，不报错**；
2. 即便带了签名，"空结果"仍会间歇出现，而且与"真的没有更多"的响应完全一致
   （都是 code=0 + items=[] + has_more=false）。所以空结果必须退避重试，
   且第一页重试耗尽时**不能**下结论说"该用户没有动态"——那可能是在撒谎。
"""
from __future__ import annotations

import json
import os
import re
import time
from datetime import datetime
from pathlib import Path

from openpyxl.utils import get_column_letter

from core import wbi, xlsx as xlsx_mod
from core.budget import BudgetExhaustedError
from core.cancel import TaskCancelledError, wait as cancel_wait
from core.xlsx_metadata import (
    FieldDefinition,
    QualityItem,
    make_metadata,
    primary_key_quality,
    write_metadata_sheets,
)
from core.xlsx_presentation import (
    TableLayout,
    append_sheet_directory,
    configure_table,
    external_link_cell,
    finish_table,
    wrap_cell,
)

DYNAMIC_FEED_URL = "https://api.bilibili.com/x/polymer/web-dynamic/v1/feed/space"

# 每页实测 12 条。仅用于给用户估算"这大概是多少条"。
PAGE_SIZE = 12
DEFAULT_MAX_PAGES = 5
MAX_PAGES_LIMIT = 100
DEFAULT_EMPTY_RETRIES = 3
# 页间业务间隔（秒）。动态接口在实测中比评论区更敏感，默认给得比评论抓取宽。
DEFAULT_SLEEP = 1.0
EMPTY_RETRY_BASE = 2.0
EMPTY_RETRY_CAP = 30.0

TYPE_LABELS = {
    "DYNAMIC_TYPE_AV": "视频投稿",
    "DYNAMIC_TYPE_DRAW": "图文",
    "DYNAMIC_TYPE_WORD": "纯文字",
    "DYNAMIC_TYPE_FORWARD": "转发",
    "DYNAMIC_TYPE_ARTICLE": "专栏",
    "DYNAMIC_TYPE_MUSIC": "音频",
}


class DynamicsUnavailable(RuntimeError):
    """第一页反复取空：无法区分"该用户没有动态"与"被限流降级"。"""


def parse_uid(text):
    """UID 纯数字，或 space.bilibili.com/<uid> 链接（含 /dynamic 等子路径）。"""
    s = (text or "").strip().strip("\"'“”")
    if not s:
        raise ValueError("请输入 UID 或用户空间链接")
    m = re.search(r"space\.bilibili\.com/(\d+)", s, re.I)
    if m:
        return int(m.group(1))
    if re.fullmatch(r"\d{1,20}", s):
        return int(s)
    raise ValueError(f"无法识别的 UID：{s[:60]}")


# ---------- 解析（纯函数，可离线测试） ----------

def to_int(value, default=0):
    """B 站的计数/时间戳常以**字符串**下发，直接当整数用会炸。"""
    if isinstance(value, bool):
        return default
    if isinstance(value, int):
        return value
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return default


def _stat_count(module_stat, key):
    node = (module_stat or {}).get(key)
    return to_int(node.get("count")) if isinstance(node, dict) else 0


def _abs_url(url):
    """动态里的跳转链接是协议相对地址（//www.bilibili.com/...），补上 https:。"""
    u = str(url or "").strip()
    if not u:
        return ""
    if u.startswith("//"):
        return "https:" + u
    return u if u.startswith(("http://", "https://")) else ""


def normalize_type(type_raw):
    return TYPE_LABELS.get(str(type_raw or ""), str(type_raw or ""))


def parse_item(item):
    """一条动态 → 扁平行。字段缺失一律降级为空/0，绝不抛异常。"""
    item = item or {}
    modules = item.get("modules") or {}
    author = modules.get("module_author") or {}
    dynamic = modules.get("module_dynamic") or {}
    major = dynamic.get("major") or {}
    archive = major.get("archive") or {}
    opus = major.get("opus") or {}
    desc = (dynamic.get("desc") or {}).get("text") or ""
    orig = item.get("orig") or {}
    orig_author = ((orig.get("modules") or {}).get("module_author") or {})

    # 正文来源随类型而变：转发/纯文字在 desc，投稿在 archive.title，
    # 图文在 opus.summary.text。按优先级取第一个非空的。
    text = (desc or archive.get("title")
            or (opus.get("summary") or {}).get("text") or opus.get("title") or "")
    pub_ts = to_int(author.get("pub_ts"))
    type_raw = str(item.get("type") or "")

    return {
        "id": str(item.get("id_str") or ""),
        "type": normalize_type(type_raw),
        "type_raw": type_raw,
        "mid": to_int(author.get("mid")),
        "author": author.get("name") or "",
        "pub_ts": pub_ts,
        "time": (datetime.fromtimestamp(pub_ts).strftime("%Y-%m-%d %H:%M")
                 if pub_ts else ""),
        "text": " ".join(str(text).split()),
        "major_type": str(major.get("type") or ""),
        "bvid": archive.get("bvid") or "",
        "url": _abs_url(archive.get("jump_url") or opus.get("jump_url")),
        "forward": _stat_count(modules.get("module_stat"), "forward"),
        "comment": _stat_count(modules.get("module_stat"), "comment"),
        "like": _stat_count(modules.get("module_stat"), "like"),
        "coin": _stat_count(modules.get("module_stat"), "coin"),
        "favorite": _stat_count(modules.get("module_stat"), "favorite"),
        "is_forward": type_raw == "DYNAMIC_TYPE_FORWARD",
        "orig_id": str(orig.get("id_str") or ""),
        "orig_author": orig_author.get("name") or "",
    }


def parse_page(payload):
    """响应体 → (items, has_more, next_offset)。结构不符预期时返回空，不抛。"""
    data = (payload or {}).get("data") or {}
    items = data.get("items") or []
    if not isinstance(items, list):
        items = []
    return items, bool(data.get("has_more")), str(data.get("offset") or "")


# ---------- 抓取 ----------

class DynamicsCrawler:
    """按 uid 翻页抓取动态。可取消；空结果会退避重试。"""

    def __init__(self, uid, out_dir, max_pages=DEFAULT_MAX_PAGES,
                 sleep=DEFAULT_SLEEP, empty_retries=DEFAULT_EMPTY_RETRIES,
                 progress=None, cancel=None, fetch=None, sleeper=None,
                 key_cache=None, budget=None):
        self.uid = int(uid)
        self.out_dir = Path(out_dir)
        self.out_dir.mkdir(parents=True, exist_ok=True)
        self.max_pages = max(1, min(MAX_PAGES_LIMIT, int(max_pages)))
        self.sleep = max(0.0, float(sleep))
        self.empty_retries = max(1, int(empty_retries))
        self.cancel = cancel or (lambda: False)
        self._sleep = sleeper or time.sleep
        self._fetch = fetch or self._session_fetch
        # 取密钥也要能取消：首次请求前会先取一次 nav，那条路同样会退避重试。
        self._keys = key_cache or wbi.WbiKeyCache(cancel=self.cancel)
        self._progress = progress or (lambda **kw: None)
        # 任务预算（core.budget.TaskBudget），每次任务新建；None = 无预算。
        # 业务请求的强制点在 BiliClient._request（fetch 透传 budget），
        # 本层的职责是页循环边界查 expired() 与落盘后 observe_records(n)。
        self.budget = budget
        self.out_path = self.out_dir / f"dynamics_{self.uid}.jsonl"
        self.rows = []
        self.stats = {"uid": self.uid, "pages": 0, "requests": 0,
                      "rows": 0, "truncated": False, "cancelled": False,
                      "candidate_records": 0, "duplicate_rows": 0,
                      "parse_failures": 0, "invalid_time": 0,
                      "missing_id": 0, "invalid_id": 0,
                      "dedup_discarded": 0, "remaining_id_conflicts": 0}

    def _p(self, **kw):
        self._progress(**kw)

    def _cancelled(self):
        return bool(self.cancel())

    def _session_fetch(self, url):
        """经 core.session 的四层风控通道发请求。

        cancel 必须透传：HTTP 层内部的退避等待受 TOTAL_WAIT_BUDGET 约束，拿不到
        取消谓词就只能睡满预算。爬取器自己那层取消检查在这之下够不着，
        不透传的话用户按取消最坏要等 90 秒。budget 同理透传：业务请求数的
        强制点在 HTTP 层入口；无预算时不附加该 kwarg，缺省调用与旧路径一致。
        """
        from core import session          # 延迟导入：本模块要能离线单独导入
        kwargs = {"cancel": self.cancel}
        if self.budget is not None:
            kwargs["budget"] = self.budget
        return session.http_get_json(url, **kwargs)

    def _load_page(self, offset):
        """取一页。空结果退避重试；返回 (items, has_more, next_offset)。"""
        params = {"host_mid": self.uid, "offset": offset, "platform": "web",
                  "timezone_offset": -480, "features": "itemOpusStyle"}
        for attempt in range(self.empty_retries):
            if self._cancelled():
                raise TaskCancelledError()
            img_key, sub_key = self._keys.get()
            url = wbi.signed_url(DYNAMIC_FEED_URL, params, img_key, sub_key)
            self.stats["requests"] += 1
            items, has_more, next_offset = parse_page(self._fetch(url))
            if items:
                return items, has_more, next_offset
            if attempt + 1 < self.empty_retries:
                wait = min(EMPTY_RETRY_BASE * (2 ** attempt), EMPTY_RETRY_CAP)
                self._p(level="warn",
                        text=f"第 {self.stats['pages'] + 1} 页返回空，"
                             f"{wait:.0f}s 后重试（{attempt + 1}/{self.empty_retries}）")
                if not cancel_wait(wait, self.cancel, sleep=self._sleep):
                    raise TaskCancelledError()
        # 连续空：密钥可能已经轮换，丢掉缓存，下次取新的
        self._keys.invalidate()
        return [], False, ""

    def _pace(self, seconds):
        if seconds <= 0:
            return True
        return cancel_wait(seconds, self.cancel, sleep=self._sleep)

    def _flush(self):
        """整表重写 + 原子替换。行数最多几千，不必做增量追加。"""
        tmp = self.out_path.with_suffix(".jsonl.tmp")
        with open(tmp, "w", encoding="utf-8") as fh:
            for row in self.rows:
                fh.write(json.dumps(row, ensure_ascii=False) + "\n")
        os.replace(tmp, self.out_path)

    def crawl(self):
        """抓取到 max_pages / 没有更多 / 用户取消 / 预算到限为止。返回 stats。"""
        seen = set()
        offset = ""
        self.rows = []
        try:
            while self.stats["pages"] < self.max_pages:
                if self._cancelled():
                    raise TaskCancelledError()
                # 预算边界：到限按正常完成收尾（不是失败、也不是取消）。
                # 取消检查在预算检查之前，cancel 与预算同时到期时按取消语义。
                if self.budget is not None and self.budget.expired():
                    self.stats["stopped_reason"] = "budget_reached"
                    self._p(level="warn",
                            text=f"已达预算上限（{self.budget.reason()}），"
                                 f"安全停止，保留已抓到的 {len(self.rows)} 条动态")
                    break
                items, has_more, next_offset = self._load_page(offset)
                if not items:
                    if self.stats["pages"] == 0:
                        raise DynamicsUnavailable(
                            f"连续 {self.empty_retries} 次取到的都是空结果。"
                            "该用户没有动态、与被限流降级，两者响应完全相同，"
                            "无法区分——请稍后重试，或换一个 UID 验证。")
                    self.stats["truncated"] = True
                    self._p(level="warn", text="后续页返回空，可能在翻页途中被截断")
                    break
                new = 0
                for item in items:
                    self.stats["candidate_records"] += 1
                    if not isinstance(item, dict):
                        self.stats["parse_failures"] += 1
                        continue
                    row = parse_item(item)
                    if row["id"] and row["id"] in seen:
                        self.stats["duplicate_rows"] += 1
                        self.stats["dedup_discarded"] += 1
                        continue
                    if not row["id"]:
                        self.stats["parse_failures"] += 1
                        if item.get("id_str") in (None, ""):
                            self.stats["missing_id"] += 1
                        else:
                            self.stats["invalid_id"] += 1
                        continue
                    raw_pub_ts = ((item.get("modules") or {}).get("module_author") or {}).get("pub_ts")
                    if raw_pub_ts not in (None, "", 0) and not row["pub_ts"]:
                        self.stats["invalid_time"] += 1
                    if row["id"]:
                        seen.add(row["id"])
                    self.rows.append(row)
                    new += 1
                self.stats["pages"] += 1
                self.stats["rows"] = len(self.rows)
                self._flush()
                if self.budget is not None:
                    self.budget.observe_records(new)
                self._p(text=f"第 {self.stats['pages']} 页：+{new} 条，"
                             f"累计 {len(self.rows)} 条动态")
                if not has_more or not next_offset:
                    break
                offset = next_offset
                if self.stats["pages"] >= self.max_pages:
                    self.stats["truncated"] = True
                    break
                if not self._pace(self.sleep):
                    raise TaskCancelledError()
        except TaskCancelledError:
            self.stats["cancelled"] = True
            self.stats["rows"] = len(self.rows)
            if self.rows:
                self._flush()
            self._p(level="warn", text=f"已取消，保留已抓到的 {len(self.rows)} 条动态")
        except BudgetExhaustedError:
            # 兜底：空结果重试等循环内部的请求被 HTTP 层硬拦。与取消分支
            # 对称地保留已抓数据，但记为预算停止——取消与预算互不冒充。
            self.stats["stopped_reason"] = "budget_reached"
            self.stats["rows"] = len(self.rows)
            if self.rows:
                self._flush()
            self._p(level="warn",
                    text=f"已达预算上限，安全停止，保留已抓到的 "
                         f"{len(self.rows)} 条动态")
        return self.stats


# ---------- 导出 ----------

_HEADERS = ("序号", "发布时间", "类型", "正文", "转发", "评论", "点赞",
            "收藏", "投币", "BV号", "链接", "动态ID")
_WIDTHS = (6, 18, 10, 60, 8, 8, 10, 8, 8, 16, 40, 22)


def export_xlsx(rows, uid, stats, path, progress=None):
    """动态明细 → Excel。rows 为空也照样出一份带说明的表。"""
    if progress:
        progress(text="正在生成 Excel…")

    def state(kind, label):
        return xlsx_mod.cell_value(None, kind, note=label)

    def value(raw, kind, label):
        if raw is None:
            return state(xlsx_mod.CellKind.MISSING, f"{label}缺失")
        return xlsx_mod.checked_cell_value(raw, kind, note=f"{label}格式异常")

    def field(mapping, key, kind, label):
        if key not in mapping:
            return state(xlsx_mod.CellKind.NOT_RETURNED, f"{label}未返回")
        if mapping[key] is None:
            return state(xlsx_mod.CellKind.MISSING, f"{label}缺失")
        if kind is xlsx_mod.CellKind.DATETIME:
            return xlsx_mod.unix_seconds_cell_value(mapping[key], note=f"{label}格式异常")
        return value(mapping[key], kind, label)

    def url_field(mapping, key, label):
        if key not in mapping:
            return state(xlsx_mod.CellKind.NOT_RETURNED, f"{label}未返回")
        if mapping[key] is None:
            return state(xlsx_mod.CellKind.MISSING, f"{label}缺失")
        return external_link_cell(sw, mapping[key])

    wb = xlsx_mod.new_workbook()
    sw = xlsx_mod.SheetWriter(wb)

    ws = wb.create_sheet("概览")
    sw.ws = ws
    sw.title_row(ws, f"用户动态导出 · UID {uid}", 4)
    if stats.get("stopped_reason") == "budget_reached":
        note = "已达上限安全停止（请求数/时长预算到限，已抓数据完整保留）"
    elif stats.get("truncated"):
        note = "注意：达到页数上限，可能还有更早的动态未抓取"
    elif not stats.get("cancelled"):
        note = "已抓到底"
    else:
        note = "任务被中途取消"
    sw.kv(ws, [
        ("UID", value(uid, xlsx_mod.CellKind.ID, "UID"), "数据来自该用户的公开动态"),
        ("动态条数", value(len(rows), xlsx_mod.CellKind.INTEGER, "动态条数"),
         f"共 {stats.get('pages', 0)} 页"
                                f"（每页约 {PAGE_SIZE} 条）"),
        ("完成情况", value(note, xlsx_mod.CellKind.TEXT, "完成情况"),
         f"取消={stats.get('cancelled', False)}"),
        ("抓取时间", value(datetime.now(xlsx_mod.ASIA_SHANGHAI).replace(tzinfo=None),
                           xlsx_mod.CellKind.DATETIME, "抓取时间"), ""),
    ])
    sw.append([None, sw.wc("口径：B站动态接口（游客通道 + WBI 签名）。"
                           "计数为接口返回的实时口径，非历史定格。",
                            font=xlsx_mod.F_CAPTION)])
    append_sheet_directory(sw, ws, (
        ("动态明细", "动态明细"),
        ("数据质量", "数据质量"),
        ("字段说明", "字段说明"),
    ))

    ws2 = wb.create_sheet("动态明细")
    sw.ws = ws2
    detail_layout = TableLayout(3, 2, len(_HEADERS) + 1)
    configure_table(ws2, detail_layout)
    sw.title_row(ws2, f"UID {uid} 的动态明细", len(_HEADERS))
    sw.header_row(ws2, list(_HEADERS))
    for i, r in enumerate(rows, 1):
        if r.get("pub_ts") not in (None, 0):
            time_cell = field(r, "pub_ts", xlsx_mod.CellKind.DATETIME, "发布时间")
        elif r.get("time"):
            # 旧 JSONL 只有本地字符串时不伪造时区或精度，保守保留文本。
            time_cell = value(r.get("time"), xlsx_mod.CellKind.TEXT, "发布时间")
        else:
            time_cell = state(xlsx_mod.CellKind.MISSING, "发布时间缺失")
        bvid_cell = (field(r, "bvid", xlsx_mod.CellKind.ID, "BV号")
                     if r.get("bvid") else state(xlsx_mod.CellKind.NOT_APPLICABLE, "BV号不适用"))
        sw.append([None,
                   value(i, xlsx_mod.CellKind.INTEGER, "序号"), time_cell,
                   field(r, "type", xlsx_mod.CellKind.TEXT, "类型"),
                   wrap_cell(sw, field(r, "text", xlsx_mod.CellKind.TEXT, "正文")),
                   field(r, "forward", xlsx_mod.CellKind.INTEGER, "转发"),
                   field(r, "comment", xlsx_mod.CellKind.INTEGER, "评论"),
                   field(r, "like", xlsx_mod.CellKind.INTEGER, "点赞"),
                   field(r, "favorite", xlsx_mod.CellKind.INTEGER, "收藏"),
                   field(r, "coin", xlsx_mod.CellKind.INTEGER, "投币"),
                   bvid_cell, url_field(r, "url", "链接"),
                   field(r, "id", xlsx_mod.CellKind.ID, "动态ID")])
    for idx, width in enumerate(_WIDTHS, start=2):
        ws2.column_dimensions[get_column_letter(idx)].width = width
    finish_table(ws2, detail_layout, len(rows))
    def q(item, raw, kind, unit, note, presentation_state=None):
        if raw is None:
            value_cell = xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE, note=note)
        else:
            value_cell = xlsx_mod.checked_cell_value(raw, kind, note=note)
        return QualityItem("动态", item, value_cell, unit, note, presentation_state)
    quality = [
        q("候选记录总数", stats.get("candidate_records"), xlsx_mod.CellKind.INTEGER, "条", "接口分页返回的候选 item 数"),
        q("实际写入动态明细数", len(rows), xlsx_mod.CellKind.INTEGER, "条", "使用真实动态 ID 去重后的有效行"),
        q("检测到的重复数", stats.get("duplicate_rows"), xlsx_mod.CellKind.INTEGER, "条", "重复动态 ID"),
        q("解析失败数", stats.get("parse_failures"), xlsx_mod.CellKind.INTEGER, "条", "非对象、缺失动态 ID 或无法归一化的 item"),
        q("非法时间数量", stats.get("invalid_time"), xlsx_mod.CellKind.INTEGER, "条", "原始 pub_ts 存在但无法解析"),
        QualityItem("接口", "接口声称数量",
                    xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_RETURNED,
                                        note="接口未返回"), "条", "接口未提供全量声明数量"),
        q("覆盖率", None, xlsx_mod.CellKind.PERCENT, "状态", "不适用：接口未返回声明数量"),
        q("是否取消", bool(stats.get("cancelled", False)), xlsx_mod.CellKind.BOOLEAN, "状态", "来自结构化 crawler stats", "stop"),
        q("是否截断", bool(stats.get("truncated", False)), xlsx_mod.CellKind.BOOLEAN, "状态", "页数或空结果边界", "warning"),
        q("是否预算到限", stats.get("stopped_reason") == "budget_reached", xlsx_mod.CellKind.BOOLEAN, "状态", "来自 stopped_reason", "stop"),
        q("是否部分成功", bool(rows) and bool(stats.get("truncated") or stats.get("cancelled") or stats.get("stopped_reason")), xlsx_mod.CellKind.BOOLEAN, "状态", "已有有效记录但任务未完整结束", "warning"),
    ]
    quality.extend(primary_key_quality(
        "主键", "动态ID", denominator=stats.get("candidate_records"),
        missing=stats.get("missing_id"), invalid=stats.get("invalid_id"),
        duplicates=stats.get("duplicate_rows"),
        dedup_discarded=stats.get("duplicate_rows"),
        remaining_conflicts=stats.get("remaining_id_conflicts"),
        source_note="来自动态 crawler 结构化统计；不扫描动态明细表",
    ))
    fields = []
    for display, stable, dtype, metric in (
        ("序号", "row_number", "整数", "本次输出顺序"), ("发布时间", "pub_ts", "日期时间", "Unix 秒转换为 Asia/Shanghai"),
        ("类型", "type", "文本", "类型枚举映射"), ("正文", "text", "文本", "用户来源文本"),
        ("转发", "forward", "整数", "接口返回计数"), ("评论", "comment", "整数", "接口返回计数"),
        ("点赞", "like", "整数", "接口返回计数"), ("收藏", "favorite", "整数", "接口返回计数"),
        ("投币", "coin", "整数", "接口返回计数"), ("BV号", "bvid", "ID", "投稿动态关联 BV"),
        ("链接", "url", "文本", "接口跳转链接"), ("动态ID", "id", "ID", "真实动态唯一 ID"),
    ):
        fields.append(FieldDefinition("动态明细", display, stable, dtype, "", "是", "动态接口/本地归一化", metric,
                                      xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE), "上游已归一化，无法区分"))
    for display, stable, dtype, metric in (
        ("UID", "uid", "ID", "任务目标 UID"),
        ("动态条数", "dynamic_count", "整数", "有效动态 ID 去重后的输出数"),
        ("完成情况", "completion", "文本", "crawler 终态说明"),
        ("抓取时间", "captured_at", "日期时间", "本地导出时间"),
    ):
        fields.append(FieldDefinition("概览", display, stable, dtype, "", "是", "动态接口/本地归一化", metric,
                                      xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE), "接口未返回或不适用"))
    metadata = make_metadata(
        tool="用户动态", report_type="用户动态导出", parameters={
            "UID": xlsx_mod.cell_value(uid, xlsx_mod.CellKind.ID),
            "页数": xlsx_mod.checked_cell_value(stats.get("pages", 0), xlsx_mod.CellKind.INTEGER),
            "请求数": xlsx_mod.checked_cell_value(stats.get("requests", 0), xlsx_mod.CellKind.INTEGER),
        }, parameter_allowlist=("UID", "页数", "请求数"), quality_items=quality, fields=fields)
    write_metadata_sheets(wb, metadata)
    xlsx_mod.save_workbook_atomic(wb, path)
    return path
