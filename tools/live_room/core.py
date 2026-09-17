# -*- coding: utf-8 -*-
"""直播间状态解析与导出（纯逻辑，网络经注入的 fetch，可离线测试）。

输入归一与 get_info 解析的正本已上提 `core/live.py`（实时监控「直播间
模式」共用同一套归一口径，见该模块 docstring）；本模块 re-export 既有
导入面——`tools.live_room.core` 上的名字一个不少，行为逐字节一致。

- GET xlive/web-room/v2/index/getRoomPlayInfo?room_id=：短号/别名 → 真实
  room_id（仅作 get_info 报业务错误时的输入归一兜底，最多用一次）。

口径（任务卡硬约束）：online 是平台「人气值」，不是精确观看人数——所有
用户可见表头与文案一律写「人气」。轮询礼仪：间隔下限 30s、上限 3600s，
轮数上限 500；每轮仅 1 次业务请求（get_info）。
"""
from __future__ import annotations

import json
from datetime import datetime

from openpyxl.utils import get_column_letter

from core import xlsx as xlsx_mod
from core.live import (GET_INFO_URL, LIVE_STATUS_LABELS, get_info_url,
                       live_status_label, parse_get_info, parse_room_input,
                       status_transition, to_int)
from core.xlsx_metadata import (
    FieldDefinition,
    QualityItem,
    make_metadata,
    primary_key_quality,
    write_metadata_sheets,
)
from core.xlsx_presentation import (
    TableLayout,
    append_sheet_directory,
    configure_table,
    finish_table,
    wrap_cell,
)

RESOLVE_URL = ("https://api.live.bilibili.com"
               "/xlive/web-room/v2/index/getRoomPlayInfo")

DEFAULT_MODE = "snapshot"
DEFAULT_ROUNDS = 10
MAX_ROUNDS = 500
DEFAULT_INTERVAL = 30
MIN_INTERVAL = 30
MAX_INTERVAL = 3600

DASH = "—"


def _s(value, dash=DASH):
    """用户可见文本里绝不放 None / 空白（参考 danmaku analysis 的 _s）。"""
    if value is None:
        return dash
    text = str(value).strip()
    return text or dash


def resolve_room_id(fetch, room_id):
    """输入归一兜底：getRoomPlayInfo → 真实 room_id（主播 uid 由 get_info 给）。

    仅供 get_info 报业务错误（短号/别名不认）时用一次；解析不出返回 0，
    由调用方决定是否回退原始错误。
    """
    url = f"{RESOLVE_URL}?room_id={int(room_id)}"
    data = (fetch(url) or {}).get("data")
    if not isinstance(data, dict):
        return 0
    return to_int(data.get("room_id"), 0)


def push_message(row, kind, prev_label=""):
    """翻转推送内容：(event_type, title, text)。

    文本只含房间号/标题/固定话术，交给 WebhookAdapter 后还会过
    webhook_payload 的白名单与 sanitize_text 双层兜底。
    """
    room = row.get("room_id")
    title = _s(row.get("title"))
    if kind == "live":
        return ("live_status", "开播提醒",
                f"房间 {room}「{title}」已开播（人气 {row.get('online', 0)}）")
    if kind == "offline":
        return ("live_status", "下播提醒",
                f"房间 {room}「{title}」已下播")
    cur_label = _s(row.get("live_status_label"))
    return ("live_status", "直播状态变化提醒",
            f"房间 {room} 直播状态变化：{_s(prev_label)} → {cur_label}")


# ---------- JSONL / Excel ----------

def jsonl_filename(room_id):
    return f"live_{room_id}.jsonl"


def append_jsonl(path, row):
    """每轮一条快照追加落盘；即时 flush，取消/到限不丢已抓轮次。"""
    with open(path, "a", encoding="utf-8") as fh:
        fh.write(json.dumps(row, ensure_ascii=False) + "\n")


_DETAIL_HEADERS = ("轮次", "时间", "直播状态", "标题", "人气", "分区",
                   "父分区", "开播时刻", "房间号", "主播UID", "标签")
_DETAIL_WIDTHS = (6, 20, 10, 50, 10, 14, 14, 20, 12, 12, 30)


def export_xlsx(rows, stats, path, progress=None):
    """快照 → Excel。单次快照只有「概览」一张表；轮询追踪另加「快照明细」。

    rows 为空不出文件（调用方保证）；stats 为 pipeline 的运行统计。
    """
    if progress:
        progress(text="正在生成 Excel…")
    rows = rows or []

    def state(kind, label):
        return xlsx_mod.cell_value(None, kind, note=label)

    def value(raw, kind, label):
        if raw is None:
            return state(xlsx_mod.CellKind.MISSING, f"{label}缺失")
        return xlsx_mod.checked_cell_value(raw, kind, note=f"{label}格式异常")

    def field(mapping, key, kind, label):
        if key not in mapping:
            return state(xlsx_mod.CellKind.NOT_RETURNED, f"{label}未返回")
        if mapping[key] is None:
            return state(xlsx_mod.CellKind.MISSING, f"{label}缺失")
        return value(mapping[key], kind, label)

    def local_time(raw, label):
        if raw is None or raw == "":
            return state(xlsx_mod.CellKind.MISSING, f"{label}缺失")
        try:
            parsed = xlsx_mod.local_text_to_excel_datetime(raw)
        except (TypeError, ValueError):
            return state(xlsx_mod.CellKind.MISSING, f"{label}格式异常")
        return value(parsed, xlsx_mod.CellKind.DATETIME, label)

    def live_time_cell(row, label):
        raw = row.get("live_time_raw")
        if raw not in (None, "", "0"):
            raw_text = str(raw).strip()
            if raw_text.isdigit():
                return xlsx_mod.unix_seconds_cell_value(raw_text, note=f"{label}格式异常")
            try:
                parsed = xlsx_mod.local_text_to_excel_datetime(raw_text)
            except (TypeError, ValueError):
                return value(row.get("live_time"), xlsx_mod.CellKind.TEXT, label)
            return value(parsed, xlsx_mod.CellKind.DATETIME, label)
        if row.get("live_time"):
            return local_time(row.get("live_time"), label)
        return state(xlsx_mod.CellKind.NOT_APPLICABLE, f"{label}不适用")

    latest = rows[-1] if rows else {}
    mode = stats.get("mode")
    track = mode == "track"

    wb = xlsx_mod.new_workbook()
    sw = xlsx_mod.SheetWriter(wb)

    ws = wb.create_sheet("概览")
    sw.ws = ws
    label = f"直播间追踪 · 房间 {latest.get('room_id', '')}"
    sw.title_row(ws, label, 4)
    if stats.get("stopped_reason") == "budget_reached":
        note = "已达上限安全停止（请求数/时长预算到限，已抓数据完整保留）"
    elif stats.get("cancelled"):
        note = "任务被中途取消（已抓轮次完整保留）"
    elif track:
        note = "已跑满设定轮数"
    else:
        note = "单次快照完成"
    area = "·".join(part for part in (
        _s(latest.get("parent_area_name")), _s(latest.get("area_name")))
        if part != DASH)
    if track:
        mode_note = (f"轮询追踪 {stats.get('rounds_requested', 0)} 轮 × "
                     f"{stats.get('interval', 0):g} 秒间隔")
    else:
        mode_note = "单次快照"
    sw.kv(ws, [
        ("房间号（真实）", field(latest, "room_id", xlsx_mod.CellKind.ID, "房间号"),
         "接口返回的真实房间号（短号/链接已归一）"),
        ("主播 UID", field(latest, "uid", xlsx_mod.CellKind.ID, "主播UID"), ""),
        ("最新标题", value(latest.get("title"), xlsx_mod.CellKind.TEXT, "最新标题"), ""),
        ("当前状态", value(latest.get("live_status_label"), xlsx_mod.CellKind.TEXT, "当前状态"),
         "0 未开播 / 1 直播中 / 2 轮播"),
        ("人气（最新）", field(latest, "online", xlsx_mod.CellKind.INTEGER, "人气"),
         "平台人气值口径，非精确观看人数"),
        ("分区", value(area or DASH, xlsx_mod.CellKind.TEXT, "分区"), ""),
        ("开播时刻", live_time_cell(latest, "开播时刻"), "接口未下发时显示占位符"),
        ("追踪模式", value(mode_note, xlsx_mod.CellKind.TEXT, "追踪模式"),
         f"推送 {stats.get('pushes', 0)} 次"),
        ("完成情况", value(note, xlsx_mod.CellKind.TEXT, "完成情况"),
         f"取消={stats.get('cancelled', False)}"),
        ("导出时间", value(datetime.now(xlsx_mod.ASIA_SHANGHAI).replace(tzinfo=None),
                            xlsx_mod.CellKind.DATETIME, "导出时间"),
         f"共 {len(rows)} 轮快照"),
    ])
    sw.append([None, sw.wc(
        "口径：「人气」为 B 站接口返回的人气值，不是精确观看人数。"
        "轮询每轮仅 1 次业务请求（get_info），间隔下限 30 秒。",
        font=xlsx_mod.F_CAPTION)])

    append_sheet_directory(
        sw,
        ws,
        (("快照明细", "快照明细"), ("数据质量", "数据质量"),
         ("字段说明", "字段说明")) if track else
        (("数据质量", "数据质量"), ("字段说明", "字段说明")),
    )

    if track:
        ws2 = wb.create_sheet("快照明细")
        sw.ws = ws2
        detail_layout = TableLayout(3, 2, len(_DETAIL_HEADERS) + 1)
        configure_table(ws2, detail_layout)
        sw.title_row(ws2, label, len(_DETAIL_HEADERS))
        sw.header_row(ws2, list(_DETAIL_HEADERS))
        for r in rows:
            sw.append([None,
                       field(r, "round", xlsx_mod.CellKind.INTEGER, "轮次"),
                       local_time(r.get("ts"), "时间"),
                       value(r.get("live_status_label"), xlsx_mod.CellKind.TEXT, "直播状态"),
                       wrap_cell(sw, value(r.get("title"), xlsx_mod.CellKind.TEXT, "标题")),
                       field(r, "online", xlsx_mod.CellKind.INTEGER, "人气"),
                       value(r.get("area_name"), xlsx_mod.CellKind.TEXT, "分区"),
                       value(r.get("parent_area_name"), xlsx_mod.CellKind.TEXT, "父分区"),
                       live_time_cell(r, "开播时刻"),
                       field(r, "room_id", xlsx_mod.CellKind.ID, "房间号"),
                       field(r, "uid", xlsx_mod.CellKind.ID, "主播UID"),
                       wrap_cell(sw, value(r.get("tags"), xlsx_mod.CellKind.TEXT, "标签"))])
        for idx, width in enumerate(_DETAIL_WIDTHS, start=2):
            ws2.column_dimensions[get_column_letter(idx)].width = width
        finish_table(ws2, detail_layout, len(rows))
    def q(item, raw, kind, unit, note, presentation_state=None):
        if raw is None:
            value_cell = xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE, note=note)
        else:
            value_cell = xlsx_mod.checked_cell_value(raw, kind, note=note)
        return QualityItem("直播", item, value_cell, unit, note, presentation_state)
    quality = [
        q("请求轮次", stats.get("rounds_requested"), xlsx_mod.CellKind.INTEGER, "轮", "任务参数"),
        q("完成轮次", stats.get("rounds_done"), xlsx_mod.CellKind.INTEGER, "轮", "真实写入快照数"),
        q("业务请求数", stats.get("requests"), xlsx_mod.CellKind.INTEGER, "次", "包含必要的归一化重试"),
        q("房间号归一化请求数", stats.get("resolve_requests"), xlsx_mod.CellKind.INTEGER, "次", "最多一次的输入归一化请求"),
        q("推送数", stats.get("pushes"), xlsx_mod.CellKind.INTEGER, "次", "只有 notify 返回成功才计数"),
        q("是否取消", bool(stats.get("cancelled", False)), xlsx_mod.CellKind.BOOLEAN, "状态", "结构化取消状态", "stop"),
        q("是否预算到限", stats.get("stopped_reason") == "budget_reached", xlsx_mod.CellKind.BOOLEAN, "状态", "结构化停止原因", "stop"),
        q("是否部分成功", bool(rows) and bool(stats.get("cancelled") or stats.get("stopped_reason")), xlsx_mod.CellKind.BOOLEAN, "状态", "有快照但未完成请求轮次", "warning"),
    ]
    key_stats = stats.get("key_stats") if isinstance(stats.get("key_stats"), dict) else {}
    key_label = "room_id / ts" if track else "room_id"
    quality.extend(primary_key_quality(
        "主键", key_label, denominator=key_stats.get("candidate_records"),
        missing=key_stats.get("missing"), invalid=key_stats.get("invalid"),
        duplicates=key_stats.get("duplicates"),
        dedup_discarded=key_stats.get("dedup_discarded"),
        remaining_conflicts=key_stats.get("remaining_conflicts"),
        source_note="来自直播流水线快照结构化统计；不扫描快照明细表",
    ))
    fields = []
    def add_field(sheet, display, stable, dtype, metric):
        fields.append(FieldDefinition(sheet, display, stable, dtype, "", "是", "直播间公开接口", metric,
                                      xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE), "接口未返回或不适用"))
    for display, stable, dtype, metric in (
        ("房间号（真实）", "room_id", "ID", "接口返回的真实房间号"),
        ("主播 UID", "uid", "ID", "接口主播 UID"), ("最新标题", "title", "文本", "最新快照标题"),
        ("当前状态", "live_status_label", "文本", "结构化直播状态映射"),
        ("人气（最新）", "online", "整数", "平台人气值，非精确观看人数"),
        ("分区", "area_name", "文本", "接口分区"), ("开播时刻", "live_time", "日期时间/文本", "接口开播时间"),
        ("追踪模式", "mode", "文本", "单次快照或轮询追踪"), ("完成情况", "completion", "文本", "任务终态说明"),
        ("导出时间", "exported_at", "日期时间", "本地导出时间"),
    ):
        add_field("概览", display, stable, dtype, metric)
    if track:
        for display, stable, dtype, metric in (
            ("轮次", "round", "整数", "请求轮次"), ("时间", "ts", "日期时间", "Asia/Shanghai 本地时间"),
            ("直播状态", "live_status_label", "文本", "结构化 live_status 映射"), ("标题", "title", "文本", "接口标题"),
            ("人气", "online", "整数", "平台人气值，非精确观看人数"), ("分区", "area_name", "文本", "接口分区"),
            ("父分区", "parent_area_name", "文本", "接口父分区"), ("开播时刻", "live_time", "日期时间/文本", "接口开播时间"),
            ("房间号", "room_id", "ID", "真实房间号"), ("主播UID", "uid", "ID", "接口主播 UID"),
            ("标签", "tags", "文本", "接口标签"),
        ):
            add_field("快照明细", display, stable, dtype, metric)
    metadata = make_metadata(
        tool="直播间", report_type="直播间追踪", parameters={
            "规范化房间号": xlsx_mod.cell_value(str(stats.get("room_id", stats.get("room_input", ""))), xlsx_mod.CellKind.ID),
            "模式": xlsx_mod.cell_value(str(stats.get("mode", DEFAULT_MODE)), xlsx_mod.CellKind.TEXT),
            "轮次": xlsx_mod.checked_cell_value(stats.get("rounds_requested", 1), xlsx_mod.CellKind.INTEGER),
            "间隔秒": xlsx_mod.checked_cell_value(stats.get("interval", 0), xlsx_mod.CellKind.DECIMAL),
        }, parameter_allowlist=("规范化房间号", "模式", "轮次", "间隔秒"),
        quality_items=quality, fields=fields)
    write_metadata_sheets(wb, metadata)
    xlsx_mod.save_workbook_atomic(wb, path)
    return path
