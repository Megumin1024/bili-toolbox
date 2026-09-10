# -*- coding: utf-8 -*-
"""本地数据修复与合并核心。

只处理调用方明确传入的本地 JSONL/XLSX；输入始终只读，全部产物先写临时文件，
源指纹复核通过后才发布。
"""
from __future__ import annotations

import json
import os
import re
import uuid
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

from core import diagnostics

from . import core as check_core


MAX_DETAIL_ROWS = 10_000
COMMENT_EXCEL_FIELDS = (
    "序号", "rpid", "层级", "用户昵称", "用户mid", "等级", "大会员",
    "性别", "评论内容", "点赞数", "楼中楼数", "发布时间", "IP属地",
)
VIDEO_EXCEL_FIELDS = (
    "排名", "BV号", "标题", "UP主", "分区", "时长(秒)", "发布时间",
    "播放", "弹幕", "评论", "点赞", "投币", "收藏", "分享",
)
_SECRET_RE = re.compile(
    r"(?i)(cookie|sessdata|bili[_ -]?jct|authorization|bearer|token|password|secret)"
)

Progress = Callable[..., None]
Cancel = Callable[[], bool]


class RepairCancelled(Exception):
    """用户取消。"""


class SourceChanged(Exception):
    """源文件在检查后发生变化。"""


@dataclass
class Detail:
    source: str
    location: str
    operation: str
    field: str = ""
    key: str = ""
    reason: str = ""

    def row(self):
        return (
            _safe(self.source, 180), _safe(self.location, 120),
            _safe(self.operation, 60), _safe(self.field, 80),
            _safe(self.key, 160), _safe(self.reason, 240),
        )


@dataclass
class RepairSummary:
    source: Path
    file_type: str
    kind: str
    status: str = "已处理"
    output: str = ""
    original_records: int = 0
    output_records: int = 0
    blanks_removed: int = 0
    duplicates_removed: int = 0
    missing_filled: int = 0
    unfixed: int = 0
    skipped_reason: str = ""
    fingerprint: dict[str, object] = field(default_factory=dict)
    unchanged: bool = True


@dataclass
class _Artifact:
    temp: Path
    final: Path
    source_paths: tuple[Path, ...]
    kind: str
    file_type: str
    summary: RepairSummary | None = None


class _Run:
    def __init__(self, out_dir: Path, details_limit: int):
        self.out_dir = out_dir
        self.run_id = uuid.uuid4().hex
        self.artifacts: list[_Artifact] = []
        self.reservations: list[Path] = []
        self.details: list[Detail] = []
        self.detail_total = 0
        self.details_limit = max(0, int(details_limit))

    def detail(self, *args, **kwargs):
        self.detail_total += 1
        if len(self.details) < self.details_limit:
            self.details.append(Detail(*args, **kwargs))

    def reserve(self, name: str, sources: Iterable[Path], kind: str,
                file_type: str, summary: RepairSummary | None = None) -> _Artifact:
        stem = Path(name).stem
        suffix = Path(name).suffix
        index = 0
        source_ids = {str(p.resolve(strict=False)).casefold() for p in sources}
        while True:
            tail = "" if index == 0 else f"_{index}"
            final = self.out_dir / f"{stem}{tail}{suffix}"
            if str(final.resolve(strict=False)).casefold() in source_ids:
                index += 1
                continue
            try:
                fd = os.open(str(final), os.O_CREAT | os.O_EXCL | os.O_WRONLY)
                os.close(fd)
                break
            except FileExistsError:
                index += 1
        temp = self.out_dir / f".{final.stem}.{self.run_id}.tmp{final.suffix}"
        artifact = _Artifact(temp, final, tuple(sources), kind, file_type, summary)
        self.artifacts.append(artifact)
        self.reservations.append(final)
        return artifact

    def cleanup(self):
        for artifact in self.artifacts:
            try:
                artifact.temp.unlink(missing_ok=True)
            except OSError:
                pass
        for path in self.reservations:
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass


def _safe(value, limit=240):
    text = diagnostics.sanitize_text("" if value is None else str(value))
    text = _SECRET_RE.sub("[已脱敏]", text)
    return text[:limit]


def _cancelled(cancel):
    if cancel():
        raise RepairCancelled


def _same_fingerprint(expected, actual):
    return all(expected.get(key) == actual.get(key)
               for key in ("path", "sha256", "size", "mtime_ns"))


def snapshot_sources(files):
    return [check_core.fingerprint_file(path) for path in files]


def verify_sources(snapshots):
    for expected in snapshots:
        try:
            actual = check_core.fingerprint_file(expected["path"])
        except (OSError, KeyError, TypeError) as exc:
            raise SourceChanged("源文件已变化，请重新检查") from exc
        if not _same_fingerprint(expected, actual):
            raise SourceChanged("源文件已变化，请重新检查")


def validate_repair_inputs(files, out_dir, snapshots, options):
    paths, output = check_core.validate_inputs(files, out_dir)
    selected = {
        "deduplicate": bool(options.get("deduplicate", True)),
        "clean_blanks": bool(options.get("clean_blanks", True)),
        "normalize_fields": bool(options.get("normalize_fields", True)),
        "merge": bool(options.get("merge", False)),
    }
    if not any(selected.values()):
        raise ValueError("请至少选择一项修复操作")
    expected = list(snapshots or [])
    expected_ids = {str(Path(item.get("path", "")).resolve(strict=False)).casefold()
                    for item in expected if isinstance(item, dict)}
    actual_ids = {str(path.resolve(strict=False)).casefold() for path in paths}
    if expected_ids != actual_ids:
        raise SourceChanged("源文件列表已变化，请重新检查")
    verify_sources(expected)
    return paths, output, expected, selected


def _kind_jsonl(path, progress, cancel):
    try:
        return check_core._scan_jsonl_kind(path, progress, cancel)  # noqa: SLF001
    except check_core.CheckCancelled as exc:
        raise RepairCancelled from exc


def _kind_excel(path):
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        return check_core._find_known_excel_sheet(workbook.sheetnames)[1]  # noqa: SLF001
    finally:
        workbook.close()


def _key_for_json(kind, item):
    if not isinstance(item, dict):
        return ("record", check_core._canonical_json(item))  # noqa: SLF001
    if kind == "已知评论 JSONL":
        key = check_core._normalise_rpid(item.get("rpid"))  # noqa: SLF001
        return ("rpid", key) if key is not None else None
    if kind == "已知视频快照 JSONL":
        bvid = item.get("bvid")
        stamp = check_core._normalise_timestamp(item.get("fetched_at"))  # noqa: SLF001
        if check_core._is_valid_bvid(bvid) and stamp is not None:  # noqa: SLF001
            return ("video", bvid.strip(), stamp)
        return None
    return ("record", check_core._canonical_json(item))  # noqa: SLF001


def _display_key(key):
    if not key:
        return ""
    if key[0] == "rpid":
        return f"rpid={key[1]}"
    if key[0] == "video":
        return f"bvid={key[1]}; fetched_at={key[2]}"
    return "完全相同记录"


def _json_compare_canonical(kind, item):
    if not isinstance(item, dict):
        return check_core._canonical_json(item)  # noqa: SLF001
    normalized = dict(item)
    if kind == "已知评论 JSONL":
        key = check_core._normalise_rpid(item.get("rpid"))  # noqa: SLF001
        if key is not None:
            normalized["rpid"] = key
    elif kind == "已知视频快照 JSONL":
        stamp = check_core._normalise_timestamp(item.get("fetched_at"))  # noqa: SLF001
        if stamp is not None:
            normalized["fetched_at"] = stamp
    return check_core._canonical_json(normalized)  # noqa: SLF001


def _json_output_fields(path, kind, cancel):
    standard = (check_core.COMMENT_FIELDS if kind == "已知评论 JSONL"
                else check_core.VIDEO_FIELDS if kind == "已知视频快照 JSONL" else ())
    if not standard:
        return ()
    extras = []
    seen = set(standard)
    with path.open("r", encoding="utf-8-sig", errors="strict") as handle:
        for raw in handle:
            _cancelled(cancel)
            if not raw.strip():
                continue
            try:
                item = json.loads(raw)
            except (json.JSONDecodeError, ValueError):
                continue
            if not isinstance(item, dict):
                continue
            for name in item:
                if name not in seen:
                    seen.add(name)
                    extras.append(name)
    return tuple(standard) + tuple(extras)


def _normalize_json(kind, item, output_fields, summary, run, line_no):
    if not isinstance(item, dict):
        return item
    standard = (check_core.COMMENT_FIELDS if kind == "已知评论 JSONL"
                else check_core.VIDEO_FIELDS if kind == "已知视频快照 JSONL" else ())
    if not output_fields:
        return item
    result = {}
    for name in output_fields:
        if name not in item:
            result[name] = None
            if name in standard:
                summary.missing_filled += 1
                run.detail(summary.source.name, f"JSONL 第 {line_no} 行", "补齐缺失字段",
                           field=name, reason="缺失字段已填充为 null")
        else:
            result[name] = item[name]
    invalid = None
    if kind == "已知评论 JSONL" and "rpid" in item:
        invalid = "rpid" if check_core._normalise_rpid(item["rpid"]) is None else None  # noqa: SLF001
    elif kind == "已知视频快照 JSONL" and "fetched_at" in item:
        invalid = ("fetched_at" if check_core._normalise_timestamp(item["fetched_at"]) is None  # noqa: SLF001
                   else None)
    if invalid:
        summary.unfixed += 1
        run.detail(summary.source.name, f"JSONL 第 {line_no} 行", "未自动修复",
                   field=invalid, reason="字段值格式异常，已原样保留")
    return result


def _repair_jsonl(path, artifact, options, run, progress, cancel):
    summary = artifact.summary
    kind = artifact.kind
    seen = {}
    output_fields = (_json_output_fields(path, kind, cancel)
                     if options["normalize_fields"] else ())
    with path.open("r", encoding="utf-8-sig", errors="strict") as source, \
            artifact.temp.open("w", encoding="utf-8", newline="\n") as target:
        for line_no, raw in enumerate(source, 1):
            _cancelled(cancel)
            if not raw.strip():
                if options["clean_blanks"]:
                    summary.blanks_removed += 1
                    run.detail(path.name, f"JSONL 第 {line_no} 行", "清理空白行")
                else:
                    target.write(raw if raw.endswith("\n") else raw + "\n")
                continue
            try:
                item = json.loads(raw)
            except (json.JSONDecodeError, ValueError):
                summary.unfixed += 1
                run.detail(path.name, f"JSONL 第 {line_no} 行", "跳过异常记录",
                           reason="JSON 语法错误或数值超出解析限制")
                continue
            summary.original_records += 1
            canonical = _json_compare_canonical(kind, item)
            key = _key_for_json(kind, item)
            if options["deduplicate"] and key is not None and key in seen:
                summary.duplicates_removed += 1
                op = "删除重复记录" if seen[key] == canonical else "冲突重复"
                run.detail(path.name, f"JSONL 第 {line_no} 行", op,
                           key=_display_key(key), reason="保留首次出现的记录")
                continue
            if key is not None:
                seen.setdefault(key, canonical)
            if options["normalize_fields"]:
                item = _normalize_json(kind, item, output_fields, summary, run, line_no)
            target.write(json.dumps(item, ensure_ascii=False, separators=(",", ":")) + "\n")
            summary.output_records += 1
            if line_no == 1 or line_no % 1000 == 0:
                progress(text=f"修复 {path.name} · 第 {line_no} 行", row=line_no)


def _excel_header(ws):
    for row_no, row in enumerate(ws.iter_rows(values_only=True), 1):
        values = tuple(row)
        if any(not check_core._is_blank(value) for value in values):  # noqa: SLF001
            offset, names, mapping = check_core._excel_headers(values)  # noqa: SLF001
            return row_no, offset, names, mapping
    return None


def _is_complex_unknown_workbook(workbook):
    nonempty = 0
    for ws in workbook.worksheets:
        if ws.merged_cells.ranges:
            return True
        found = False
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    return True
                if not check_core._is_blank(cell.value):  # noqa: SLF001
                    found = True
        nonempty += int(found)
    return nonempty > 1


def _excel_key(kind, mapping, row):
    if kind == "已知评论 Excel" and "rpid" in mapping:
        value = row[mapping["rpid"]] if mapping["rpid"] < len(row) else None
        key = check_core._normalise_rpid(value)  # noqa: SLF001
        return ("rpid", key) if key is not None else None
    if kind == "已知视频 Excel" and "BV号" in mapping:
        bvid = row[mapping["BV号"]] if mapping["BV号"] < len(row) else None
        if not check_core._is_valid_bvid(bvid):  # noqa: SLF001
            return None
        stamp_index = mapping.get("抓取时间")
        if stamp_index is not None:
            stamp = check_core._normalise_timestamp(  # noqa: SLF001
                row[stamp_index] if stamp_index < len(row) else None)
            return ("video", bvid.strip(), stamp) if stamp is not None else None
        return ("bvid", bvid.strip())
    return ("record", check_core._canonical_row(row))  # noqa: SLF001


def _excel_compare_canonical(kind, mapping, row, offset):
    values = list(row)
    if kind == "已知评论 Excel" and "rpid" in mapping:
        index = mapping["rpid"]
        key = check_core._normalise_rpid(values[index] if index < len(values) else None)  # noqa: SLF001
        if key is not None and index < len(values):
            values[index] = key
    elif kind == "已知视频 Excel":
        stamp_index = mapping.get("抓取时间")
        if stamp_index is not None and stamp_index < len(values):
            stamp = check_core._normalise_timestamp(values[stamp_index])  # noqa: SLF001
            if stamp is not None:
                values[stamp_index] = stamp
    return check_core._canonical_row(values[offset:])  # noqa: SLF001


def _repair_excel(path, artifact, options, run, progress, cancel):
    summary = artifact.summary
    workbook = load_workbook(path, data_only=False, keep_links=False)
    try:
        selected, kind = check_core._find_known_excel_sheet(workbook.sheetnames)  # noqa: SLF001
        if selected is None and _is_complex_unknown_workbook(workbook):
            summary.status = "已跳过"
            summary.skipped_reason = "未知 Excel 结构复杂（公式、合并单元格或多个数据工作表）"
            run.detail(path.name, "工作簿", "跳过文件", reason=summary.skipped_reason)
            return False
        data_sheets = [workbook[selected]] if selected else list(workbook.worksheets)
        for ws in data_sheets:
            _cancelled(cancel)
            header = _excel_header(ws)
            if header is None:
                continue
            header_row, offset, names, mapping = header
            standard = (COMMENT_EXCEL_FIELDS if kind == "已知评论 Excel"
                        else VIDEO_EXCEL_FIELDS if kind == "已知视频 Excel" else ())
            if standard and options["normalize_fields"]:
                extras = [name for name in names if name and name not in standard]
                ordered = list(standard) + extras
                source_rows = [tuple(cell.value for cell in row)
                               for row in ws.iter_rows(min_row=header_row + 1)]
                ws.delete_rows(header_row, max(1, ws.max_row - header_row + 1))
                for column, name in enumerate(ordered, offset + 1):
                    ws.cell(header_row, column, name)
                mapping = {name: offset + index for index, name in enumerate(ordered)}
                old_mapping = {name: offset + index for index, name in enumerate(names) if name}
                for row in source_rows:
                    values = [None] * offset
                    for name in ordered:
                        index = old_mapping.get(name)
                        if index is None:
                            values.append(None)
                            summary.missing_filled += 1
                        else:
                            values.append(row[index] if index < len(row) else None)
                    ws.append(values)
                names = ordered
            mapping = {name: offset + index for index, name in enumerate(names) if name}
            seen = {}
            delete_rows = []
            for row_no in range(header_row + 1, ws.max_row + 1):
                _cancelled(cancel)
                row = tuple(ws.cell(row_no, col).value for col in range(1, ws.max_column + 1))
                if not any(not check_core._is_blank(v) for v in row[offset:]):  # noqa: SLF001
                    if options["clean_blanks"]:
                        delete_rows.append(row_no)
                        summary.blanks_removed += 1
                        run.detail(path.name, f"工作表 {ws.title} 第 {row_no} 行", "清理空白行")
                    continue
                summary.original_records += 1
                key = _excel_key(kind, mapping, row)
                canonical = _excel_compare_canonical(kind, mapping, row, offset)
                invalid_field = None
                if kind == "已知评论 Excel" and "rpid" in mapping:
                    value = row[mapping["rpid"]] if mapping["rpid"] < len(row) else None
                    invalid_field = ("rpid" if check_core._normalise_rpid(value) is None  # noqa: SLF001
                                     else None)
                elif kind == "已知视频 Excel" and "BV号" in mapping:
                    value = row[mapping["BV号"]] if mapping["BV号"] < len(row) else None
                    invalid_field = ("BV号" if not check_core._is_valid_bvid(value)  # noqa: SLF001
                                     else None)
                if invalid_field:
                    summary.unfixed += 1
                    run.detail(path.name, f"工作表 {ws.title} 第 {row_no} 行",
                               "未自动修复", field=invalid_field,
                               reason="字段值格式异常，已原样保留")
                if options["deduplicate"] and key is not None and key in seen:
                    delete_rows.append(row_no)
                    summary.duplicates_removed += 1
                    op = "删除重复记录" if seen[key] == canonical else "冲突重复"
                    run.detail(path.name, f"工作表 {ws.title} 第 {row_no} 行", op,
                               key=_display_key(key), reason="保留首次出现的记录")
                else:
                    if key is not None:
                        seen.setdefault(key, canonical)
                    summary.output_records += 1
            for row_no in reversed(delete_rows):
                ws.delete_rows(row_no)
            progress(text=f"修复 {path.name} · {ws.title}")
        workbook.save(artifact.temp)
        return True
    finally:
        workbook.close()


def _validate_artifact(artifact):
    if artifact.file_type == "JSONL":
        with artifact.temp.open("r", encoding="utf-8", errors="strict") as handle:
            for raw in handle:
                if raw.strip():
                    json.loads(raw)
    else:
        workbook = load_workbook(artifact.temp, read_only=True, data_only=False)
        workbook.close()


def _stable_json_union(artifacts, kind, cancel):
    fields = []
    known = set()
    standard = (check_core.COMMENT_FIELDS if kind == "已知评论 JSONL"
                else check_core.VIDEO_FIELDS if kind == "已知视频快照 JSONL" else ())
    for name in standard:
        fields.append(name)
        known.add(name)
    for artifact in artifacts:
        with artifact.temp.open("r", encoding="utf-8", errors="strict") as source:
            for raw in source:
                _cancelled(cancel)
                if not raw.strip():
                    continue
                try:
                    item = json.loads(raw)
                except (json.JSONDecodeError, ValueError) as exc:
                    raise ValueError("修复副本包含无法解析的 JSON") from exc
                if not isinstance(item, dict):
                    continue
                for name in item:
                    if name not in known:
                        known.add(name)
                        fields.append(name)
    return tuple(fields)


def _merge_jsonl(artifacts, merged, run, cancel, options):
    seen = {} if options["deduplicate"] else None
    fields = (_stable_json_union(artifacts, artifacts[0].kind, cancel)
              if options["normalize_fields"] else ())
    standard = set(check_core.COMMENT_FIELDS if artifacts[0].kind == "已知评论 JSONL"
                   else check_core.VIDEO_FIELDS)
    with merged.temp.open("w", encoding="utf-8", newline="\n") as target:
        for artifact in artifacts:
            with artifact.temp.open("r", encoding="utf-8", errors="strict") as source:
                for line_no, raw in enumerate(source, 1):
                    _cancelled(cancel)
                    if not raw.strip():
                        if not options["clean_blanks"]:
                            target.write(raw if raw.endswith("\n") else raw + "\n")
                        continue
                    item = json.loads(raw)
                    key = _key_for_json(artifact.kind, item)
                    canonical = _json_compare_canonical(artifact.kind, item)
                    if seen is not None and key is not None and key in seen:
                        artifact.summary.duplicates_removed += 1
                        op = "合并去重" if seen[key] == canonical else "冲突重复"
                        run.detail(artifact.summary.source.name, f"修复副本第 {line_no} 行", op,
                                   key=_display_key(key), reason="合并时保留文件列表中首次记录")
                        continue
                    if seen is not None and key is not None:
                        seen[key] = canonical
                    if fields and isinstance(item, dict):
                        normalized = {}
                        for name in fields:
                            if name in item:
                                normalized[name] = item[name]
                            else:
                                normalized[name] = None
                                artifact.summary.missing_filled += 1
                                if name not in standard:
                                    run.detail(
                                        artifact.summary.source.name,
                                        f"修复副本第 {line_no} 行", "合并补齐字段",
                                        field=name, reason="合并组字段并集缺少该扩展字段")
                        target.write(json.dumps(
                            normalized, ensure_ascii=False, separators=(",", ":")) + "\n")
                    else:
                        target.write(raw if raw.endswith("\n") else raw + "\n")


def _unique_excel_sheet_title(workbook, source_title, reserved=()):
    cleaned = re.sub(r"[:\\/?*\[\]]", "_", str(source_title))[:31] or "工作表"
    used = {sheet.title.casefold() for sheet in workbook.worksheets}
    used.update(str(title).casefold() for title in reserved)
    if cleaned.casefold() not in used:
        return cleaned
    for index in range(2, 10_000):
        suffix = f"_{index}"
        candidate = cleaned[:31 - len(suffix)] + suffix
        if candidate.casefold() not in used:
            return candidate
    raise ValueError("无法为 Excel 工作表生成唯一名称")


def _excel_sheet_name_map(workbook, source, source_data_name, target_data_name):
    mapping = {source_data_name: target_data_name}
    reserved = {target_data_name}
    for source_ws in source.worksheets:
        if source_ws.title == source_data_name:
            continue
        target_title = _unique_excel_sheet_title(workbook, source_ws.title, reserved)
        mapping[source_ws.title] = target_title
        reserved.add(target_title)
    return mapping


def _rewrite_excel_formula(value, sheet_map):
    if not isinstance(value, str) or not value.startswith("="):
        return value
    rewritten = value
    for source_title, target_title in sorted(
            sheet_map.items(), key=lambda item: len(item[0]), reverse=True):
        if source_title == target_title:
            continue
        source_quoted = source_title.replace("'", "''")
        target_quoted = target_title.replace("'", "''")
        quoted = re.compile(rf"'{re.escape(source_quoted)}'!")
        rewritten = quoted.sub(
            lambda _match: f"'{target_quoted}'!", rewritten)
        unquoted = re.compile(
            rf"(?<![A-Za-z0-9_']){re.escape(source_title)}!")
        rewritten = unquoted.sub(
            lambda _match: f"{target_title}!", rewritten)
    return rewritten


def _copy_excel_nondata_sheet(workbook, source_ws, target_title, sheet_map, cancel):
    target_ws = workbook.create_sheet(target_title)
    for row in source_ws.iter_rows():
        _cancelled(cancel)
        for cell in row:
            target_ws.cell(
                row=cell.row, column=cell.column,
                value=_rewrite_excel_formula(cell.value, sheet_map))
    for index, dimension in source_ws.row_dimensions.items():
        target = target_ws.row_dimensions[index]
        target.height = dimension.height
        target.hidden = dimension.hidden
        target.outlineLevel = dimension.outlineLevel
        target.collapsed = dimension.collapsed
    for key, dimension in source_ws.column_dimensions.items():
        target = target_ws.column_dimensions[key]
        target.width = dimension.width
        target.hidden = dimension.hidden
        target.outlineLevel = dimension.outlineLevel
        target.collapsed = dimension.collapsed
    return target_ws


def _append_excel_nondata_sheets(workbook, artifacts, target_data_name, cancel):
    sheet_maps = {}
    for artifact in artifacts:
        _cancelled(cancel)
        source = load_workbook(artifact.temp, data_only=False, keep_links=False)
        try:
            selected, _ = check_core._find_known_excel_sheet(source.sheetnames)  # noqa: SLF001
            data_name = selected or source.sheetnames[0]
            sheet_map = _excel_sheet_name_map(
                workbook, source, data_name, target_data_name)
            sheet_maps[artifact.temp] = sheet_map
            for source_ws in source.worksheets:
                _cancelled(cancel)
                if source_ws.title == data_name:
                    continue
                _copy_excel_nondata_sheet(
                    workbook, source_ws, sheet_map[source_ws.title], sheet_map, cancel)
        finally:
            source.close()
    return sheet_maps


def _merge_excel(artifacts, merged, run, cancel, options):
    base = load_workbook(artifacts[0].temp, data_only=False, keep_links=False)
    try:
        selected, kind = check_core._find_known_excel_sheet(base.sheetnames)  # noqa: SLF001
        data_name = selected or base.sheetnames[0]
        sheet_maps = _append_excel_nondata_sheets(
            base, artifacts[1:], data_name, cancel)
        ws = base[data_name]
        header = _excel_header(ws)
        if header is None:
            base.save(merged.temp)
            return
        header_row, offset, names, mapping = header
        normalize = options["normalize_fields"] and kind in {
            "已知评论 Excel", "已知视频 Excel"
        }
        if normalize:
            standard = (COMMENT_EXCEL_FIELDS if kind == "已知评论 Excel"
                        else VIDEO_EXCEL_FIELDS)
            union_names = list(standard)
            union_seen = set(union_names)
            schemas = []
            for source_artifact in artifacts:
                _cancelled(cancel)
                other = load_workbook(
                    source_artifact.temp, read_only=True, data_only=False, keep_links=False)
                try:
                    other_selected, _ = check_core._find_known_excel_sheet(  # noqa: SLF001
                        other.sheetnames)
                    other_ws = other[other_selected or other.sheetnames[0]]
                    other_header = _excel_header(other_ws)
                    schemas.append(other_header)
                    if other_header:
                        for name in other_header[2]:
                            if name and name not in union_seen:
                                union_seen.add(name)
                                union_names.append(name)
                finally:
                    other.close()
            for column, name in enumerate(union_names, offset + 1):
                ws.cell(header_row, column).value = name
            names = union_names
            mapping = {name: offset + index for index, name in enumerate(names)}
        seen = {} if options["deduplicate"] else None
        for row_no in range(header_row + 1, ws.max_row + 1):
            row = tuple(ws.cell(row_no, c).value for c in range(1, ws.max_column + 1))
            key = _excel_key(kind, mapping, row)
            if seen is not None and key is not None:
                seen[key] = _excel_compare_canonical(kind, mapping, row, offset)
        for artifact in artifacts[1:]:
            _cancelled(cancel)
            sheet_map = sheet_maps.get(artifact.temp, {})
            other = load_workbook(artifact.temp, read_only=True, data_only=False, keep_links=False)
            try:
                other_selected, _ = check_core._find_known_excel_sheet(other.sheetnames)  # noqa: SLF001
                other_ws = other[other_selected or other.sheetnames[0]]
                other_header = _excel_header(other_ws)
                if other_header is None:
                    continue
                other_header_row, other_offset, other_names, other_mapping = other_header
                for row in other_ws.iter_rows(
                        min_row=other_header_row + 1, values_only=False):
                    _cancelled(cancel)
                    row = tuple(cell.value for cell in row)
                    row_is_blank = not any(
                        not check_core._is_blank(v) for v in row[other_offset:]  # noqa: SLF001
                    )
                    if row_is_blank and options["clean_blanks"]:
                        continue
                    key = _excel_key(kind, other_mapping, row)
                    canonical = _excel_compare_canonical(
                        kind, other_mapping, row, other_offset)
                    if seen is not None and key is not None and key in seen:
                        artifact.summary.duplicates_removed += 1
                        run.detail(artifact.summary.source.name, f"工作表 {other_ws.title}",
                                   "合并去重" if seen[key] == canonical else "冲突重复",
                                   key=_display_key(key), reason="合并时保留文件列表中首次记录")
                        continue
                    if seen is not None and key is not None:
                        seen[key] = canonical
                    if normalize:
                        values = [None] * offset
                        by_name = {name: row[index] if index < len(row) else None
                                   for name, index in other_mapping.items()}
                        values.extend(by_name.get(name) for name in names)
                        if not row_is_blank:
                            for name in names:
                                if name not in other_mapping:
                                    artifact.summary.missing_filled += 1
                                    run.detail(
                                        artifact.summary.source.name,
                                        f"工作表 {other_ws.title}", "合并补齐字段",
                                        field=name, reason="合并组列并集缺少该扩展列")
                    else:
                        values = list(row)
                    values = [_rewrite_excel_formula(value, sheet_map)
                              for value in values]
                    ws.append(values)
            finally:
                other.close()
        base.save(merged.temp)
    finally:
        base.close()


def _merge_groups(run, options, progress, cancel):
    if not options["merge"]:
        return []
    groups = {}
    for artifact in run.artifacts:
        if artifact.summary is None or artifact.summary.status != "已处理":
            continue
        if artifact.file_type == "Excel":
            workbook = load_workbook(artifact.temp, read_only=True, data_only=False)
            try:
                ws = workbook[workbook.sheetnames[0]]
                header = _excel_header(ws)
                signature = tuple(header[2]) if header else ()
            finally:
                workbook.close()
            if not artifact.kind.startswith("未知") and options["normalize_fields"]:
                signature = "known"
        elif artifact.kind.startswith("未知"):
            signature = _jsonl_signature(artifact.temp)
            if signature is None:
                signature = ("incompatible", str(artifact.temp))
        else:
            signature = "known"
        groups.setdefault((artifact.file_type, artifact.kind, signature), []).append(artifact)
    merged_outputs = []
    for (file_type, kind, _signature), artifacts in groups.items():
        if len(artifacts) < 2:
            continue
        _cancelled(cancel)
        if kind == "已知评论 JSONL":
            name = "评论数据_合并修复副本.jsonl"
        elif kind == "已知视频快照 JSONL":
            name = "视频快照_合并修复副本.jsonl"
        elif kind == "已知评论 Excel":
            name = "评论数据_合并修复副本.xlsx"
        elif kind == "已知视频 Excel":
            name = "视频快照_合并修复副本.xlsx"
        else:
            name = f"未知数据_合并修复副本{'.jsonl' if file_type == 'JSONL' else '.xlsx'}"
        merged = run.reserve(name, [p for a in artifacts for p in a.source_paths], kind, file_type)
        if file_type == "JSONL":
            _merge_jsonl(artifacts, merged, run, cancel, options)
        else:
            _merge_excel(artifacts, merged, run, cancel, options)
        _validate_artifact(merged)
        merged_outputs.append(str(merged.final))
        progress(text=f"已生成合并候选：{merged.final.name}")
    return merged_outputs


def _jsonl_signature(path, cancel=None):
    cancel = cancel or (lambda: False)
    signature = None
    try:
        with path.open("r", encoding="utf-8", errors="strict") as handle:
            for raw in handle:
                _cancelled(cancel)
                if not raw.strip():
                    continue
                try:
                    item = json.loads(raw)
                except (json.JSONDecodeError, ValueError, TypeError):
                    return None
                if not isinstance(item, dict):
                    return None
                keys = tuple(item.keys())
                if signature is None:
                    signature = keys
                elif keys != signature:
                    return None
    except (OSError, UnicodeError):
        return None
    return signature


def _cell(ws, value, header=False):
    safe = value if value is None or isinstance(value, (int, float, bool)) else _safe(value)
    cell = WriteOnlyCell(ws, value=safe)
    if header:
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1B2A4A")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    else:
        cell.font = Font(name="Microsoft YaHei", color="37352F")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    return cell


def _write_manifest(artifact, summaries, run, options, merged_outputs, cancel):
    wb = Workbook(write_only=True)
    try:
        overview = wb.create_sheet("修复概览")
        overview.append([_cell(overview, v, True) for v in ("指标", "数值", "说明")])
        rows = (
            ("输入文件数", len(summaries), "本次检查快照中的文件"),
            ("已处理文件数", sum(s.status == "已处理" for s in summaries), "成功生成修复副本"),
            ("跳过文件数", sum(s.status == "已跳过" for s in summaries), "损坏或复杂未知文件"),
            ("原始记录数", sum(s.original_records for s in summaries), "可解析数据记录"),
            ("输出记录数", sum(s.output_records for s in summaries), "各文件修复副本记录合计"),
            ("清理空白行", sum(s.blanks_removed for s in summaries), "仅数据区或 JSONL 物理空白行"),
            ("删除重复记录", sum(s.duplicates_removed for s in summaries), "保留首次记录"),
            ("补齐缺失字段", sum(s.missing_filled for s in summaries), "null 或空白单元格"),
            ("未自动修复", sum(s.unfixed for s in summaries), "异常值或语法错误"),
            ("输出文件数", sum(s.status == "已处理" for s in summaries) + len(merged_outputs), "不含清单"),
            ("已生成合并文件", "是" if merged_outputs else "否", "仅兼容组且至少两个文件"),
            ("修复选项", json.dumps(options, ensure_ascii=False, sort_keys=True), "本次执行参数"),
            ("明细总数", run.detail_total, "统计包含超出展示上限的明细"),
            ("明细展示上限", run.details_limit, "超出部分不写入清单"),
        )
        for row in rows:
            overview.append([_cell(overview, v) for v in row])

        files_ws = wb.create_sheet("文件汇总")
        headers = ("源文件名", "文件类型", "识别结构", "状态", "SHA-256", "大小", "mtime_ns",
                   "结束时源文件未变化", "输出文件名", "原始记录", "输出记录", "空白清理",
                   "重复清理", "缺失字段补齐", "未自动修复", "跳过原因")
        files_ws.append([_cell(files_ws, v, True) for v in headers])
        for summary in summaries:
            fp = summary.fingerprint
            values = (summary.source.name, summary.file_type, summary.kind, summary.status,
                      fp.get("sha256", ""), fp.get("size", 0), fp.get("mtime_ns", 0),
                      "是" if summary.unchanged else "否", Path(summary.output).name if summary.output else "",
                      summary.original_records, summary.output_records, summary.blanks_removed,
                      summary.duplicates_removed, summary.missing_filled, summary.unfixed,
                      summary.skipped_reason)
            files_ws.append([_cell(files_ws, v) for v in values])

        details_ws = wb.create_sheet("修复明细")
        details_ws.append([_cell(details_ws, v, True) for v in
                           ("源文件", "行号/工作表", "操作", "字段", "安全键", "原因")])
        for detail in run.details:
            _cancelled(cancel)
            details_ws.append([_cell(details_ws, v) for v in detail.row()])
        wb.save(artifact.temp)
    finally:
        wb.close()


def run_repair(files, out_dir, snapshots, *, deduplicate=True, clean_blanks=True,
               normalize_fields=True, merge=False, progress=None, cancel=None,
               max_detail_rows=MAX_DETAIL_ROWS):
    progress = progress or (lambda **_kwargs: None)
    cancel = cancel or (lambda: False)
    options = {
        "deduplicate": deduplicate,
        "clean_blanks": clean_blanks,
        "normalize_fields": normalize_fields,
        "merge": merge,
    }
    try:
        paths, output, expected, options = validate_repair_inputs(
            files, out_dir, snapshots, options)
    except SourceChanged:
        return {"copies": [], "merged": [], "manifest": "", "dir": "",
                "files": [], "stats": {"cancelled": False, "status": "interrupted",
                                         "reason": "源文件已变化，请重新检查"}}
    run = _Run(output, max_detail_rows)
    summaries = []
    try:
        for index, path in enumerate(paths, 1):
            _cancelled(cancel)
            fp = next(item for item in expected
                      if str(Path(item["path"]).resolve(strict=False)).casefold()
                      == str(path.resolve(strict=False)).casefold())
            file_type = "JSONL" if path.suffix.lower() == ".jsonl" else "Excel"
            summary = RepairSummary(path, file_type, "识别中", fingerprint=fp)
            summaries.append(summary)
            try:
                kind = (_kind_jsonl(path, progress, cancel) if file_type == "JSONL"
                        else _kind_excel(path))
                summary.kind = kind
                artifact = run.reserve(f"{path.stem}_修复副本{path.suffix}", paths,
                                       kind, file_type, summary)
                summary.output = str(artifact.final)
                if file_type == "JSONL":
                    _repair_jsonl(path, artifact, options, run, progress, cancel)
                    success = True
                else:
                    success = _repair_excel(path, artifact, options, run, progress, cancel)
                if success:
                    _validate_artifact(artifact)
                else:
                    artifact.temp.unlink(missing_ok=True)
                    artifact.final.unlink(missing_ok=True)
                    run.reservations.remove(artifact.final)
                    summary.output = ""
            except (OSError, UnicodeError, zipfile.BadZipFile, ValueError, KeyError,
                    IndexError, EOFError) as exc:
                artifact = next((item for item in run.artifacts if item.summary is summary), None)
                if artifact is not None:
                    artifact.temp.unlink(missing_ok=True)
                    artifact.final.unlink(missing_ok=True)
                    if artifact.final in run.reservations:
                        run.reservations.remove(artifact.final)
                summary.status = "已跳过"
                summary.output = ""
                summary.skipped_reason = "文件损坏或无法安全读取"
                run.detail(path.name, "文件", "跳过文件", reason=summary.skipped_reason)
                progress(text=f"跳过 {path.name}：{type(exc).__name__}", level="warn")
            progress(text=f"已处理 {path.name}", done=index, total=len(paths))

        merged_outputs = _merge_groups(run, options, progress, cancel)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        manifest = run.reserve(f"数据修复清单_{timestamp}.xlsx", paths, "修复清单", "Excel")
        _write_manifest(manifest, summaries, run, options, merged_outputs, cancel)
        _validate_artifact(manifest)
        _cancelled(cancel)
        verify_sources(expected)
        for summary in summaries:
            summary.unchanged = True
        for artifact in run.artifacts:
            if artifact.temp.exists():
                os.replace(artifact.temp, artifact.final)
        return {
            "copies": [s.output for s in summaries if s.output],
            "merged": merged_outputs,
            "manifest": str(manifest.final),
            "dir": str(output),
            "files": [{
                "source": str(s.source), "type": s.file_type, "kind": s.kind,
                "status": s.status, "output": s.output,
                "original_records": s.original_records, "output_records": s.output_records,
                "blanks_removed": s.blanks_removed,
                "duplicates_removed": s.duplicates_removed,
                "missing_filled": s.missing_filled, "unfixed": s.unfixed,
                "skipped_reason": s.skipped_reason,
            } for s in summaries],
            "stats": {
                "status": "completed", "cancelled": False,
                "processed_files": sum(s.status == "已处理" for s in summaries),
                "skipped_files": sum(s.status == "已跳过" for s in summaries),
                "outputs": sum(bool(s.output) for s in summaries) + len(merged_outputs) + 1,
                "details": run.detail_total,
                "details_truncated": run.detail_total > run.details_limit,
            },
        }
    except RepairCancelled:
        run.cleanup()
        return {"copies": [], "merged": [], "manifest": "", "dir": "",
                "files": [], "stats": {"cancelled": True, "status": "cancelled"}}
    except SourceChanged:
        run.cleanup()
        return {"copies": [], "merged": [], "manifest": "", "dir": "",
                "files": [], "stats": {"cancelled": False, "status": "interrupted",
                                         "reason": "源文件已变化，请重新检查"}}
    except Exception:
        run.cleanup()
        raise
