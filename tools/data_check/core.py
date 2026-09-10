# -*- coding: utf-8 -*-
"""本地 Excel / JSONL 检查、统计与脱敏报告生成。

本模块只访问调用方明确传入的本地文件；不会修改输入文件，也不包含网络调用。
"""
from __future__ import annotations

import json
import hashlib
import os
import re
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Callable, Iterable
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from core import diagnostics


MAX_ISSUE_DETAILS = 10_000
JSON_KIND_SCAN_LIMIT = 100
MAX_POSITIVE_INT_DIGITS = 20
REPORT_PREFIX = "数据检查报告_"

COMMENT_FIELDS = (
    "rpid", "parent", "root", "is_main", "uname", "mid", "sex", "level",
    "vip", "message", "like", "rcount", "ctime", "location", "is_top",
)
VIDEO_FIELDS = (
    "bvid", "aid", "title", "owner", "owner_mid", "tname", "pubdate",
    "duration", "view", "danmaku", "reply", "favorite", "coin", "share",
    "like", "fetched_at",
)

COMMENT_EXCEL_HEADERS = {"rpid": "rpid", "message": "评论内容"}
VIDEO_EXCEL_HEADERS = {"bvid": "BV号", "title": "标题"}

_COMMENT_KIND_FEATURES = {"message", "mid", "is_main", "uname", "parent", "root"}
_COMMENT_KIND_MARKERS = {"rpid"} | _COMMENT_KIND_FEATURES
_VIDEO_KIND_MARKERS = {"bvid", "fetched_at", "title", "view", "owner"}
_BV_RE = re.compile(r"^BV[0-9A-Za-z]+$")
_POSITIVE_INT_RE = re.compile(r"^[0-9]+$")
_POSITIVE_INT_UPPER_BOUND = 10 ** MAX_POSITIVE_INT_DIGITS

Progress = Callable[..., None]
Cancel = Callable[[], bool]


class CheckCancelled(Exception):
    """内部取消信号；由 run_check 转换为 TaskPage 可识别的结果。"""


@dataclass
class Issue:
    file: str
    location: str
    category: str
    field: str
    message: str
    key: str = ""

    def as_dict(self) -> dict[str, str]:
        return {
            "file": self.file,
            "location": self.location,
            "category": self.category,
            "field": self.field,
            "message": self.message,
            "key": self.key,
        }


@dataclass
class FileSummary:
    file: str
    file_type: str
    structure: str
    records: int = 0
    duplicates: int = 0
    missing: int = 0
    formats: int = 0
    blanks: int = 0
    damaged: int = 0
    conclusion: str = "检查完成，未发现问题"

    def as_dict(self) -> dict[str, object]:
        return {
            "file": self.file,
            "file_type": self.file_type,
            "structure": self.structure,
            "records": self.records,
            "duplicates": self.duplicates,
            "missing": self.missing,
            "formats": self.formats,
            "blanks": self.blanks,
            "damaged": self.damaged,
            "conclusion": self.conclusion,
        }


class _Accumulator:
    def __init__(self, max_details: int = MAX_ISSUE_DETAILS):
        self.max_details = max(0, int(max_details))
        self.issues: list[Issue] = []
        self.total_issues = 0
        self.counts = {
            "duplicates": 0,
            "missing": 0,
            "formats": 0,
            "blanks": 0,
            "damaged": 0,
        }
        self.truncated = False

    def add(self, issue: Issue, summary: FileSummary) -> None:
        self.total_issues += 1
        counter = {
            "重复记录": "duplicates",
            "缺失字段": "missing",
            "空白内容": "blanks",
            "文件损坏": "damaged",
        }.get(issue.category, "formats")
        self.counts[counter] += 1
        if len(self.issues) < self.max_details:
            self.issues.append(issue)
        else:
            self.truncated = True

        if counter == "duplicates":
            summary.duplicates += 1
        elif counter == "missing":
            summary.missing += 1
        elif counter == "blanks":
            summary.blanks += 1
        elif counter == "damaged":
            summary.damaged += 1
        else:
            summary.formats += 1


def _safe_text(value, limit: int = 240) -> str:
    """报告中只保留短的、经过现有脱敏逻辑处理的文本。"""
    text = diagnostics.sanitize_text("" if value is None else str(value))
    return text[:limit]


def _safe_name(path: Path) -> str:
    return _safe_text(path.name or str(path), 180)


def fingerprint_file(path: str | Path) -> dict[str, object]:
    """返回修复阶段使用的只读源文件指纹。"""
    source = Path(path).resolve(strict=False)
    digest = hashlib.sha256()
    with source.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    stat = source.stat()
    return {
        "path": str(source),
        "sha256": digest.hexdigest(),
        "size": stat.st_size,
        "mtime_ns": stat.st_mtime_ns,
    }


def _is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _normalise_positive_int(value) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return value if 0 < value < _POSITIVE_INT_UPPER_BOUND else None
    if isinstance(value, str):
        text = value.strip()
        if len(text) > MAX_POSITIVE_INT_DIGITS or not _POSITIVE_INT_RE.fullmatch(text):
            return None
        try:
            number = int(text)
        except (ValueError, OverflowError):
            return None
        return number if 0 < number < _POSITIVE_INT_UPPER_BOUND else None
    return None


def _normalise_rpid(value) -> int | None:
    return _normalise_positive_int(value)


def _is_valid_bvid(value) -> bool:
    return isinstance(value, str) and bool(_BV_RE.fullmatch(value.strip()))


def _normalise_timestamp(value) -> int | None:
    return _normalise_positive_int(value)


def _canonical_json(value) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _normalise_excel_value(value):
    if isinstance(value, datetime):
        return {"__type__": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"__type__": "date", "value": value.isoformat()}
    if isinstance(value, time):
        return {"__type__": "time", "value": value.isoformat()}
    if isinstance(value, timedelta):
        return {
            "__type__": "timedelta",
            "days": value.days,
            "seconds": value.seconds,
            "microseconds": value.microseconds,
        }
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return {
        "__type__": f"{type(value).__module__}.{type(value).__qualname__}",
        "value": _safe_text(value),
    }


def _canonical_row(values: Iterable[object]) -> str:
    return _canonical_json([_normalise_excel_value(value) for value in values])


def _issue(summary: FileSummary, accumulator: _Accumulator, *, location: str,
           category: str, field: str = "", message: str, key: str = "") -> None:
    accumulator.add(
        Issue(
            file=summary.file,
            location=_safe_text(location, 120),
            category=_safe_text(category, 40),
            field=_safe_text(field, 80),
            message=_safe_text(message, 240),
            key=_safe_text(key, 160),
        ),
        summary,
    )


def _json_kind_evidence(keys: set[str]) -> tuple[bool, bool]:
    comment = "rpid" in keys and bool(keys & _COMMENT_KIND_FEATURES)
    video = {"bvid", "fetched_at"}.issubset(keys)
    return comment, video


def _scan_jsonl_kind(path: Path, progress: Progress, cancel: Cancel) -> str:
    comment_votes = 0
    video_votes = 0
    comment_compatible = 0
    video_compatible = 0
    object_count = 0
    with path.open("r", encoding="utf-8-sig", errors="strict") as handle:
        for line_no, raw_line in enumerate(handle, 1):
            if cancel():
                raise CheckCancelled
            if not raw_line.strip():
                continue
            try:
                item = json.loads(raw_line)
            except (json.JSONDecodeError, ValueError):
                continue
            if not isinstance(item, dict):
                continue
            keys = {str(key) for key in item}
            comment, video = _json_kind_evidence(keys)
            comment_votes += int(comment)
            video_votes += int(video)
            comment_compatible += int(bool(keys & _COMMENT_KIND_MARKERS))
            video_compatible += int(bool(keys & _VIDEO_KIND_MARKERS))
            object_count += 1
            if line_no == 1 or line_no % 500 == 0:
                progress(
                    text=f"识别 {_safe_name(path)} · JSONL 第 {line_no} 行",
                    row=line_no,
                )
            if object_count >= JSON_KIND_SCAN_LIMIT:
                break
    if (comment_votes and not video_votes
            and comment_compatible * 2 >= object_count):
        return "已知评论 JSONL"
    if (video_votes and not comment_votes
            and video_compatible * 2 >= object_count):
        return "已知视频快照 JSONL"
    return "未知结构，仅完成通用检查"


def _jsonl_issue_for_record(summary: FileSummary, accumulator: _Accumulator,
                            kind: str, item: object, line_no: int,
                            seen: set[object], baseline_keys: set[str] | None) -> None:
    location = f"JSONL 第 {line_no} 行"
    if not isinstance(item, dict):
        _issue(summary, accumulator, location=location, category="非对象记录",
               field="顶层", message="顶层 JSON 值不是对象，未按业务记录检查。")
        return

    summary.records += 1
    keys = {str(key) for key in item}
    if kind == "已知评论 JSONL":
        for field in COMMENT_FIELDS:
            if field not in item:
                _issue(summary, accumulator, location=location, category="缺失字段",
                       field=field, message=f"评论记录缺少字段“{field}”。")
        rpid_key = _normalise_rpid(item.get("rpid")) if "rpid" in item else None
        if "rpid" in item and rpid_key is None:
            _issue(summary, accumulator, location=location, category="异常格式",
                   field="rpid", message="rpid 为空或不是正整数。")
        elif rpid_key is not None and rpid_key in seen:
            _issue(summary, accumulator, location=location, category="重复记录",
                   field="rpid", message="同一文件中后续记录使用了重复 rpid。",
                   key=f"rpid={rpid_key}")
        elif rpid_key is not None:
            seen.add(rpid_key)
        if "message" in item and _is_blank(item.get("message")):
            _issue(summary, accumulator, location=location, category="空白内容",
                   field="message", message="评论内容为空、为 null 或仅包含空白字符。")
        return

    if kind == "已知视频快照 JSONL":
        for field in VIDEO_FIELDS:
            if field not in item:
                _issue(summary, accumulator, location=location, category="缺失字段",
                       field=field, message=f"视频快照缺少字段“{field}”。")
        bvid = item.get("bvid")
        if "bvid" in item and _is_blank(bvid):
            _issue(summary, accumulator, location=location, category="空白内容",
                   field="bvid", message="BV号为空或仅包含空白字符。")
        elif "bvid" in item and not _is_valid_bvid(bvid):
            _issue(summary, accumulator, location=location, category="异常格式",
                   field="bvid", message="BV号不是以 BV 开头的有效标识。")
        if "title" in item and _is_blank(item.get("title")):
            _issue(summary, accumulator, location=location, category="空白内容",
                   field="title", message="标题为空或仅包含空白字符。")
        if "fetched_at" in item and _is_blank(item.get("fetched_at")):
            _issue(summary, accumulator, location=location, category="空白内容",
                   field="fetched_at", message="抓取时间为空。")
        elif "fetched_at" in item and _normalise_timestamp(item.get("fetched_at")) is None:
            _issue(summary, accumulator, location=location, category="异常格式",
                   field="fetched_at", message="抓取时间不是有效的时间值。")
        bvid_key = bvid.strip() if _is_valid_bvid(bvid) else None
        fetched_at = item.get("fetched_at")
        fetched_at_key = _normalise_timestamp(fetched_at)
        key = (bvid_key, fetched_at_key)
        if bvid_key is not None and fetched_at_key is not None:
            if key in seen:
                _issue(summary, accumulator, location=location, category="重复记录",
                       field="bvid,fetched_at",
                       message="同一文件中后续快照重复使用了相同 (bvid, fetched_at)。",
                       key=f"bvid={bvid_key}; fetched_at={fetched_at_key}")
            else:
                seen.add(key)
        return

    canonical = _canonical_json(item)
    if canonical in seen:
        _issue(summary, accumulator, location=location, category="重复记录",
               field="整条对象", message="未知 JSONL 中出现完全相同的对象。")
    else:
        seen.add(canonical)
    if baseline_keys is not None and keys != baseline_keys:
        _issue(summary, accumulator, location=location, category="结构不一致",
               field="字段集合", message="未知 JSONL 的对象字段集合与首条对象不一致。")


def _check_jsonl(path: Path, accumulator: _Accumulator, progress: Progress,
                 cancel: Cancel) -> FileSummary:
    summary = FileSummary(_safe_name(path), "JSONL", "识别中")
    seen: set[object] = set()
    baseline_keys: set[str] | None = None
    try:
        if cancel():
            raise CheckCancelled
        kind = _scan_jsonl_kind(path, progress, cancel)
        summary.structure = kind
        saw_physical_line = False
        with path.open("r", encoding="utf-8-sig", errors="strict") as handle:
            for line_no, raw_line in enumerate(handle, 1):
                saw_physical_line = True
                if cancel():
                    raise CheckCancelled
                if not raw_line.strip():
                    _issue(summary, accumulator, location=f"JSONL 第 {line_no} 行",
                           category="空白内容", field="整行",
                           message="JSONL 中存在空白行。")
                    continue
                try:
                    item = json.loads(raw_line)
                except (json.JSONDecodeError, ValueError):
                    _issue(summary, accumulator, location=f"JSONL 第 {line_no} 行",
                           category="JSON语法错误", field="整行",
                           message="该行不是完整、有效的 JSON。")
                    continue
                if isinstance(item, dict) and baseline_keys is None:
                    baseline_keys = {str(key) for key in item}
                _jsonl_issue_for_record(
                    summary, accumulator, kind, item, line_no, seen, baseline_keys
                )
                if line_no == 1 or line_no % 500 == 0:
                    progress(
                        text=f"检查 {_safe_name(path)} · JSONL 第 {line_no} 行",
                        row=line_no,
                    )
        if not saw_physical_line:
            _issue(summary, accumulator, location="文件", category="空白内容",
                   field="整文件", message="JSONL 文件内容为空。")
    except CheckCancelled:
        raise
    except (OSError, UnicodeError):
        summary.structure = "无法读取"
        _issue(summary, accumulator, location="文件", category="文件损坏",
               message="文件无法按 UTF-8 JSONL 完整读取，已跳过其余内容。")
        return summary
    summary.conclusion = _conclusion(summary)
    return summary


def _first_nonempty_row(rows: Iterable[tuple[int, tuple[object, ...]]]):
    for row_no, row in rows:
        if any(not _is_blank(value) for value in row):
            return row_no, row
    return None


def _trim_leading_empty(values: tuple[object, ...]) -> tuple[int, tuple[object, ...]]:
    for index, value in enumerate(values):
        if not _is_blank(value):
            return index, values[index:]
    return len(values), ()


def _excel_headers(row: tuple[object, ...]) -> tuple[int, list[str], dict[str, int]]:
    offset, trimmed = _trim_leading_empty(row)
    names = ["" if _is_blank(value) else _safe_text(value, 120).strip() for value in trimmed]
    mapping = {name: offset + index for index, name in enumerate(names) if name}
    return offset, names, mapping


def _row_value(row: tuple[object, ...], index: int | None):
    if index is None or index >= len(row):
        return None
    return row[index]


def _find_known_excel_sheet(sheetnames: list[str]) -> tuple[str | None, str]:
    if "全量评论" in sheetnames:
        return "全量评论", "已知评论 Excel"
    if "视频总表" in sheetnames:
        return "视频总表", "已知视频 Excel"
    return None, "未知结构，仅完成通用检查"


def _check_known_excel_sheet(ws, summary: FileSummary, accumulator: _Accumulator,
                             kind: str, progress: Progress, cancel: Cancel) -> None:
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        _issue(summary, accumulator, location=f"工作表 {ws.title}",
               category="空工作表", field="表头", message="记录明细工作表为空。")
        return

    offset, _names, mapping = _excel_headers(tuple(header))
    required = COMMENT_EXCEL_HEADERS if kind == "已知评论 Excel" else VIDEO_EXCEL_HEADERS
    for field, header_name in required.items():
        if header_name not in mapping:
            _issue(summary, accumulator, location=f"工作表 {ws.title} 表头",
                   category="缺失字段", field=header_name,
                   message=f"记录明细工作表缺少关键表头“{header_name}”。")

    seen: set[object] = set()
    for row_no, row in enumerate(rows, 2):
        if cancel():
            raise CheckCancelled
        row = tuple(row)
        if not any(not _is_blank(value) for value in row[offset:]):
            continue
        summary.records += 1
        location = f"工作表 {ws.title} 第 {row_no} 行"
        if kind == "已知评论 Excel":
            if "rpid" in mapping:
                rpid = _row_value(row, mapping["rpid"])
                rpid_key = _normalise_rpid(rpid)
                if rpid_key is None:
                    _issue(summary, accumulator, location=location, category="异常格式",
                           field="rpid", message="rpid 为空或不是正整数。")
                elif rpid_key in seen:
                    _issue(summary, accumulator, location=location, category="重复记录",
                           field="rpid", message="同一工作表中后续记录使用了重复 rpid。",
                           key=f"rpid={rpid_key}")
                else:
                    seen.add(rpid_key)
            if "评论内容" in mapping:
                message = _row_value(row, mapping["评论内容"])
                if _is_blank(message):
                    _issue(summary, accumulator, location=location, category="空白内容",
                           field="评论内容", message="评论内容为空或仅包含空白字符。")
        else:
            bvid = _row_value(row, mapping["BV号"]) if "BV号" in mapping else None
            if "BV号" in mapping and _is_blank(bvid):
                _issue(summary, accumulator, location=location, category="空白内容",
                       field="BV号", message="BV号为空或仅包含空白字符。")
            elif "BV号" in mapping and not _is_valid_bvid(bvid):
                _issue(summary, accumulator, location=location, category="异常格式",
                       field="BV号", message="BV号不是以 BV 开头的有效标识。")
            elif "BV号" in mapping:
                bvid_key = bvid.strip()
                if bvid_key in seen:
                    _issue(summary, accumulator, location=location, category="重复记录",
                           field="BV号", message="同一工作表中后续记录使用了重复 BV号。",
                           key=f"bvid={bvid_key}")
                else:
                    seen.add(bvid_key)
            if "标题" in mapping:
                title = _row_value(row, mapping["标题"])
                if _is_blank(title):
                    _issue(summary, accumulator, location=location, category="空白内容",
                           field="标题", message="标题为空或仅包含空白字符。")
        if row_no == 2 or row_no % 500 == 0:
            progress(text=f"检查 {summary.file} · {ws.title} 第 {row_no} 行", row=row_no)


def _check_unknown_excel_sheet(ws, summary: FileSummary, accumulator: _Accumulator,
                               progress: Progress, cancel: Cancel) -> None:
    rows = ws.iter_rows(values_only=True)
    header_info = _first_nonempty_row(enumerate(rows, 1))
    if header_info is None:
        _issue(summary, accumulator, location=f"工作表 {ws.title}",
               category="空工作表", field="表头", message="工作表完全为空。")
        return
    header_row_no, header = header_info
    offset, _names, _mapping = _excel_headers(tuple(header))
    seen: set[str] = set()
    for row_no, row in enumerate(rows, header_row_no + 1):
        if cancel():
            raise CheckCancelled
        row = tuple(row)
        if not any(not _is_blank(value) for value in row[offset:]):
            continue
        summary.records += 1
        canonical = _canonical_row(row[offset:])
        if canonical in seen:
            _issue(summary, accumulator, location=f"工作表 {ws.title} 第 {row_no} 行",
                   category="重复记录", field="整行",
                   message="未知 Excel 中出现完全相同的数据行。")
        else:
            seen.add(canonical)
        if row_no == header_row_no + 1 or row_no % 500 == 0:
            progress(text=f"检查 {summary.file} · {ws.title} 第 {row_no} 行", row=row_no)


def _check_excel(path: Path, accumulator: _Accumulator, progress: Progress,
                 cancel: Cancel) -> FileSummary:
    summary = FileSummary(_safe_name(path), "Excel", "识别中")
    workbook = None
    try:
        workbook = load_workbook(
            filename=str(path), read_only=True, data_only=True, keep_links=False
        )
        sheetnames = list(workbook.sheetnames)
        if not sheetnames:
            summary.structure = "空工作簿"
            _issue(summary, accumulator, location="工作簿", category="空工作簿",
                   field="工作表", message="工作簿没有可读取的工作表。")
            return summary
        selected_sheet, kind = _find_known_excel_sheet(sheetnames)
        summary.structure = kind
        if selected_sheet is not None:
            _check_known_excel_sheet(workbook[selected_sheet], summary, accumulator,
                                     kind, progress, cancel)
        else:
            for sheet_name in sheetnames:
                if cancel():
                    raise CheckCancelled
                _check_unknown_excel_sheet(
                    workbook[sheet_name], summary, accumulator, progress, cancel
                )
    except CheckCancelled:
        raise
    except (OSError, UnicodeError, ValueError, TypeError, KeyError, IndexError,
            EOFError, zipfile.BadZipFile, InvalidFileException, ParseError,
            RuntimeError):
        summary.structure = "无法读取"
        _issue(summary, accumulator, location="文件", category="文件损坏",
               message="Excel 文件无法读取，已跳过其内容。")
        return summary
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except OSError:
                pass
    summary.conclusion = _conclusion(summary)
    return summary


def _conclusion(summary: FileSummary) -> str:
    if summary.damaged:
        return "文件损坏，未能完整读取"
    total = summary.duplicates + summary.missing + summary.formats + summary.blanks
    if total:
        return f"发现 {total} 项问题"
    return "检查完成，未发现问题"


def validate_inputs(files: Iterable[str], out_dir: str) -> tuple[list[Path], Path]:
    """校验页面参数；只检查输入文件，不读取其内容。"""
    if isinstance(files, (str, Path)):
        files = [files]
    raw_files = [str(item).strip() for item in (files or []) if str(item).strip()]
    if not raw_files:
        raise ValueError("请至少选择一个 .xlsx 或 .jsonl 文件")
    result: list[Path] = []
    seen: set[str] = set()
    for raw in raw_files:
        path = Path(raw).expanduser()
        if not path.exists() or not path.is_file():
            raise ValueError(f"输入文件不存在或不是普通文件：{_safe_text(path.name or raw)}")
        if path.suffix.lower() not in {".xlsx", ".jsonl"}:
            raise ValueError("只支持 .xlsx 和 .jsonl 文件（扩展名不区分大小写）")
        identity = str(path.resolve(strict=False)).casefold()
        if identity not in seen:
            result.append(path.resolve(strict=False))
            seen.add(identity)
    out_text = str(out_dir or "").strip()
    if not out_text:
        raise ValueError("请设置检查报告输出目录")
    target = Path(out_text).expanduser()
    try:
        target.mkdir(parents=True, exist_ok=True)
        if not target.is_dir():
            raise ValueError("检查报告输出路径不是目录")
        with tempfile.NamedTemporaryFile(
            mode="w", encoding="utf-8", prefix=".data-check-write-",
            suffix=".tmp", dir=str(target), delete=False,
        ) as probe:
            probe.write("ok")
        Path(probe.name).unlink(missing_ok=True)
    except ValueError:
        raise
    except OSError as exc:
        raise ValueError("检查报告输出目录无法创建或写入") from exc
    return result, target.resolve(strict=False)


def _reserve_report_path(out_dir: Path, files: list[Path]) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    source_ids = {str(path.resolve(strict=False)).casefold() for path in files}
    index = 0
    while True:
        suffix = "" if index == 0 else f"_{index}"
        candidate = out_dir / f"{REPORT_PREFIX}{timestamp}{suffix}.xlsx"
        if str(candidate.resolve(strict=False)).casefold() in source_ids:
            index += 1
            continue
        try:
            fd = os.open(str(candidate), os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.close(fd)
            return candidate
        except FileExistsError:
            index += 1


def _cell(ws, value, *, header: bool = False):
    safe_value = value if isinstance(value, (int, float, bool)) or value is None else diagnostics.sanitize_text(value)
    cell = WriteOnlyCell(ws, value=safe_value)
    if header:
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1B2A4A")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    else:
        cell.font = Font(name="Microsoft YaHei", color="37352F")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    return cell


def _write_report(path: Path, files: list[FileSummary], issues: list[Issue],
                  stats: dict[str, object], cancel: Cancel) -> None:
    workbook = Workbook(write_only=True)
    workbook.properties.creator = "BiliToolbox"
    temp_path = path.with_suffix(path.suffix + ".tmp")
    try:
        overview = workbook.create_sheet("检查概览")
        overview.append([_cell(overview, value, header=True) for value in ("指标", "数值", "说明")])
        overview_rows = [
            ("文件数量", stats["file_count"], "本次选择的本地文件数量"),
            ("可读取文件数量", stats["readable_files"], "成功完成本地读取的文件数量"),
            ("损坏文件数量", stats["damaged_files"], "文件级读取失败或 UTF-8 解码失败"),
            ("有效记录数", stats["record_count"], "有效 JSON 对象或非空 Excel 数据行"),
            ("重复数量", stats["duplicates"], "按已识别结构的重复键或完全相同记录统计"),
            ("缺失字段数量", stats["missing"], "按记录/表头缺失事件统计"),
            ("异常格式数量", stats["formats"], "含语法错误、非对象、结构不一致和格式异常"),
            ("空白内容数量", stats["blanks"], "空白行或规定字段为空"),
            ("总问题数量", stats["total_issues"], "包括文件损坏、重复、缺失、格式和空白问题"),
        ]
        if stats.get("details_truncated"):
            overview_rows.append(
                ("问题明细说明", f"仅展示前 {stats['detail_limit']:,} 条",
                 f"实际问题共 {stats['total_issues']:,} 条，统计值仍为完整结果")
            )
        for row in overview_rows:
            overview.append([_cell(overview, value) for value in row])

        summary_sheet = workbook.create_sheet("文件汇总")
        summary_headers = (
            "输入文件名", "文件类型", "识别出的数据结构", "记录数", "重复", "缺失字段",
            "异常格式", "空白内容", "文件损坏", "检查结论",
        )
        summary_sheet.append([_cell(summary_sheet, value, header=True) for value in summary_headers])
        for item in files:
            row = item.as_dict()
            summary_sheet.append([_cell(summary_sheet, row[key]) for key in (
                "file", "file_type", "structure", "records", "duplicates", "missing",
                "formats", "blanks", "damaged", "conclusion",
            )])

        details = workbook.create_sheet("问题明细")
        detail_headers = ("文件", "工作表或 JSONL 行号", "问题类别", "字段",
                          "安全的问题说明", "重复键/定位")
        details.append([_cell(details, value, header=True) for value in detail_headers])
        for item in issues:
            if cancel():
                raise CheckCancelled
            details.append([_cell(details, value) for value in (
                item.file, item.location, item.category, item.field, item.message, item.key,
            )])
        workbook.save(str(temp_path))
    finally:
        try:
            workbook.close()
        except OSError:
            pass


def _write_report_atomically(final_path: Path, files: list[FileSummary], issues: list[Issue],
                             stats: dict[str, object], cancel: Cancel) -> None:
    temp_path = final_path.with_suffix(final_path.suffix + ".tmp")
    try:
        _write_report(final_path, files, issues, stats, cancel)
        if cancel():
            raise CheckCancelled
        os.replace(str(temp_path), str(final_path))
    except CheckCancelled:
        temp_path.unlink(missing_ok=True)
        final_path.unlink(missing_ok=True)
        raise
    except Exception:
        temp_path.unlink(missing_ok=True)
        final_path.unlink(missing_ok=True)
        raise


def run_check(files: Iterable[str], out_dir: str, progress: Progress | None = None,
              cancel: Cancel | None = None,
              max_issue_details: int = MAX_ISSUE_DETAILS) -> dict[str, object]:
    """检查所有文件并生成报告；文件级损坏不会使整个任务失败。"""
    progress = progress or (lambda **_kwargs: None)
    cancel = cancel or (lambda: False)
    file_paths, output_dir = validate_inputs(files, out_dir)
    accumulator = _Accumulator(max_issue_details)
    summaries: list[FileSummary] = []
    for index, path in enumerate(file_paths, 1):
        if cancel():
            return _cancelled_result(summaries, accumulator)
        progress(text=f"开始检查 {path.name}", done=index - 1, total=len(file_paths))
        try:
            summary = (
                _check_jsonl(path, accumulator, progress, cancel)
                if path.suffix.lower() == ".jsonl"
                else _check_excel(path, accumulator, progress, cancel)
            )
        except CheckCancelled:
            return _cancelled_result(summaries, accumulator)
        summary.conclusion = _conclusion(summary)
        summaries.append(summary)
        progress(text=f"完成检查 {path.name} · {summary.conclusion}",
                 done=index, total=len(file_paths))

    stats = {
        "file_count": len(file_paths),
        "readable_files": sum(1 for item in summaries if not item.damaged),
        "damaged_files": sum(1 for item in summaries if item.damaged),
        "record_count": sum(item.records for item in summaries),
        "duplicates": accumulator.counts["duplicates"],
        "missing": accumulator.counts["missing"],
        "formats": accumulator.counts["formats"],
        "blanks": accumulator.counts["blanks"],
        "total_issues": accumulator.total_issues,
        "details_truncated": accumulator.truncated,
        "detail_limit": accumulator.max_details,
        "cancelled": False,
    }
    if cancel():
        return _cancelled_result(summaries, accumulator)
    final_path = _reserve_report_path(output_dir, file_paths)
    try:
        _write_report_atomically(final_path, summaries, accumulator.issues, stats, cancel)
    except CheckCancelled:
        return _cancelled_result(summaries, accumulator)
    progress(text=f"检查报告已生成：{final_path.name}", done=len(file_paths), total=len(file_paths))
    return {
        "report": str(final_path),
        "dir": str(output_dir),
        "source_files": [fingerprint_file(path) for path in file_paths],
        "files": [item.as_dict() for item in summaries],
        "issues": [item.as_dict() for item in accumulator.issues],
        "stats": stats,
    }


def _cancelled_result(summaries: list[FileSummary], accumulator: _Accumulator) -> dict[str, object]:
    return {
        "report": "",
        "dir": "",
        "files": [item.as_dict() for item in summaries],
        "issues": [item.as_dict() for item in accumulator.issues],
        "stats": {
            "cancelled": True,
            "file_count": len(summaries),
            "record_count": sum(item.records for item in summaries),
            "total_issues": accumulator.total_issues,
        },
    }
