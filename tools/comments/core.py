# -*- coding: utf-8 -*-
"""评论抓取核心：gRPC 全量抓取（主楼+楼中楼、断点续传）+ 精确分析 + Excel。

- Crawler 的 gRPC metadata 由 core.session.grpc_metadata() 注入（预热 cookie
  伪装）
- 可选实验通道 grpc_tls（Chrome TLS 指纹 gRPC），失败自动回退 grpcio
- 可选任务预算（core.budget.TaskBudget）：gRPC 每页业务请求记账一次，到限按
  正常完成收尾，断点保留可续传
- Excel 导出基于 core.xlsx 共享基建
- 本模块自建 gRPC 通道、不走 BiliClient，因此**自己过全局闸门**
  （core.gate.shared_gate）：这是全项目请求速率最高的一条流。
"""
import json
import random
import re
import time
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import grpc

from core import xlsx as xlsx_mod
from core.budget import BudgetExhaustedError
from core.cancel import wait as cancel_wait
from core.gate import shared_gate
from core.xlsx_metadata import (
    FieldDefinition,
    QualityItem,
    classify_declared_count,
    make_metadata,
    percent_cell,
    primary_key_quality,
    write_metadata_sheets,
)
from core.xlsx_presentation import (
    TableLayout,
    append_sheet_directory,
    configure_table,
    finish_table,
    wrap_cell,
)

from . import grpc_tls
from . import reply_pb2, reply_pb2_grpc

UA_APP = "Mozilla/5.0 BiliDroid/8.16.0 (bbcall@126.com)"
GRPC_HOST = "grpc.biliapi.net:443"

try:  # 有 openpyxl 才能导出
    from openpyxl.utils import get_column_letter
    HAVE_XLSX = True
except ImportError:  # noqa: F841
    HAVE_XLSX = False


def random_jitter():
    return random.random() * 0.1


class TaskCancelled(Exception):
    """表示调用方主动取消当前任务。"""


def _normalize_vip(raw):
    """把 gRPC 的 VIP 状态映射为明确的布尔语义。"""
    if isinstance(raw, bool):
        return raw, None
    if isinstance(raw, int) and raw in (0, 1):
        return bool(raw), None
    if isinstance(raw, int):
        return None, "大会员状态未知（仅支持整数 0/1）"
    return None, "大会员状态未知（仅支持 bool 或整数 0/1）"


def _vip_is_true(raw):
    """只把明确的 True 或协议状态 1 计入大会员统计。"""
    normalized, _note = _normalize_vip(raw)
    return normalized is True


_UID_TEXT_RE = re.compile(r"^[0-9]{1,20}$")


def _normalize_uid(raw):
    """把评论行 mid 安全归一化为不经过 float 的 UID 文本。"""
    if isinstance(raw, bool):
        return None
    if isinstance(raw, int):
        text = str(raw)
    elif isinstance(raw, str):
        text = raw.strip()
        if not _UID_TEXT_RE.fullmatch(text):
            return None
    else:
        return None
    if not _UID_TEXT_RE.fullmatch(text):
        return None
    normalized = text.lstrip("0")
    return normalized if normalized else None


def _build_uid_rows(rows):
    """按评论行首次出现顺序构造 UID 汇总，并统计被排除的 mid。"""
    entries = {}
    ordered = []
    missing = invalid = 0
    for row in rows:
        if "mid" not in row or row.get("mid") is None:
            missing += 1
            continue
        uid = _normalize_uid(row.get("mid"))
        if uid is None:
            invalid += 1
            continue
        entry = entries.get(uid)
        if entry is None:
            entry = {"uid": uid, "row": row, "count": 0}
            entries[uid] = entry
            ordered.append(entry)
        entry["count"] += 1
    return ordered, missing, invalid


class _CommentMetadataWorkbookView:
    """让共享元数据写入器在首次写行前配置评论报告的日期列宽。"""

    def __init__(self, workbook):
        self._workbook = workbook

    @property
    def sheetnames(self):
        return self._workbook.sheetnames

    def create_sheet(self, title=None, index=None):
        worksheet = self._workbook.create_sheet(title, index)
        if worksheet.title == "数据质量":
            worksheet.column_dimensions["C"].width = 21
            worksheet.column_dimensions["D"].width = 21
        elif worksheet.title == "字段说明":
            worksheet.column_dimensions["C"].width = 21
        return worksheet

    def close(self):
        return self._workbook.close()


# ============================ gRPC 抓取 ============================

def norm_grpc(r, root_rpid=None):
    member = r.member
    ctrl = r.reply_control
    return {
        "rpid": r.id,
        "parent": r.parent,
        "root": r.root or (root_rpid or 0),
        "is_main": root_rpid is None,
        "uname": member.name,
        "mid": member.mid,
        "sex": member.sex,
        "level": member.level,
        "vip": member.vip_status,
        "message": r.content.message,
        "like": r.like,
        "rcount": r.count,
        "ctime": r.ctime,
        "location": ctrl.location,
        "is_top": bool(ctrl.is_up_top or ctrl.is_admin_top),
    }


class Crawler:
    """全量评论抓取（gRPC 游客通道，可取消、可限页、断点续传）。"""

    # 类级缺省：让绕过 __init__ 的构造方式（离线测试用 __new__ 手工装配）
    # 读 budget 时也能安全拿到 None，不改变任何真实构造路径。
    budget = None

    def __init__(self, oid, rtype, out_dir, sleep=0.2, max_pages=0,
                 progress=None, cancel=None, metadata=None, use_tls_grpc=False,
                 gate=None, sleeper=None, budget=None):
        self.oid, self.rtype = int(oid), int(rtype)
        self.out_dir = Path(out_dir)
        self.out_dir.mkdir(parents=True, exist_ok=True)
        # sleep 是**业务间隔**（秒/页，用户可调）；sleeper 是等待实现，仅供注入。
        self.sleep = sleep
        self._sleep_fn = sleeper or time.sleep
        self.max_pages = int(max_pages or 0)  # 0=不限
        self.progress = progress or (lambda **k: None)
        self.cancel = cancel or (lambda: False)
        # 任务预算（core.budget.TaskBudget），None = 无上限：整条请求链的调用
        # 与没有预算时逐字节一致。
        self.budget = budget
        # 全局限速 + 熔断：本模块自建 gRPC 通道绕过 BiliClient，须自己过闸门
        # 才能与其它工具共享同一份背压（生产路径共用进程级单例）。
        self.gate = shared_gate() if gate is None else gate
        self.out_path = self.out_dir / "comments.jsonl"
        self.ckpt_path = self.out_dir / "checkpoint.json"
        self.ckpt = self._load_ckpt()
        # ckpt 中的字段是持久化的上次中断标记；以下字段只表示本次 crawl()。
        self._run_aborted = False
        self._run_cancelled = False
        self._run_error = None
        self._run_budget_stopped = False
        self.channel = grpc.secure_channel(GRPC_HOST, grpc.ssl_channel_credentials())
        self.stub = reply_pb2_grpc.ReplyStub(self.channel)
        self.md = metadata or [("user-agent", UA_APP)]
        self._tls_ok = bool(use_tls_grpc)

    # ---- 断点 ----
    def _load_ckpt(self):
        if self.ckpt_path.exists():
            try:
                data = json.loads(self.ckpt_path.read_text(encoding="utf-8"))
                if isinstance(data, dict):
                    data.setdefault("aborted", False)
                    data.setdefault("cancelled", False)
                    return data
            except ValueError:
                pass
        return {"phase": "main", "cursor_next": 0, "main_done": 0, "sub_done": 0,
                "pending_roots": [], "next_root_index": 0, "sub_cursor_next": 0,
                "pages": 0, "aborted": False, "cancelled": False}

    def _save_ckpt(self):
        tmp = self.ckpt_path.with_suffix(".tmp")
        tmp.write_text(json.dumps(self.ckpt, ensure_ascii=False), encoding="utf-8")
        tmp.replace(self.ckpt_path)

    def _call_tls(self, fn, req):
        if fn is self.stub.MainList:
            path, cls = grpc_tls.PATH_MAIN, reply_pb2.MainListReply
        else:
            path, cls = grpc_tls.PATH_DETAIL, reply_pb2.DetailListReply
        data = grpc_tls.call(path, req.SerializeToString(), self.md)
        return cls.FromString(data)

    def _call(self, fn, req, tries=4):
        last = None
        for i in range(tries):
            if self.cancel():
                raise TaskCancelled()
            if self.budget is not None and i == 0:
                # 预算强制点：cancel 之后、gate.acquire 之前；检查通过立即计数
                # （与 core.client._request 的 attempt==0 同一口径）。一次 _call
                # 调用 = 一页业务请求，内部退避重试不重复检查/记账；到限抛
                # BudgetExhaustedError，不被重试/退避逻辑捕获，不改变 gate 状态。
                self.budget.check_request()
                self.budget.observe_request()
            if not self.gate.acquire(self.cancel, on_wait=self._log_gate_wait):
                raise TaskCancelled()
            try:
                if self._tls_ok:
                    reply = self._call_tls(fn, req)
                else:
                    reply = fn(req, metadata=self.md, timeout=20)
                self.gate.record_success()
                return reply
            except grpc.RpcError as e:
                last = e
                # 只有明确的服务端限流（RESOURCE_EXHAUSTED）才算"平台在拦我们"，
                # 用来开全局熔断；其余（UNAVAILABLE/DEADLINE_EXCEEDED 等）是传输层
                # 故障，交给本方法自己的退避重试，不能污染熔断状态。
                if e.code() == grpc.StatusCode.RESOURCE_EXHAUSTED:
                    self.gate.record_block(f"grpc:{e.code()}")
                else:
                    self.gate.record_neutral()
                wait = min(2 ** i * 2, 30)
                self.progress(level="warn", text=f"gRPC {e.code()}，{wait:.0f}s 后重试")
                # 退避改为可中断：旧实现是 time.sleep 完才检查取消，用户按下停止
                # 后最长仍要空等 30s。
                if not cancel_wait(wait, self.cancel, sleep=self._sleep_fn):
                    raise TaskCancelled()
            except grpc_tls.GrpcTlsError as e:
                self._tls_ok = False  # 实验通道失败即回退 grpcio，不消耗后续重试
                self.gate.record_neutral()
                self.progress(level="warn",
                              text=f"Chrome-TLS gRPC 通道不可用({e})，已回退 grpcio 通道")
        raise RuntimeError(f"gRPC 连续失败: {last}")

    def _pace(self, seconds):
        """业务间隔等待；被取消返回 False（不再靠 sleep 完再检查）。"""
        return cancel_wait(seconds, self.cancel, sleep=self._sleep_fn)

    def _log_gate_wait(self, seconds, reason):
        """闸门即将暂停时写进进度——否则抓取会静默卡住，用户不知道发生了什么。"""
        self.progress(level="warn",
                      text=f"全局闸门：{reason}，暂停 {seconds:.0f}s（全局限速/熔断）")

    def reset(self):
        """清空断点（新一轮抓取）。"""
        self.ckpt = {"phase": "main", "cursor_next": 0, "main_done": 0, "sub_done": 0,
                     "pending_roots": [], "next_root_index": 0, "sub_cursor_next": 0,
                     "pages": 0, "aborted": False, "cancelled": False}
        self._run_aborted = False
        self._run_cancelled = False
        self._run_error = None
        self._run_budget_stopped = False
        self._save_ckpt()

    def _mark_aborted(self, cancelled=False, error=None):
        """同时保存断点，并区分本次运行状态与断点中的上次中断标记。"""
        self._run_aborted = True
        self._run_cancelled = bool(cancelled)
        self._run_error = str(error) if error is not None else None
        self.ckpt["aborted"] = True
        self.ckpt["cancelled"] = bool(cancelled)
        self._save_ckpt()

    def _mark_budget_stopped(self):
        """预算到限：按正常完成收尾——不设 aborted/cancelled 标记，断点照常
        保留（下次运行从当前游标续传）。取消优先于预算，取消路径不会走到这里。
        """
        self._run_budget_stopped = True
        if self.budget is not None:
            done = self.ckpt["main_done"] + self.ckpt["sub_done"]
            self.progress(level="warn",
                          text=f"已达预算上限（{self.budget.reason()}），"
                               f"安全停止：已抓 {done:,} 条评论，"
                               "断点已保留，可续传")
        self._save_ckpt()

    def crawl(self):
        """执行两阶段抓取，返回统计 dict。中断(取消/限页)时标记 aborted。"""
        # 上一次取消/限页/异常只属于旧运行，不能阻止本次从断点继续。
        self._run_aborted = False
        self._run_cancelled = False
        self._run_error = None
        self._run_budget_stopped = False
        self.ckpt["aborted"] = False
        self.ckpt["cancelled"] = False
        with open(self.out_path, "a", encoding="utf-8") as out:
            self._phase_main(out)
            # 预算到限与中断同样收尾：不再进入下一阶段（否则楼中楼首个请求
            # 会再次触发预算检查，把同一条警告刷两遍）。
            if not self._run_aborted and not self._run_budget_stopped:
                self._phase_sub(out)
        self.channel.close()
        self._save_ckpt()
        if self._run_cancelled:
            status = "cancelled"
        elif self._run_error:
            status = "error"
        elif self._run_aborted:
            status = "interrupted"
        else:
            status = "completed"
        stats = {"main": self.ckpt["main_done"], "sub": self.ckpt["sub_done"],
                 "pages": self.ckpt["pages"],
                 "aborted": self._run_aborted,
                 "cancelled": self._run_cancelled,
                 "status": status,
                 "error": self._run_error}
        if self._run_budget_stopped:
            # 只在真的预算到限时才出现该键：budget=None 路径的返回结构与
            # 引入预算前逐字节一致。
            stats["stopped_reason"] = "budget_reached"
        return stats

    def _phase_main(self, out):
        ck = self.ckpt
        if ck["phase"] != "main":
            return
        self.progress(phase="主楼", text=f"从游标 {ck['cursor_next']} 续传")
        while True:
            if self.cancel():
                self._mark_aborted(cancelled=True)
                return
            req = reply_pb2.MainListReq(oid=self.oid, type=self.rtype,
                                        cursor=reply_pb2.CursorReq(next=ck["cursor_next"],
                                                                   mode=2))
            try:
                resp = self._call(self.stub.MainList, req)
            except TaskCancelled:
                self._mark_aborted(cancelled=True)
                return
            except BudgetExhaustedError:
                self._mark_budget_stopped()
                return
            except RuntimeError as e:
                self.progress(level="error", text=f"主楼阶段终止: {e}")
                self._mark_aborted(error=f"主楼阶段终止：{e}")
                return
            for r in resp.replies:
                out.write(json.dumps(norm_grpc(r), ensure_ascii=False) + "\n")
                ck["main_done"] += 1
                if r.count > 0:
                    ck["pending_roots"].append([r.id, r.count])
            ck["pages"] += 1
            is_end = resp.cursor.isEnd
            ck["cursor_next"] = resp.cursor.next
            self.progress(phase="主楼", pages=ck["pages"], done=ck["main_done"],
                          next=ck["cursor_next"], is_end=is_end,
                          pending=len(ck["pending_roots"]))
            if self.max_pages and ck["pages"] >= self.max_pages:
                self._mark_aborted()
                return
            if is_end:
                ck["phase"] = "sub"
                self._save_ckpt()
                return
            if not resp.replies:
                ck["phase"] = "sub"
                self._save_ckpt()
                return
            if not self._pace(self.sleep + random_jitter()):
                self._mark_aborted(cancelled=True)
                return

    def _phase_sub(self, out):
        ck = self.ckpt
        idx = ck.get("next_root_index", 0)
        pending = ck.get("pending_roots", [])
        while idx < len(pending):
            if self.cancel():
                self._mark_aborted(cancelled=True)
                return
            rpid, _count = pending[idx]
            sub_next = ck.get("sub_cursor_next", 0)
            while True:
                req = reply_pb2.DetailListReq(oid=self.oid, type=self.rtype, root=rpid,
                                              cursor=reply_pb2.CursorReq(next=sub_next,
                                                                         mode=2))
                try:
                    resp = self._call(self.stub.DetailList, req)
                except TaskCancelled:
                    self._mark_aborted(cancelled=True)
                    return
                except BudgetExhaustedError:
                    self._mark_budget_stopped()
                    return
                except RuntimeError as e:
                    self.progress(level="error", text=f"楼中楼 {rpid} 终止: {e}")
                    self._mark_aborted(error=f"楼中楼 {rpid} 终止：{e}")
                    return
                for sub in (resp.root.replies if resp.root else []):
                    out.write(json.dumps(norm_grpc(sub, root_rpid=rpid),
                                         ensure_ascii=False) + "\n")
                    ck["sub_done"] += 1
                ck["pages"] += 1
                sub_next = resp.cursor.next
                ck["sub_cursor_next"] = sub_next
                self.progress(phase="楼中楼", pages=ck["pages"], done=ck["sub_done"],
                              root_idx=idx + 1, root_total=len(pending))
                if resp.cursor.isEnd:
                    break
                if not self._pace(self.sleep + random_jitter()):
                    self._mark_aborted(cancelled=True)
                    return
            idx += 1
            ck["next_root_index"] = idx
            ck["sub_cursor_next"] = 0
            if self.max_pages and ck["pages"] >= self.max_pages:
                self._mark_aborted()
                return
            if not self._pace(self.sleep):
                self._mark_aborted(cancelled=True)
                return
        self._save_ckpt()


# ============================ 分析 ============================

EMOJI_RE = re.compile(r"\[([^\[\]]{1,12})\]")
STOP = set(("的 了 是 我 你 他 她 它 我们 你们 他们 这那 就是 也 都 还有 但是 因为 所以 如果 现在 "
            "什么 怎么 一个 这个 那个 自己 没有 可以 不是 已经 知道 觉得 时候 直接 还是 只是 真的").split())


def analyze(rows, meta):
    """返回 (报告文本, 概览指标 dict)。"""
    def safe_time(value):
        try:
            return xlsx_mod.unix_seconds_to_excel_datetime(value)
        except (TypeError, ValueError):
            return None

    mains = [r for r in rows if r.get("is_main")]
    subs = [r for r in rows if not r.get("is_main")]
    n = len(rows)
    mids = defaultdict(lambda: {"n": 0, "uname": "", "like": 0})
    for r in rows:
        m = mids[r.get("mid")]
        m["n"] += 1
        m["uname"] = r.get("uname")
        m["like"] += r.get("like") or 0
    ts_list = [safe_time(r.get("ctime")) for r in rows
               if r.get("ctime") is not None]
    ts_list = sorted(item for item in ts_list if item is not None)
    by_day = Counter(item.strftime("%m-%d") for item in ts_list)
    by_hour = Counter(item.hour for item in ts_list)
    likes_all = sorted((r["like"] for r in rows), reverse=True)
    zero_like = sum(1 for v in likes_all if v == 0)
    vip = sum(1 for r in rows if _vip_is_true(r.get("vip")))
    levels = Counter(r.get("level") for r in rows)
    emo, uni = Counter(), Counter()
    build_cnt = apology = defense = stop_pay = genshin = spam = 0
    dup = Counter((r["message"] or "").strip() for r in rows)
    SPAM_WORDS = ("代练", "修两层", "接单", "上号", "低价", "变昔涟", "神游")
    for r in rows:
        msg = r["message"] or ""
        for e in set(EMOJI_RE.findall(msg)):
            emo[e] += 1
        for ch in set(c for c in msg if ord(c) > 0x1F000):
            uni[ch] += 1
        build_cnt += ("🧱" in msg or "🔨" in msg)
        apology += ("道歉" in msg)
        defense += ("这么多喷崩铁" in msg or "有什么错" in msg)
        stop_pay += ("停氪" in msg or "退款" in msg)
        genshin += ("原神" in msg)
        spam += any(w in msg for w in SPAM_WORDS)
    toks = Counter()
    for r in rows:
        text = EMOJI_RE.sub("", r["message"] or "")
        for seg in re.findall(r"[\u4e00-\u9fff]{2,}", text):
            for size in (2, 3, 4):
                for i in range(len(seg) - size + 1):
                    tk = seg[i:i + size]
                    if tk not in STOP and not any(c in tk for c in "的一是在不了有和人这"):
                        toks[tk] += 1
    top_words = []
    for tk, c in toks.most_common(150):
        if any(o != tk and len(o) > len(tk) and tk in o and oc >= c * 0.6
               for o, oc in toks.most_common(150)):
            continue
        top_words.append((tk, c))
        if len(top_words) >= 20:
            break
    hours_span = ((max(ts_list) - min(ts_list)).total_seconds() / 3600
                  if ts_list else 0)
    claimed_status, claimed = classify_declared_count(
        meta.get("claimed_comment_count"),
        present="claimed_comment_count" in meta,
    )
    claimed_numeric = claimed_status == "valid"
    coverage_text = (
        f"{n / claimed * 100:.1f}%" if claimed_numeric and claimed
        else "不适用" if claimed_status == "invalid" else "-"
    )

    kpi = {
        "claimed": claimed if claimed_numeric else None, "fetched": n,
        "coverage": coverage_text,
        "coverage_ratio": n / claimed if claimed_numeric and claimed else None,
        "main": len(mains), "sub": len(subs), "users": len(mids),
        "per_user": f"{n / max(len(mids), 1):.2f}",
        # 0 行评论（如启动即取消）也要能出报告：n=0 的占比口径降级为 "-"
        "vip_pct": f"{vip / n * 100:.1f}%" if n else "-",
        "vip_ratio": vip / n if n else None,
        "zero_like_pct": f"{zero_like / n * 100:.1f}%" if n else "-",
        "zero_like_ratio": zero_like / n if n else None,
        "hours": f"{hours_span:.0f}", "build": build_cnt,
        "build_pct": f"{build_cnt / n * 100:.1f}%" if n else "-",
        "build_ratio": build_cnt / n if n else None,
        "apology": apology, "defense": defense, "stop_pay": stop_pay,
        "genshin": genshin, "spam": spam,
        "by_day": sorted(by_day.items()), "by_hour": sorted(by_hour.items()),
        "top_words": top_words, "emo_top": emo.most_common(10),
        "uni_top": uni.most_common(10),
        "levels": sorted(levels.items(), key=lambda x: (x[0] is None, x[0])),
    }

    top_likes = sorted(rows, key=lambda r: -(r["like"] or 0))[:10]
    top_users = sorted(mids.items(), key=lambda kv: -kv[1]["n"])[:10]
    top_dups = [(m, c) for m, c in dup.most_common(6) if c >= 3][:5]

    L = []
    a = L.append
    a("# 评论精确分析\n")
    pub_dt = (safe_time(meta["pub_ts"])
              if meta.get("pub_ts") is not None else None)
    a(f"- 对象：{meta.get('author', '')} 《{meta.get('title', '')}》"
      f"（{pub_dt.strftime('%Y-%m-%d %H:%M') if pub_dt else '-'} 发布）")
    claimed_text = (
        "接口未返回" if claimed_status == "not_returned"
        else f"{claimed:,}" if claimed_numeric
        else "声明数量格式异常"
    )
    a(f"- B站声称评论总数：{claimed_text}；抓取 {n:,} 条（{kpi['coverage']}）"
      f" = 主楼 {len(mains):,} + 楼中楼 {len(subs):,}\n")
    a("## 用户结构")
    a(f"- 独立用户 {len(mids):,}（人均 {kpi['per_user']}）；大会员 {kpi['vip_pct']}%；"
      f"0赞占比 {kpi['zero_like_pct']}")
    a("- 等级：" + "，".join(f"Lv{k}×{v}" for k, v in kpi["levels"]))
    a("- 刷屏 Top：" + "，".join(f"{m['uname']}({m['n']}条)" for _, m in top_users[:5]) + "\n")
    a("## 时间")
    a("- 按日：" + "，".join(f"{d}×{c:,}" for d, c in kpi["by_day"]))
    a("- 高峰：" + "，".join(f"{h}点" for h, _ in sorted(by_hour.items(), key=lambda x: -x[1])[:3]) + "\n")
    a("## 点赞 Top10")
    for r in top_likes:
        tag = "主" if r.get("is_main") else "楼"
        a(f"- [{tag}] {r['uname']}（赞{r['like']:,}）：{(r['message'] or '')[:60]}")
    a("\n## 表情")
    a("- 小表情：" + "，".join(f"[{k}]{v}条" for k, v in kpi["emo_top"]))
    a("- emoji：" + " ".join(f"{k}{v}条" for k, v in kpi["uni_top"]) + "\n")
    a("## 高频词 Top20")
    a("；".join(f"{k}({v:,})" for k, v in top_words) + "\n")
    a("## 话题与阵营")
    a(f"- 盖楼抗议（🧱🔨）：{build_cnt:,} 条（{kpi['build_pct']}）；要求道歉 {apology:,}；"
      f"护官反呛 {defense:,}；停氪/退款 {stop_pay:,}；提及原神 {genshin:,}；疑似引流 {spam:,}\n")
    a("## 重复复读 Top")
    for m, c in top_dups:
        a(f"- ×{c}：{m[:60]}")
    report = "\n".join(L)
    return report, kpi


# ============================ Excel 导出 ============================

def export_xlsx(rows, meta, kpi, out_path, progress=None, stats=None):
    if not xlsx_mod or not HAVE_XLSX:
        raise RuntimeError("缺少 openpyxl（pip install openpyxl）")

    def state(kind, label):
        return xlsx_mod.cell_value(None, kind, note=label)

    def field(mapping, key, kind, label):
        if key not in mapping:
            return state(xlsx_mod.CellKind.NOT_RETURNED, f"{label}未返回")
        if mapping[key] is None:
            return state(xlsx_mod.CellKind.MISSING, f"{label}缺失")
        if kind is xlsx_mod.CellKind.DATETIME:
            return xlsx_mod.unix_seconds_cell_value(mapping[key], note=f"{label}格式异常")
        return xlsx_mod.checked_cell_value(mapping[key], kind, note=f"{label}格式异常")

    def vip_field(mapping):
        if "vip" not in mapping:
            return state(xlsx_mod.CellKind.NOT_RETURNED, "大会员未返回")
        raw = mapping["vip"]
        if raw is None:
            return state(xlsx_mod.CellKind.MISSING, "大会员缺失")
        normalized, note = _normalize_vip(raw)
        if note:
            return state(xlsx_mod.CellKind.MISSING, note)
        return xlsx_mod.cell_value(normalized, xlsx_mod.CellKind.BOOLEAN)

    def value(value, kind, label, number_format=None):
        if value is None:
            return state(xlsx_mod.CellKind.MISSING, f"{label}缺失")
        return xlsx_mod.checked_cell_value(value, kind, number_format=number_format,
                                           note=f"{label}格式异常")

    claimed_status, claimed_value = classify_declared_count(
        meta.get("claimed_comment_count"),
        present="claimed_comment_count" in meta,
    )
    if claimed_status == "not_returned":
        claimed_overview_cell = state(xlsx_mod.CellKind.NOT_RETURNED, "接口未返回")
    elif claimed_status == "invalid":
        claimed_overview_cell = state(xlsx_mod.CellKind.MISSING, "声明数量格式异常")
    else:
        claimed_overview_cell = xlsx_mod.checked_cell_value(
            claimed_value, xlsx_mod.CellKind.INTEGER, note="声明数量格式异常"
        )

    def sort_count(row, key):
        candidate = row.get(key)
        return candidate if isinstance(candidate, (int, float)) and not isinstance(candidate, bool) else 0

    def safe_time(raw, label):
        if raw is None:
            return state(xlsx_mod.CellKind.MISSING, f"{label}缺失")
        return xlsx_mod.unix_seconds_cell_value(raw, note=f"{label}格式异常")

    n = len(rows)
    wb = xlsx_mod.new_workbook()
    sw = xlsx_mod.SheetWriter(wb)
    W = 8

    ts_points = []
    for row in rows:
        raw = row.get("ctime")
        if raw is None:
            continue
        try:
            converted = xlsx_mod.unix_seconds_to_excel_datetime(raw)
        except (TypeError, ValueError):
            continue
        if converted is not None:
            ts_points.append(converted)
    ts_note = (f"{min(ts_points):%m-%d %H:%M}"
               f" ~ {max(ts_points):%m-%d %H:%M}"
               ) if ts_points else "-"

    # Sheet1 统计概览
    ws = wb.create_sheet("统计概览")
    sw.ws = ws
    ws.column_dimensions["C"].width = 21
    sw.title_row(ws, f"评论精确分析 | {meta.get('author', '')} 《{meta.get('title', '')}》", W)
    pub_cell = field(meta, "pub_ts", xlsx_mod.CellKind.DATETIME, "发布时间")
    sw.kv(ws, [
        ("发布时间", pub_cell if "pub_ts" in meta and meta.get("pub_ts") is not None
         else state(xlsx_mod.CellKind.NOT_RETURNED, "发布时间未返回"), meta.get("author", "")),
        ("声称评论总数", claimed_overview_cell, "含楼中楼"),
        ("实际抓取", value(n, xlsx_mod.CellKind.INTEGER, "实际抓取"),
         f"覆盖率 {kpi['coverage']}（差额=已删除/折叠/不可见）"),
        ("主楼 / 楼中楼", value(kpi["main"], xlsx_mod.CellKind.INTEGER, "主楼数"),
         f"楼中楼 {kpi['sub']:,}"),
        ("独立用户", value(kpi["users"], xlsx_mod.CellKind.INTEGER, "独立用户"),
         f"人均 {kpi['per_user']} 条"),
        ("时间跨度", value(kpi["hours"], xlsx_mod.CellKind.DECIMAL, "时间跨度",
                           '0" 小时"'), ts_note),
        ("0赞占比", value(kpi.get("zero_like_ratio"), xlsx_mod.CellKind.PERCENT, "0赞占比"),
         f"大会员 {kpi['vip_pct']}（长尾结构与老玩家占比）"),
        ("盖楼抗议（🧱🔨）", value(kpi["build"], xlsx_mod.CellKind.INTEGER, "盖楼抗议数"),
         f"占比 {kpi['build_pct']}（复读接龙抗议规模）"),
        ("要求道歉 / 护官反呛", value(kpi["apology"], xlsx_mod.CellKind.INTEGER, "要求道歉数"),
         f"护官反呛 {kpi['defense']:,}（阵营对喷比值）"),
        ("停氪退款 / 提及原神 / 疑似引流", value(kpi["stop_pay"], xlsx_mod.CellKind.INTEGER,
                                                   "停氪退款数"),
         f"提及原神 {kpi['genshin']:,} / 疑似引流 {kpi['spam']:,}"),
    ], W)
    sw.header_row(ws, ["排名", "高频词", "出现次数"])
    for i, (tk, c) in enumerate(kpi["top_words"]):
        sw.append([None, value(i + 1, xlsx_mod.CellKind.INTEGER, "排名"),
                   value(tk, xlsx_mod.CellKind.TEXT, "高频词"),
                   value(c, xlsx_mod.CellKind.INTEGER, "出现次数")])
    sw.append([None])
    sw.header_row(ws, ["排名", "B站小表情", "评论数", "", "排名", "emoji", "评论数"])
    for i in range(10):
        e1 = f"[{kpi['emo_top'][i][0]}]" if i < len(kpi["emo_top"]) else ""
        v1 = kpi["emo_top"][i][1] if i < len(kpi["emo_top"]) else None
        e2 = kpi["uni_top"][i][0] if i < len(kpi["uni_top"]) else ""
        v2 = kpi["uni_top"][i][1] if i < len(kpi["uni_top"]) else None
        sw.append([None, value(i + 1, xlsx_mod.CellKind.INTEGER, "排名"),
                   value(e1, xlsx_mod.CellKind.TEXT, "小表情"),
                   value(v1, xlsx_mod.CellKind.INTEGER, "评论数"), None,
                   value(i + 1, xlsx_mod.CellKind.INTEGER, "排名"),
                   value(e2, xlsx_mod.CellKind.TEXT, "emoji"),
                   value(v2, xlsx_mod.CellKind.INTEGER, "评论数")])
    sw.append([None])
    sw.append([None, sw.wc("口径：游客 gRPC 通道全量抓取；情感/主题为规则词典法粗判。",
                            font=xlsx_mod.F_CAPTION)])
    append_sheet_directory(sw, ws, (
        ("时间分布", "分布统计"),
        ("点赞 Top100", "点赞Top100"),
        ("全量评论", "全量评论"),
        ("UID列表", "UID列表"),
        ("数据质量", "数据质量"),
        ("字段说明", "字段说明"),
    ))

    # Sheet2 按日/按小时
    ws2 = wb.create_sheet("分布统计")
    sw.ws = ws2
    sw.title_row(ws2, "时间分布", 4)
    sw.header_row(ws2, ["日期", "评论条数", "占比"])
    for d, c in kpi["by_day"]:
        sw.append([None, value(d, xlsx_mod.CellKind.TEXT, "日期"),
                   value(c, xlsx_mod.CellKind.INTEGER, "评论条数"),
                   value(c / n if n else None, xlsx_mod.CellKind.PERCENT, "占比")])
    sw.append([None])
    sw.header_row(ws2, ["小时", "评论条数", "占比"])
    for h, c in kpi["by_hour"]:
        sw.append([None, value(f"{h:02d}:00-{h:02d}:59", xlsx_mod.CellKind.TEXT, "小时"),
                   value(c, xlsx_mod.CellKind.INTEGER, "评论条数"),
                   value(c / n if n else None, xlsx_mod.CellKind.PERCENT, "占比")])

    # Sheet3 点赞Top100
    ws3 = wb.create_sheet("点赞Top100")
    sw.ws = ws3
    top_layout = TableLayout(3, 2, 7)
    configure_table(ws3, top_layout)
    sw.title_row(ws3, "点赞 Top100", 6)
    sw.header_row(ws3, ["排名", "层级", "用户", "点赞", "内容", "楼中楼数"])
    for i, r in enumerate(sorted(rows, key=lambda row: -sort_count(row, "like"))[:100]):
        sw.append([None, value(i + 1, xlsx_mod.CellKind.INTEGER, "排名"),
                   value("主楼" if r.get("is_main") else "楼中楼", xlsx_mod.CellKind.TEXT, "层级"),
                   field(r, "uname", xlsx_mod.CellKind.TEXT, "用户昵称"),
                   field(r, "like", xlsx_mod.CellKind.INTEGER, "点赞数"),
                    wrap_cell(sw, value((r.get("message") or "")[:120],
                                        xlsx_mod.CellKind.TEXT, "评论内容")),
                   field(r, "rcount", xlsx_mod.CellKind.INTEGER, "楼中楼数")])
    finish_table(ws3, top_layout, min(n, 100))

    # Sheet4 全量评论
    ws4 = wb.create_sheet("全量评论")
    sw.ws = ws4
    headers = ["序号", "rpid", "层级", "用户昵称", "用户mid", "等级", "大会员",
               "性别", "评论内容", "点赞数", "楼中楼数", "发布时间", "IP属地"]
    ws4.column_dimensions["A"].width = 3
    for i, h in enumerate(headers):
        ws4.column_dimensions[get_column_letter(i + 2)].width = {
            "序号": 8, "rpid": 16, "层级": 8, "用户昵称": 20, "用户mid": 13, "等级": 6,
            "大会员": 8, "性别": 6, "评论内容": 60, "点赞数": 10, "楼中楼数": 10,
            "发布时间": 21, "IP属地": 12}.get(h, 12)
    detail_layout = TableLayout(1, 2, len(headers) + 1)
    configure_table(ws4, detail_layout)
    sw.header_row(ws4, headers)
    for i, r in enumerate(rows):
        if progress and i % 20000 == 0:
            progress(text=f"写入 Excel {i}/{n}")
        location = r.get("location")
        location = location.replace("IP属地：", "") if isinstance(location, str) else location
        sw.append([None, value(i + 1, xlsx_mod.CellKind.INTEGER, "序号"),
                   field(r, "rpid", xlsx_mod.CellKind.ID, "rpid"),
                   value("主楼" if r.get("is_main") else "楼中楼", xlsx_mod.CellKind.TEXT, "层级"),
                   field(r, "uname", xlsx_mod.CellKind.TEXT, "用户昵称"),
                   field(r, "mid", xlsx_mod.CellKind.ID, "用户mid"),
                   field(r, "level", xlsx_mod.CellKind.INTEGER, "等级"),
                   vip_field(r),
                   field(r, "sex", xlsx_mod.CellKind.TEXT, "性别"),
                   wrap_cell(sw, field(r, "message", xlsx_mod.CellKind.TEXT, "评论内容")),
                   field(r, "like", xlsx_mod.CellKind.INTEGER, "点赞数"),
                   field(r, "rcount", xlsx_mod.CellKind.INTEGER, "楼中楼数"),
                   safe_time(r.get("ctime"), "发布时间"),
                   value(location, xlsx_mod.CellKind.TEXT, "IP属地")])
    finish_table(ws4, detail_layout, n)
    uid_rows, missing_mid, invalid_mid = _build_uid_rows(rows)
    uid_ws = wb.create_sheet("UID列表")
    sw.ws = uid_ws
    uid_ws.column_dimensions["A"].width = 3
    uid_ws.column_dimensions["B"].width = 22
    uid_ws.column_dimensions["C"].width = 20
    uid_ws.column_dimensions["D"].width = 10
    uid_layout = TableLayout(3, 2, 4)
    configure_table(uid_ws, uid_layout)
    sw.title_row(uid_ws, "UID列表", 4)
    sw.header_row(uid_ws, ["UID", "用户昵称", "评论数"])
    for entry in uid_rows:
        sw.append([
            None,
            xlsx_mod.cell_value(entry["uid"], xlsx_mod.CellKind.ID),
            field(entry["row"], "uname", xlsx_mod.CellKind.TEXT, "用户昵称"),
            value(entry["count"], xlsx_mod.CellKind.INTEGER, "评论数"),
        ])
    finish_table(uid_ws, uid_layout, len(uid_rows))

    stats = stats or {}
    def q(category, item, value, unit, note, presentation_state=None):
        return QualityItem(category, item, value, unit, note, presentation_state)

    def q_state(raw, label):
        if raw is None:
            return xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE, note=label)
        return xlsx_mod.checked_cell_value(raw, xlsx_mod.CellKind.INTEGER, note=label)

    invalid_time = 0
    times = []
    for row in rows:
        raw = row.get("ctime")
        if raw is None:
            continue
        try:
            parsed = xlsx_mod.unix_seconds_to_excel_datetime(raw)
        except (TypeError, ValueError):
            invalid_time += 1
            continue
        if parsed is not None:
            times.append(parsed)

    def parsed_time_value(value, label):
        if value is None:
            return xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_RETURNED, note=label)
        return xlsx_mod.checked_cell_value(value, xlsx_mod.CellKind.DATETIME, note=label)

    if claimed_status == "not_returned":
        claimed_cell = xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_RETURNED,
                                            note="接口未返回")
        coverage_cell = xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_RETURNED,
                                             note="接口未返回")
        coverage_note = "接口未返回声明数量"
    elif claimed_status == "invalid":
        claimed_cell = xlsx_mod.cell_value(None, xlsx_mod.CellKind.MISSING,
                                            note="声明数量格式异常")
        coverage_cell = xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE,
                                             note="声明数量格式异常，覆盖率不适用")
        coverage_note = "声明数量格式异常，覆盖率不适用"
    else:
        claimed_cell = xlsx_mod.checked_cell_value(
            claimed_value, xlsx_mod.CellKind.INTEGER, note="声明数量格式异常")
        coverage_cell = percent_cell(n, claimed_value,
                                     note="声明数量为 0，覆盖率不适用")
        coverage_note = "按接口声明数量与实际唯一 rpid 比较"

    quality = [
        q("记录", "候选记录总数", q_state(stats.get("jsonl_candidate_records"), "上游未传入候选记录统计"), "条", "JSONL 重读阶段的非空候选行"),
        q("记录", "实际写入主明细表的有效记录数", xlsx_mod.checked_cell_value(n, xlsx_mod.CellKind.INTEGER), "条", "按 rpid 去重并成功写入全量评论"),
        q("记录", "检测到的重复数", q_state(stats.get("jsonl_duplicate_rows"), "上游未传入重复统计"), "条", "JSONL 重读阶段重复 rpid"),
        q("记录", "解析失败数", q_state(stats.get("jsonl_parse_failures"), "上游未传入解析失败统计"), "条", "JSONL 无法解析的候选行"),
        q("记录", "缺失 rpid 数", q_state(stats.get("jsonl_missing_rpid"), "上游未传入缺失键统计"), "条", "无法恢复的唯一键缺失不猜测"),
        q("字段", "发布时间非法值数量", xlsx_mod.checked_cell_value(invalid_time, xlsx_mod.CellKind.INTEGER), "条", "无法转换为 Asia/Shanghai 时间", "error"),
        q("字段", "有效时间最早值", parsed_time_value(min(times) if times else None, "无可用时间"), "时间", "仅统计可转换的 ctime"),
        q("字段", "有效时间最晚值", parsed_time_value(max(times) if times else None, "无可用时间"), "时间", "仅统计可转换的 ctime"),
        q("UID", "有效 UID 数", xlsx_mod.checked_cell_value(len(uid_rows), xlsx_mod.CellKind.INTEGER), "个", "按评论行 mid 去重，按首次出现顺序输出"),
        q("UID", "缺失 mid 数", xlsx_mod.checked_cell_value(missing_mid, xlsx_mod.CellKind.INTEGER), "条", "mid 缺失或为 None，不进入 UID列表"),
        q("UID", "非法 mid 数", xlsx_mod.checked_cell_value(invalid_mid, xlsx_mod.CellKind.INTEGER), "条", "mid 不是 1～20 位正整数，或无法安全转换，不进入 UID列表"),
        q("UID", "排除 mid 总数", xlsx_mod.checked_cell_value(missing_mid + invalid_mid, xlsx_mod.CellKind.INTEGER), "条", "缺失与非法 mid 的合计"),
        q("接口", "接口声称评论数", claimed_cell, "条", "缺失、0、正数保持不同语义"),
        q("记录", "实际获得数量", xlsx_mod.checked_cell_value(n, xlsx_mod.CellKind.INTEGER), "条", "主明细表有效唯一 rpid 数"),
        q("覆盖", "覆盖率", coverage_cell, "百分比", coverage_note),
        q("任务", "是否取消", xlsx_mod.checked_cell_value(bool(stats.get("cancelled", False)), xlsx_mod.CellKind.BOOLEAN), "状态", "来自结构化 crawler stats", "stop"),
        q("任务", "是否截断", xlsx_mod.checked_cell_value(bool(stats.get("aborted", False)), xlsx_mod.CellKind.BOOLEAN), "状态", "取消、限页或异常中断由 aborted 区分", "warning"),
        q("任务", "是否预算到限", xlsx_mod.checked_cell_value(stats.get("stopped_reason") == "budget_reached", xlsx_mod.CellKind.BOOLEAN), "状态", "来自 stopped_reason", "stop"),
        q("任务", "是否部分成功", xlsx_mod.checked_cell_value(bool(rows) and bool(stats.get("aborted") or stats.get("error")), xlsx_mod.CellKind.BOOLEAN), "状态", "已有有效报告但抓取阶段未完整结束", "warning"),
    ]
    quality.extend(primary_key_quality(
        "主键", "rpid", denominator=stats.get("jsonl_candidate_records"),
        missing=stats.get("jsonl_missing_rpid"),
        invalid=stats.get("jsonl_invalid_rpid"),
        duplicates=stats.get("jsonl_duplicate_rows"),
        dedup_discarded=stats.get("jsonl_duplicate_rows"),
        remaining_conflicts=stats.get("jsonl_remaining_rpid_conflicts"),
        source_note="来自评论 JSONL 重读统计；不扫描 XLSX 明细表",
    ))
    fields = []
    def fd(sheet, display, stable, dtype, metric, missing="字段缺失", example=None):
        fields.append(FieldDefinition(sheet, display, stable, dtype, "", "是", "评论抓取/本地归一化", metric,
                                      example or xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE), missing))
    for display, stable, dtype, metric in (
        ("rpid", "rpid", "ID", "按 rpid 唯一"), ("层级", "is_main", "布尔", "主楼/楼中楼"),
        ("用户昵称", "uname", "文本", "原始用户显示名"), ("用户mid", "mid", "ID", "用户标识"),
        ("等级", "level", "整数", "接口返回等级"), ("大会员", "vip", "布尔", "接口返回会员状态"),
        ("性别", "sex", "文本", "接口返回性别"), ("评论内容", "message", "文本", "用户来源文本"),
        ("点赞数", "like", "整数", "接口返回计数"), ("楼中楼数", "rcount", "整数", "接口返回计数"),
        ("发布时间", "ctime", "日期时间", "Unix 秒转换为 Asia/Shanghai"), ("IP属地", "location", "文本", "接口返回属地"),
    ):
        fd("全量评论", display, stable, dtype, metric)
    fd("全量评论", "序号", "row_number", "整数", "本次输出顺序")
    for display, stable, dtype, metric in (
        ("指标", "metric", "文本", "概览稳定指标标签"),
        ("数值", "value", "数值", "概览指标值；状态值为空时看单元格批注"),
        ("说明", "note", "文本", "概览指标说明"),
        ("排名", "rank", "整数", "各统计块内的本次排名"),
        ("高频词", "top_word", "文本", "本地 n-gram 高频词"),
        ("出现次数", "occurrences", "整数", "高频词出现次数"),
        ("B站小表情", "bili_emoji", "文本", "评论中的 B 站小表情"),
        ("评论数", "count", "整数", "表情对应的评论条数"),
        ("emoji", "unicode_emoji", "文本", "评论中的 Unicode emoji"),
        ("", "reserved", "文本", "并排统计块的布局占位列"),
    ):
        fd("统计概览", display, stable, dtype, metric)
    for display, stable, dtype, metric in (
        ("日期", "date", "文本", "按评论发布时间归属日期"),
        ("评论条数", "count", "整数", "按日期或小时的评论条数"),
        ("占比", "ratio", "百分比", "评论条数/有效评论总数；总数为 0 时不适用"),
        ("小时", "hour", "文本", "按评论发布时间的小时区间"),
    ):
        fd("分布统计", display, stable, dtype, metric)
    for display, stable, dtype, metric in (
        ("发布时间", "published_at", "日期时间", "接口返回发布时间"),
        ("声称评论总数", "claimed_comment_count", "整数", "接口声明值；缺失与 0 保持不同语义"),
        ("实际抓取", "fetched_count", "整数", "有效 rpid 评论数"),
        ("主楼 / 楼中楼", "main_sub_counts", "文本", "主楼与楼中楼数量"),
        ("独立用户", "unique_users", "整数", "按用户 mid 去重"),
        ("时间跨度", "time_span_hours", "小数", "有效评论时间跨度"),
        ("0赞占比", "zero_like_ratio", "百分比", "0 赞评论数/有效评论数"),
        ("盖楼抗议（🧱🔨）", "build_count", "整数", "规则词命中评论数"),
        ("要求道歉 / 护官反呛", "apology_defense_counts", "文本", "规则词命中计数"),
        ("停氪退款 / 提及原神 / 疑似引流", "topic_counts", "文本", "规则词命中计数"),
    ):
        fd("统计概览", display, stable, dtype, metric)
    for display, stable, dtype, metric in (
        ("排名", "rank", "整数", "按点赞数降序的本次 Top100 顺序"),
        ("层级", "is_main", "文本", "主楼或楼中楼"),
        ("用户", "uname", "文本", "原始用户显示名"),
        ("点赞", "like", "整数", "接口返回点赞计数"),
        ("内容", "message", "文本", "用户来源文本"),
        ("楼中楼数", "rcount", "整数", "接口返回楼中楼计数"),
    ):
        fd("点赞Top100", display, stable, dtype, metric)
    fields.extend((
        FieldDefinition(
            "UID列表", "UID", "uid", "ID", "", "是", "评论行的 member.mid / mid",
            "按首次出现顺序去重；每个 UID 一行，评论数为有效 mid 的评论条数",
            xlsx_mod.cell_value("12345678901234567890", xlsx_mod.CellKind.ID),
            "缺失或非法 mid 不生成 UID",
        ),
        FieldDefinition(
            "UID列表", "用户昵称", "uname", "文本", "", "是", "首次出现该 UID 的评论行",
            "保留该 UID 首次出现时的用户昵称",
            xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE),
            "字段缺失",
        ),
        FieldDefinition(
            "UID列表", "评论数", "comment_count", "整数", "条", "否", "评论行本地分组",
            "按有效 mid 统计评论行数量",
            xlsx_mod.cell_value(1, xlsx_mod.CellKind.INTEGER),
            "不适用",
        ),
    ))
    metadata = make_metadata(
        tool="评论", report_type="评论分析", parameters={
            "目标类型": xlsx_mod.cell_value(meta.get("target_type", "未知"), xlsx_mod.CellKind.TEXT),
            "规范化 oid": xlsx_mod.cell_value(str(meta.get("normalized_oid", "")), xlsx_mod.CellKind.ID),
            "分页上限": xlsx_mod.cell_value(max(0, int(meta.get("max_pages", 0) or 0)), xlsx_mod.CellKind.INTEGER),
            "TLS gRPC": xlsx_mod.cell_value(bool(meta.get("use_tls_grpc", False)), xlsx_mod.CellKind.BOOLEAN),
        }, parameter_allowlist=("目标类型", "规范化 oid", "分页上限", "TLS gRPC"),
        quality_items=quality, fields=fields)
    write_metadata_sheets(_CommentMetadataWorkbookView(wb), metadata)
    xlsx_mod.save_workbook_atomic(wb, out_path)
