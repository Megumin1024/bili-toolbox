# -*- coding: utf-8 -*-
"""本地报告中心的数据读取、对比和导出核心。

本模块只访问调用方提供的本地路径。它不导入网络客户端、采集流水线、
监控服务或 TaskRunner；JSONL 流式读取，XLSX 使用 openpyxl read_only 模式。
"""
from __future__ import annotations

import hashlib
import json
import math
import os
import re
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator

from openpyxl import load_workbook

from core import diagnostics, task_history
from core import output as output_mod
from core import xlsx as xlsx_mod
from core.redact import EXPORT_SENSITIVE_KEY_RE as _SENSITIVE_KEY_RE
from core.xlsx_metadata import (
    FIELDS_SHEET,
    FieldDefinition,
    QualityItem,
    QUALITY_SHEET,
    contains_export_sensitive_text,
    is_complete_metadata_workbook,
    make_metadata,
    primary_key_quality,
    write_metadata_sheets,
)
from core.xlsx_presentation import (
    TableLayout,
    append_sheet_directory,
    configure_table,
    finish_table,
    wrap_cell,
)


KIND_COMMENT = "comment"
KIND_VIDEO = "video_snapshot"
KIND_MONITOR = "monitor_history"
KIND_UNKNOWN = "unknown"

FORMAT_JSONL = "JSONL"
FORMAT_XLSX = "XLSX"

STATUS_READABLE = "readable"
STATUS_INCOMPLETE = "incomplete"
STATUS_DAMAGED = "damaged"
STATUS_UNREADABLE = "unreadable"
STATUS_UNCOMPARABLE = "uncomparable"

PREVIEW_LIMIT = 200
JSON_DETECT_LIMIT = 100
SUPPORTED_SUFFIXES = {".jsonl", ".xlsx"}

# The existing producers use int(time.time()). Excel exports use
# datetime.fromtimestamp(), which is a naive value in the application local
# zone. The product's data contract is Asia/Shanghai, so the internal
# representation keeps that named +08:00 zone explicit on every machine.
LOCAL_TZ = timezone(timedelta(hours=8), "Asia/Shanghai")

TIME_FIELDS = {
    KIND_COMMENT: "ctime",
    KIND_VIDEO: "fetched_at",
    KIND_MONITOR: "ts",
}
KEY_FIELDS = {
    KIND_COMMENT: ("rpid",),
    KIND_VIDEO: ("bvid", "fetched_at"),
    KIND_MONITOR: ("bvid", "ts"),
}
KIND_LABELS = {
    KIND_COMMENT: "评论",
    KIND_VIDEO: "视频快照",
    KIND_MONITOR: "监控历史",
    KIND_UNKNOWN: "未知结构",
}
STATUS_LABELS = {
    STATUS_READABLE: "可读取",
    STATUS_INCOMPLETE: "字段不完整",
    STATUS_DAMAGED: "损坏",
    STATUS_UNREADABLE: "不可读取",
    STATUS_UNCOMPARABLE: "不可用于对比",
}

NUMERIC_CANDIDATES = {
    KIND_COMMENT: ("like", "rcount", "level", "mid"),
    KIND_VIDEO: ("view", "danmaku", "reply", "favorite", "coin", "share", "like", "duration"),
    KIND_MONITOR: ("view", "danmaku", "reply", "favorite", "coin", "share", "like", "online"),
}

COMMENT_HEADER_MAP = {
    "rpid": "rpid", "用户昵称": "uname", "用户mid": "mid", "等级": "level",
    "大会员": "vip", "性别": "sex", "评论内容": "message", "点赞数": "like",
    "楼中楼数": "rcount", "发布时间": "ctime", "IP属地": "location",
}
VIDEO_HEADER_MAP = {
    "BV号": "bvid", "标题": "title", "UP主": "owner", "分区": "tname",
    "时长(秒)": "duration", "发布时间": "pubdate", "播放": "view",
    "弹幕": "danmaku", "评论": "reply", "点赞": "like", "投币": "coin",
    "收藏": "favorite", "分享": "share",
}
# 敏感键判定已收拢为 core.redact.EXPORT_SENSITIVE_KEY_RE（原串搬移，见顶部导入）。

Cancel = Callable[[], bool]
Progress = Callable[..., None]


class ReportCenterError(RuntimeError):
    """报告中心可向界面展示的本地数据错误。"""


class SourceChangedError(ReportCenterError):
    """源文件在导出期间发生变化。"""


class ComparisonCancelled(ReportCenterError):
    """用户取消了后台对比。"""


@dataclass(frozen=True)
class Fingerprint:
    path: str
    sha256: str
    size: int
    mtime_ns: int

    def as_dict(self) -> dict[str, Any]:
        return {"path": self.path, "sha256": self.sha256, "size": self.size, "mtime_ns": self.mtime_ns}


@dataclass
class Source:
    path: str
    file_name: str = ""
    fmt: str = ""
    kind: str = KIND_UNKNOWN
    status: str = STATUS_UNREADABLE
    message: str = ""
    exists: bool = False
    size: int = 0
    mtime_ns: int = 0
    sha256: str = ""
    record_count: int = 0
    time_start: datetime | None = None
    time_end: datetime | None = None
    time_field: str = ""
    fields: tuple[str, ...] = ()
    metrics: tuple[str, ...] = ()
    key_fields: tuple[str, ...] = ()
    can_compare: bool = False
    duplicate_keys: int = 0
    key_missing: int | None = None
    key_invalid: int | None = None
    key_dedup_discarded: int | None = None
    key_remaining_conflicts: int | None = None
    invalid_time: int = 0
    issues: list[str] = field(default_factory=list)
    sheet_name: str = ""
    header_row: int = 0
    header_map: dict[str, int] = field(default_factory=dict)
    original_headers: tuple[str, ...] = ()
    discovered_from: str = ""
    preview_rows: list[dict[str, Any]] = field(default_factory=list, repr=False)

    @property
    def kind_label(self) -> str:
        return KIND_LABELS.get(self.kind, KIND_LABELS[KIND_UNKNOWN])

    @property
    def status_label(self) -> str:
        return STATUS_LABELS.get(self.status, self.status)

    @property
    def time_range_label(self) -> str:
        if self.time_start is None or self.time_end is None:
            return "暂无可靠时间"
        return f"{format_datetime(self.time_start)} ~ {format_datetime(self.time_end)}"

    def fingerprint(self) -> Fingerprint | None:
        if not self.exists or not self.sha256:
            return None
        return Fingerprint(self.path, self.sha256, self.size, self.mtime_ns)

    def as_dict(self, redact_path: bool = False) -> dict[str, Any]:
        shown_path = diagnostics.sanitize_text(self.path) if redact_path else self.path
        return {
            "path": shown_path,
            "file_name": diagnostics.sanitize_text(self.file_name),
            "format": self.fmt,
            "kind": self.kind,
            "kind_label": self.kind_label,
            "status": self.status,
            "status_label": self.status_label,
            "message": diagnostics.sanitize_text(self.message),
            "exists": self.exists,
            "size": self.size,
            "mtime_ns": self.mtime_ns,
            "sha256": self.sha256,
            "records": self.record_count,
            "time_start": format_datetime(self.time_start) if self.time_start else None,
            "time_end": format_datetime(self.time_end) if self.time_end else None,
            "time_field": self.time_field,
            "fields": list(self.fields),
            "metrics": list(self.metrics),
            "key_fields": list(self.key_fields),
            "can_compare": self.can_compare,
            "issues": [diagnostics.sanitize_text(item) for item in self.issues],
            "discovered_from": diagnostics.sanitize_text(self.discovered_from),
        }


@dataclass
class ReadResult:
    source: Source
    preview_rows: list[dict[str, Any]] = field(default_factory=list)


@dataclass
class MetricChange:
    field: str
    a_value: float | int | None
    b_value: float | int | None
    absolute_change: float | int | None
    percent_change: float | None
    compared_records: int = 0
    missing_a: int = 0
    missing_b: int = 0
    note: str = ""

    def as_dict(self) -> dict[str, Any]:
        return {"field": self.field, "a_value": self.a_value, "b_value": self.b_value,
                "absolute_change": self.absolute_change, "percent_change": self.percent_change,
                "compared_records": self.compared_records, "missing_a": self.missing_a,
                "missing_b": self.missing_b, "note": diagnostics.sanitize_text(self.note)}


@dataclass
class ComparisonResult:
    left: Source
    right: Source
    compatible: bool
    matching_mode: str
    message: str
    a_records: int = 0
    b_records: int = 0
    common_records: int | None = None
    only_a_records: int | None = None
    only_b_records: int | None = None
    changed_records: int | None = None
    metric_changes: list[MetricChange] = field(default_factory=list)
    time_range_a: tuple[datetime | None, datetime | None] = (None, None)
    time_range_b: tuple[datetime | None, datetime | None] = (None, None)
    missing_fields_a: tuple[str, ...] = ()
    missing_fields_b: tuple[str, ...] = ()
    row_changes: list[dict[str, Any]] = field(default_factory=list)
    fingerprint_a: Fingerprint | None = None
    fingerprint_b: Fingerprint | None = None

    def as_dict(self, redact_paths: bool = True) -> dict[str, Any]:
        return {
            "left": self.left.as_dict(redact_paths), "right": self.right.as_dict(redact_paths),
            "compatible": self.compatible, "matching_mode": self.matching_mode,
            "message": diagnostics.sanitize_text(self.message), "a_records": self.a_records,
            "b_records": self.b_records, "common_records": self.common_records,
            "only_a_records": self.only_a_records, "only_b_records": self.only_b_records,
            "changed_records": self.changed_records,
            "metric_changes": [item.as_dict() for item in self.metric_changes],
            "time_range_a": [format_datetime(item) if item else None for item in self.time_range_a],
            "time_range_b": [format_datetime(item) if item else None for item in self.time_range_b],
            "missing_fields_a": list(self.missing_fields_a), "missing_fields_b": list(self.missing_fields_b),
            "fingerprint_a": self.fingerprint_a.as_dict() if self.fingerprint_a else None,
            "fingerprint_b": self.fingerprint_b.as_dict() if self.fingerprint_b else None,
            "row_changes": [sanitize_public_row(item) for item in self.row_changes],
        }


@dataclass
class PeriodStats:
    start: datetime
    end: datetime
    sample_count: int = 0
    valid_count: int = 0
    missing_data: int = 0
    missing_time: int = 0
    start_value: float | int | None = None
    end_value: float | int | None = None
    average: float | None = None
    maximum: float | int | None = None
    minimum: float | int | None = None

    @property
    def absolute_change(self):
        if self.start_value is None or self.end_value is None:
            return None
        return self.end_value - self.start_value

    @property
    def percent_change(self):
        if self.start_value is None or self.end_value is None:
            return None
        if self.start_value == 0:
            return 0.0 if self.end_value == 0 else None
        return (self.end_value - self.start_value) / self.start_value * 100

    def as_dict(self) -> dict[str, Any]:
        return {"start": format_datetime(self.start), "end": format_datetime(self.end),
                "sample_count": self.sample_count, "valid_count": self.valid_count,
                "missing_data": self.missing_data, "missing_time": self.missing_time,
                "start_value": self.start_value, "end_value": self.end_value,
                "average": self.average, "maximum": self.maximum, "minimum": self.minimum,
                "absolute_change": self.absolute_change, "percent_change": self.percent_change}


@dataclass
class PeriodComparisonResult:
    source: Source
    metric: str
    period_a: PeriodStats
    period_b: PeriodStats
    values_a: list[tuple[datetime, float | int]] = field(default_factory=list)
    values_b: list[tuple[datetime, float | int]] = field(default_factory=list)
    message: str = ""
    fingerprint: Fingerprint | None = None

    def as_dict(self, redact_path: bool = True) -> dict[str, Any]:
        return {"source": self.source.as_dict(redact_path), "metric": self.metric,
                "period_a": self.period_a.as_dict(), "period_b": self.period_b.as_dict(),
                "values_a": [[format_datetime(t), v] for t, v in self.values_a],
                "values_b": [[format_datetime(t), v] for t, v in self.values_b],
                "message": diagnostics.sanitize_text(self.message),
                "fingerprint": self.fingerprint.as_dict() if self.fingerprint else None}


def _cancelled(cancel: Cancel | None) -> bool:
    return bool(cancel and cancel())


def _safe_error(exc: Exception) -> str:
    return diagnostics.sanitize_text(f"{type(exc).__name__}: {exc}")[:500]


def format_datetime(value: datetime | None) -> str:
    if value is None:
        return ""
    return _to_internal_datetime(value).isoformat(timespec="seconds")


def _to_internal_datetime(value: datetime) -> datetime:
    if value.tzinfo is None:
        return value.replace(tzinfo=LOCAL_TZ)
    return value.astimezone(LOCAL_TZ)


def parse_user_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return _to_internal_datetime(value)
    if isinstance(value, date):
        return datetime.combine(value, time.min, tzinfo=LOCAL_TZ)
    text = str(value or "").strip()
    if not text:
        return None
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return None
    return _to_internal_datetime(parsed)


def parse_epoch_seconds(value: Any) -> datetime | None:
    """Only accept plausible Unix seconds; milliseconds/microseconds are rejected."""
    if isinstance(value, bool) or value is None:
        return None
    try:
        if isinstance(value, str):
            text = value.strip()
            if len(text) > xlsx_mod.MAX_NUMERIC_TEXT_LENGTH:
                return None
            number = Decimal(text)
        elif isinstance(value, int):
            if value < 100_000_000 or value >= 10_000_000_000:
                return None
            number = Decimal(value)
        elif isinstance(value, float):
            if not math.isfinite(value):
                return None
            number = Decimal(repr(value))
        else:
            return None
    except (InvalidOperation, TypeError, ValueError, OverflowError):
        return None
    if not number.is_finite() or number < 100_000_000 or number >= 10_000_000_000:
        return None
    try:
        converted = xlsx_mod.unix_seconds_to_excel_datetime(value)
        return converted.replace(tzinfo=LOCAL_TZ) if converted is not None else None
    except (TypeError, OverflowError, OSError, ValueError):
        return None


def parse_source_time(value: Any, *, epoch: bool) -> datetime | None:
    if isinstance(value, datetime):
        return _to_internal_datetime(value)
    if isinstance(value, date):
        return datetime.combine(value, time.min, tzinfo=LOCAL_TZ)
    if epoch:
        return parse_epoch_seconds(value)
    return parse_user_datetime(value)


def numeric_value(value: Any) -> int | float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return value if math.isfinite(float(value)) else None
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if not text or text.startswith("="):
            return None
        try:
            number = float(text)
        except ValueError:
            return None
        if not math.isfinite(number):
            return None
        return int(number) if number.is_integer() else number
    return None


def _canonical_header(value: Any) -> str:
    return str(value or "").strip()


def _kind_from_keys(keys: set[str]) -> str:
    if "rpid" in keys and bool(keys & {"message", "uname", "parent", "root"}):
        return KIND_COMMENT
    if "bvid" in keys and "fetched_at" in keys:
        return KIND_VIDEO
    if "bvid" in keys and "ts" in keys:
        return KIND_MONITOR
    return KIND_UNKNOWN


def _kind_from_headers(headers: set[str], sheet_name: str = "") -> tuple[str, dict[str, str]]:
    if sheet_name == "全量评论" or ("rpid" in headers and "评论内容" in headers):
        return KIND_COMMENT, COMMENT_HEADER_MAP
    if sheet_name == "视频总表" or ("BV号" in headers and "标题" in headers):
        return KIND_VIDEO, VIDEO_HEADER_MAP
    if "bvid" in headers and "ts" in headers:
        return KIND_MONITOR, {item: item for item in MONITOR_FIELDS}
    return KIND_UNKNOWN, {}


def _path(path: str | Path) -> Path:
    return Path(path).expanduser().resolve(strict=False)


def snapshot_source(path: str | Path) -> Fingerprint:
    source = _path(path)
    digest = hashlib.sha256()
    with source.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    stat = source.stat()
    return Fingerprint(str(source), digest.hexdigest(), stat.st_size, stat.st_mtime_ns)


def _metadata(path: str | Path) -> Source:
    source = _path(path)
    fmt = FORMAT_JSONL if source.suffix.lower() == ".jsonl" else FORMAT_XLSX if source.suffix.lower() == ".xlsx" else source.suffix.lstrip(".").upper()
    return Source(path=str(source), file_name=source.name, fmt=fmt, exists=source.is_file())


def _detect_jsonl(path: Path, cancel: Cancel | None = None,
                  progress: Progress | None = None) -> tuple[str, set[str], list[str]]:
    votes: dict[str, int] = {KIND_COMMENT: 0, KIND_VIDEO: 0, KIND_MONITOR: 0}
    fields: set[str] = set()
    with path.open("r", encoding="utf-8-sig", errors="strict") as handle:
        for line_no, raw in enumerate(handle, 1):
            if _cancelled(cancel):
                break
            if line_no > JSON_DETECT_LIMIT:
                break
            if not raw.strip():
                continue
            try:
                item = json.loads(raw)
            except (json.JSONDecodeError, ValueError):
                continue
            if not isinstance(item, dict):
                continue
            keys = {str(key) for key in item}
            fields.update(keys)
            kind = _kind_from_keys(keys)
            if kind in votes:
                votes[kind] += 1
            if progress and line_no % 500 == 0:
                progress(text=f"识别 {path.name} · {line_no:,} 行", row=line_no)
    kind = max(votes, key=votes.get) if any(votes.values()) else KIND_UNKNOWN
    return kind, fields, []


def _xlsx_layout(path: Path, cancel: Cancel | None = None,
                 progress: Progress | None = None) -> tuple[str, str, int, dict[str, int], tuple[str, ...], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        preferred = None
        for name in ("全量评论", "视频总表"):
            if name in workbook.sheetnames:
                preferred = workbook[name]
                break
        sheets = [preferred] if preferred is not None else list(workbook.worksheets)
        metadata_rows = {
            name: tuple(
                tuple(row)
                for row in workbook[name].iter_rows(max_row=40, values_only=True)
            )
            for name in (QUALITY_SHEET, FIELDS_SHEET)
            if name in workbook.sheetnames
        }
        complete_metadata = is_complete_metadata_workbook(metadata_rows)
        for worksheet in sheets:
            if _cancelled(cancel):
                break
            if complete_metadata and worksheet.title in (QUALITY_SHEET, FIELDS_SHEET):
                continue
            header_row_no = 0
            header_values: tuple[Any, ...] = ()
            for index, row in enumerate(worksheet.iter_rows(values_only=True), 1):
                if _cancelled(cancel):
                    break
                if any(value not in (None, "") for value in row):
                    header_row_no = index
                    header_values = tuple(row)
                    break
                if index >= 30:
                    break
            if progress and header_row_no:
                progress(text=f"识别 {path.name} · 工作表 {worksheet.title}", row=header_row_no)
            if _cancelled(cancel):
                break
            if not header_row_no:
                continue
            headers = tuple(_canonical_header(value) for value in header_values)
            kind, mapping = _kind_from_headers({value for value in headers if value}, worksheet.title)
            if kind == KIND_UNKNOWN and preferred is not None:
                continue
            header_map: dict[str, int] = {}
            for index, header in enumerate(headers):
                canonical = mapping.get(header)
                if canonical:
                    header_map[canonical] = index
            return kind, worksheet.title, header_row_no, header_map, headers, []
        return KIND_UNKNOWN, "", 0, {}, (), ["没有找到可识别的工作表表头"]
    finally:
        workbook.close()


def _schema(source: Source, cancel: Cancel | None = None,
            progress: Progress | None = None) -> Source:
    path = Path(source.path)
    if not source.exists:
        source.status = STATUS_UNREADABLE
        source.message = "文件不存在"
        return source
    try:
        if source.fmt == FORMAT_JSONL:
            kind, fields, issues = _detect_jsonl(path, cancel, progress)
            source.kind = kind
            source.fields = tuple(sorted(fields))
            source.issues.extend(issues)
        elif source.fmt == FORMAT_XLSX:
            kind, sheet, header_row, header_map, headers, issues = _xlsx_layout(path, cancel, progress)
            source.kind = kind
            source.sheet_name = sheet
            source.header_row = header_row
            source.header_map = header_map
            source.original_headers = headers
            source.fields = tuple(sorted(header_map))
            source.issues.extend(issues)
        else:
            source.status = STATUS_UNCOMPARABLE
            source.message = "仅支持 JSONL 和 XLSX"
            return source
    except Exception as exc:
        source.status = STATUS_DAMAGED
        source.message = _safe_error(exc)
        source.issues.append(source.message)
        return source
    source.time_field = TIME_FIELDS.get(source.kind, "")
    source.key_fields = KEY_FIELDS.get(source.kind, ())
    return source


def _convert_xlsx_value(canonical: str, value: Any) -> Any:
    if canonical == "is_main":
        return str(value or "").strip() in {"主楼", "main", "True", "true", "1"}
    if canonical == "vip":
        return value in {True, 1, "是", "true", "True", "1"}
    return value


def _iter_jsonl(source: Source, cancel: Cancel | None = None) -> Iterator[tuple[dict[str, Any], int]]:
    with Path(source.path).open("r", encoding="utf-8-sig", errors="strict") as handle:
        for line_no, raw in enumerate(handle, 1):
            if _cancelled(cancel):
                return
            if not raw.strip():
                continue
            try:
                item = json.loads(raw)
            except (json.JSONDecodeError, ValueError) as exc:
                source.issues.append(f"JSONL 第 {line_no} 行：{_safe_error(exc)}")
                source.status = STATUS_DAMAGED
                continue
            if isinstance(item, dict):
                yield dict(item), line_no


def _iter_xlsx(source: Source, cancel: Cancel | None = None) -> Iterator[tuple[dict[str, Any], int]]:
    workbook = load_workbook(source.path, read_only=True, data_only=False)
    try:
        worksheet = workbook[source.sheet_name]
        for row_no, row in enumerate(worksheet.iter_rows(min_row=source.header_row + 1, values_only=True), source.header_row + 1):
            if _cancelled(cancel):
                return
            values = tuple(row)
            if not any(value not in (None, "") for value in values):
                continue
            raw = {header: values[index] if index < len(values) else None
                   for index, header in enumerate(source.original_headers) if header}
            item = {canonical: _convert_xlsx_value(canonical, values[index] if index < len(values) else None)
                    for canonical, index in source.header_map.items()}
            item["_raw_fields"] = raw
            yield item, row_no
    finally:
        workbook.close()


def _iter_records(source: Source, cancel: Cancel | None = None) -> Iterator[tuple[dict[str, Any], int]]:
    if source.fmt == FORMAT_JSONL:
        yield from _iter_jsonl(source, cancel)
    elif source.fmt == FORMAT_XLSX and source.sheet_name:
        yield from _iter_xlsx(source, cancel)


def _public_fields(row: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in row.items() if not key.startswith("_")}


def sanitize_public_row(row: dict[str, Any]) -> dict[str, Any]:
    raw = row.get("_raw_fields") if isinstance(row, dict) else None
    values = raw if isinstance(raw, dict) else _public_fields(row)
    safe: dict[str, Any] = {}
    for key, value in values.items():
        key_text = str(key)
        if _SENSITIVE_KEY_RE.search(key_text):
            continue
        if contains_export_sensitive_text({key_text: value}):
            continue
        safe[key_text] = _json_value(value)
    return safe


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return _to_internal_datetime(value).isoformat(timespec="seconds")
    if isinstance(value, (date, time)):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return diagnostics.sanitize_text(value)


def _xlsx_safe_value(value: Any) -> Any:
    """只清洗控制字符；公式防护由 ``SheetWriter`` 的 TEXT 类型完成。"""
    if isinstance(value, str):
        return xlsx_mod.clean_text(value)
    if isinstance(value, (datetime, date, time, int, float, bool)) or value is None:
        return value
    return _xlsx_safe_value(diagnostics.sanitize_text(value))


_REPORT_ID_FIELDS = {
    KIND_COMMENT: {"rpid", "mid", "用户mid"},
    KIND_VIDEO: {"bvid", "aid", "owner_mid", "mid"},
    KIND_MONITOR: {"bvid", "aid", "owner_mid", "mid", "room_id", "uid"},
}
_REPORT_INTEGER_FIELDS = {
    KIND_COMMENT: {"like", "rcount", "level", "点赞数", "楼中楼数", "等级"},
    KIND_VIDEO: {"duration", "view", "danmaku", "reply", "favorite", "coin", "share", "like"},
    KIND_MONITOR: {"view", "danmaku", "reply", "favorite", "coin", "share", "like", "online"},
}
_REPORT_BOOLEAN_FIELDS = {KIND_COMMENT: {"vip", "大会员"}}
_REPORT_DATETIME_FIELDS = {
    KIND_COMMENT: {"ctime", "发布时间"},
    KIND_VIDEO: {"pubdate", "fetched_at", "发布时间"},
    KIND_MONITOR: {"ts"},
}
_REPORT_LONG_TEXT_FIELDS = {
    "message", "评论内容", "title", "标题", "text", "正文", "content", "内容",
    "note", "说明", "reason", "原因",
}


def _report_canonical_field(source: Source, field_name: str) -> str:
    maps = {
        KIND_COMMENT: COMMENT_HEADER_MAP,
        KIND_VIDEO: VIDEO_HEADER_MAP,
    }
    mapping = maps.get(source.kind, {})
    return mapping.get(field_name, field_name)


def _report_kind(source: Source, field_name: str, value: Any) -> xlsx_mod.CellKind:
    canonical = _report_canonical_field(source, field_name)
    if field_name in _REPORT_ID_FIELDS.get(source.kind, set()) or canonical in _REPORT_ID_FIELDS.get(source.kind, set()):
        return xlsx_mod.CellKind.ID
    if field_name in _REPORT_INTEGER_FIELDS.get(source.kind, set()) or canonical in _REPORT_INTEGER_FIELDS.get(source.kind, set()):
        return xlsx_mod.CellKind.INTEGER
    if field_name in _REPORT_BOOLEAN_FIELDS.get(source.kind, set()) or canonical in _REPORT_BOOLEAN_FIELDS.get(source.kind, set()):
        return xlsx_mod.CellKind.BOOLEAN
    if field_name in _REPORT_DATETIME_FIELDS.get(source.kind, set()) or canonical in _REPORT_DATETIME_FIELDS.get(source.kind, set()):
        return xlsx_mod.CellKind.DATETIME
    if isinstance(value, bool):
        return xlsx_mod.CellKind.BOOLEAN
    if isinstance(value, int):
        return xlsx_mod.CellKind.INTEGER
    if isinstance(value, float):
        return xlsx_mod.CellKind.DECIMAL
    if isinstance(value, (datetime, date)):
        return xlsx_mod.CellKind.DATETIME
    return xlsx_mod.CellKind.TEXT


def _report_cell(source: Source, field_name: str, value: Any, *, present=True):
    if not present:
        return xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_RETURNED,
                                   note=f"{field_name}未返回")
    if value is None:
        return xlsx_mod.cell_value(None, xlsx_mod.CellKind.MISSING,
                                   note=f"{field_name}缺失")
    canonical = _report_canonical_field(source, field_name)
    kind = _report_kind(source, field_name, value)
    if kind is xlsx_mod.CellKind.DATETIME:
        if source.fmt == FORMAT_JSONL and canonical in {"ctime", "pubdate", "fetched_at"}:
            return xlsx_mod.unix_seconds_cell_value(value, note=f"{field_name}格式异常")
        if canonical == "ts" and not isinstance(value, (datetime, date)):
            try:
                parsed = xlsx_mod.local_text_to_excel_datetime(value)
            except (TypeError, ValueError):
                return xlsx_mod.checked_cell_value(value, xlsx_mod.CellKind.TEXT,
                                                   note=f"{field_name}格式异常")
            return xlsx_mod.checked_cell_value(parsed, kind, note=f"{field_name}格式异常")
        if isinstance(value, date) and not isinstance(value, datetime):
            value = datetime.combine(value, time.min)
    return xlsx_mod.checked_cell_value(_xlsx_safe_value(value), kind,
                                       note=f"{field_name}格式异常")


def _generic_cell(value, label="字段"):
    if value is None:
        return xlsx_mod.cell_value(None, xlsx_mod.CellKind.MISSING, note=f"{label}缺失")
    if contains_export_sensitive_text(value):
        return xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE,
                                   note="已排除敏感值")
    if isinstance(value, datetime):
        kind = xlsx_mod.CellKind.DATETIME
    elif isinstance(value, date):
        value = datetime.combine(value, time.min)
        kind = xlsx_mod.CellKind.DATETIME
    elif isinstance(value, bool):
        kind = xlsx_mod.CellKind.BOOLEAN
    elif isinstance(value, int):
        kind = xlsx_mod.CellKind.INTEGER
    elif isinstance(value, float):
        kind = xlsx_mod.CellKind.DECIMAL
    else:
        kind = xlsx_mod.CellKind.TEXT
    return xlsx_mod.checked_cell_value(_xlsx_safe_value(value), kind,
                                       note=f"{label}格式异常")


def _percent_point_cell(value, label="百分比变化"):
    """将报告内部的百分数点转换为 Excel 的 0~1 比例值。"""
    if value is None:
        return xlsx_mod.cell_value(None, xlsx_mod.CellKind.MISSING,
                                   note=f"{label}不可计算")
    return xlsx_mod.checked_cell_value(value / 100, xlsx_mod.CellKind.PERCENT,
                                       note=f"{label}格式异常")


def _summary_cell(value, label):
    if isinstance(value, xlsx_mod.CellValue):
        return value
    return _generic_cell(_json_value(value), label)


def _xlsx_safe_row(row: dict[str, Any]) -> dict[str, Any]:
    raw = row.get("_raw_fields") if isinstance(row, dict) else None
    values = raw if isinstance(raw, dict) else _public_fields(row)
    safe: dict[str, Any] = {}
    for key, value in values.items():
        key_text = str(key)
        if _SENSITIVE_KEY_RE.search(key_text):
            continue
        if contains_export_sensitive_text({key_text: value}):
            continue
        safe[key_text] = _xlsx_safe_value(value)
    return safe


def _row_time(source: Source, row: dict[str, Any]) -> datetime | None:
    if not source.time_field:
        return None
    return parse_source_time(row.get(source.time_field), epoch=source.fmt == FORMAT_JSONL)


def _row_key_state(source: Source, row: dict[str, Any]) -> tuple[tuple[Any, ...] | None, str]:
    if not source.key_fields:
        return None, "not_applicable"
    if any(field not in row or row.get(field) in (None, "")
           for field in source.key_fields):
        return None, "missing"
    values: list[Any] = []
    for field_name in source.key_fields:
        value = row.get(field_name)
        if field_name in {"ctime", "fetched_at", "ts"}:
            value = _row_time(source, row)
            if value is None:
                return None, "invalid"
            value = int(value.timestamp())
        elif field_name == "rpid":
            try:
                value = int(value)
            except (TypeError, ValueError, OverflowError):
                return None, "invalid"
        else:
            value = str(value).strip()
            if not value:
                return None, "missing"
        values.append(value)
    return tuple(values), "valid"


def _row_key(source: Source, row: dict[str, Any]) -> tuple[Any, ...] | None:
    return _row_key_state(source, row)[0]


def _required_fields(source: Source) -> set[str]:
    if source.kind == KIND_COMMENT:
        return {"rpid", "ctime"}
    if source.kind == KIND_VIDEO:
        return {"bvid", "fetched_at"}
    if source.kind == KIND_MONITOR:
        return {"bvid", "ts"}
    return set()


def _source_metrics(source: Source, seen_numeric: dict[str, bool]) -> tuple[str, ...]:
    return tuple(name for name in NUMERIC_CANDIDATES.get(source.kind, ()) if seen_numeric.get(name))


def read_source(source_or_path: Source | str | Path, preview_limit: int = PREVIEW_LIMIT,
                cancel: Cancel | None = None, progress: Progress | None = None) -> ReadResult:
    source = source_or_path if isinstance(source_or_path, Source) else _metadata(source_or_path)
    if not source.exists:
        source.status = STATUS_UNREADABLE
        source.message = "文件不存在"
        return ReadResult(source)
    if _cancelled(cancel):
        source.status = STATUS_INCOMPLETE
        source.message = "用户取消读取"
        return ReadResult(source)
    source = _schema(source, cancel, progress)
    if _cancelled(cancel):
        source.status = STATUS_INCOMPLETE
        source.message = "用户取消读取"
        return ReadResult(source)
    if source.status in {STATUS_DAMAGED, STATUS_UNCOMPARABLE} and not source.sheet_name and source.fmt == FORMAT_XLSX:
        return ReadResult(source)
    preview: list[dict[str, Any]] = []
    fields = set(source.fields)
    time_values: list[datetime] = []
    numeric_seen: dict[str, bool] = {}
    keys: set[tuple[Any, ...]] = set()
    count = 0
    invalid_time = 0
    duplicate_keys = 0
    key_missing = 0
    key_invalid = 0
    source.status = STATUS_READABLE
    cancelled = False
    try:
        for row, _row_no in _iter_records(source, cancel):
            if _cancelled(cancel):
                source.status = STATUS_INCOMPLETE
                source.message = "用户取消读取"
                cancelled = True
                break
            count += 1
            fields.update(_public_fields(row))
            if len(preview) < max(0, int(preview_limit)):
                preview.append(row)
            current_time = _row_time(source, row)
            if source.time_field and row.get(source.time_field) not in (None, ""):
                if current_time is None:
                    invalid_time += 1
                else:
                    time_values.append(current_time)
            key, key_state = _row_key_state(source, row)
            if key_state == "missing":
                key_missing += 1
            elif key_state == "invalid":
                key_invalid += 1
            if key is not None:
                if key in keys:
                    duplicate_keys += 1
                keys.add(key)
            for name, value in row.items():
                if not name.startswith("_") and numeric_value(value) is not None:
                    numeric_seen[name] = True
            if progress and (count == 1 or count % 500 == 0):
                progress(text=f"读取 {source.file_name} · {count:,} 条", row=count)
        if _cancelled(cancel):
            cancelled = True
    except (OSError, UnicodeError, ValueError, TypeError) as exc:
        source.status = STATUS_DAMAGED
        source.message = _safe_error(exc)
        source.issues.append(source.message)
    except Exception as exc:
        source.status = STATUS_DAMAGED
        source.message = _safe_error(exc)
        source.issues.append(source.message)

    source.record_count = count
    source.fields = tuple(sorted(fields))
    source.time_start = min(time_values) if time_values else None
    source.time_end = max(time_values) if time_values else None
    source.invalid_time = invalid_time
    source.duplicate_keys = duplicate_keys
    source.key_missing = key_missing if source.key_fields else None
    source.key_invalid = key_invalid if source.key_fields else None
    source.key_dedup_discarded = None
    source.key_remaining_conflicts = duplicate_keys if source.key_fields else None
    source.metrics = _source_metrics(source, numeric_seen)
    if cancelled:
        source.status = STATUS_INCOMPLETE
        source.message = "用户取消读取"
        source.can_compare = False
        source.preview_rows = preview
        return ReadResult(source, preview)
    missing_required = sorted(_required_fields(source) - set(source.fields))
    if source.status == STATUS_READABLE:
        if source.kind == KIND_UNKNOWN:
            source.status = STATUS_UNCOMPARABLE
            source.message = "无法识别评论、视频快照或监控历史字段"
        elif missing_required or invalid_time:
            source.status = STATUS_INCOMPLETE
            source.message = "缺少稳定业务键或可靠时间字段"
        elif duplicate_keys:
            source.status = STATUS_INCOMPLETE
            source.message = "存在重复稳定业务键，无法保证逐条匹配"
    source.can_compare = (
        source.kind in {KIND_COMMENT, KIND_VIDEO, KIND_MONITOR}
        and not missing_required and invalid_time == 0 and source.status == STATUS_READABLE
    )
    if source.kind in {KIND_COMMENT, KIND_VIDEO, KIND_MONITOR} and not source.can_compare and source.status == STATUS_READABLE:
        source.status = STATUS_UNCOMPARABLE
        source.message = "字段不足，不能安全逐条匹配"
    try:
        fp = snapshot_source(source.path)
        source.sha256, source.size, source.mtime_ns = fp.sha256, fp.size, fp.mtime_ns
    except (OSError, ValueError) as exc:
        source.status = STATUS_UNREADABLE
        source.message = _safe_error(exc)
    source.preview_rows = preview
    return ReadResult(source, preview)


def inspect_source(path: str | Path) -> Source:
    return read_source(path).source


def default_monitor_roots(cfg: dict[str, Any] | None = None) -> list[Path]:
    values = cfg or {}
    roots = [Path.cwd() / "data", output_mod.app_base_dir() / "data"]
    if values.get("out_dir"):
        roots.append(Path(values["out_dir"]) / "监控数据")
    roots.append(output_mod.app_base_dir() / "监控数据")
    unique: list[Path] = []
    seen: set[str] = set()
    for root in roots:
        key = str(root.resolve(strict=False)).casefold()
        if key not in seen:
            seen.add(key)
            unique.append(root)
    return unique


def _iter_candidate_paths(paths: Iterable[str | Path], cancel: Cancel | None = None,
                          progress: Progress | None = None) -> Iterator[Path]:
    for value in paths or ():
        if _cancelled(cancel):
            return
        path = _path(value)
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES:
            if progress:
                progress(text=f"发现候选文件：{path.name}", path=str(path))
            yield path
        elif path.is_dir():
            try:
                for child in path.rglob("*"):
                    if _cancelled(cancel):
                        return
                    if child.is_file() and child.suffix.lower() in SUPPORTED_SUFFIXES:
                        if progress:
                            progress(text=f"发现候选文件：{child.name}", path=str(child))
                        yield child
            except OSError:
                continue


def _iter_discovery_candidates(history_records: Iterable[dict[str, Any]] | None,
                               manual_files: Iterable[str | Path] | None,
                               manual_dirs: Iterable[str | Path] | None,
                               monitor_roots: Iterable[str | Path] | None,
                               cancel: Cancel | None = None,
                               progress: Progress | None = None) -> Iterator[tuple[Path, str]]:
    records = task_history.load_history() if history_records is None else history_records
    for record in records:
        if _cancelled(cancel):
            return
        outputs = record.get("outputs", []) if isinstance(record, dict) else []
        for value in outputs:
            if _cancelled(cancel):
                return
            path = _path(value)
            if path.suffix.lower() in SUPPORTED_SUFFIXES:
                yield path, "任务历史"
    for value in manual_files or ():
        if _cancelled(cancel):
            return
        path = _path(value)
        if path.suffix.lower() in SUPPORTED_SUFFIXES:
            yield path, "手动添加"
    for path in _iter_candidate_paths(manual_dirs or (), cancel, progress):
        yield path, "目录扫描"
    for path in _iter_candidate_paths(monitor_roots or (), cancel, progress):
        if _cancelled(cancel):
            return
        if path.name.lower().startswith("history_"):
            yield path, "监控历史"


def discover_sources(history_records: Iterable[dict[str, Any]] | None = None,
                     manual_files: Iterable[str | Path] | None = None,
                     manual_dirs: Iterable[str | Path] | None = None,
                     monitor_roots: Iterable[str | Path] | None = None,
                     cancel: Cancel | None = None,
                     progress: Progress | None = None) -> list[Source]:
    seen: set[str] = set()
    results: list[Source] = []
    for path, origin in _iter_discovery_candidates(
        history_records, manual_files, manual_dirs, monitor_roots, cancel, progress
    ):
        if _cancelled(cancel):
            break
        key = str(path).casefold()
        if key in seen:
            continue
        seen.add(key)
        result = read_source(path, cancel=cancel, progress=progress)
        if _cancelled(cancel) or result.source.message == "用户取消读取":
            break
        source = result.source
        source.discovered_from = origin
        results.append(source)
        if progress:
            progress(text=f"已读取 {source.file_name}", done=len(results), total=None)
    results.sort(key=lambda item: (item.kind_label, item.file_name.casefold(), item.path.casefold()))
    return results


def _source_input(value: Source | ReadResult | str | Path,
                  cancel: Cancel | None = None,
                  progress: Progress | None = None) -> Source:
    if isinstance(value, ReadResult):
        if _cancelled(cancel):
            raise ComparisonCancelled("用户取消对比")
        return value.source
    result = read_source(value.path if isinstance(value, Source) else value,
                         cancel=cancel, progress=progress)
    if _cancelled(cancel) or result.source.message == "用户取消读取":
        raise ComparisonCancelled("用户取消对比")
    if isinstance(value, Source):
        result.source.discovered_from = value.discovered_from
    return result.source


def _iter_all(source: Source, cancel: Cancel | None = None) -> Iterator[dict[str, Any]]:
    for row, _row_no in _iter_records(source, cancel):
        yield row


def _percent_change(a_value: float | int, b_value: float | int) -> tuple[float | None, str]:
    if a_value == 0:
        return (0.0, "") if b_value == 0 else (None, "A 值为 0，百分比变化不可计算")
    return float((b_value - a_value) / a_value * 100), ""


def _aggregate_metric(rows: list[dict[str, Any]], field_name: str) -> tuple[float | int | None, int, int]:
    values = [numeric_value(row.get(field_name)) for row in rows]
    valid = [value for value in values if value is not None]
    return (sum(valid) if valid else None, len(valid), len(values) - len(valid))


def _row_diff_fields(a: dict[str, Any], b: dict[str, Any]) -> list[str]:
    keys = set(_public_fields(a)) | set(_public_fields(b))
    changed: list[str] = []
    for name in sorted(keys):
        if _json_value(a.get(name)) != _json_value(b.get(name)):
            changed.append(name)
    return changed


def _raise_if_comparison_cancelled(cancel: Cancel | None) -> None:
    if _cancelled(cancel):
        raise ComparisonCancelled("用户取消对比")


def compare_sources(left: Source | ReadResult | str | Path,
                    right: Source | ReadResult | str | Path,
                    row_limit: int = 500,
                    cancel: Cancel | None = None,
                    progress: Progress | None = None) -> ComparisonResult:
    source_a = _source_input(left, cancel=cancel, progress=progress)
    source_b = _source_input(right, cancel=cancel, progress=progress)
    compatible = source_a.kind == source_b.kind and source_a.kind in {KIND_COMMENT, KIND_VIDEO, KIND_MONITOR}
    message = "" if compatible else f"类型不兼容：{source_a.kind_label} 不能与 {source_b.kind_label} 比较"
    result = ComparisonResult(
        left=source_a, right=source_b, compatible=compatible,
        matching_mode="key" if compatible and source_a.can_compare and source_b.can_compare else "summary",
        message=message, a_records=source_a.record_count, b_records=source_b.record_count,
        time_range_a=(source_a.time_start, source_a.time_end), time_range_b=(source_b.time_start, source_b.time_end),
        missing_fields_a=tuple(sorted(set(source_b.fields) - set(source_a.fields))),
        missing_fields_b=tuple(sorted(set(source_a.fields) - set(source_b.fields))),
        fingerprint_a=source_a.fingerprint(), fingerprint_b=source_b.fingerprint(),
    )
    rows_a: list[dict[str, Any]] = []
    rows_b: list[dict[str, Any]] = []
    for index, row in enumerate(_iter_all(source_a, cancel), 1):
        rows_a.append(row)
        if progress and index % 500 == 0:
            progress(text=f"对比 A：{index:,} 条", row=index)
    _raise_if_comparison_cancelled(cancel)
    for index, row in enumerate(_iter_all(source_b, cancel), 1):
        rows_b.append(row)
        if progress and index % 500 == 0:
            progress(text=f"对比 B：{index:,} 条", row=index)
    _raise_if_comparison_cancelled(cancel)
    if not compatible:
        return result
    if result.matching_mode != "key":
        result.message = result.message or "无法确定完整稳定业务键，仅生成汇总对比；无法逐条匹配"
        for metric_index, name in enumerate(sorted(set(source_a.fields) & set(source_b.fields)), 1):
            _raise_if_comparison_cancelled(cancel)
            if name in set(source_a.key_fields) or name in {"message", "title", "owner", "uname"}:
                continue
            a_value, a_valid, missing_a = _aggregate_metric(rows_a, name)
            b_value, b_valid, missing_b = _aggregate_metric(rows_b, name)
            if not (a_valid or b_valid):
                continue
            absolute = b_value - a_value if a_value is not None and b_value is not None else None
            percent, note = _percent_change(a_value, b_value) if absolute is not None else (None, "缺少可比较数值")
            result.metric_changes.append(MetricChange(name, a_value, b_value, absolute, percent,
                                                       min(a_valid, b_valid), missing_a, missing_b, note))
        return result

    map_a: dict[tuple[Any, ...], dict[str, Any]] = {}
    map_b: dict[tuple[Any, ...], dict[str, Any]] = {}
    for row_index, row in enumerate(rows_a, 1):
        _raise_if_comparison_cancelled(cancel)
        key = _row_key(source_a, row)
        if key is not None:
            map_a[key] = row
        if progress and row_index % 500 == 0:
            progress(text=f"建立 A 索引：{row_index:,} 条", row=row_index)
    for row_index, row in enumerate(rows_b, 1):
        _raise_if_comparison_cancelled(cancel)
        key = _row_key(source_b, row)
        if key is not None:
            map_b[key] = row
        if progress and row_index % 500 == 0:
            progress(text=f"建立 B 索引：{row_index:,} 条", row=row_index)
    keys_a, keys_b = set(map_a), set(map_b)
    common = sorted(keys_a & keys_b, key=str)
    result.common_records = len(common)
    result.only_a_records = len(keys_a - keys_b)
    result.only_b_records = len(keys_b - keys_a)
    result.changed_records = 0
    for key_index, key in enumerate(common, 1):
        _raise_if_comparison_cancelled(cancel)
        changed = _row_diff_fields(map_a[key], map_b[key])
        if changed:
            result.changed_records += 1
            if len(result.row_changes) < row_limit:
                result.row_changes.append({"key": list(key), "changed_fields": changed})
    key_fields = set(source_a.key_fields)
    for name in sorted(set(source_a.fields) & set(source_b.fields)):
        if name in key_fields or name in {"message", "title", "owner", "uname"}:
            continue
        a_values: list[float | int] = []
        b_values: list[float | int] = []
        missing_a = missing_b = 0
        for key_index, key in enumerate(common, 1):
            _raise_if_comparison_cancelled(cancel)
            a_value = numeric_value(map_a[key].get(name))
            b_value = numeric_value(map_b[key].get(name))
            if a_value is None:
                missing_a += 1
            else:
                a_values.append(a_value)
            if b_value is None:
                missing_b += 1
            else:
                b_values.append(b_value)
        if not a_values and not b_values:
            continue
        a_total = sum(a_values) if a_values else None
        b_total = sum(b_values) if b_values else None
        absolute = b_total - a_total if a_total is not None and b_total is not None else None
        percent, note = _percent_change(a_total, b_total) if absolute is not None else (None, "缺少可比较数值")
        result.metric_changes.append(MetricChange(name, a_total, b_total, absolute, percent,
                                                   min(len(a_values), len(b_values)), missing_a, missing_b, note))
    return result


def _period_stats(source: Source, rows: Iterable[dict[str, Any]], start: datetime, end: datetime,
                  metric: str, cancel: Cancel | None = None,
                  progress: Progress | None = None) -> tuple[PeriodStats, list[tuple[datetime, float | int]]]:
    result = PeriodStats(start, end)
    values: list[tuple[datetime, float | int]] = []
    for row_index, row in enumerate(rows, 1):
        _raise_if_comparison_cancelled(cancel)
        current = _row_time(source, row)
        if current is None:
            result.missing_time += 1
            continue
        if not (start <= current <= end):
            continue
        result.sample_count += 1
        value = numeric_value(row.get(metric))
        if value is None:
            result.missing_data += 1
            continue
        result.valid_count += 1
        values.append((current, value))
        if progress and row_index % 500 == 0:
            progress(text=f"统计 {metric}：{row_index:,} 条", row=row_index)
    values.sort(key=lambda item: item[0])
    if values:
        numeric_values = [value for _time, value in values]
        result.start_value, result.end_value = numeric_values[0], numeric_values[-1]
        result.average = sum(numeric_values) / len(numeric_values)
        result.maximum, result.minimum = max(numeric_values), min(numeric_values)
    return result, values


def compare_periods(source_value: Source | ReadResult | str | Path,
                    period_a: tuple[datetime | str, datetime | str],
                    period_b: tuple[datetime | str, datetime | str], metric: str,
                    cancel: Cancel | None = None,
                    progress: Progress | None = None) -> PeriodComparisonResult:
    source = _source_input(source_value, cancel=cancel, progress=progress)
    start_a, end_a = (parse_user_datetime(item) for item in period_a)
    start_b, end_b = (parse_user_datetime(item) for item in period_b)
    if not start_a or not end_a or not start_b or not end_b:
        raise ValueError("时间段必须是带日期和时间的有效值")
    if start_a > end_a or start_b > end_b:
        raise ValueError("时间段开始时间不能晚于结束时间")
    if not source.time_field:
        raise ValueError("来源没有可靠时间字段，不能进行时间段对比")
    if metric not in source.metrics:
        raise ValueError(f"指标“{metric}”不存在或不是可计算数值")
    stats_a, values_a = _period_stats(source, _iter_all(source, cancel), start_a, end_a, metric,
                                      cancel=cancel, progress=progress)
    _raise_if_comparison_cancelled(cancel)
    stats_b, values_b = _period_stats(source, _iter_all(source, cancel), start_b, end_b, metric,
                                      cancel=cancel, progress=progress)
    _raise_if_comparison_cancelled(cancel)
    return PeriodComparisonResult(source, metric, stats_a, stats_b, values_a, values_b,
                                  fingerprint=source.fingerprint())


def filter_sources(sources: Iterable[Source], query: str = "", kind: str = "", fmt: str = "",
                   status: str = "", start: datetime | None = None, end: datetime | None = None) -> list[Source]:
    query = str(query or "").strip().casefold()
    output: list[Source] = []
    for source in sources:
        haystack = " ".join([source.file_name, source.path, source.kind_label, *source.fields]).casefold()
        if query and query not in haystack:
            continue
        if kind and source.kind != kind:
            continue
        if fmt and source.fmt != fmt:
            continue
        if status and source.status != status:
            continue
        if start and source.time_end and source.time_end < start:
            continue
        if end and source.time_start and source.time_start > end:
            continue
        output.append(source)
    return output


def search_source_matches(sources: Iterable[Source], query: str,
                          cancel: Cancel | None = None,
                          progress: Progress | None = None) -> dict[str, str]:
    """返回来源路径到命中类型的映射：metadata 或 content。

    记录内容始终流式扫描全量记录，因此不会把预览上限误当成搜索范围。
    """
    needle = str(query or "").strip().casefold()
    values = list(sources)
    if not needle:
        return {source.path.casefold(): "metadata" for source in values}
    matched: dict[str, str] = {}
    for source_index, source in enumerate(values, 1):
        if _cancelled(cancel):
            return matched
        metadata = " ".join((source.file_name, source.path, source.kind_label, *source.fields)).casefold()
        metadata_hit = needle in metadata
        content_hit = False
        for scanned, row in enumerate(_iter_all(source, cancel), 1):
            if _cancelled(cancel):
                return matched
            text = json.dumps(sanitize_public_row(row), ensure_ascii=False).casefold()
            if needle in text:
                content_hit = True
            if progress and scanned % 500 == 0:
                progress(text=f"搜索 {source.file_name} · {scanned:,} 条", row=scanned)
        if metadata_hit:
            matched[source.path.casefold()] = "metadata"
        elif content_hit:
            matched[source.path.casefold()] = "content"
        if progress:
            progress(text=f"已搜索 {source.file_name}", done=source_index, total=len(values))
    return matched


def search_sources(sources: Iterable[Source], query: str,
                   cancel: Cancel | None = None,
                   progress: Progress | None = None) -> list[Source]:
    """后台流式搜索来源元数据和全量记录，不受预览上限限制。"""
    values = list(sources)
    matches = search_source_matches(values, query, cancel=cancel, progress=progress)
    return [source for source in values if source.path.casefold() in matches]


def _unique_output_path(directory: Path, stem: str, suffix: str, forbidden: set[str] | None = None) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    forbidden = {str(Path(item).resolve(strict=False)).casefold() for item in (forbidden or set())}
    candidate = directory / f"{stem}{suffix}"
    index = 1
    while candidate.exists() or str(candidate.resolve(strict=False)).casefold() in forbidden:
        candidate = directory / f"{stem}_{index}{suffix}"
        index += 1
    return candidate


def _atomic_text(path: Path, lines: Iterable[str]) -> None:
    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile("w", encoding="utf-8", prefix=".report-", suffix=".tmp",
                                        dir=str(path.parent), delete=False) as handle:
            temp_path = Path(handle.name)
            for line in lines:
                handle.write(line)
            handle.flush()
            os.fsync(handle.fileno())
        temp_path.replace(path)
    finally:
        if temp_path and temp_path.exists():
            temp_path.unlink(missing_ok=True)


def _atomic_xlsx(path: Path, build: Callable[[], Any], before_replace=None) -> None:
    workbook = None
    try:
        workbook = build()
        xlsx_mod.save_workbook_atomic(workbook, path, before_replace=before_replace)
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass


def _verify_unchanged(before: dict[str, Fingerprint]) -> None:
    for path, expected in before.items():
        try:
            current = snapshot_source(path)
        except OSError as exc:
            raise SourceChangedError("源文件无法再次读取") from exc
        if current != expected:
            raise SourceChangedError("源文件在导出期间发生变化")


def _verify_comparison_fingerprints(result: ComparisonResult | PeriodComparisonResult) -> None:
    expected = ([result.fingerprint_a, result.fingerprint_b]
                if isinstance(result, ComparisonResult)
                else [result.fingerprint])
    for fingerprint in expected:
        if fingerprint is None:
            raise SourceChangedError("源文件已变化，请重新生成对比")
        try:
            current = snapshot_source(fingerprint.path)
        except OSError as exc:
            raise SourceChangedError("源文件已变化，请重新生成对比") from exc
        if current != fingerprint:
            raise SourceChangedError("源文件已变化，请重新生成对比")


def _export_context(sources: Iterable[Source]) -> dict[str, Fingerprint]:
    return {source.path: snapshot_source(source.path) for source in sources}


def _write_rows_xlsx(path: Path, source: Source, rows: Iterable[dict[str, Any]],
                     summary: dict[str, Any], before_replace=None) -> None:
    written_count = {"value": 0}
    output_key_conflicts = {"value": 0}
    def build():
        workbook = xlsx_mod.new_workbook()
        sw = xlsx_mod.SheetWriter(workbook)
        data_sheet = workbook.create_sheet("数据")
        sw.ws = data_sheet
        fields = list(source.original_headers) if source.original_headers else list(source.fields)
        fields = [field for field in fields if not _SENSITIVE_KEY_RE.search(str(field))]
        data_layout = TableLayout(1, 1, max(1, len(fields)))
        configure_table(data_sheet, data_layout)
        sw.append([_xlsx_safe_value(field) for field in fields])
        output_keys = set()
        for row in rows:
            key = _row_key(source, row)
            if key is not None:
                if key in output_keys:
                    output_key_conflicts["value"] += 1
                output_keys.add(key)
            values = _xlsx_safe_row(row)
            row_cells = []
            for field in fields:
                cell = _report_cell(source, field, values.get(field), present=field in values)
                row_cells.append(wrap_cell(sw, cell) if field in _REPORT_LONG_TEXT_FIELDS else cell)
            sw.append(row_cells)
            written_count["value"] += 1
        finish_table(data_sheet, data_layout, written_count["value"])
        report = workbook.create_sheet("报告汇总")
        sw.ws = report
        summary_layout = TableLayout(1, 1, 2)
        configure_table(report, summary_layout)
        sw.append([_generic_cell("字段", "报告汇总"), _generic_cell("值", "报告汇总")])
        for key, value in summary.items():
                sw.append([_generic_cell(str(key), "报告字段"), _generic_cell(value, str(key))])
        finish_table(report, summary_layout, len(summary))
        quality = []
        def q(item, raw, kind, unit, note):
            value = (xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE, note=note)
                     if raw is None else xlsx_mod.checked_cell_value(raw, kind, note=note))
            return QualityItem("报告中心", item, value, unit, note)
        quality.extend([
            q("来源记录总数", source.record_count, xlsx_mod.CellKind.INTEGER, "条", "Source 已观测记录数"),
            q("实际写入数据表记录数", written_count["value"], xlsx_mod.CellKind.INTEGER, "条", "筛选生成器真实写入数"),
            q("检测到的重复键数量", source.duplicate_keys, xlsx_mod.CellKind.INTEGER, "条", "Source 统计；不从输出表猜测"),
            q("无法解析时间数量", source.invalid_time, xlsx_mod.CellKind.INTEGER, "条", "Source 统计"),
            q("可观测问题数量", len(source.issues), xlsx_mod.CellKind.INTEGER, "条", "已脱敏的问题摘要数量"),
            q("接口声称数量", None, xlsx_mod.CellKind.INTEGER, "条", "本地报告中心不适用"),
            q("覆盖率", None, xlsx_mod.CellKind.PERCENT, "状态", "本地报告中心不适用"),
            q("是否取消", False, xlsx_mod.CellKind.BOOLEAN, "状态", "导出完成"),
            q("是否部分成功", source.status != STATUS_READABLE, xlsx_mod.CellKind.BOOLEAN, "状态", "来自 Source.status"),
        ])
        key_label = " / ".join(source.key_fields) if source.key_fields else "主键"
        quality.extend(primary_key_quality(
            "主键", key_label,
            denominator=source.record_count if source.key_fields else None,
            missing=source.key_missing,
            invalid=source.key_invalid,
            duplicates=source.duplicate_keys if source.key_fields else None,
            dedup_discarded=0 if source.key_fields else None,
            remaining_conflicts=output_key_conflicts["value"] if source.key_fields else None,
            source_note="来自 Source 结构化统计；导出器不扫描 XLSX 工作表",
        ))
        fields_meta = []
        known_fields = set(source.fields)
        for field_name in fields:
            known = field_name in known_fields and (
                field_name in _REPORT_ID_FIELDS.get(source.kind, set())
                or field_name in _REPORT_INTEGER_FIELDS.get(source.kind, set())
                or field_name in _REPORT_BOOLEAN_FIELDS.get(source.kind, set())
                or field_name in _REPORT_DATETIME_FIELDS.get(source.kind, set())
            )
            dtype = _report_kind(source, field_name, None).value if known else "源文件未声明/未知"
            field_source = "报告中心已知映射" if known else "源文件未声明/未知"
            metric = "按 Source 已知字段处理" if known else "源文件未声明/未知"
            fields_meta.append(FieldDefinition(
                "数据", str(field_name), str(field_name), dtype, "", "是", field_source,
                metric, xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE),
                "字段缺失或源文件未声明"))
        for key, value in summary.items():
            fields_meta.append(FieldDefinition(
                "报告汇总", str(key), f"summary.{key}",
                "布尔" if isinstance(value, bool) else
                "整数" if isinstance(value, int) and not isinstance(value, bool) else
                "小数" if isinstance(value, float) else "文本",
                "", "是", "报告中心结构化筛选结果", "本次导出摘要",
                xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE),
                "未生成或不适用"))
        fields_meta.extend((
            FieldDefinition("报告汇总", "字段", "summary_key", "文本", "", "否",
                            "报告中心结构化筛选结果", "摘要键", xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE), "不适用"),
            FieldDefinition("报告汇总", "值", "summary_value", "文本/数值", "", "是",
                            "报告中心结构化筛选结果", "摘要值；敏感值已排除", xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE), "不适用或已排除敏感值"),
        ))
        metadata = make_metadata(
            tool="报告中心", report_type="筛选导出",
            parameters={
                "来源 basename": xlsx_mod.cell_value(source.file_name, xlsx_mod.CellKind.TEXT),
                "来源类型": xlsx_mod.cell_value(source.kind_label, xlsx_mod.CellKind.TEXT),
                "是否包含筛选条件": xlsx_mod.cell_value(bool(summary.get("是否包含筛选条件", False)), xlsx_mod.CellKind.BOOLEAN),
            },
            parameter_allowlist=("来源 basename", "来源类型", "是否包含筛选条件"),
            quality_items=quality, fields=fields_meta)
        write_metadata_sheets(workbook, metadata)
        return workbook
    _atomic_xlsx(path, build, before_replace=before_replace)


def _write_summary_xlsx(path: Path, sheets: list[tuple[str, list[list[Any]]]],
                        before_replace=None) -> None:
    def build():
        workbook = xlsx_mod.new_workbook()
        sw = xlsx_mod.SheetWriter(workbook)
        for index, (name, rows) in enumerate(sheets):
            worksheet = workbook.create_sheet(name[:31])
            sw.ws = worksheet
            width = max(1, len(rows[0]) if rows else 1)
            layout = TableLayout(1, 1, width)
            configure_table(worksheet, layout)
            for row in rows:
                sw.append([_summary_cell(value, str(name)) for value in row])
            finish_table(worksheet, layout, max(0, len(rows) - 1))
            if index == 0 and name in ("对比汇总", "时间段汇总"):
                append_sheet_directory(
                    sw,
                    worksheet,
                    tuple((sheet_name[:31], sheet_name[:31]) for sheet_name, _ in sheets)
                    + (("数据质量", "数据质量"), ("字段说明", "字段说明")),
                    prefix_columns=0,
                )
        output_rows = sum(max(0, len(rows) - 1) for _name, rows in sheets)
        quality = [
            QualityItem("报告中心", "输出业务表数量",
                        xlsx_mod.checked_cell_value(len(sheets), xlsx_mod.CellKind.INTEGER), "张", "本次导出的业务结果表"),
            QualityItem("报告中心", "实际输出记录数",
                        xlsx_mod.checked_cell_value(output_rows, xlsx_mod.CellKind.INTEGER), "行", "按导出表头之后的真实行数"),
            QualityItem("报告中心", "接口声称数量",
                        xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE, note="本地报告中心不适用"), "条", "本地比较结果没有接口声明数量"),
            QualityItem("报告中心", "覆盖率",
                        xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE, note="本地报告中心不适用"), "状态", "不计算外部接口覆盖率"),
            QualityItem("报告中心", "是否取消",
                        xlsx_mod.checked_cell_value(False, xlsx_mod.CellKind.BOOLEAN), "状态", "导出完成"),
        ]
        fields = []
        for name, rows in sheets:
            headers = rows[0] if rows else []
            for header in headers:
                text = str(header)
                stable = {
                    "字段": "field", "值": "value", "A值": "a_value", "B值": "b_value",
                    "绝对变化": "absolute_change", "百分比变化": "percent_change",
                    "可比较记录": "compared_records", "A缺失": "missing_a", "B缺失": "missing_b",
                    "业务键": "business_key", "变化字段": "changed_fields", "时间段": "period",
                    "时间": "time", "样本数": "sample_count", "有效数值数": "valid_count",
                    "缺失数据数": "missing_data", "缺失时间数": "missing_time",
                    "起始值": "start_value", "结束值": "end_value", "平均值": "average",
                    "最大值": "maximum", "最小值": "minimum",
                }.get(text, f"report.{name[:20]}.{text}")
                fields.append(FieldDefinition(
                    name[:31], text, stable, "源文件未声明/未知", "", "是",
                    "报告中心计算结果", "报告中心输出字段；源文件口径不适用",
                    xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE),
                    "字段缺失或不适用"))
        metadata = make_metadata(
            tool="报告中心", report_type="本地比较报告",
            parameters={
                "业务表数量": xlsx_mod.cell_value(len(sheets), xlsx_mod.CellKind.INTEGER),
                "导出类型": xlsx_mod.cell_value("本地比较", xlsx_mod.CellKind.TEXT),
            }, parameter_allowlist=("业务表数量", "导出类型"),
            quality_items=quality, fields=fields)
        write_metadata_sheets(workbook, metadata)
        return workbook
    _atomic_xlsx(path, build, before_replace=before_replace)


def export_filtered(source_value: Source | ReadResult, out_dir: str | Path, fmt: str,
                    query: str = "", forbidden_paths: Iterable[str | Path] | None = None,
                    progress: Progress | None = None, cancel: Cancel | None = None) -> Path:
    source = _source_input(source_value)
    before = _export_context([source])
    suffix = ".jsonl" if fmt.upper() == FORMAT_JSONL else ".xlsx"
    output = _unique_output_path(_path(out_dir), "报告中心_筛选结果", suffix,
                                 {str(item) for item in (forbidden_paths or ())})
    needle = str(query or "").strip().casefold()

    def matching_rows() -> Iterator[dict[str, Any]]:
        for index, row in enumerate(_iter_all(source, cancel), 1):
            if _cancelled(cancel):
                raise ReportCenterError("用户取消导出")
            public = sanitize_public_row(row)
            if needle and needle not in json.dumps(public, ensure_ascii=False).casefold():
                continue
            if progress and index % 500 == 0:
                progress(text=f"导出筛选结果 · {index:,} 条", row=index)
            yield row

    if suffix == ".jsonl":
        _atomic_text(output, (json.dumps(sanitize_public_row(row), ensure_ascii=False) + "\n" for row in matching_rows()))
    else:
        _write_rows_xlsx(
            output, source, matching_rows(),
            {"来源类型": source.kind_label, "来源文件": source.file_name,
             "是否包含筛选条件": bool(needle)},
            before_replace=lambda: _verify_unchanged(before),
        )
    _verify_unchanged(before)
    return output


def _comparison_lines(result: ComparisonResult) -> Iterator[str]:
    yield json.dumps({"report_type": "file_comparison", "summary": {
        "a_records": result.a_records, "b_records": result.b_records, "common_records": result.common_records,
        "only_a_records": result.only_a_records, "only_b_records": result.only_b_records,
        "changed_records": result.changed_records, "matching_mode": result.matching_mode, "message": result.message,
        "fingerprint_a": result.fingerprint_a.as_dict() if result.fingerprint_a else None,
        "fingerprint_b": result.fingerprint_b.as_dict() if result.fingerprint_b else None,
    }, "time_range_a": [format_datetime(item) if item else None for item in result.time_range_a],
        "time_range_b": [format_datetime(item) if item else None for item in result.time_range_b],
        "missing_fields_a": list(result.missing_fields_a), "missing_fields_b": list(result.missing_fields_b)}, ensure_ascii=False) + "\n"
    for item in result.metric_changes:
        yield json.dumps({"record_type": "metric_change", **item.as_dict()}, ensure_ascii=False) + "\n"
    for item in result.row_changes:
        yield json.dumps({"record_type": "row_change", **sanitize_public_row(item)}, ensure_ascii=False) + "\n"


def export_file_comparison(result: ComparisonResult, out_dir: str | Path, fmt: str) -> Path:
    _verify_comparison_fingerprints(result)
    before = _export_context([result.left, result.right])
    suffix = ".jsonl" if fmt.upper() == FORMAT_JSONL else ".xlsx"
    output = _unique_output_path(_path(out_dir), "报告中心_文件对比", suffix, {result.left.path, result.right.path})
    if suffix == ".jsonl":
        _atomic_text(output, _comparison_lines(result))
    else:
        summary = result.as_dict(redact_paths=True)
        metric_rows = [["字段", "A值", "B值", "绝对变化", "百分比变化", "可比较记录", "A缺失", "B缺失", "说明"]]
        metric_rows.extend([[item.field, item.a_value, item.b_value, item.absolute_change,
                             _percent_point_cell(item.percent_change),
                             item.compared_records, item.missing_a, item.missing_b, item.note] for item in result.metric_changes])
        row_rows = [["业务键", "变化字段"]]
        row_rows.extend([[" / ".join(map(str, item.get("key", []))), ", ".join(item.get("changed_fields", []))] for item in result.row_changes])
        summary_rows = [["字段", "值"]] + [[key, value] for key, value in summary.items()]
        _write_summary_xlsx(
            output, [("对比汇总", summary_rows), ("指标变化", metric_rows),
                     ("记录变化", row_rows)],
            before_replace=lambda: _verify_unchanged(before),
        )
    _verify_unchanged(before)
    return output


def _period_lines(result: PeriodComparisonResult) -> Iterator[str]:
    yield json.dumps({"report_type": "period_comparison", **result.as_dict()}, ensure_ascii=False) + "\n"
    for label, values in (("A", result.values_a), ("B", result.values_b)):
        for current, value in values:
            yield json.dumps({"record_type": "sample", "period": label, "time": format_datetime(current),
                              "metric": result.metric, "value": value}, ensure_ascii=False) + "\n"


def export_period_comparison(result: PeriodComparisonResult, out_dir: str | Path, fmt: str) -> Path:
    _verify_comparison_fingerprints(result)
    before = _export_context([result.source])
    suffix = ".jsonl" if fmt.upper() == FORMAT_JSONL else ".xlsx"
    output = _unique_output_path(_path(out_dir), "报告中心_时间段对比", suffix, {result.source.path})
    if suffix == ".jsonl":
        _atomic_text(output, _period_lines(result))
    else:
        summary = [["字段", "时间段A", "时间段B"], ["样本数", result.period_a.sample_count, result.period_b.sample_count],
                   ["有效数值数", result.period_a.valid_count, result.period_b.valid_count],
                   ["缺失数据数", result.period_a.missing_data, result.period_b.missing_data],
                   ["缺失时间数", result.period_a.missing_time, result.period_b.missing_time],
                   ["起始值", result.period_a.start_value, result.period_b.start_value],
                   ["结束值", result.period_a.end_value, result.period_b.end_value], ["平均值", result.period_a.average, result.period_b.average],
                   ["最大值", result.period_a.maximum, result.period_b.maximum], ["最小值", result.period_a.minimum, result.period_b.minimum],
                   ["绝对变化", result.period_a.absolute_change, result.period_b.absolute_change],
                   ["百分比变化", _percent_point_cell(result.period_a.percent_change),
                    _percent_point_cell(result.period_b.percent_change)]]
        samples = [["时间段", "时间", result.metric]]
        samples.extend([["A", format_datetime(current), value] for current, value in result.values_a])
        samples.extend([["B", format_datetime(current), value] for current, value in result.values_b])
        _write_summary_xlsx(
            output, [("时间段汇总", summary), ("样本序列", samples)],
            before_replace=lambda: _verify_unchanged(before),
        )
    _verify_unchanged(before)
    return output
