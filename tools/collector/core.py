# -*- coding: utf-8 -*-
"""采集核心：快照采集 + 分析 + Excel。

HTTP 请求统一走 core.session 风控栈；来源展开（收藏夹/合集/系列/txt）见
core.links.expand_source；Excel 基于 core.xlsx 共享基建。
"""
import json
import random
import time
from collections import Counter
from datetime import datetime
from pathlib import Path

from core import session
from core.budget import BudgetExhaustedError
from core.cancel import TaskCancelledError, wait as wait_or_cancel
from core.risk import RiskChallengeError  # noqa: F401  供 pipeline 引用
from core import xlsx as xlsx_mod
from core.xlsx_metadata import (
    FieldDefinition,
    QualityItem,
    make_metadata,
    primary_key_quality,
    write_metadata_sheets,
)
from core.xlsx_presentation import (
    TableLayout,
    append_sheet_directory,
    configure_table,
    finish_table,
    wrap_cell,
)

try:
    from openpyxl.utils import get_column_letter
    HAVE_XLSX = True
except ImportError:  # noqa: F841
    HAVE_XLSX = False


# ============================ 采集 ============================

def fetch_view(bvid, cancel=None, budget=None):
    """取一份视频公开数据快照。budget 非 None 时透传任务预算
    （强制点在 HTTP 层入口），为 None 时调用与旧路径逐字节一致。"""
    kwargs = {"cancel": cancel}
    if budget is not None:
        kwargs["budget"] = budget
    d = session.http_get_json(
        f"https://api.bilibili.com/x/web-interface/view?bvid={bvid}", **kwargs)
    if d.get("code") != 0:
        raise ValueError(f"{bvid}: code={d.get('code')} {d.get('message')}")
    v = d["data"]
    st = v.get("stat") or {}
    return {
        "bvid": v.get("bvid"),
        "aid": v.get("aid"),
        "title": v.get("title", ""),
        "owner": (v.get("owner") or {}).get("name", ""),
        "owner_mid": (v.get("owner") or {}).get("mid"),
        "tname": v.get("tname", ""),
        "pubdate": v.get("pubdate"),
        "duration": v.get("duration"),
        "view": st.get("view"), "danmaku": st.get("danmaku"), "reply": st.get("reply"),
        "favorite": st.get("favorite"), "coin": st.get("coin"),
        "share": st.get("share"), "like": st.get("like"),
        "fetched_at": int(time.time()),
    }


def collect_snapshot(bvids, sleep=0.3, progress=None, cancel=None,
                     snapshot_path=None, budget=None):
    """逐视频采集一份快照；返回 (成功列表, 失败列表[(bvid, err)])。

    cancel 会一路透传到请求内部：不仅在本循环间生效，也能打断 BiliClient 的
    退避等待。取消属主动行为，不计入失败列表——干净收尾并返回已采到的部分。

    budget 为任务预算（core.budget.TaskBudget），None = 无预算。视频循环
    边界查 expired()；请求在 HTTP 层入口被预算硬拦时同样按正常收尾处理
    （break，不计入失败列表），落盘后以快照条数记账 observe_records(1)。
    """
    ok, fail = [], []
    total = len(bvids)
    for i, bv in enumerate(bvids, 1):
        if cancel and cancel():
            break
        if budget is not None and budget.expired():
            break
        try:
            kwargs = {"cancel": cancel}
            if budget is not None:
                kwargs["budget"] = budget
            snap = fetch_view(bv, **kwargs)
            ok.append(snap)
            if snapshot_path:
                with open(snapshot_path, "a", encoding="utf-8") as fh:
                    fh.write(json.dumps(snap, ensure_ascii=False) + "\n")
            if budget is not None:
                budget.observe_records(1)
        except RiskChallengeError as e:
            e.resume_index = i - 1  # 断点：下一轮从此视频续采（0-based）
            raise
        except TaskCancelledError:
            break                   # 退避途中被取消：不是失败，直接收尾
        except BudgetExhaustedError:
            break                   # 预算硬拦：正常完成语义，不是失败
        except Exception as e:  # noqa: BLE001
            fail.append((bv, str(e)[:80]))
        if progress:
            progress(done=i, total=total, ok=len(ok), fail=len(fail),
                     text=f"采集 {i}/{total}: {bv}（成功{len(ok)} 失败{len(fail)}）")
        if not wait_or_cancel(sleep + random.random() * 0.15, cancel):
            break                   # 视频间隔也可被打断
    return ok, fail


# ============================ 分析 ============================

def analyze_snapshot(snaps):
    n = len(snaps)
    views = sorted(s["view"] for s in snaps
                   if isinstance(s.get("view"), int) and not isinstance(s.get("view"), bool))
    eng = sorted(s["like"] / s["view"] for s in snaps
                 if isinstance(s.get("view"), int) and s["view"] > 0
                 and isinstance(s.get("like"), int)
                 and not isinstance(s.get("like"), bool))
    days = Counter()
    for snapshot in snaps:
        pubdate = snapshot.get("pubdate")
        if pubdate is None:
            continue
        try:
            parsed_pubdate = xlsx_mod.unix_seconds_to_excel_datetime(pubdate)
        except (TypeError, ValueError, OverflowError, OSError):
            continue
        days[parsed_pubdate.strftime("%Y")] += 1
    kpi = {
        "count": n, "total_view": sum(views),
        "median_view": views[len(views) // 2] if views else 0,
        "median_engage": f"{eng[len(eng) // 2] * 100:.2f}%" if eng else "-",
        "median_engage_ratio": eng[len(eng) // 2] if eng else None,
        "top_view": snaps[0] if snaps else None,
        "years": sorted(days.items()),
    }
    return kpi


def analyze_growth(snapshots_by_bvid):
    """监控模式：每个 bvid 首末快照对比 → 增速榜。"""
    growth = []
    for bvid, snaps in snapshots_by_bvid.items():
        if len(snaps) < 2:
            continue
        first, last = snaps[0], snaps[-1]
        required = ("fetched_at", "view", "like", "coin", "favorite")
        if any(key not in first or key not in last
               or not isinstance(first[key], (int, float))
               or isinstance(first[key], bool)
               or not isinstance(last[key], (int, float))
               or isinstance(last[key], bool)
               for key in required):
            continue
        dt_h = max((last["fetched_at"] - first["fetched_at"]) / 3600, 1 / 60)
        d_view = last["view"] - first["view"]
        growth.append({
            "bvid": bvid, "title": last.get("title", ""), "owner": last.get("owner", ""),
            "first_view": first["view"], "last_view": last["view"],
            "d_view": d_view,
            "d_like": last["like"] - first["like"],
            "d_coin": last["coin"] - first["coin"],
            "d_fav": last["favorite"] - first["favorite"],
            "hours": round(dt_h, 2),
            "view_per_hour": round(d_view / dt_h, 1),
        })
    growth.sort(key=lambda g: -g["view_per_hour"])
    return growth


# ============================ Excel 导出 ============================

def export_xlsx(snaps, growth, meta, out_path, progress=None):
    if not HAVE_XLSX:
        raise RuntimeError("缺少 openpyxl（pip install openpyxl）")

    def state(kind, label):
        return xlsx_mod.cell_value(None, kind, note=label)

    def field(mapping, key, kind, label, number_format=None):
        if key not in mapping:
            return state(xlsx_mod.CellKind.NOT_RETURNED, f"{label}未返回")
        if mapping[key] is None:
            return state(xlsx_mod.CellKind.MISSING, f"{label}缺失")
        if kind is xlsx_mod.CellKind.DATETIME:
            return xlsx_mod.unix_seconds_cell_value(mapping[key], note=f"{label}格式异常")
        return xlsx_mod.checked_cell_value(mapping[key], kind,
                                           number_format=number_format,
                                           note=f"{label}格式异常")

    def value(raw, kind, label, number_format=None):
        if raw is None:
            return state(xlsx_mod.CellKind.MISSING, f"{label}缺失")
        return xlsx_mod.checked_cell_value(raw, kind, number_format=number_format,
                                           note=f"{label}格式异常")

    def metric(raw):
        return raw if isinstance(raw, (int, float)) and not isinstance(raw, bool) else 0

    n = len(snaps)
    wb = xlsx_mod.new_workbook()
    sw = xlsx_mod.SheetWriter(wb)

    views = sorted(s["view"] for s in snaps
                   if isinstance(s.get("view"), (int, float))
                   and not isinstance(s.get("view"), bool))
    eng = sorted((s["like"] / s["view"] for s in snaps
                  if isinstance(s.get("view"), int) and s["view"] > 0
                  and isinstance(s.get("like"), int)))
    total_view = sum(views) if views else None

    # Sheet1 概览
    ws = wb.create_sheet("采集概览")
    sw.ws = ws
    sw.title_row(ws, f"视频批量采集报告 | {meta.get('title', '')} 共{n}个视频", 6)
    sw.kv(ws, [
        ("采集时间", xlsx_mod.checked_cell_value(
            datetime.now(xlsx_mod.ASIA_SHANGHAI).replace(tzinfo=None),
            xlsx_mod.CellKind.DATETIME, note="采集时间格式异常"), meta.get("source", "")),
        ("视频数量", value(n, xlsx_mod.CellKind.INTEGER, "视频数量"),
         f"成功率 {meta.get('ok', n)}/{meta.get('attempted', n)}"),
        ("播放总量", value(total_view, xlsx_mod.CellKind.INTEGER, "播放总量"),
         f"中位数 {views[len(views) // 2] if views else 0:,}"),
        ("互动率中位", value(eng[len(eng) // 2] if eng else None,
                              xlsx_mod.CellKind.PERCENT, "互动率中位"), "点赞/播放"),
        ("监控模式", value(meta.get("monitor", "否"), xlsx_mod.CellKind.TEXT, "监控模式"),
         f"轮次 {meta.get('rounds', 1)}，间隔 {meta.get('interval_min', '-')} 分钟"),
        ("增速榜第一", value(growth[0]["title"][:30] if growth else "-",
                              xlsx_mod.CellKind.TEXT, "增速榜第一"),
         f"{growth[0]['view_per_hour']:,}/小时" if growth else ""),
    ])
    sw.append([None, sw.wc("口径：游客通道公开数据采集；播放量等计数为B站页面口径。",
                            font=xlsx_mod.F_CAPTION)])
    append_sheet_directory(sw, ws, (
        ("视频总表", "视频总表"),
        ("增速榜", "增速榜"),
        ("数据质量", "数据质量"),
        ("字段说明", "字段说明"),
    ))

    # Sheet2 视频总表（按播放降序）
    ws2 = wb.create_sheet("视频总表")
    sw.ws = ws2
    headers = ["排名", "BV号", "标题", "UP主", "分区", "时长(秒)", "发布时间",
               "播放", "弹幕", "评论", "点赞", "投币", "收藏", "分享"]
    ws2.column_dimensions["A"].width = 3
    for i, h in enumerate(headers):
        ws2.column_dimensions[get_column_letter(i + 2)].width = {
            "排名": 6, "BV号": 14, "标题": 46, "UP主": 18, "分区": 12, "时长(秒)": 9,
            "发布时间": 17, "播放": 10, "弹幕": 9, "评论": 9, "点赞": 10,
            "投币": 9, "收藏": 9, "分享": 9}.get(h, 10)
    detail_layout = TableLayout(1, 2, len(headers) + 1)
    configure_table(ws2, detail_layout)
    sw.header_row(ws2, headers)
    for i, s in enumerate(sorted(snaps, key=lambda x: -metric(x.get("view")))):
        if progress and i % 5000 == 0:
            progress(text=f"写入 Excel {i}/{n}")
        sw.append([None, value(i + 1, xlsx_mod.CellKind.INTEGER, "排名"),
                   field(s, "bvid", xlsx_mod.CellKind.ID, "BV号"),
                   wrap_cell(sw, field(s, "title", xlsx_mod.CellKind.TEXT, "标题")),
                   field(s, "owner", xlsx_mod.CellKind.TEXT, "UP主"),
                   field(s, "tname", xlsx_mod.CellKind.TEXT, "分区"),
                   field(s, "duration", xlsx_mod.CellKind.INTEGER, "时长",
                         '#,##0" 秒"'),
                   field(s, "pubdate", xlsx_mod.CellKind.DATETIME, "发布时间"),
                   field(s, "view", xlsx_mod.CellKind.INTEGER, "播放"),
                   field(s, "danmaku", xlsx_mod.CellKind.INTEGER, "弹幕"),
                   field(s, "reply", xlsx_mod.CellKind.INTEGER, "评论"),
                   field(s, "like", xlsx_mod.CellKind.INTEGER, "点赞"),
                   field(s, "coin", xlsx_mod.CellKind.INTEGER, "投币"),
                   field(s, "favorite", xlsx_mod.CellKind.INTEGER, "收藏"),
                   field(s, "share", xlsx_mod.CellKind.INTEGER, "分享")])
    finish_table(ws2, detail_layout, n)

    # Sheet3 增速榜（监控模式）
    ws3 = wb.create_sheet("增速榜")
    sw.ws = ws3
    growth_layout = TableLayout(3, 2, 10)
    configure_table(ws3, growth_layout)
    sw.title_row(ws3, "增速榜（监控期增量，按播放/小时降序）", 9)
    sw.header_row(ws3, ["排名", "BV号", "标题", "UP主", "期初播放", "期末播放",
                        "播放增量", "点赞增量", "播放/小时"])
    for i, g in enumerate(growth or []):
        sw.append([None, value(i + 1, xlsx_mod.CellKind.INTEGER, "排名"),
                   field(g, "bvid", xlsx_mod.CellKind.ID, "BV号"),
                   wrap_cell(sw, field(g, "title", xlsx_mod.CellKind.TEXT, "标题")),
                   field(g, "owner", xlsx_mod.CellKind.TEXT, "UP主"),
                   field(g, "first_view", xlsx_mod.CellKind.INTEGER, "期初播放"),
                   field(g, "last_view", xlsx_mod.CellKind.INTEGER, "期末播放"),
                   field(g, "d_view", xlsx_mod.CellKind.INTEGER, "播放增量"),
                   field(g, "d_like", xlsx_mod.CellKind.INTEGER, "点赞增量"),
                   field(g, "view_per_hour", xlsx_mod.CellKind.DECIMAL, "播放/小时")])
    if not growth:
        sw.append([None, sw.wc("单次快照模式无增速数据（使用定时追踪模式可采集时序增量）",
                                 font=xlsx_mod.F_CAPTION)])
    finish_table(ws3, growth_layout, len(growth or []))
    def q(item, raw, kind, unit, note, presentation_state=None):
        value = (xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE, note=note)
                 if raw is None else xlsx_mod.checked_cell_value(raw, kind, note=note))
        return QualityItem("采集", item, value, unit, note, presentation_state)

    quality = [
        q("候选视频总数", meta.get("attempted"), xlsx_mod.CellKind.INTEGER, "个", "本轮实际尝试的视频数量"),
        q("实际写入视频总表的有效记录数", n, xlsx_mod.CellKind.INTEGER, "个", "按视频总表业务结构写入"),
        q("本轮成功数", meta.get("ok"), xlsx_mod.CellKind.INTEGER, "个", "来自结构化本轮统计"),
        q("本轮失败数", meta.get("failed"), xlsx_mod.CellKind.INTEGER, "个", "来自结构化本轮统计"),
        q("检测到的重复数", None, xlsx_mod.CellKind.INTEGER, "条", "当前生产管线未提供重复判定，不能扫描猜测"),
        q("有效时间非法值数量", None, xlsx_mod.CellKind.INTEGER, "条", "当前导出上下文未传入该统计"),
        q("是否取消", meta.get("cancelled", False), xlsx_mod.CellKind.BOOLEAN, "状态", "来自结构化任务状态", "stop"),
        q("是否预算到限", meta.get("stopped_reason") == "budget_reached", xlsx_mod.CellKind.BOOLEAN, "状态", "来自 stopped_reason", "stop"),
        q("是否部分成功", bool(snaps) and (bool(meta.get("failed", 0))
          or bool(meta.get("round_incomplete"))
          or bool(meta.get("cancelled"))
          or meta.get("stopped_reason") == "budget_reached"),
          xlsx_mod.CellKind.BOOLEAN, "状态", "仅在存在有效视频记录且任务未完整结束时为真", "warning"),
        q("终态可观测", meta.get("terminal_known"), xlsx_mod.CellKind.BOOLEAN, "状态", "未传入时标记上游状态已丢失"),
    ]
    key_stats = meta.get("key_stats") if isinstance(meta.get("key_stats"), dict) else {}
    quality.extend(primary_key_quality(
        "主键", "bvid", denominator=key_stats.get("candidate_records"),
        missing=key_stats.get("missing"), invalid=key_stats.get("invalid"),
        duplicates=key_stats.get("duplicates"),
        dedup_discarded=key_stats.get("dedup_discarded"),
        remaining_conflicts=key_stats.get("remaining_conflicts"),
        source_note="采集器结构化快照统计；不扫描视频总表",
    ))
    fields = []
    def add_field(sheet, display, stable, dtype, metric):
        fields.append(FieldDefinition(sheet, display, stable, dtype, "", "是", "视频公开接口", metric,
                                      xlsx_mod.cell_value(None, xlsx_mod.CellKind.NOT_APPLICABLE), "接口未返回"))
    for display, stable, dtype, metric in (
        ("排名", "row_number", "整数", "按播放降序"), ("BV号", "bvid", "ID", "视频唯一标识"),
        ("标题", "title", "文本", "接口返回标题"), ("UP主", "owner", "文本", "接口返回作者名"),
        ("分区", "tname", "文本", "接口返回分区"), ("时长(秒)", "duration", "整数", "视频时长秒数"),
        ("发布时间", "pubdate", "日期时间", "Unix 秒转换为 Asia/Shanghai"), ("播放", "view", "整数", "接口公开计数"),
        ("弹幕", "danmaku", "整数", "接口公开计数"), ("评论", "reply", "整数", "接口公开计数"),
        ("点赞", "like", "整数", "接口公开计数"), ("投币", "coin", "整数", "接口公开计数"),
        ("收藏", "favorite", "整数", "接口公开计数"), ("分享", "share", "整数", "接口公开计数"),
    ):
        add_field("视频总表", display, stable, dtype, metric)
    for display, stable, dtype, metric in (
        ("采集时间", "collected_at", "日期时间", "本次导出生成时间"),
        ("视频数量", "video_count", "整数", "本次快照有效视频数"),
        ("播放总量", "total_view", "整数", "有效播放值之和"),
        ("互动率中位", "median_engagement_rate", "百分比", "点赞/播放的中位数"),
        ("监控模式", "monitor_mode", "文本", "任务是否启用监控"),
        ("增速榜第一", "top_growth_title", "文本", "增速榜首视频标题"),
    ):
        add_field("采集概览", display, stable, dtype, metric)
    for display, stable, dtype, metric in (
        ("排名", "row_number", "整数", "按播放/小时降序"),
        ("BV号", "bvid", "ID", "视频唯一标识"),
        ("标题", "title", "文本", "接口返回标题"),
        ("UP主", "owner", "文本", "接口返回作者名"),
        ("期初播放", "first_view", "整数", "监控期首个观测播放量"),
        ("期末播放", "last_view", "整数", "监控期末个观测播放量"),
        ("播放增量", "d_view", "整数", "期末减期初"),
        ("点赞增量", "d_like", "整数", "期末减期初"),
        ("播放/小时", "view_per_hour", "小数", "播放增量/观测小时"),
    ):
        add_field("增速榜", display, stable, dtype, metric)
    metadata = make_metadata(
        tool="采集器", report_type="视频数据报告", parameters={
            "来源数量": xlsx_mod.cell_value(meta.get("source_count", 0), xlsx_mod.CellKind.INTEGER),
            "视频数量": xlsx_mod.cell_value(meta.get("attempted", n), xlsx_mod.CellKind.INTEGER),
            "监控开关": xlsx_mod.cell_value(str(meta.get("monitor", "否")), xlsx_mod.CellKind.TEXT),
            "轮次": xlsx_mod.cell_value(meta.get("rounds", 1), xlsx_mod.CellKind.INTEGER),
            "间隔分钟": xlsx_mod.checked_cell_value(meta.get("interval_min", 0), xlsx_mod.CellKind.DECIMAL),
        }, parameter_allowlist=("来源数量", "视频数量", "监控开关", "轮次", "间隔分钟"),
        quality_items=quality, fields=fields)
    write_metadata_sheets(wb, metadata)
    xlsx_mod.save_workbook_atomic(wb, out_path)
